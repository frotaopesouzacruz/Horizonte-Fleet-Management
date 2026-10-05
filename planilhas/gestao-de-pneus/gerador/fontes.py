"""Localiza as planilhas de origem e extrai delas os dados de referência usados na geração.

Nada operacional fica fixo no código. Daqui saem, a cada geração:
  * custos de referência por medida (compras do Rodopar 978 e reformas da Análise CPK);
  * De-Para Filial + Unidade → Operação (aba De Para da planilha Controle de Pneus);
  * data da alocação de frotas mais recente (Base de Fidelização);
  * recapadoras usadas nos lotes (histórico de ressolagem);
  * caminho do Rodopar 10 gravado na consulta Power Query.
"""
import base64, collections, fnmatch, functools, io, os, re, statistics, struct, urllib.parse, warnings, zipfile
import datetime as dt
import openpyxl
import config

warnings.filterwarnings('ignore')

# chave: (padrão do nome do arquivo, descrição) — o padrão ignora maiúsculas/minúsculas
PADROES = {
    'rodopar': ('*rodopar*10*.xlsx', 'exportação Rodopar 10 (base mestre de pneus)'),
    'controle': ('*controle*pneus*.xlsx', 'Controle de Pneus OPE (histórico de ressolagem, De Para, alocação de frotas)'),
    'recapagem': ('*demanda*recapagem*.xlsx', 'Demanda de Recapagem (observações do lote mais recente)'),
    'cpk': ('*cpk*.xlsx', 'Gestão de CPK (cadastro de frotas, compras Rodopar 978, custos de reforma)'),
}
PREFIXO_UPLOAD = re.compile(r'^[0-9a-f]{8}-', re.I)  # prefixo que alguns sistemas acrescentam a anexos


@functools.lru_cache(None)
def arquivo(chave):
    padrao, descricao = PADROES[chave]
    if not os.path.isdir(config.FONTES):
        raise SystemExit(f'Pasta de fontes não encontrada: {config.FONTES} (defina GP_FONTES)')
    achados = [os.path.join(config.FONTES, n) for n in os.listdir(config.FONTES)
               if fnmatch.fnmatch(n.lower(), padrao) and not n.startswith('~$')]
    if not achados:
        raise SystemExit(f'Fonte não encontrada: {descricao}\n  procurei "{padrao}" em {config.FONTES}')
    escolhido = max(achados, key=os.path.getmtime)
    if len(achados) > 1:
        print(f'aviso: {len(achados)} arquivos para "{padrao}"; usando o mais recente: {os.path.basename(escolhido)}')
    return escolhido


def nome_limpo(path):
    return PREFIXO_UPLOAD.sub('', os.path.basename(path))


def arquivo_fonte_inicial():
    """Texto da coluna "Arquivo Fonte" na carga inicial (a consulta grava o caminho real a cada atualização)."""
    return f'Carga inicial: {nome_limpo(arquivo("rodopar"))} ({config.CARGA_INICIAL:%d/%m/%Y})'


@functools.lru_cache(None)
def _wb(chave, somente_leitura=True):
    return openpyxl.load_workbook(arquivo(chave), read_only=somente_leitura, data_only=True)


def _cabecalho(h):
    return re.sub(r'\s+', ' ', str(h)).strip() if h is not None else None


def _medida(v):
    import model  # import tardio: model importa este módulo
    return model.fn_medida(v)


# ------------------------------------------------------------------ custos de referência
@functools.lru_cache(None)
def custos_novos():
    """Pneu novo: preço médio da compra mais recente de cada medida (aba "Custos Pneus - Rodopar 978" do arquivo de CPK).
    Compras com data posterior à carga são ignoradas (data digitada errada)."""
    ws = _wb('cpk')['Custos Pneus - Rodopar 978']
    linhas = ws.iter_rows(values_only=True)
    idx = {_cabecalho(h): i for i, h in enumerate(next(linhas))}
    i_dim, i_data, i_preco = idx['Dimensão'], idx['Data da Compra'], idx['Preço']
    i_marca, i_modelo = idx['Marca'], idx['Modelo']
    por_medida = collections.defaultdict(list)
    futuras = collections.Counter()
    for r in linhas:
        med, data, preco = _medida(r[i_dim]), r[i_data], r[i_preco]
        if med is None or not isinstance(data, dt.datetime) or not isinstance(preco, (int, float)) or preco <= 0:
            continue
        if data > config.CARGA_INICIAL:
            futuras[med] += 1
            continue
        por_medida[med].append((data, float(preco), f'{str(r[i_marca]).strip().title()} {str(r[i_modelo]).strip()}'))
    out = {}
    for med, compras in por_medida.items():
        ultima = max(c[0] for c in compras)
        lote = [c for c in compras if c[0] == ultima]
        produto = collections.Counter(c[2] for c in lote).most_common(1)[0][0]
        out[med] = dict(valor=round(statistics.mean(c[1] for c in lote), 2), data=ultima, produto=produto,
                        n=len(lote), futuras=futuras.get(med, 0))
    return out


@functools.lru_cache(None)
def custos_ressolagem():
    """Ressolagem: valor mais frequente de "Custo Reforma" (2ª a 4ª vida) na aba "Análise CPK"; empate → maior valor."""
    ws = _wb('cpk')['Análise CPK']
    linhas = list(ws.iter_rows(values_only=True))
    hr = next(i for i, r in enumerate(linhas) if 'N.Fogo' in r and 'Dimensão' in r)
    hdr = [_cabecalho(h) for h in linhas[hr]]
    i_dim = hdr.index('Dimensão')
    i_custos = [i for i, h in enumerate(hdr) if h and h.startswith('Custo Reforma')]
    valores = collections.defaultdict(collections.Counter)
    for r in linhas[hr + 1:]:
        med = _medida(r[i_dim])
        if med is None:
            continue
        for i in i_custos:
            v = r[i]
            if isinstance(v, (int, float)) and v > 0:
                valores[med][float(v)] += 1
    out = {}
    for med, cont in valores.items():
        valor, n = max(cont.items(), key=lambda kv: (kv[1], kv[0]))
        out[med] = dict(valor=round(valor, 2), n=n, total=sum(cont.values()))
    return out


@functools.lru_cache(None)
def medidas():
    """Tabela de medidas: regras de config.MEDIDAS_REGRAS + custos de referência calculados das fontes.
    Formato: (Medida, Utilização, Vida Máxima, Limite Específico, Custo Novo, Custo Ressolagem, Referência)."""
    novos, ress = custos_novos(), custos_ressolagem()
    out = []
    for med, util, vmax, lim, obs in config.MEDIDAS_REGRAS:
        ref, cn, cr = [], None, None
        if med in novos:
            c = novos[med]
            cn = c['valor']
            ref.append(f'Novo: preço médio da última compra no Rodopar 978 ({c["data"]:%d/%m/%Y}, {c["produto"]}, '
                       f'{c["n"]} pneu{"s" if c["n"] > 1 else ""}).')
            if c['futuras']:
                ref.append(f'{c["futuras"]} compra(s) com data futura ignorada(s).')
        else:
            ref.append('Novo: sem compra registrada no Rodopar 978.')
        if vmax <= 1:
            ref.append('Medida não ressolada (descarte na 1ª vida).')
        elif med in ress:
            c = ress[med]
            cr = c['valor']
            ref.append(f'Ressolagem: valor mais frequente na Análise CPK ({c["n"]} de {c["total"]} reformas).')
        else:
            ref.append('Sem histórico de custo de ressolagem.')
        if obs:
            ref.append(obs)
        out.append((med, util, vmax, lim, cn, cr, ' '.join(ref)))
    return out


# ------------------------------------------------------------------ cadastros auxiliares
@functools.lru_cache(None)
def centros_custo():
    """De-Para Filial Pneu + Cód Unidade → Operação (aba "De Para" da planilha Controle de Pneus)."""
    ws = _wb('controle')['De Para']
    linhas = list(ws.iter_rows(values_only=True))
    for hr, r in enumerate(linhas):
        hdr = [_cabecalho(h) for h in r]
        if 'Filial Pneu' in hdr and 'Operação' in hdr:
            break
    else:
        raise SystemExit('Aba "De Para": cabeçalho com "Filial Pneu" e "Operação" não encontrado')
    i_fil, i_uni, i_op = hdr.index('Filial Pneu'), hdr.index('Cód Unidade'), hdr.index('Operação')
    out = set()
    for r in linhas[hr + 1:]:
        fil, uni, op = r[i_fil], r[i_uni], r[i_op]
        if isinstance(fil, (int, float)) and isinstance(uni, (int, float)) and op:
            out.add((int(fil), int(uni), str(op).strip()))
    return sorted(out)


@functools.lru_cache(None)
def data_alocacao():
    """Data da alocação de frotas usada na aba Frotas: a mais recente até o dia da carga na aba Base_Dados_Fidelização
    (a base costuma vir preenchida para o mês inteiro, inclusive dias futuros)."""
    ws = _wb('controle', False)['Base_Dados_Fidelização']
    datas = {ws.cell(r, 8).value.date() for r in range(10, ws.max_row + 1) if isinstance(ws.cell(r, 8).value, dt.datetime)}
    ate_hoje = [d for d in datas if d <= config.CARGA_INICIAL.date()]
    return max(ate_hoje) if ate_hoje else min(datas)


def _parecido(a, b):
    """Grafias que diferem em no máximo uma letra (ex.: "Recatac" × "Recatec")."""
    a, b = a.casefold(), b.casefold()
    if a == b:
        return True
    if abs(len(a) - len(b)) > 1:
        return False
    if len(a) == len(b):
        return sum(x != y for x, y in zip(a, b)) <= 1
    if len(a) > len(b):
        a, b = b, a
    return any(b[:i] + b[i + 1:] == a for i in range(len(b)))


def recapadoras(historico):
    """Lista sugerida de recapadoras: nomes usados nos lotes; variações de grafia ficam com a forma mais usada.
    Retorna (nomes, variações) — variações = [(grafia encontrada, forma adotada)]."""
    cont = collections.Counter(str(f['Recapadora']).strip() for f in historico if f.get('Recapadora'))
    nomes, variacoes = [], []
    for nome, _ in cont.most_common():
        igual = next((n for n in nomes if _parecido(nome, n)), None)
        if igual is None:
            nomes.append(nome)
        elif nome != igual:
            variacoes.append((nome, igual))
    return sorted(nomes, key=str.casefold), variacoes


# ------------------------------------------------------------------ caminho do Rodopar 10
def _consultas_m(path):
    """Código M (Power Query) de uma pasta de trabalho, lido do DataMashup em customXml."""
    z = zipfile.ZipFile(path)
    for n in z.namelist():
        if not (n.startswith('customXml/item') and n.endswith('.xml')):
            continue
        raw = z.read(n)
        texto = raw.decode('utf-16') if raw[:2] in (b'\xff\xfe', b'\xfe\xff') else raw.decode('utf-8', 'ignore')
        m = re.search(r'<DataMashup[^>]*>(.*)</DataMashup>', texto, re.S)
        if not m:
            continue
        blob = base64.b64decode(m.group(1))
        tam = struct.unpack('<I', blob[4:8])[0]
        pacote = zipfile.ZipFile(io.BytesIO(blob[8:8 + tam]))
        return pacote.read('Formulas/Section1.m').decode('utf-8-sig')
    return ''


@functools.lru_cache(None)
def caminho_rodopar():
    """(caminho, origem) gravados na etapa CaminhoArquivo da consulta.
    1. variável GP_CAMINHO_RODOPAR; 2. "Rodopar 10.xlsx" na mesma pasta do SharePoint em que a planilha Controle de Pneus
    se publica (link lido das consultas dela); 3. exemplo de caminho local, a ajustar na primeira atualização."""
    if os.environ.get('GP_CAMINHO_RODOPAR'):
        return os.environ['GP_CAMINHO_RODOPAR'], 'definido na geração'
    for url in re.findall(r'"(https?://[^"]+)"', _consultas_m(arquivo('controle'))):
        pasta, nome = url.rsplit('/', 1)
        nome = urllib.parse.unquote(nome).casefold()
        if 'controle' in nome and 'pneus' in nome:
            return pasta + '/Rodopar%2010.xlsx', 'sharepoint'
    return r'C:\Pneus\Rodopar 10.xlsx', 'exemplo'
