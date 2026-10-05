"""Abas Análises (visões gerenciais) e Auditoria (qualidade da base)."""
from openpyxl.formatting.rule import DataBarRule, FormulaRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from views_common import *
from build_views import MED, MEDINFO

def fixed_block(ws, row, c1, title, headers, lines, last_col=None, first_w=None):
    """lines: lista de (rótulo|fórmula_rotulo, rótulo_esperado, [(fórmula, esperado, nf), ...])"""
    section(ws, f'{col(c1)}{row}', title, last_col or (c1 + len(headers) - 1), row_h=18)
    for i, h in enumerate(headers):
        put(ws, f'{col(c1+i)}{row+1}', h)
    table_header_style(ws, row + 1, c1, c1 + len(headers) - 1, h=18)
    r = row + 1
    for lab, lab_exp, cells in lines:
        r += 1
        if isinstance(lab, str) and lab.startswith('='):
            put(ws, f'{col(c1)}{r}', lab, f=font(9.5, True), a=al('left'), expected=lab_exp)
        else:
            put(ws, f'{col(c1)}{r}', lab, f=font(9.5, True), a=al('left'))
        for j, (fm, ex, nf) in enumerate(cells):
            put(ws, f'{col(c1+1+j)}{r}', fm, f=font(9.5), a=al('center'), nf=nf, expected=ex)
        for ci in range(c1, c1 + len(headers)):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
    return r

def dyn_block(ws, row, c1, title, headers, key_formula, key_vals, metric_formulas, metric_vals, fmts, ncols_title=None):
    """Bloco dinâmico: chave(s) via ÚNICO/CLASSIFICAR + métricas que derramam junto (referência #)."""
    section(ws, f'{col(c1)}{row}', title, c1 + len(headers) - 1, row_h=18)
    for i, h in enumerate(headers):
        put(ws, f'{col(c1+i)}{row+1}', h)
    table_header_style(ws, row + 1, c1, c1 + len(headers) - 1, h=18)
    anchor = f'{col(c1)}{row+2}'
    dyn(ws, anchor, key_formula, key_vals)
    nkey = len(key_vals[0]) if key_vals else 1
    for j, (fm, vals) in enumerate(zip(metric_formulas, metric_vals)):
        dyn(ws, f'{col(c1+nkey+j)}{row+2}', fm.replace('{A}', f'_xlfn.ANCHORARRAY({anchor})'), [[v] for v in vals])
    n = len(key_vals)
    for r in range(row + 2, row + 2 + max(n, 1) + 25):
        for i, (fmt, a_) in enumerate(fmts):
            c = ws.cell(r, c1 + i)
            c.number_format = fmt; c.alignment = al(a_, 'center'); c.font = font(9.5, i < nkey)
    rng = f'{col(c1)}{row+2}:{col(c1+len(headers)-1)}{row+400}'
    a0 = f'${col(c1)}{row+2}'
    ws.conditional_formatting.add(rng, lowp(ws, FormulaRule(formula=[f'AND({a0}<>"",MOD(ROW(),2)=0)'], fill=fill(C['surface2']))))
    ws.conditional_formatting.add(rng, lowp(ws, FormulaRule(formula=[f'{a0}<>""'], border=Border(bottom=Side(style='thin', color=C['border'])))))
    return row + 2

def uniq_sorted(values):
    seen = []
    for v in values:
        if v not in seen: seen.append(v)
    return sorted(seen, key=coll)

# ======================================================================= ANÁLISES
def build_analises(wb, ws, rows, frotas_calc_rows):
    sheet_base(ws, C['blue'])
    widths(ws, {'B': 25, 'C': 11, 'D': 11, 'E': 11, 'F': 12, 'G': 3, 'H': 24, 'I': 18, 'J': 11, 'K': 11, 'L': 11, 'M': 12,
                'N': 12, 'O': 3, 'P': 9, 'Q': 10, 'R': 14, 'S': 20, 'T': 9, 'U': 11, 'V': 9, 'W': 11, 'X': 9, 'Y': 11})
    header(ws, 'Análises gerenciais', 'Distribuição da frota de pneus por faixa de profundidade, situação, vida, eixo, '
           'classificação, operação, local, fabricante/modelo e veículo.', 25)
    ativos = sel(rows, Ativo='SIM')
    lc, lr, lp, lf = P['pLimCritico'], P['pLimRessolagem'], P['pLimCompra'], P['pLimFaixa']
    fx = M.fixed_pt
    def mmc(lo, hi, extra=None):
        cs = []
        for r in ativos:
            mm = r['Menor Milimetragem']
            if not M.isnum(mm): continue
            if lo is not None and mm < lo: continue
            if hi is not None and mm >= hi: continue
            if extra and not extra(r): continue
            cs.append(r)
        return len(cs)
    na = len(ativos)
    faixas = [
        ('="< "&FIXED(pLimCritico,2)&" mm"', f'< {fx(lc)} mm', None, 'pLimCritico', None, lc),
        ('=FIXED(pLimCritico,2)&" a < "&FIXED(pLimRessolagem,2)&" mm"', f'{fx(lc)} a < {fx(lr)} mm', 'pLimCritico', 'pLimRessolagem', lc, lr),
        ('=FIXED(pLimRessolagem,2)&" a < "&FIXED(pLimCompra,2)&" mm"', f'{fx(lr)} a < {fx(lp)} mm', 'pLimRessolagem', 'pLimCompra', lr, lp),
        ('=FIXED(pLimCompra,2)&" a < "&FIXED(pLimFaixa,2)&" mm"', f'{fx(lp)} a < {fx(lf)} mm', 'pLimCompra', 'pLimFaixa', lp, lf),
        ('="≥ "&FIXED(pLimFaixa,2)&" mm"', f'≥ {fx(lf)} mm', 'pLimFaixa', None, lf, None),
    ]
    lines = []
    for lab, labx, plo, phi, vlo, vhi in faixas:
        crit = ''
        if plo: crit += f',tbPneus[Menor Milimetragem],">="&{plo}'
        if phi: crit += f',tbPneus[Menor Milimetragem],"<"&{phi}'
        t = mmc(vlo if plo else None, vhi if phi else None)
        u = mmc(vlo if plo else None, vhi if phi else None, lambda r: r['Situação Pneu'] == 'USO')
        e = mmc(vlo if plo else None, vhi if phi else None, lambda r: r['Situação Pneu'] == 'ESTOQUE')
        lines.append((lab, labx, [
            (f'=COUNTIFS(tbPneus[Ativo],"SIM"{crit})', t, '#,##0'),
            (f'=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[Situação Pneu],"USO"{crit})', u, '#,##0'),
            (f'=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[Situação Pneu],"ESTOQUE"{crit})', e, '#,##0'),
            (f'=IFERROR(COUNTIFS(tbPneus[Ativo],"SIM"{crit})/COUNTIF(tbPneus[Ativo],"SIM"),0)', t / na if na else 0, '0.0%')]))
    sm = sum(1 for r in ativos if r['Faixa MM'] == 'Sem medição')
    lines.append(('Sem medição', None, [
        ('=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[Faixa MM],"Sem medição")', sm, '#,##0'),
        ('=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[Faixa MM],"Sem medição",tbPneus[Situação Pneu],"USO")',
         sum(1 for r in ativos if r['Faixa MM'] == 'Sem medição' and r['Situação Pneu'] == 'USO'), '#,##0'),
        ('=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[Faixa MM],"Sem medição",tbPneus[Situação Pneu],"ESTOQUE")',
         sum(1 for r in ativos if r['Faixa MM'] == 'Sem medição' and r['Situação Pneu'] == 'ESTOQUE'), '#,##0'),
        ('=IFERROR(COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[Faixa MM],"Sem medição")/COUNTIF(tbPneus[Ativo],"SIM"),0)', sm / na if na else 0, '0.0%')]))
    r = fixed_block(ws, 6, 2, 'Faixa de profundidade — pneus ativos', ['Faixa (Menor MM)', 'Total', 'Em uso', 'Estoque', '% ativos'], lines)
    faixa_rows = (8, r)
    # cores das faixas
    tones = [C['danger'], C['danger_soft'], C['warning_soft'], C['success_soft'], C['success_soft']]
    for k, tn in enumerate(tones):
        ws.cell(8 + k, 2).fill = fill(tn)
        if k == 0: ws.cell(8 + k, 2).font = font(9.5, True, 'FFFFFF')
    # Situação
    sits = [s[0] for s in M.SITUACOES]
    lines = []
    tot_all = len(rows)
    for k, s in enumerate(sits):
        lab = f'=INDEX(tbSituacoes[Situação Pneu],{k+1})'
        t = cnt(rows, **{'Situação__Pneu': s})
        v1 = sum(1 for x in rows if x['Situação Pneu'] == s and x['N. Vida'] == 1)
        vr = sum(1 for x in rows if x['Situação Pneu'] == s and M.isnum(x['N. Vida']) and x['N. Vida'] >= 2)
        lines.append((lab, s, [
            (f'=COUNTIF(tbPneus[Situação Pneu],B{r+4+k})', t, '#,##0'),
            (f'=COUNTIFS(tbPneus[Situação Pneu],B{r+4+k},tbPneus[N. Vida],1)', v1, '#,##0'),
            (f'=COUNTIFS(tbPneus[Situação Pneu],B{r+4+k},tbPneus[N. Vida],">=2")', vr, '#,##0'),
            (f'=IFERROR(COUNTIF(tbPneus[Situação Pneu],B{r+4+k})/ROWS(tbPneus[N.Fogo]),0)', t / tot_all, '0.0%')]))
    r = fixed_block(ws, r + 2, 2, 'Situação do pneu — toda a base', ['Situação', 'Total', '1ª vida', 'Ressolados', '% base'], lines)
    # Vida
    lines = []
    for v, lab in [(1, '1ª vida'), (2, '2ª vida'), (3, '3ª vida'), (4, '4ª vida ou mais')]:
        if v < 4:
            crit = f'tbPneus[N. Vida],{v}'; f_ = lambda x, v=v: x['N. Vida'] == v
        else:
            crit = 'tbPneus[N. Vida],">=4"'; f_ = lambda x: M.isnum(x['N. Vida']) and x['N. Vida'] >= 4
        a_ = [x for x in ativos if f_(x)]
        lines.append((lab, None, [
            (f'=COUNTIFS(tbPneus[Ativo],"SIM",{crit})', len(a_), '#,##0'),
            (f'=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[Situação Pneu],"USO",{crit})', sum(x['Situação Pneu'] == 'USO' for x in a_), '#,##0'),
            (f'=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[Situação Pneu],"ESTOQUE",{crit})', sum(x['Situação Pneu'] == 'ESTOQUE' for x in a_), '#,##0'),
            (f'=COUNTIFS(tbPneus[Demanda Ressolagem],"SIM",{crit})', sum(x['Demanda Ressolagem'] == 'SIM' for x in a_), '#,##0')]))
    r = fixed_block(ws, r + 2, 2, 'Vida do pneu — pneus ativos', ['Vida', 'Ativos', 'Em uso', 'Estoque', 'Dem. ressol.'], lines)
    # Eixo
    lines = []
    uso = sel(rows, **{'Situação__Pneu': 'USO'})
    for e in ['Dianteiro', 'Traseiro', 'Estepe']:
        a_ = [x for x in uso if x['Eixo'] == e]
        lines.append((e, None, [
            (f'=COUNTIFS(tbPneus[Situação Pneu],"USO",tbPneus[Eixo],"{e}")', len(a_), '#,##0'),
            (f'=COUNTIFS(tbPneus[Situação Pneu],"USO",tbPneus[Eixo],"{e}",tbPneus[N. Vida],1)', sum(x['N. Vida'] == 1 for x in a_), '#,##0'),
            (f'=COUNTIFS(tbPneus[Situação Pneu],"USO",tbPneus[Eixo],"{e}",tbPneus[Demanda Compra],"SIM")', sum(x['Demanda Compra'] == 'SIM' for x in a_), '#,##0'),
            (f'=COUNTIFS(tbPneus[Situação Pneu],"USO",tbPneus[Eixo],"{e}",tbPneus[Demanda Ressolagem],"SIM")', sum(x['Demanda Ressolagem'] == 'SIM' for x in a_), '#,##0')]))
    r = fixed_block(ws, r + 2, 2, 'Eixo — pneus em uso', ['Eixo', 'Pneus', '1ª vida', 'Dem. compra', 'Dem. ressol.'], lines)
    # Classificação
    lines = []
    for cl in M.CLASSES:
        a_ = [x for x in rows if x['Classificação Desgaste'] == cl]
        lines.append((cl, None, [
            (f'=COUNTIF(tbPneus[Classificação Desgaste],"{cl}")', len(a_), '#,##0'),
            (f'=COUNTIFS(tbPneus[Classificação Desgaste],"{cl}",tbPneus[Situação Pneu],"USO")', sum(x['Situação Pneu'] == 'USO' for x in a_), '#,##0'),
            (f'=COUNTIFS(tbPneus[Classificação Desgaste],"{cl}",tbPneus[Situação Pneu],"ESTOQUE")', sum(x['Situação Pneu'] == 'ESTOQUE' for x in a_), '#,##0'),
            (f'=IFERROR(COUNTIF(tbPneus[Classificação Desgaste],"{cl}")/ROWS(tbPneus[N.Fogo]),0)', len(a_) / tot_all, '0.0%')]))
    r0 = r + 2
    r = fixed_block(ws, r0, 2, 'Classificação de desgaste — toda a base', ['Classificação', 'Total', 'Em uso', 'Estoque', '% base'], lines)
    class_cf(ws, f'B{r0+2}:B{r}')
    # Aferição
    lines = []
    for s_ in ['Em dia', 'Próx. do vencimento', 'Vencida', 'Sem medição']:
        a_ = [x for x in ativos if x['Status Aferição'] == s_]
        lines.append((s_, None, [
            (f'=COUNTIF(tbPneus[Status Aferição],"{s_}")', len(a_), '#,##0'),
            (f'=COUNTIFS(tbPneus[Status Aferição],"{s_}",tbPneus[Situação Pneu],"USO")', sum(x['Situação Pneu'] == 'USO' for x in a_), '#,##0'),
            (f'=COUNTIFS(tbPneus[Status Aferição],"{s_}",tbPneus[Situação Pneu],"ESTOQUE")', sum(x['Situação Pneu'] == 'ESTOQUE' for x in a_), '#,##0'),
            (f'=IFERROR(COUNTIF(tbPneus[Status Aferição],"{s_}")/COUNTIF(tbPneus[Ativo],"SIM"),0)', len(a_) / na, '0.0%')]))
    r0 = r + 2
    r = fixed_block(ws, r0, 2, 'Atualidade da medição — pneus ativos', ['Status da aferição', 'Total', 'Em uso', 'Estoque', '% ativos'], lines)
    text_cf(ws, f'B{r0+2}:B{r}', [('Vencida', 'danger'), ('Próx. do vencimento', 'warning'), ('Em dia', 'success')])
    put(ws, f'B{r+1}', '="Prazo: aviso acima de "&pAfAviso&" dias; vencida acima de "&pAfPrazo&" dias desde a última medição."',
        f=font(8, False, C['muted'], italic=True), expected=f"Prazo: aviso acima de {P['pAfAviso']} dias; vencida acima de {P['pAfPrazo']} dias desde a última medição.")
    # ---------------- blocos dinâmicos (coluna H em diante)
    # Operação
    ops = uniq_sorted([x['Operação'] for x in ativos])
    kf = '=_xlfn._xlws.SORT(_xlfn.UNIQUE(_xlfn._xlws.FILTER(tbPneus[Operação],tbPneus[Ativo]="SIM","—")))'
    mf = ['=COUNTIFS(tbPneus[Operação],{A},tbPneus[Ativo],"SIM")',
          '=COUNTIFS(tbPneus[Operação],{A},tbPneus[Situação Pneu],"USO")',
          '=COUNTIFS(tbPneus[Operação],{A},tbPneus[Situação Pneu],"ESTOQUE")',
          '=COUNTIFS(tbPneus[Operação],{A},tbPneus[Demanda Ressolagem],"SIM")',
          '=COUNTIFS(tbPneus[Operação],{A},tbPneus[Demanda Compra],"SIM")']
    mv = [[cnt(rows, Operação=o, Ativo='SIM') for o in ops],
          [cnt(rows, Operação=o, **{'Situação__Pneu': 'USO'}) for o in ops],
          [cnt(rows, Operação=o, **{'Situação__Pneu': 'ESTOQUE'}) for o in ops],
          [cnt(rows, Operação=o, **{'Demanda__Ressolagem': 'SIM'}) for o in ops],
          [cnt(rows, Operação=o, **{'Demanda__Compra': 'SIM'}) for o in ops]]
    fm = [('@', 'left')] + [('#,##0', 'center')] * 5
    dyn_block(ws, 6, 8, 'Por operação (centro de custo) — pneus ativos', ['Operação', 'Ativos', 'Em uso', 'Estoque', 'Dem. ressol.', 'Dem. compra'],
              kf, [[o] for o in ops], mf, mv, fm)
    # Local de operação (em uso)
    row_loc = 6 + 2 + 14
    locs = uniq_sorted([x['Local de Operação'] for x in uso])
    kf = '=_xlfn._xlws.SORT(_xlfn.UNIQUE(_xlfn._xlws.FILTER(tbPneus[Local de Operação],tbPneus[Situação Pneu]="USO","—")))'
    mf = ['=COUNTIFS(tbPneus[Local de Operação],{A},tbPneus[Situação Pneu],"USO")',
          '=COUNTIFS(tbPneus[Local de Operação],{A},tbPneus[Situação Pneu],"USO",tbPneus[Demanda Ressolagem],"SIM")',
          '=COUNTIFS(tbPneus[Local de Operação],{A},tbPneus[Situação Pneu],"USO",tbPneus[Demanda Compra],"SIM")',
          '=COUNTIFS(tbPneus[Local de Operação],{A},tbPneus[Situação Pneu],"USO",tbPneus[Status Aferição],"Vencida")']
    mv = [[sum(1 for x in uso if x['Local de Operação'] == l) for l in locs],
          [sum(1 for x in uso if x['Local de Operação'] == l and x['Demanda Ressolagem'] == 'SIM') for l in locs],
          [sum(1 for x in uso if x['Local de Operação'] == l and x['Demanda Compra'] == 'SIM') for l in locs],
          [sum(1 for x in uso if x['Local de Operação'] == l and x['Status Aferição'] == 'Vencida') for l in locs]]
    fm = [('@', 'left')] + [('#,##0', 'center')] * 4
    dyn_block(ws, row_loc, 8, 'Por local de operação (base do veículo) — pneus em uso',
              ['Local de operação', 'Pneus em uso', 'Dem. ressol.', 'Dem. compra', 'Afer. vencida'], kf, [[l] for l in locs], mf, mv, fm)
    # Fabricante e modelo
    row_mod = row_loc + 2 + 34
    pairs = []
    for x in ativos:
        p_ = (x['Marca'], x['Modelo Pneu'])
        if p_ not in pairs: pairs.append(p_)
    pairs = sorted(pairs, key=lambda p_: (coll(p_[0]), coll(p_[1])))
    kf = '=_xlfn._xlws.SORT(_xlfn.UNIQUE(_xlfn._xlws.FILTER(tbPneus[[Marca]:[Modelo Pneu]],tbPneus[Ativo]="SIM","—")),{1,2},{1,1})'
    A1 = 'INDEX({A},0,1)'; A2 = 'INDEX({A},0,2)'
    mf = [f'=COUNTIFS(tbPneus[Marca],{A1},tbPneus[Modelo Pneu],{A2},tbPneus[Ativo],"SIM")',
          f'=COUNTIFS(tbPneus[Marca],{A1},tbPneus[Modelo Pneu],{A2},tbPneus[Situação Pneu],"USO")',
          f'=COUNTIFS(tbPneus[Marca],{A1},tbPneus[Modelo Pneu],{A2},tbPneus[Situação Pneu],"ESTOQUE")',
          f'=COUNTIFS(tbPneus[Marca],{A1},tbPneus[Modelo Pneu],{A2},tbPneus[Demanda Ressolagem],"SIM")',
          f'=COUNTIFS(tbPneus[Marca],{A1},tbPneus[Modelo Pneu],{A2},tbPneus[Demanda Compra],"SIM")']
    def pc(p_, **k):
        return sum(1 for x in rows if x['Marca'] == p_[0] and x['Modelo Pneu'] == p_[1] and all(x.get(a.replace('__', ' ')) == b for a, b in k.items()))
    mv = [[pc(p_, Ativo='SIM') for p_ in pairs], [pc(p_, **{'Situação__Pneu': 'USO'}) for p_ in pairs],
          [pc(p_, **{'Situação__Pneu': 'ESTOQUE'}) for p_ in pairs], [pc(p_, **{'Demanda__Ressolagem': 'SIM'}) for p_ in pairs],
          [pc(p_, **{'Demanda__Compra': 'SIM'}) for p_ in pairs]]
    fm = [('@', 'left'), ('@', 'left')] + [('#,##0', 'center')] * 5
    dyn_block(ws, row_mod, 8, 'Por fabricante e modelo — pneus ativos',
              ['Fabricante', 'Modelo', 'Ativos', 'Em uso', 'Estoque', 'Dem. ressol.', 'Dem. compra'], kf, [list(p_) for p_ in pairs], mf, mv, fm)
    # Veículo: banda própria (P..Y), sem nada abaixo — cresce livremente
    fr_cols = ['Frota', 'Placa', 'Tipo', 'Local de Operação', 'Pneus Aplicados', 'Em Demanda Ressolagem', 'Críticos',
               'Em Demanda Compra', 'Menor MM', 'Aferições Vencidas']
    hdrs = ['Frota', 'Placa', 'Tipo', 'Local de operação', 'Pneus', 'Dem. ressol.', 'Críticos', 'Dem. compra', 'Menor MM', 'Afer. venc.']
    section(ws, 'P6', 'Por veículo — pneus aplicados (ordenado por criticidade)', 25, row_h=18)
    fmts = [('@', 'center'), ('@', 'center'), ('@', 'left'), ('@', 'left'), ('0', 'center'), ('0', 'center'), ('0', 'center'),
            ('0', 'center'), ('0.00', 'center'), ('0', 'center')]
    veic = [f for f in frotas_calc_rows if f['Pneus Aplicados'] > 0]
    list_block(ws, 7, 16, hdrs, fmts, 'tbFrotas[Pneus Aplicados]>0', fr_cols, [('Pontuação', -1)], veic,
               'Nenhum veículo com pneus aplicados.', table='tbFrotas', max_rows=400, header_h=18)
    red = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['danger_fg']), fill=fill(C['danger_soft']))
    amb = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['warning_fg']), fill=fill(C['warning_soft']))
    ws.conditional_formatting.add('U8:V408', Rule(type='cellIs', operator='greaterThan', formula=['0'], dxf=red))
    ws.conditional_formatting.add('W8:W408', Rule(type='cellIs', operator='greaterThan', formula=['0'], dxf=amb))
    ws.conditional_formatting.add('Y8:Y408', Rule(type='cellIs', operator='greaterThan', formula=['0'], dxf=amb))
    mm_cf(ws, 'X8:X408', 'X8')
    ws.freeze_panes = 'A5'
    return dict(faixa_rows=faixa_rows)

# ======================================================================= AUDITORIA
def build_auditoria(wb, ws, rows):
    from openpyxl.workbook.defined_name import DefinedName
    sheet_base(ws, C['blue'])
    widths(ws, {'B': 9, 'C': 10, 'D': 11, 'E': 9, 'F': 7, 'G': 7, 'H': 7, 'I': 7, 'J': 8, 'K': 18, 'L': 9, 'M': 7,
                'N': 9, 'O': 6, 'P': 58, 'Q': 3, 'R': 30})
    header(ws, 'Auditoria da base', 'Verificações automáticas de qualidade do cadastro Rodopar 10 e registro dos tratamentos '
           'aplicados pelo Power Query. Corrija na origem (Rodopar) e atualize.', 16)
    n = len(rows)
    with_inc = sum(1 for r in rows if r['Qtd Inconsistências'] > 0)
    occ = sum(r['Qtd Inconsistências'] for r in rows)
    high_cols = [a[0] for a in M.AUD if a[2] == 'Alta']
    high = sum(1 for r in rows if any(r[c] == 1 for c in high_cols))
    ativos = sel(rows, Ativo='SIM')
    venc = sum(1 for r in ativos if r['Status Aferição'] == 'Vencida')
    q = 1 - with_inc / n
    kpi(ws, 6, 2, 3, 'QUALIDADE DA BASE', '=IFERROR(1-COUNTIF(tbPneus[Qtd Inconsistências],">0")/ROWS(tbPneus[N.Fogo]),0)', q,
        nf='0.0%', tone='success' if q >= 0.9 else 'warning', note='registros sem inconsistência')
    kpi(ws, 6, 5, 5, 'REGISTROS COM INCONSISTÊNCIA', '=COUNTIF(tbPneus[Qtd Inconsistências],">0")', with_inc, tone='warning',
        note_formula='="de "&ROWS(tbPneus[N.Fogo])&" pneus"', note_expected=f'de {n} pneus')
    kpi(ws, 6, 10, 2, 'OCORRÊNCIAS', '=SUM(tbPneus[Qtd Inconsistências])', occ, tone='default', note='total de alertas')
    hsum = '+'.join(f'(tbPneus[{c}]=1)' for c in high_cols)
    kpi(ws, 6, 12, 4, 'AFETAM AS REGRAS', f'=SUMPRODUCT(--(({hsum})>0))', high, tone='danger', note='severidade alta')
    kpi(ws, 6, 16, 1, 'AFERIÇÃO VENCIDA (ATIVOS)', '=COUNTIF(tbPneus[Status Aferição],"Vencida")', venc, tone='warning',
        note_formula='="mais de "&pAfPrazo&" dias sem medição"', note_expected=f"mais de {P['pAfPrazo']} dias sem medição")
    section(ws, 'B10', 'Verificações de cadastro', 16)
    hr = 11
    for coord, h in [('B', 'Nº'), ('C', 'Verificação'), ('J', 'Severidade'), ('K', 'Escopo'), ('L', 'Ocorrências'), ('N', '% da base'),
                     ('P', 'Como corrigir')]:
        put(ws, f'{coord}{hr}', h)
    table_header_style(ws, hr, 2, 16, h=24, dividers=False)
    for ci in range(3, 10):
        ws.cell(hr, ci).alignment = Alignment(horizontal='centerContinuous', vertical='center')
    r = hr
    for k, (cname, lab, sev, esc, fix) in enumerate(M.AUD, start=1):
        r += 1
        oc = sum(x[cname] for x in rows)
        put(ws, f'B{r}', k, f=font(9), a=al('center'))
        put(ws, f'C{r}', lab, f=font(9.5, True), a=al('left', 'center'))
        put(ws, f'J{r}', sev, f=font(9), a=al('center'))
        put(ws, f'K{r}', esc, f=font(9), a=al('center'))
        put(ws, f'L{r}', f'=SUM(tbPneus[{cname}])', f=font(10, True), a=al('center'), nf='#,##0', expected=oc)
        ws[f'M{r}'].alignment = Alignment(horizontal='centerContinuous', vertical='center')
        ws[f'L{r}'].alignment = Alignment(horizontal='centerContinuous', vertical='center')
        put(ws, f'N{r}', f'=IFERROR(L{r}/ROWS(tbPneus[N.Fogo]),0)', f=font(9), a=al('center'), nf='0.0%', expected=oc / n)
        ws[f'O{r}'].alignment = Alignment(horizontal='centerContinuous', vertical='center')
        ws[f'N{r}'].alignment = Alignment(horizontal='centerContinuous', vertical='center')
        put(ws, f'P{r}', fix, f=font(9), a=al('left', 'center', wrap=True))
        put(ws, f'R{r}', cname, f=font(8, False, C['muted']), a=al('left'))
        ws.row_dimensions[r].height = 24
        for ci in range(2, 17):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
    last = r
    put(ws, f'R{hr}', 'Coluna na Base Tratada', f=font(8, True, C['muted']))
    wb.defined_names['rngAudVerificacao'] = DefinedName('rngAudVerificacao', attr_text=f"'Auditoria'!$C${hr+1}:$C${last}")
    wb.defined_names['rngAudColuna'] = DefinedName('rngAudColuna', attr_text=f"'Auditoria'!$R${hr+1}:$R${last}")
    wb.defined_names['rngAudOcorrencias'] = DefinedName('rngAudOcorrencias', attr_text=f"'Auditoria'!$L${hr+1}:$L${last}")
    text_cf(ws, f'J{hr+1}:J{last}', [('Alta', 'danger'), ('Média', 'warning'), ('Baixa', 'neutral')])
    red = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['danger_fg']), fill=fill(C['danger_soft']))
    ok = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['success_fg']), fill=fill(C['success_soft']))
    ws.conditional_formatting.add(f'L{hr+1}:L{last}', Rule(type='cellIs', operator='greaterThan', formula=['0'], dxf=red))
    ws.conditional_formatting.add(f'L{hr+1}:L{last}', Rule(type='cellIs', operator='equal', formula=['0'], dxf=ok))
    # tratamentos PQ
    tr = last + 2
    section(ws, f'B{tr}', 'Tratamentos aplicados pelo Power Query (padronização automática)', 16)
    put(ws, f'B{tr+1}', 'Tratamento'); put(ws, f'L{tr+1}', 'Registros'); put(ws, f'P{tr+1}', 'O que foi feito')
    table_header_style(ws, tr + 1, 2, 16, h=22, dividers=False)
    ws.cell(tr + 1, 2).alignment = al('left', 'center')
    trat = [('Calibragem convertida de data para número', 'Valores de calibragem gravados com formato de data (ex.: 10/03/1900) voltaram a ser número (70 PSI).'),
            ('Medida padronizada', 'Dimensão reescrita no padrão "175/70 R14" (ex.: "175/70R14" → "175/70 R14").'),
            ('Situação padronizada', 'Situação convertida para maiúsculas e sem espaços (ex.: " Baixado " → "BAIXADO").'),
            ('Número armazenado como texto convertido', 'Milimetragem/N. Vida digitados como texto (ex.: "2,75") convertidos para número.'),
            ('Espaços extras removidos', 'Espaços antes/depois removidos de marca, modelo, posição, status e frota.')]
    for k, (t, d) in enumerate(trat):
        rr = tr + 2 + k
        put(ws, f'B{rr}', t, f=font(9.5, True))
        v = sum(1 for x in rows if x['Tratamento Aplicado'] and t in x['Tratamento Aplicado'])
        put(ws, f'L{rr}', f'=COUNTIF(tbPneus[Tratamento Aplicado],"*{t}*")', f=font(10, True), a=Alignment(horizontal='centerContinuous', vertical='center'),
            nf='#,##0', expected=v)
        ws[f'M{rr}'].alignment = Alignment(horizontal='centerContinuous', vertical='center')
        put(ws, f'P{rr}', d, f=font(8.5, False, C['muted']), a=al('left', 'center', wrap=True))
        ws.row_dimensions[rr].height = 22
        for ci in range(2, 17):
            ws.cell(rr, ci).border = Border(bottom=Side(style='thin', color=C['border']))
    # lista de registros
    lr = tr + 2 + len(trat) + 2
    section(ws, f'B{lr}', 'Registros com inconsistência (rastreabilidade)', 16)
    helper_col = 'T'
    put(ws, f'{helper_col}10', '(Todas)', f=font(8, False, C['muted']))
    for k, a in enumerate(M.AUD, start=1):
        put(ws, f'{helper_col}{10+k}', a[1], f=font(8, False, C['muted']))
    ws.column_dimensions[helper_col].hidden = True
    wb.defined_names['lstFiltroAuditoria'] = DefinedName('lstFiltroAuditoria', attr_text=f"'Auditoria'!${helper_col}$10:${helper_col}${10+len(M.AUD)}")
    put(ws, f'B{lr+1}', 'Verificação:', f=font(9, True, C['muted']), a=al('left', 'center'))
    put(ws, f'C{lr+1}', '(Todas)', f=font(10, True, C['blue_deep']), fl=fill(C['gold_soft']), a=al('left', 'center'), b=BOX)
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type='list', formula1='=lstFiltroAuditoria', allow_blank=False, showErrorMessage=True)
    ws.add_data_validation(dv); dv.add(f'C{lr+1}')
    wb.defined_names['fAudCheck'] = DefinedName('fAudCheck', attr_text=f"'Auditoria'!$C${lr+1}")
    put(ws, f'K{lr+1}', '=IF(fAudCheck="(Todas)",COUNTIF(tbPneus[Qtd Inconsistências],">0"),INDEX(rngAudOcorrencias,MATCH(fAudCheck,rngAudVerificacao,0)))&" registros listados"',
        f=font(9, True, C['blue_deep']), expected=f'{with_inc} registros listados')
    hdrs = ['N.Fogo', 'Situação', 'Medida', 'Menor MM', 'Sulco 1', 'Sulco 2', 'Sulco 3', 'Sulco 4', 'N. Vida', 'Status Rodopar',
            'Frota', 'DOT', 'Km Real', 'Qtd', 'Inconsistências']
    srcs = ['N.Fogo', 'Situação Pneu', 'Dimensão', 'Menor Milimetragem', 'Sulco 1', 'Sulco 2', 'Sulco 3', 'Sulco 4', 'N. Vida',
            'Status', 'N. Frota', 'Dot', 'Km Real', 'Qtd Inconsistências', 'Inconsistências']
    fmts = [('0', 'center'), ('@', 'center'), ('@', 'center'), ('0.00', 'center'), ('0.00', 'center'), ('0.00', 'center'),
            ('0.00', 'center'), ('0.00', 'center'), ('0', 'center'), ('@', 'left'), ('0;-0;;@', 'center'), ('0', 'center'),
            ('#,##0', 'center'), ('0', 'center'), ('@', 'left')]
    cond = ('IF(fAudCheck="(Todas)",tbPneus[Qtd Inconsistências]>0,INDEX(tbPneus,0,_xlfn.XMATCH(INDEX(rngAudColuna,'
            'MATCH(fAudCheck,rngAudVerificacao,0)),tbPneus[#Headers]))=1)')
    inc = [r_ for r_ in rows if r_['Qtd Inconsistências'] > 0]
    list_block(ws, lr + 3, 2, hdrs, fmts, cond, srcs, [('Qtd Inconsistências', -1), ('N.Fogo', 1)], inc,
               'Nenhum registro com inconsistência.', max_rows=1500)
    ws.freeze_panes = 'A5'
    return dict(q=q, with_inc=with_inc, occ=occ, high=high, venc=venc)
