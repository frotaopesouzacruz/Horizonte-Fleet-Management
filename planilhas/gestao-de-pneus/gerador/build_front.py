"""Abas de frente: Início, Dashboard e Documentação."""
from openpyxl.drawing.image import Image as XLImage
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DFont
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import DataPoint
from openpyxl.formatting.rule import Rule, FormulaRule
from openpyxl.styles.differential import DifferentialStyle
from views_common import *
from build_views import MED, MEDINFO
import fontes, urllib.parse

def frase_caminho():
    caminho, origem = fontes.caminho_rodopar()
    if origem == 'sharepoint':
        pasta = urllib.parse.unquote(caminho.rsplit('/', 2)[-2])
        return f'O caminho padrão aponta para a pasta {pasta} do SharePoint da operação.'
    if origem == 'exemplo':
        return 'O caminho gravado é apenas um exemplo: ajuste-o antes da primeira atualização.'
    return 'O caminho padrão foi definido na geração do arquivo; confirme-o na primeira atualização.'

def logo(ws, anchor, height_px):
    img = XLImage(F_LOGO)
    ratio = img.width / img.height
    img.height = height_px
    img.width = int(round(height_px * ratio))
    ws.add_image(img, anchor)

def txt(size=900, color='4A5872', bold=False):
    cp = CharacterProperties(latin=DFont(typeface=FONT), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])

def style_chart(ch, title=None):
    ch.style = 2
    ch.title = None
    ch.legend.position = 'b'
    ch.legend.txPr = txt(850, '4A5872')
    ch.x_axis.txPr = txt(850, '4A5872')
    ch.y_axis.txPr = txt(800, '8B96A9')
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E6EAF1'))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill='C5CDDA'))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.y_axis.numFmt = '0'
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))

def color_series(s, hexc):
    s.graphicalProperties.solidFill = hexc
    s.graphicalProperties.line.solidFill = hexc

def labels(s, color='1F2937', size=800, pos='outEnd'):
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = False; s.dLbls.showCatName = False; s.dLbls.showLegendKey = False; s.dLbls.showPercent = False
    s.dLbls.position = pos
    s.dLbls.txPr = txt(size, color, True)

# ======================================================================= DASHBOARD
def build_dashboard(wb, ws, rows, ctx):
    sheet_base(ws, C['blue'], zoom=85)
    for i in range(2, 18):
        ws.column_dimensions[col(i)].width = 10.6
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 8
    logo(ws, 'B2', 50)
    put(ws, 'D2', 'Gestão de Pneus  |  Dashboard', f=font(18, True, C['blue']), a=al('left', 'center'))
    put(ws, 'D3', 'Operações Souza Cruz · planejamento de ressolagem, compra e estoque a partir da base oficial Rodopar 10',
        f=font(9, False, C['muted']), a=al('left', 'center'))
    put(ws, 'M2', '=MAX(tbPneus[Atualizado em])', f=font(9, True, C['blue_deep']), a=al('right', 'center'),
        nf='"Dados Rodopar 10 atualizados em "dd/mm/yyyy hh:mm', expected=CARGA_INICIAL)
    ws.merge_cells('M2:Q2')
    put(ws, 'Q3', '="Base Rodopar 10: "&ROWS(tbPneus[N.Fogo])&" registros  ·  qualidade "&FIXED(Auditoria!$B$7*100,1)&"%"',
        f=font(9, False, C['muted']), a=al('right', 'center'),
        expected=f"Base Rodopar 10: {len(rows)} registros  ·  qualidade {M.fixed_pt(ctx['aud']['q']*100,1)}%")
    for ci in range(2, 18):
        ws.cell(4, ci).border = Border(bottom=Side(style='medium', color=C['gold']))
    at = sel(rows, Ativo='SIM')
    # --------- faixa 1: base
    section(ws, 'B6', 'Base de pneus', 17, row_h=18)
    n_tot = len(rows); n_at = len(at)
    n_uso = cnt(rows, **{'Situação__Pneu': 'USO'}); n_est = cnt(rows, **{'Situação__Pneu': 'ESTOQUE'})
    n_v1 = sum(1 for r in at if r['N. Vida'] == 1); n_vr = sum(1 for r in at if M.isnum(r['N. Vida']) and r['N. Vida'] >= 2)
    n_bx = cnt(rows, **{'Situação__Pneu': 'BAIXADO'}); n_ds = cnt(rows, **{'Situação__Pneu': 'DESCARTE'})
    cards = [
        ('TOTAL CADASTRADO', '=ROWS(tbPneus[N.Fogo])', n_tot, 'default', 'registros na Rodopar 10'),
        ('PNEUS ATIVOS', '=COUNTIF(tbPneus[Ativo],"SIM")', n_at, 'blue', 'exceto baixados/descartados'),
        ('APLICADOS (EM USO)', '=COUNTIF(tbPneus[Situação Pneu],"USO")', n_uso, 'default', 'montados em veículos'),
        ('EM ESTOQUE', '=COUNTIF(tbPneus[Situação Pneu],"ESTOQUE")', n_est, 'default', 'ver aba Estoque'),
        ('1ª VIDA (ATIVOS)', '=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[N. Vida],1)', n_v1, 'default', 'nunca ressolados'),
        ('RESSOLADOS (ATIVOS)', '=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[N. Vida],">=2")', n_vr, 'default', 'N. Vida ≥ 2'),
        ('BAIXADOS', '=COUNTIF(tbPneus[Situação Pneu],"BAIXADO")', n_bx, 'neutral', 'fora de operação'),
        ('DESCARTADOS', '=COUNTIF(tbPneus[Situação Pneu],"DESCARTE")', n_ds, 'neutral', 'fora de operação'),
    ]
    for k, (lab, fm, ex, tone, note) in enumerate(cards):
        kpi(ws, 7, 2 + 2 * k, 2, lab, fm, ex, tone=tone, note=note)
    # --------- faixa 2: planejamento
    section(ws, 'B11', 'Planejamento', 17, row_h=18)
    rs, cp = ctx['res'], ctx['compra']
    cards = [
        ('DEMANDA DE RESSOLAGEM', '=COUNTIF(tbPneus[Demanda Ressolagem],"SIM")', rs['n_dem'], 'danger', None,
         '="Menor MM < "&FIXED(pLimRessolagem,2)&" mm"', f"Menor MM < {M.fixed_pt(P['pLimRessolagem'])} mm"),
        ('A RESSOLAR', '=COUNTIF(tbPneus[Destino Sugerido],"RESSOLAR")', rs['n_res'], 'blue', 'dentro do limite de vida', None, None),
        ('DESCARTAR (LIMITE DE VIDA)', '=COUNTIF(tbPneus[Destino Sugerido],"DESCARTAR (LIMITE DE VIDA)")', rs['n_des'], 'neutral',
         'vida máxima atingida', None, None),
        ('DEMANDA PROJETADA DE COMPRA', '=SUM(tbResCompra[Demanda Projetada])', cp['tot']['Demanda Projetada'], 'blue', None,
         '="1ª vida < "&FIXED(pLimCompra,2)&" mm"', f"1ª vida < {M.fixed_pt(P['pLimCompra'])} mm"),
        ('ESTOQUE DISPONÍVEL', '=SUM(tbResCompra[Estoque Disponível])', cp['tot']['Estoque Disponível'], 'success', 'cobre a compra', None, None),
        ('NECESSIDADE LÍQUIDA', '=SUM(tbResCompra[Necessidade Líquida])', cp['tot']['Necessidade Líquida'], 'danger', 'pneus a comprar', None, None),
        ('INVESTIMENTO (COMPRA)', '=SUM(tbResCompra[Investimento Estimado])', cp['tot']['Investimento Estimado'], 'gold', 'estimativa', None, None),
        ('CUSTO (RESSOLAGEM)', '=SUM(tbResRessolagem[Custo Estimado])', rs['custo'], 'gold', 'estimativa', None, None),
    ]
    for k, (lab, fm, ex, tone, note, nf_, ne) in enumerate(cards):
        nf = '"R$" #,##0' if 'INVEST' in lab or 'CUSTO' in lab else '#,##0'
        kpi(ws, 12, 2 + 2 * k, 2, lab, fm, ex, nf=nf, tone=tone, note=note, note_formula=nf_, note_expected=ne)
    # --------- faixa 3: condição e qualidade
    section(ws, 'B16', 'Condição da frota e qualidade da base', 17, row_h=18)
    lp, lr, lc = P['pLimCompra'], P['pLimRessolagem'], P['pLimCritico']
    n_lt4 = sum(1 for r in at if M.isnum(r['Menor Milimetragem']) and r['Menor Milimetragem'] < lp)
    n_lt275 = sum(1 for r in at if M.isnum(r['Menor Milimetragem']) and r['Menor Milimetragem'] < lr)
    n_cri = rs['n_cri']
    n_venc = ctx['aud']['venc']
    n_vei = sum(1 for f in ctx['frotas'] if f['Pneus Aplicados'] > 0 and (f['Em Demanda Ressolagem'] + f['Em Demanda Compra']) > 0)
    n_med = sum(1 for s in cp['summary'] if s['Necessidade Líquida'] > 0)
    cards = [
        ('ABAIXO DE 4,00 MM', '=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[Menor Milimetragem],"<"&pLimCompra)', n_lt4, 'warning', 'pneus ativos',
         None, None, '="ABAIXO DE "&FIXED(pLimCompra,2)&" MM"', f"ABAIXO DE {M.fixed_pt(lp)} MM"),
        ('ABAIXO DE 2,75 MM', '=COUNTIFS(tbPneus[Ativo],"SIM",tbPneus[Menor Milimetragem],"<"&pLimRessolagem)', n_lt275, 'danger', 'pneus ativos',
         None, None, '="ABAIXO DE "&FIXED(pLimRessolagem,2)&" MM"', f"ABAIXO DE {M.fixed_pt(lr)} MM"),
        ('CRÍTICOS EM USO', '=COUNTIF(tbPneus[Classificação Desgaste],"CRÍTICO")', n_cri, 'danger', None,
         '="abaixo de "&FIXED(pLimCritico,2)&" mm (TWI)"', f"abaixo de {M.fixed_pt(lc)} mm (TWI)", None, None),
        ('AFERIÇÃO VENCIDA', '=COUNTIF(tbPneus[Status Aferição],"Vencida")', n_venc, 'warning', None,
         '="ativos > "&pAfPrazo&" dias sem medir"', f"ativos > {P['pAfPrazo']} dias sem medir", None, None),
        ('VEÍCULOS COM DEMANDA', '=SUMPRODUCT((tbFrotas[Pneus Aplicados]>0)*((tbFrotas[Em Demanda Ressolagem]+tbFrotas[Em Demanda Compra])>0))',
         n_vei, 'warning', 'ressolagem ou compra', None, None, None, None),
        ('MEDIDAS COM NECESSIDADE', '=COUNTIF(tbResCompra[Necessidade Líquida],">0")', n_med, 'warning', None,
         '="de "&ROWS(tbResCompra[Medida])&" medidas"', f"de {len(MED)} medidas", None, None),
        ('QUALIDADE DA BASE', '=Auditoria!$B$7', ctx['aud']['q'], 'success' if ctx['aud']['q'] >= 0.9 else 'warning',
         'sem inconsistência', None, None, None, None),
        ('INCONSISTÊNCIAS', '=SUM(tbPneus[Qtd Inconsistências])', ctx['aud']['occ'], 'neutral', 'ver aba Auditoria', None, None, None, None),
    ]
    for k, (lab, fm, ex, tone, note, nf_, ne, lab_f, lab_e) in enumerate(cards):
        nf = '0.0%' if 'QUALIDADE' in lab else '#,##0'
        kpi(ws, 17, 2 + 2 * k, 2, lab, fm, ex, nf=nf, tone=tone, note=note, note_formula=nf_, note_expected=ne)
        if lab_f:
            put(ws, f'{col(2+2*k)}17', lab_f, f=font(8, True, C['muted']),
                a=Alignment(horizontal='centerContinuous', vertical='bottom'), expected=lab_e)
    # --------- gráficos
    section(ws, 'B21', 'Compra por medida — demanda × estoque × necessidade', 9, row_h=18)
    section(ws, 'J21', 'Ressolagem por medida — ressolar × descartar', 17, row_h=18)
    wc = wb['Compra']; wr = wb['Ressolagem']; wa = wb['Análises']
    n = len(MED)
    ch = BarChart(); ch.type = 'col'; ch.grouping = 'clustered'; ch.gapWidth = 60; ch.overlap = -10
    for colx, name, color in [(5, 'Demanda projetada', C['blue']), (8, 'Estoque disponível', C['cyan']), (9, 'Necessidade líquida', C['gold'])]:
        ref = Reference(wc, min_col=colx, min_row=11, max_row=11 + n)
        s = Series(ref, title_from_data=True)
        color_series(s, color); labels(s)
        s.dLbls.numFmt = '0;-0;;'
        ch.series.append(s)
    ch.set_categories(Reference(wc, min_col=3, min_row=12, max_row=11 + n))
    style_chart(ch); ch.width = 17.2; ch.height = 7.6
    ws.add_chart(ch, 'B22')
    ch2 = BarChart(); ch2.type = 'col'; ch2.grouping = 'clustered'; ch2.gapWidth = 80; ch2.overlap = -10
    for colx, name, color in [(6, 'Ressolar', C['blue']), (7, 'Descartar', C['neutral'])]:
        ref = Reference(wr, min_col=colx, min_row=11, max_row=11 + n)
        s = Series(ref, title_from_data=True)
        color_series(s, color); labels(s)
        s.dLbls.numFmt = '0;-0;;'
        ch2.series.append(s)
    ch2.set_categories(Reference(wr, min_col=3, min_row=12, max_row=11 + n))
    style_chart(ch2); ch2.width = 17.2; ch2.height = 7.6
    ws.add_chart(ch2, 'J22')
    for r in range(22, 38):
        ws.row_dimensions[r].height = 14.5
    # --------- ranking de pressão de reposição
    section(ws, 'B39', 'Pressão de reposição por medida (ordenado pela necessidade líquida)', 9, row_h=18)
    hdr = [('B', 'Medida'), ('D', 'Utilização'), ('E', 'Demanda'), ('F', 'Estoque'), ('G', 'Necessidade'), ('H', 'Pressão'), ('I', 'Investimento')]
    for c_, h in hdr:
        put(ws, f'{c_}40', h)
    table_header_style(ws, 40, 2, 9, h=22, dividers=False)
    ws['B40'].alignment = al('left', 'center')
    summ = cp['summary']
    order = sorted(range(len(summ)), key=lambda i: (summ[i]['Necessidade Líquida'] * 1e6 + summ[i]['Demanda Projetada'] * 1e3 + (1000 - (12 + i)) / 1000), reverse=True)
    for k in range(n):
        r = 41 + k
        s = summ[order[k]]
        idx = f'MATCH(LARGE(tbResCompra[Ordem],{k+1}),tbResCompra[Ordem],0)'
        put(ws, f'B{r}', f'=INDEX(tbResCompra[Medida],{idx})', f=font(10, True), a=al('left'), expected=s['Medida'])
        put(ws, f'D{r}', f'=INDEX(tbResCompra[Utilização],{idx})', f=font(9.5), a=al('center'), expected=s['Utilização'])
        put(ws, f'E{r}', f'=INDEX(tbResCompra[Demanda Projetada],{idx})', f=font(9.5), a=al('center'), nf='#,##0', expected=s['Demanda Projetada'])
        put(ws, f'F{r}', f'=INDEX(tbResCompra[Estoque Disponível],{idx})', f=font(9.5), a=al('center'), nf='#,##0', expected=s['Estoque Disponível'])
        put(ws, f'G{r}', f'=INDEX(tbResCompra[Necessidade Líquida],{idx})', f=font(10, True), a=al('center'), nf='#,##0', expected=s['Necessidade Líquida'])
        put(ws, f'H{r}', f'=INDEX(tbResCompra[Pressão de Reposição],{idx})', f=font(9.5), a=al('center'), nf='0%', expected=s['Pressão de Reposição'])
        put(ws, f'I{r}', f'=INDEX(tbResCompra[Investimento Estimado],{idx})', f=font(9.5), a=al('center'), nf='"R$" #,##0', expected=s['Investimento Estimado'])
        for ci in range(2, 10):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
    from openpyxl.formatting.rule import DataBarRule
    ws.conditional_formatting.add(f'G41:G{40+n}', DataBarRule(start_type='num', start_value=0, end_type='max', color='E9A3A3', showValue=True))
    link(ws, f'B{41+n}', 'Ver planejamento de compra ▸', "'Compra'!A1")
    # --------- faixa de milimetragem (gráfico)
    section(ws, 'J39', 'Pneus ativos por faixa de Menor Milimetragem', 17, row_h=18)
    ch3 = BarChart(); ch3.type = 'col'; ch3.gapWidth = 40
    fr0, fr1 = ctx['an']['faixa_rows']
    ref = Reference(wa, min_col=3, min_row=7, max_row=7 + 5)
    s = Series(ref, title_from_data=True)
    color_series(s, C['blue']); labels(s)
    cols5 = [C['danger'], 'E26B6B', C['gold'], '6CB98F', C['success']]
    for i_, c5 in enumerate(cols5):
        dp = DataPoint(idx=i_)
        dp.graphicalProperties.solidFill = c5
        dp.graphicalProperties.line.solidFill = c5
        s.dPt.append(dp)
    ch3.series.append(s)
    ch3.set_categories(Reference(wa, min_col=2, min_row=8, max_row=12))
    style_chart(ch3); ch3.legend = None; ch3.width = 17.2; ch3.height = 6.2
    ws.add_chart(ch3, 'J40')
    # --------- veículos
    vr = 54
    section(ws, f'B{vr}', 'Veículos com pneus em demanda (10 mais críticos)', 17, row_h=18)
    hdr = [('B', 'Frota'), ('C', 'Placa'), ('D', 'Tipo'), ('F', 'Local de operação'), ('I', 'Pneus'), ('J', 'Dem. ressolagem'),
           ('K', 'Críticos'), ('L', 'Dem. compra'), ('M', 'Menor MM'), ('N', 'Aferição venc.'), ('O', 'Gestor')]
    for c_, h in hdr:
        put(ws, f'{c_}{vr+1}', h)
    table_header_style(ws, vr + 1, 2, 17, h=26, dividers=False)
    for c_ in ('D', 'F', 'O'):
        ws[f'{c_}{vr+1}'].alignment = al('left', 'center')
    fr_sorted = sorted(ctx['frotas'], key=lambda f: f['Pontuação'], reverse=True)
    for k in range(10):
        r = vr + 2 + k
        f = fr_sorted[k] if k < len(fr_sorted) else None
        ok = f is not None and f['Pontuação'] >= 100
        idx = f'MATCH(LARGE(tbFrotas[Pontuação],{k+1}),tbFrotas[Pontuação],0)'
        g = f'IF(LARGE(tbFrotas[Pontuação],{k+1})<100,"",INDEX(tbFrotas[{{c}}],{idx}){{amp}})'
        for c_, cname, nf, a_, b_ in [('B', 'Frota', '@', 'center', True), ('C', 'Placa', '@', 'center', False), ('D', 'Tipo', '@', 'left', False),
                                      ('F', 'Local de Operação', '@', 'left', False), ('I', 'Pneus Aplicados', '0', 'center', False),
                                      ('J', 'Em Demanda Ressolagem', '0', 'center', True), ('K', 'Críticos', '0', 'center', True),
                                      ('L', 'Em Demanda Compra', '0', 'center', True), ('M', 'Menor MM', '0.00', 'center', False),
                                      ('N', 'Aferições Vencidas', '0', 'center', False), ('O', 'Gestor', '@', 'left', False)]:
            is_txt = cname in ('Tipo', 'Local de Operação', 'Gestor', 'Placa', 'Frota')
            if ok:
                ex = (f[cname] or '') if is_txt else (f[cname] if f[cname] is not None else 0)
            else:
                ex = ''
            put(ws, f'{c_}{r}', '=' + g.replace('{c}', cname).replace('{amp}', '&""' if is_txt else ''), f=font(9.5, b_),
                a=al(a_, 'center'), nf=nf, expected=ex)
        for ci in range(2, 18):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
    red = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['danger_fg']), fill=fill(C['danger_soft']))
    amb = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['warning_fg']), fill=fill(C['warning_soft']))
    ws.conditional_formatting.add(f'J{vr+2}:K{vr+11}', Rule(type='cellIs', operator='greaterThan', formula=['0'], dxf=red))
    ws.conditional_formatting.add(f'L{vr+2}:L{vr+11}', Rule(type='cellIs', operator='greaterThan', formula=['0'], dxf=amb))
    ws.conditional_formatting.add(f'N{vr+2}:N{vr+11}', Rule(type='cellIs', operator='greaterThan', formula=['0'], dxf=amb))
    mm_cf(ws, f'M{vr+2}:M{vr+11}', f'M{vr+2}')
    link(ws, f'B{vr+12}', 'Ver todos os veículos (Análises) ▸', "'Análises'!P6")
    put(ws, f'B{vr+14}', 'Rastreabilidade: cada número acima vem da Base Tratada (Rodopar 10). As listas nominais estão nas abas Compra, '
        'Ressolagem, Estoque e Auditoria (filtros por medida e situação).', f=font(8, False, C['muted'], italic=True))
    ws.freeze_panes = 'A5'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

# ======================================================================= INÍCIO
def build_inicio(wb, ws, rows, ctx):
    sheet_base(ws, C['navy'], zoom=95)
    widths(ws, {'B': 3, 'C': 34, 'D': 50, 'E': 4, 'F': 4, 'G': 34, 'H': 52})
    ws.row_dimensions[1].height = 10
    logo(ws, 'C2', 74)
    ws.row_dimensions[2].height = 22; ws.row_dimensions[3].height = 22; ws.row_dimensions[4].height = 16; ws.row_dimensions[5].height = 10
    put(ws, 'D2', 'Gestão de Pneus', f=font(22, True, C['blue']), a=al('left', 'center'))
    put(ws, 'D3', 'Planejamento de ressolagem, compra e estoque', f=font(12, False, C['text2']), a=al('left', 'center'))
    put(ws, 'D4', 'Operações Souza Cruz · Horizonte Logística · fonte oficial: Rodopar 10', f=font(9, False, C['muted']), a=al('left', 'center'))
    for ci in range(2, 9):
        ws.cell(6, ci).border = Border(top=Side(style='medium', color=C['gold']))
    # situação dos dados
    section(ws, 'C8', 'Situação dos dados', 4)
    lines = [
        ('Última atualização (Atualizar Tudo)', '=MAX(tbPneus[Atualizado em])', CARGA_INICIAL, 'dd/mm/yyyy hh:mm'),
        ('Registros carregados da Rodopar 10', '=ROWS(tbPneus[N.Fogo])', len(rows), '#,##0'),
        ('Alteração mais recente no Rodopar', '=MAX(tbPneus[Data Última Alteração])', max(r['Data Última Alteração'] for r in rows), 'dd/mm/yyyy hh:mm'),
        ('Medição mais recente', '=MAX(tbPneus[Dt. Medição])', max(r['Dt. Medição'] for r in rows), 'dd/mm/yyyy hh:mm'),
        ('Qualidade da base', '=Auditoria!$B$7', ctx['aud']['q'], '0.0%'),
        ('Arquivo de origem', '=INDEX(tbPneus[Arquivo Fonte],1)', fontes.arquivo_fonte_inicial(), '@'),
    ]
    for k, (lab, fm, ex, nf) in enumerate(lines):
        r = 9 + k
        put(ws, f'C{r}', lab, f=font(9.5, False, C['text2']), a=al('left', 'center'))
        put(ws, f'D{r}', fm, f=font(10.5, True, C['text']), a=al('left', 'center', wrap=(k == 5)), nf=nf, expected=ex)
        ws.row_dimensions[r].height = 20 if k < 5 else 30
        for ci in (3, 4):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
    # como atualizar
    section(ws, 'G8', 'Como atualizar', 8)
    steps = [('1', 'Exporte o relatório Rodopar 10 e salve-o como "Rodopar 10.xlsx" no local configurado.'),
             ('2', 'Abra este arquivo e clique em Dados ▸ Atualizar Tudo (Ctrl+Alt+F5).'),
             ('3', 'Pronto: base tratada, planejamentos, auditoria e dashboard são recalculados.')]
    for k, (n_, t) in enumerate(steps):
        r = 9 + k
        put(ws, f'G{r}', f'{n_}   {t}', f=font(9.5, k == 1, C['text']), a=al('left', 'center', wrap=True))
        ws.merge_cells(f'G{r}:H{r}')
    put(ws, 'G13', 'Primeira utilização — apontar o arquivo de origem', f=font(9.5, True, C['blue_deep']))
    put(ws, 'G14', 'Dados ▸ Obter Dados ▸ Iniciar Editor do Power Query ▸ consulta "Pneus_Rodopar10" ▸ etapa "CaminhoArquivo": '
        'substitua o texto pelo caminho do arquivo (ex.: C:\\Pneus\\Rodopar 10.xlsx ou o link do SharePoint) ▸ Fechar e Carregar. '
        + frase_caminho(), f=font(9, False, C['text2']), a=al('left', 'top', wrap=True))
    ws.merge_cells('G14:H16')
    for r in (14, 15, 16): ws.row_dimensions[r].height = 18
    # navegação
    section(ws, 'C18', 'Navegação', 8)
    nav = [('Dashboard', 'Visão gerencial: indicadores, gráficos, medidas e veículos que demandam atenção.'),
           ('Compra', 'Demanda projetada de compra, cobertura pelo estoque, necessidade líquida e investimento.'),
           ('Ressolagem', 'Pneus abaixo do limite, prioridade (CRÍTICO / ENVIAR / ACOMPANHAR / NORMAL) e destino.'),
           ('Fluxo Ressolagem', 'Controle dos lotes: identificado → retirado → enviado → análise → ressolado → aplicado.'),
           ('Estoque', 'Estoque por medida e condição: novos, ressolados, utilizável, comprometido e cobertura.'),
           ('Análises', 'Visões por faixa, situação, vida, eixo, operação, local, fabricante/modelo e veículo.'),
           ('Auditoria', 'Qualidade da base: inconsistências de cadastro e tratamentos aplicados.'),
           ('Parâmetros', 'Limites das regras e tabelas de referência (medidas, situações, centros de custo).'),
           ('Frotas', 'Cadastro N. Frota → placa, tipo e base de operação.'),
           ('Base Tratada', 'Resultado do Power Query + colunas de regra por pneu (memória de cálculo).'),
           ('Documentação', 'Regras, arquitetura, dicionário, reconciliação com a planilha anterior e testes.')]
    for k, (sh, d) in enumerate(nav):
        r = 19 + k
        link(ws, f'C{r}', f'{sh}  ▸', f"'{sh}'!A1", f=font(10, True, C['blue'], underline=None))
        put(ws, f'D{r}', d, f=font(9, False, C['text2']), a=al('left', 'center'))
        ws.merge_cells(f'D{r}:H{r}')
        ws.row_dimensions[r].height = 18
        for ci in range(3, 9):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
    # regras vigentes
    r0 = 19 + len(nav) + 1
    section(ws, f'C{r0}', 'Regras vigentes (lidas da aba Parâmetros)', 8)
    rules = [
        ('Demanda de ressolagem', '="Menor Milimetragem < "&FIXED(pLimRessolagem,2)&" mm  E  Situação diferente de Baixado e Descarte"',
         f"Menor Milimetragem < {M.fixed_pt(P['pLimRessolagem'])} mm  E  Situação diferente de Baixado e Descarte"),
        ('Demanda projetada de compra', '="Menor Milimetragem < "&FIXED(pLimCompra,2)&" mm  E  N. Vida = "&pVidaCompra&"  E  Situação diferente de Baixado e Descarte (estoque incluído: "&pIncluirEstoque&")"',
         f"Menor Milimetragem < {M.fixed_pt(P['pLimCompra'])} mm  E  N. Vida = {P['pVidaCompra']}  E  Situação diferente de Baixado e Descarte (estoque incluído: {P['pIncluirEstoque']})"),
        ('Estoque disponível (cobertura)', '="Situação ESTOQUE  E  N. Vida = "&pVidaCompra&"  E  Menor Milimetragem ≥ "&FIXED(pLimCompra,2)&" mm"',
         f"Situação ESTOQUE  E  N. Vida = {P['pVidaCompra']}  E  Menor Milimetragem ≥ {M.fixed_pt(P['pLimCompra'])} mm"),
        ('Necessidade líquida', 'Demanda projetada − estoque disponível, por medida (mínimo zero)', None),
        ('Classificação', '="CRÍTICO: em uso < "&FIXED(pLimCritico,2)&" mm · ENVIAR PARA RESSOLAGEM: < "&FIXED(pLimRessolagem,2)&" mm · ACOMPANHAR: < "&FIXED(pLimCompra,2)&" mm · NORMAL: demais"',
         f"CRÍTICO: em uso < {M.fixed_pt(P['pLimCritico'])} mm · ENVIAR PARA RESSOLAGEM: < {M.fixed_pt(P['pLimRessolagem'])} mm · ACOMPANHAR: < {M.fixed_pt(P['pLimCompra'])} mm · NORMAL: demais"),
        ('Destino sugerido', 'Na demanda de ressolagem: DESCARTAR (LIMITE DE VIDA) quando N. Vida ≥ vida máxima da medida; senão RESSOLAR', None),
    ]
    for k, (lab, fm, ex) in enumerate(rules):
        r = r0 + 1 + k
        put(ws, f'C{r}', lab, f=font(9.5, True, C['text']), a=al('left', 'center'))
        put(ws, f'D{r}', fm, f=font(9, False, C['text2']), a=al('left', 'center', wrap=True), expected=ex)
        ws.merge_cells(f'D{r}:H{r}')
        ws.row_dimensions[r].height = 20
        for ci in range(3, 9):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
    # legenda
    r1 = r0 + len(rules) + 2
    section(ws, f'C{r1}', 'Legenda', 8)
    leg = [(C['danger'], 'FFFFFF', 'CRÍTICO', 'Pneu em uso abaixo do limite legal — retirada imediata'),
           (C['danger_soft'], C['danger_fg'], 'ENVIAR PARA RESSOLAGEM', 'Abaixo do limite de ressolagem — necessidade imediata'),
           (C['warning_soft'], C['warning_fg'], 'ACOMPANHAR', 'Aproximando-se do limite (entre os limites de ressolagem e de compra)'),
           (C['success_soft'], C['success_fg'], 'NORMAL', 'Condição adequada'),
           (C['neutral_soft'], C['neutral_fg'], 'INATIVO', 'Baixado / descarte — registro não operacional'),
           (C['gold_soft'], C['blue_deep'], 'Célula editável', 'Parâmetros, cadastros, filtros e registro do fluxo de ressolagem')]
    for k, (bg, fg, lab, d) in enumerate(leg):
        r = r1 + 1 + k
        put(ws, f'C{r}', lab, f=font(9, True, fg), fl=fill(bg), a=al('center', 'center'))
        put(ws, f'D{r}', d, f=font(9, False, C['text2']), a=al('left', 'center'))
        ws.row_dimensions[r].height = 18
    put(ws, f'C{r1+len(leg)+2}', 'Versão 1.0 · outubro/2026 · construída a partir de Rodopar 10, Controle de Pneus OPE Souza Cruz, '
        'Demandas de Compra/Recapagem e Gestão de CPK.', f=font(8, False, C['muted'], italic=True))

# ======================================================================= DOCUMENTAÇÃO
def build_doc(wb, ws, rows, ctx, recon, tests):
    db = M.data_base(rows)
    n_cob = sum(r['Cobertura Compra'] == 'SIM' for r in rows)
    sheet_base(ws, C['navy'], zoom=95)
    widths(ws, {'B': 32, 'C': 16, 'D': 16, 'E': 13, 'F': 16, 'G': 34, 'H': 46})
    header(ws, 'Documentação técnica', 'Fonte, arquitetura, regras, dicionário de colunas, reconciliação com a planilha anterior, '
           'testes de validação e limitações.', 8)
    r = 6
    def block(title, items, h=None):
        nonlocal r
        section(ws, f'B{r}', title, 8)
        r += 1
        for lab, text in items:
            put(ws, f'B{r}', lab, f=font(9.5, True, C['text']), a=al('left', 'top', wrap=True))
            if isinstance(text, tuple):
                put(ws, f'C{r}', text[0], f=font(9, False, C['text2']), a=al('left', 'top', wrap=True), expected=text[1])
            else:
                put(ws, f'C{r}', text, f=font(9, False, C['text2']), a=al('left', 'top', wrap=True))
            ws.merge_cells(f'C{r}:H{r}')
            txtv = text[1] if isinstance(text, tuple) else text
            nlines = max(1, int(len(str(txtv)) / 125) + 1 + str(txtv).count('\n'))
            ws.row_dimensions[r].height = max(16, 13 * nlines + 3)
            for ci in range(2, 9):
                ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
            r += 1
        r += 1
    block('1. Fonte dos dados', [
        ('Fonte oficial', 'Rodopar 10 (cadastro/base operacional de pneus) — base mestre. A estrutura original não é alterada: a planilha '
                          'apenas lê o arquivo exportado.'),
        ('Leitura', 'Consulta Power Query "Pneus_Rodopar10": lê a aba "Planilha1" (ou a primeira aba) do arquivo configurado na etapa '
                    '"CaminhoArquivo" (caminho local ou link do SharePoint). Cabeçalhos são normalizados (quebras de linha e espaços duplos '
                    'removidos: "Cód⏎Unidade" → "Cód Unidade", "N.  Frota" → "N. Frota") e as 34 colunas são selecionadas pelo nome, '
                    'em ordem fixa — colunas novas na exportação são ignoradas; colunas ausentes ficam vazias e aparecem na Auditoria.'),
        ('Campos-chave', 'Menor Milimetragem, N. Vida, Situação Pneu, Dimensão (medida), N.Fogo (identificação), N. Frota (veículo), '
                         'Posição, Filial Pneu + Cód Unidade (operação), Status (Novo/Usado/Recauchutado), Dt. Medição.'),
        ('Fontes complementares', 'Cadastro de frotas (placa/tipo/base): Base Cadastro Frotas do arquivo Gestão de CPK + alocação de '
                                  f'{fontes.data_alocacao():%d/%m/%Y} da Base de Fidelização. Custos de referência: Rodopar 978 e Análise CPK. Histórico de '
                                  'ressolagem: aba Ressolagem (Controle_Demanda) da planilha Controle de Pneus OPE Souza Cruz.'),
    ])
    block('2. Arquitetura', [
        ('Fluxo', 'Rodopar 10 ▸ Power Query (tratamento) ▸ Base Tratada (tbPneus) ▸ Regras (colunas calculadas que leem a aba '
                  'Parâmetros) ▸ Controles (Compra, Ressolagem, Fluxo, Estoque, Auditoria) ▸ Dashboard.'),
        ('Atualização', '"Atualizar Tudo" executa a consulta; as colunas de regra e todos os resumos se ajustam automaticamente ao '
                        'novo número de linhas (tabelas estruturadas, sem intervalos fixos). As listas nominais usam FILTRO/CLASSIFICARPOR '
                        'e crescem sozinhas.'),
        ('Desempenho', 'Uma única leitura do arquivo; regras por linha com CONT.SES/ÍNDICE+CORRESP sobre colunas de tabela (sem colunas '
                       'inteiras, sem PROCV encadeado e sem funções voláteis nas regras — a data de referência é a hora da atualização).'),
    ])
    block('3. Regras de negócio', [
        ('Demanda de ressolagem', ('="Menor Milimetragem < "&FIXED(pLimRessolagem,2)&" mm (estritamente menor; 2,75 não entra) E Situação Pneu diferente de BAIXADO e DESCARTE (comparação sem diferenciar maiúsculas e espaços). Limite específico por medida opcional em Parâmetros (vazio = limite geral)."',
                                   f"Menor Milimetragem < {M.fixed_pt(P['pLimRessolagem'])} mm (estritamente menor; 2,75 não entra) E Situação Pneu diferente de BAIXADO e DESCARTE (comparação sem diferenciar maiúsculas e espaços). Limite específico por medida opcional em Parâmetros (vazio = limite geral).")),
        ('Demanda projetada de compra', ('="Menor Milimetragem < "&FIXED(pLimCompra,2)&" mm (estritamente menor; 4,00 não entra) E N. Vida = "&pVidaCompra&" E Situação diferente de BAIXADO e DESCARTE. Pneus em ESTOQUE incluídos: "&pIncluirEstoque&"."',
                                         f"Menor Milimetragem < {M.fixed_pt(P['pLimCompra'])} mm (estritamente menor; 4,00 não entra) E N. Vida = {P['pVidaCompra']} E Situação diferente de BAIXADO e DESCARTE. Pneus em ESTOQUE incluídos: {P['pIncluirEstoque']}.")),
        ('Estoque disponível (cobertura)', 'Pneus em ESTOQUE com N. Vida igual à vida considerada para compra e Menor Milimetragem ≥ limite '
                                           'de compra. É exatamente a lista de estoque usada na planilha "Demanda Compra de Pneus" '
                                           f'({n_cob} pneus na base de {db:%d/%m/%Y}). Pneus ressolados não abatem a compra: 100% das posições dianteiras usam '
                                           'pneus de 1ª vida.'),
        ('Necessidade líquida', 'Por medida: MÁXIMO(0; Demanda projetada − Estoque disponível). Investimento = necessidade × custo de '
                                'referência da medida (Parâmetros).'),
        ('Classificação de desgaste', 'CRÍTICO: em uso e abaixo do limite legal (TWI 1,6 mm — CONTRAN 558/1980). ENVIAR PARA RESSOLAGEM: '
                                      'demais pneus da demanda de ressolagem. ACOMPANHAR: entre o limite de ressolagem e o de compra. NORMAL: '
                                      'acima do limite de compra. INATIVO: baixado/descarte. SEM AFERIÇÃO: sem Menor Milimetragem.'),
        ('Destino sugerido', 'Para pneus em demanda de ressolagem: DESCARTAR (LIMITE DE VIDA) quando N. Vida ≥ vida máxima da medida '
                             '(Van 3, Fiorino 1, Caminhão 3); caso contrário RESSOLAR. Reproduz a prática dos lotes Ago/Out-2026 '
                             '("Descartar por Limite de Vida do Pneu").'),
        ('Fluxo de ressolagem', 'Status por lote: Identificado ▸ Retirado do veículo ▸ Enviado ao fornecedor ▸ Em análise ▸ Aprovado | '
                                'Reprovado ▸ Ressolado ▸ Retornado ao estoque ▸ Aplicado (ou Descartar). Situação no fluxo de cada pneu = '
                                'último lote em que aparece: Não programado, Em andamento, Encerrado – baixar no Rodopar, Concluído.'),
        ('Aferição', 'Dias desde a última medição na data da atualização: até 20 = Em dia; 21–25 = Próx. do vencimento; acima de 25 = '
                     'Vencida (critério da planilha anterior).'),
    ])
    # dicionário
    section(ws, f'B{r}', '4. Dicionário das colunas calculadas (Base Tratada)', 8)
    r += 1
    dic = [
        ('Ativo', 'SIM se a Situação não estiver marcada como inativa em Parâmetros (BAIXADO/DESCARTE).'),
        ('Utilização', 'Van / Fiorino / Caminhão, pela medida (Parâmetros › Medidas).'),
        ('Placa / Tipo Veículo', 'Pelo N. Frota na aba Frotas; "(sem cadastro)" quando a frota não existe no cadastro.'),
        ('Local de Operação', 'Base do veículo (pneus em uso), "Estoque FL <filial>" (estoque) ou a própria situação.'),
        ('Operação', 'Centro de custo pelo De-Para Filial Pneu + Cód Unidade.'),
        ('Eixo', 'Dianteiro (ED*), Traseiro (ET*), Estepe (EST*), a partir da Posição.'),
        ('Faixa MM', 'Faixa de profundidade da Menor Milimetragem (limites crítico, ressolagem, compra e atenção).'),
        ('Limite Ressolagem Aplicado', 'Limite específico da medida, se houver; senão o limite geral.'),
        ('Demanda Ressolagem / Demanda Compra', 'Regras 3.1 e 3.2 (SIM/NÃO).'),
        ('Classificação Desgaste / Destino Sugerido', 'Regras 3.5 e 3.6.'),
        ('Estoque Utilizável / Cobertura Compra', 'Estoque ≥ limite de ressolagem / estoque de 1ª vida ≥ limite de compra.'),
        ('Dias desde Medição / Status Aferição', 'Idade da última medição na data da atualização e respectivo status.'),
        ('Lote Ressolagem / Status Fluxo / Situação Fluxo', 'Último lote do pneu no Fluxo de Ressolagem e sua situação.'),
        ('Prioridade', 'Ordem de atendimento usada nas listas (1 = CRÍTICO … 6 = INATIVO).'),
        ('AUD … / Qtd Inconsistências / Inconsistências', 'Verificações de auditoria (1 = ocorrência) e resumo textual por pneu.'),
    ]
    for lab, d in dic:
        put(ws, f'B{r}', lab, f=font(9, True), a=al('left', 'top', wrap=True))
        put(ws, f'C{r}', d, f=font(9, False, C['text2']), a=al('left', 'top', wrap=True))
        ws.merge_cells(f'C{r}:H{r}')
        ws.row_dimensions[r].height = 26 if len(lab) > 30 else 16
        for ci in range(2, 9):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
        r += 1
    r += 1
    # reconciliação
    section(ws, f'B{r}', f'5. Reconciliação com a planilha atual (base Rodopar 10 de {db:%d/%m/%Y})', 8)
    r += 1
    hh = ['Indicador', 'Planilha anterior', 'Nova planilha', 'Diferença', 'Causa identificada', None, 'Conclusão']
    for i, h in enumerate(hh):
        if h: put(ws, f'{col(2+i)}{r}', h)
    table_header_style(ws, r, 2, 8, h=24)
    r += 1
    for row in recon:
        put(ws, f'B{r}', row[0], f=font(9, True), a=al('left', 'top', wrap=True))
        put(ws, f'C{r}', row[1], f=font(9), a=al('center', 'top', wrap=True))
        put(ws, f'D{r}', row[2], f=font(9, True, C['blue_deep']), a=al('center', 'top', wrap=True))
        put(ws, f'E{r}', row[3], f=font(9), a=al('center', 'top', wrap=True))
        put(ws, f'F{r}', row[4], f=font(8.5, False, C['text2']), a=al('left', 'top', wrap=True))
        ws.merge_cells(f'F{r}:G{r}')
        put(ws, f'H{r}', row[5], f=font(8.5, False, C['text2']), a=al('left', 'top', wrap=True))
        nl = max(len(row[4]) / 55, len(row[5]) / 60, len(str(row[0])) / 30, 1)
        ws.row_dimensions[r].height = 12 * (int(nl) + 1) + 4
        for ci in range(2, 9):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
        r += 1
    r += 1
    # testes
    section(ws, f'B{r}', '6. Testes de validação com registros reais (resultado recalculado ao vivo)', 8)
    r += 1
    hh = ['Caso (escopo)', 'N.Fogo', 'Menor MM', 'N. Vida', 'Situação', 'Obtido', 'Esperado / resultado']
    for i, h in enumerate(hh):
        put(ws, f'{col(2+i)}{r}', h)
    table_header_style(ws, r, 2, 8, h=24)
    r += 1
    for t in tests:
        caso, fogo, campo, esperado, row_ = t
        m = f'MATCH(C{r},tbPneus[N.Fogo],0)'
        put(ws, f'B{r}', caso, f=font(9, True), a=al('left', 'center', wrap=True))
        put(ws, f'C{r}', fogo, f=font(9.5, True, C['blue_deep']), a=al('center'))
        put(ws, f'D{r}', f'=INDEX(tbPneus[Menor Milimetragem],{m})', f=font(9), a=al('center'), nf='0.00', expected=row_['Menor Milimetragem'])
        put(ws, f'E{r}', f'=INDEX(tbPneus[N. Vida],{m})', f=font(9), a=al('center'), expected=row_['N. Vida'])
        put(ws, f'F{r}', f'=INDEX(tbPneus[Situação Pneu],{m})', f=font(9), a=al('center'), expected=row_['Situação Pneu'])
        lab = 'Ressolagem' if campo == 'Demanda Ressolagem' else ('Compra' if campo == 'Demanda Compra' else 'Destino')
        put(ws, f'G{r}', f'="{lab} = "&INDEX(tbPneus[{campo}],{m})', f=font(9, True), a=al('center'),
            expected=f"{lab} = {row_[campo]}")
        put(ws, f'H{r}', f'=IF(INDEX(tbPneus[{campo}],{m})="{esperado}","✔ OK — esperado {lab} = {esperado}","✖ VERIFICAR — esperado {lab} = {esperado}")',
            f=font(9, True), a=al('left', 'center'),
            expected=(f"✔ OK — esperado {lab} = {esperado}" if row_[campo] == esperado else f"✖ VERIFICAR — esperado {lab} = {esperado}"))
        ws.row_dimensions[r].height = 30
        for ci in range(2, 9):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
        r += 1
    ok = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['success_fg']), fill=fill(C['success_soft']))
    bad = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['danger_fg']), fill=fill(C['danger_soft']))
    ws.conditional_formatting.add(f'H{r-len(tests)}:H{r-1}', Rule(type='beginsWith', operator='beginsWith', text='✔',
                                                                    formula=[f'LEFT(H{r-len(tests)},1)="✔"'], dxf=ok))
    ws.conditional_formatting.add(f'H{r-len(tests)}:H{r-1}', Rule(type='beginsWith', operator='beginsWith', text='✖',
                                                                    formula=[f'LEFT(H{r-len(tests)},1)="✖"'], dxf=bad))
    put(ws, f'B{r}', f'Os casos usam pneus reais da base de {db:%d/%m/%Y}. Se a situação de um desses pneus mudar no Rodopar, o resultado '
        'acompanha os novos dados (o teste valida a regra com os valores atuais da linha).', f=font(8, False, C['muted'], italic=True))
    r += 2
    block('7. Manutenção', [
        ('Alterar limites', 'Edite os valores amarelos em Parâmetros (ex.: 2,75 → 3,00). Todas as regras, resumos, listas e o dashboard '
                            'são recalculados — nenhuma fórmula precisa ser editada.'),
        ('Nova medida', 'Inclua uma linha em Parâmetros › Medidas (utilização, vida máxima, custos) e estenda as tabelas de resumo das abas '
                        'Compra, Ressolagem e Estoque arrastando a alça do canto inferior da tabela (uma linha por medida). A Auditoria '
                        'acusa "Medida não cadastrada" enquanto isso não for feito.'),
        ('Nova frota / mudança de base', 'Atualize a aba Frotas (colunas amarelas). Frota deve ser digitada como texto.'),
        ('Novo lote de ressolagem', 'Na aba Fluxo Ressolagem, acrescente as linhas logo abaixo da tabela (ela se expande): nº do lote, '
                                    'referência, N.Fogo, status e recapadora. Use a lista Ressolagem › "Não programado" como ponto de partida.'),
        ('Caminho do arquivo', 'Power Query › consulta Pneus_Rodopar10 › etapa CaminhoArquivo.'),
    ])
    block('8. Limitações conhecidas', [
        ('Excel', 'Requer Excel 365 / 2021 ou superior (FILTRO, CLASSIFICARPOR, ÚNICO, SEQUÊNCIA, CORRESPX, LET). Versões anteriores '
                  'exibem #NOME? nas listas nominais; resumos e indicadores continuam funcionando.'),
        ('Primeira atualização', frase_caminho() + ' Confirme o local do arquivo e, se solicitado, entre com a conta corporativa. '
                                 f'Os dados atuais vêm da carga inicial ({fontes.nome_limpo(fontes.arquivo("rodopar"))}, '
                                 f'{CARGA_INICIAL:%d/%m/%Y}).'),
        ('Cadastro de frotas', f'É uma fotografia de {fontes.data_alocacao():%d/%m/%Y} (CPK + Fidelização). Mudanças de base/veículo precisam ser atualizadas na aba Frotas.'),
        ('Fluxo de ressolagem', 'A Rodopar 10 não traz o fluxo de envio/retorno; o controle de lotes é manual. O histórico de movimentação '
                                'do Rodopar 607 (usado no arquivo de CPK) pode automatizar etapas no futuro.'),
        ('Custos', 'Custos de referência são estimativas (última compra registrada / histórico de reforma) e devem ser revisados com cotações.'),
        ('Vida máxima', 'Van = 3 e Fiorino = 1 seguem a prática dos lotes; Caminhão = 3 vem da Vida Prev. do Rodopar 978 (validar com a operação).'),
        ('Fora do escopo', 'Calibragem, consertos, alinhamento/balanceamento, aferições internas e controle de descartes (MTR/CDF/NF) '
                           'continuam na planilha Controle de Pneus OPE Souza Cruz; esta ferramenta é dedicada ao planejamento.'),
        ('Recálculo', 'O arquivo recalcula ao abrir (cálculo completo); o Excel pode perguntar se deseja salvar ao fechar.'),
    ])
    block('9. Histórico', [
        ('v1.0 — 05/10/2026', 'Versão inicial: consulta Power Query da Rodopar 10, regras parametrizadas, planejamento de compra e ressolagem, '
                              'fluxo de lotes com histórico migrado (569 registros), estoque, análises, auditoria, dashboard e documentação.'),
    ])
