"""Abas de referência e dados: Parâmetros, Frotas, Fluxo Ressolagem, Base Tratada."""
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from xlh import *
from config import *
import model as M
import fontes

INPUT_FILL = fill(C['gold_soft'])
INPUT_FONT = font(10, False, C['blue_deep'])

def R(c, t='tbPneus'):
    return f"{t}[[#This Row],[{c}]]"

def defname(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)

# ------------------------------------------------------------------ PARÂMETROS
def build_parametros(wb, ws):
    sheet_base(ws, C['gold'])
    widths(ws, {'B': 36, 'C': 13, 'D': 11, 'E': 14, 'F': 14, 'G': 14, 'H': 64})
    header(ws, 'Parâmetros e tabelas de referência', 'Limites das regras, medidas, situações e centros de custo. '
           'Células em amarelo-claro são editáveis; todas as regras da planilha leem estes valores.', 8)
    section(ws, 'B6', 'Parâmetros das regras', 8)
    hdr = ['Parâmetro', 'Valor', 'Unidade', 'Como é usado', None, None, 'Origem do critério']
    for i, h in enumerate(hdr):
        if h: put(ws, f'{col(2+i)}7', h)
    table_header_style(ws, 7, 2, 8, h=22)
    ws.merge_cells('E7:G7')
    r = 8
    for name, label, val, unit, uso, origem, typ in PARAMS:
        put(ws, f'B{r}', label, f=font(10, True), a=al('left', 'center', wrap=True))
        c = put(ws, f'C{r}', val, f=font(11, True, C['blue_deep']), fl=INPUT_FILL, a=al('center'),
                nf={'dec': '0.00', 'int': '0', 'sn': '@'}[typ])
        c.border = BOX
        put(ws, f'D{r}', unit, f=font(9, False, C['muted']), a=al('center'))
        put(ws, f'E{r}', uso, f=font(9), a=al('left', 'center', wrap=True))
        ws.merge_cells(f'E{r}:G{r}')
        put(ws, f'H{r}', origem, f=font(9, False, C['muted']), a=al('left', 'center', wrap=True))
        ws.row_dimensions[r].height = 40
        for ci in range(2, 9):
            cc = ws.cell(r, ci)
            cc.border = Border(bottom=Side(style='thin', color=C['border'])) if ci != 3 else BOX
        defname(wb, name, f"'Parâmetros'!$C${r}")
        if typ == 'dec':
            dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='100', allow_blank=False,
                                showErrorMessage=True, errorTitle='Valor inválido', error='Informe um número entre 0 e 100.')
        elif typ == 'int':
            dv = DataValidation(type='whole', operator='between', formula1='0', formula2='365', allow_blank=False,
                                showErrorMessage=True, errorTitle='Valor inválido', error='Informe um número inteiro.')
        else:
            dv = DataValidation(type='list', formula1='"SIM,NÃO"', allow_blank=False, showErrorMessage=True,
                                errorTitle='Valor inválido', error='Escolha SIM ou NÃO.')
        ws.add_data_validation(dv); dv.add(f'C{r}')
        r += 1
    last_param_row = r - 1
    # ---------------- tbMedidas
    r += 1
    section(ws, f'B{r}', 'Medidas (dimensões) — utilização, limite de vida e custos de referência', 8)
    r += 1
    mcols = ['Medida', 'Utilização', 'Vida Máxima', 'Limite Ressolagem Específico (mm)', 'Custo Pneu Novo (R$)',
             'Custo Ressolagem (R$)', 'Referência do Custo']
    for i, h in enumerate(mcols):
        put(ws, f'{col(2+i)}{r}', h)
    hdr_row = r
    for m in fontes.medidas():
        r += 1
        vals = [m[0], m[1], m[2], m[3], m[4], m[5], m[6]]
        for i, v in enumerate(vals):
            c = put(ws, f'{col(2+i)}{r}', v, f=INPUT_FONT if i < 6 else font(9, False, C['muted']),
                    a=al('center' if i in (1, 2, 3) else ('right' if i in (4, 5) else 'left'), 'center', wrap=(i == 6)))
            if i < 6: c.fill = INPUT_FILL
            if i in (4, 5): c.number_format = '"R$" #,##0.00'
            if i == 3: c.number_format = '0.00'
        ws.row_dimensions[r].height = max(30, 12 * -(-len(m[6] or '') // 95) + 4)  # ~95 caracteres por linha
    add_table(ws, 'tbMedidas', f'B{hdr_row}:H{r}', mcols)
    table_header_style(ws, hdr_row, 2, 8, h=32)
    put(ws, f'B{r+1}', 'Vida Máxima: número de vidas após o qual o pneu em demanda é indicado para descarte (Van: 3 — padrão observado '
        'nos lotes "Descartar por Limite de Vida"; Fiorino: 1 — não ressola; Caminhão: 3 — Vida Prev. do Rodopar 978, validar). '
        'Limite específico: deixe vazio para usar o limite geral de ressolagem.', f=font(8, False, C['muted'], italic=True),
        a=al('left', 'top', wrap=True))
    ws.merge_cells(f'B{r+1}:H{r+1}'); ws.row_dimensions[r+1].height = 30
    dvm = DataValidation(type='list', formula1='"Van,Fiorino,Caminhão"', allow_blank=True)
    ws.add_data_validation(dvm); dvm.add(f'C{hdr_row+1}:C{hdr_row+40}')
    # ---------------- tbSituacoes
    r += 3
    section(ws, f'B{r}', 'Situações do pneu (Rodopar)', 8)
    r += 1
    scol = ['Situação Pneu', 'Inativo', 'Descrição']
    for i, h in enumerate(scol):
        put(ws, f'{col(2+i)}{r}', h)
    sh = r
    for s in SITUACOES:
        r += 1
        put(ws, f'B{r}', s[0], f=INPUT_FONT, fl=INPUT_FILL)
        put(ws, f'C{r}', s[2], f=INPUT_FONT, fl=INPUT_FILL, a=al('center'))
        put(ws, f'D{r}', s[1], f=font(9))
    add_table(ws, 'tbSituacoes', f'B{sh}:D{r}', scol)
    table_header_style(ws, sh, 2, 4, h=22)
    put(ws, f'F{sh}', 'Inativo = SIM exclui o pneu das demandas (regra: Situação diferente de Baixado e Descarte). '
        'A comparação ignora maiúsculas/minúsculas e espaços (padronização no Power Query).',
        f=font(8, False, C['muted'], italic=True), a=al('left', 'top', wrap=True))
    ws.merge_cells(f'F{sh}:H{sh+2}')
    dvs = DataValidation(type='list', formula1='"SIM,NÃO"', allow_blank=False)
    ws.add_data_validation(dvs); dvs.add(f'C{sh+1}:C{sh+20}')
    # ---------------- tbCentrosCusto
    r += 2
    section(ws, f'B{r}', 'Centros de custo — De-Para Filial + Unidade → Operação', 8)
    r += 1
    ccol = ['Operação', 'Filial Pneu', 'Cód Unidade', 'Chave']
    for i, h in enumerate(ccol):
        put(ws, f'{col(2+i)}{r}', h)
    ch = r
    for fil, uni, op in fontes.centros_custo():
        r += 1
        put(ws, f'B{r}', op, f=INPUT_FONT, fl=INPUT_FILL)
        put(ws, f'C{r}', fil, f=INPUT_FONT, fl=INPUT_FILL, a=al('center'))
        put(ws, f'D{r}', uni, f=INPUT_FONT, fl=INPUT_FILL, a=al('center'))
        put(ws, f'E{r}', f'={R("Filial Pneu","tbCentrosCusto")}&"-"&{R("Cód Unidade","tbCentrosCusto")}', f=font(9, False, C['muted']),
            a=al('center'), expected=f'{fil}-{uni}')
    add_table(ws, 'tbCentrosCusto', f'B{ch}:E{r}', ccol,
              calc={'Chave': f'{R("Filial Pneu","tbCentrosCusto")}&"-"&{R("Cód Unidade","tbCentrosCusto")}'})
    table_header_style(ws, ch, 2, 5, h=22)
    put(ws, f'F{ch}', 'Mesmo De-Para da planilha anterior (aba De Para, colunas L:O). Combinações ausentes aparecem como '
        '"Não mapeada" e são listadas na Auditoria.', f=font(8, False, C['muted'], italic=True), a=al('left', 'top', wrap=True))
    ws.merge_cells(f'F{ch}:H{ch+2}')
    ws.freeze_panes = 'A5'
    return last_param_row

# ------------------------------------------------------------------ FROTAS
FROTA_IN = ['Frota', 'Placa', 'Tipo', 'Marca', 'Modelo', 'Carroceria', 'Tipo Operação', 'Local de Operação', 'Gestor', 'Status Frota']
FROTA_CALC = ['Pneus Aplicados', 'Em Demanda Ressolagem', 'Críticos', 'Em Demanda Compra', 'Menor MM', 'Aferições Vencidas', 'Pontuação']

def frota_formulas():
    F = lambda c: R(c, 'tbFrotas')
    base = f'tbPneus[N. Frota],{F("Frota")},tbPneus[Situação Pneu],"USO"'
    return {
        'Pneus Aplicados': f'COUNTIFS({base})',
        'Em Demanda Ressolagem': f'COUNTIFS({base},tbPneus[Demanda Ressolagem],"SIM")',
        'Críticos': f'COUNTIFS({base},tbPneus[Classificação Desgaste],"CRÍTICO")',
        'Em Demanda Compra': f'COUNTIFS({base},tbPneus[Demanda Compra],"SIM")',
        'Menor MM': f'IF({F("Pneus Aplicados")}=0,"",_xlfn.MINIFS(tbPneus[Menor Milimetragem],{base}))',
        'Aferições Vencidas': f'COUNTIFS({base},tbPneus[Status Aferição],"Vencida")',
        'Pontuação': (f'IF({F("Pneus Aplicados")}=0,0,{F("Críticos")}*1000000+{F("Em Demanda Ressolagem")}*10000'
                      f'+{F("Em Demanda Compra")}*100+MAX(0,20-N({F("Menor MM")}))+ROW()/10000000)'),
    }

def build_frotas(wb, ws, frotas_rows):
    sheet_base(ws, C['gold'])
    widths(ws, {'B': 9, 'C': 10, 'D': 15, 'E': 14, 'F': 16, 'G': 12, 'H': 18, 'I': 22, 'J': 25, 'K': 13,
                'L': 10, 'M': 12, 'N': 9, 'O': 11, 'P': 9, 'Q': 11, 'R': 11})
    header(ws, 'Cadastro de frotas', 'Vínculo N. Frota (Rodopar) → placa, tipo e base de operação. Fonte: Base Cadastro '
           f'Frotas (arquivo Gestão de CPK) + alocação de {fontes.data_alocacao():%d/%m/%Y} da Base de Fidelização. Atualize as '
           'colunas amarelas quando houver troca de veículo ou de base.', 18)
    hr = 6
    cols = FROTA_IN + FROTA_CALC
    for i, h in enumerate(cols):
        put(ws, f'{col(2+i)}{hr}', h)
    fml = frota_formulas()
    r = hr
    for i, f in enumerate(frotas_rows):
        r += 1
        for j, cname in enumerate(FROTA_IN):
            c = put(ws, f'{col(2+j)}{r}', f.get(cname), f=font(9, False, C['blue_deep']), a=al('left' if j > 1 else 'center'))
            if j == 0: c.number_format = '@'
        for j, cname in enumerate(FROTA_CALC):
            jj = len(FROTA_IN) + j
            exp = f[cname] if cname != 'Pontuação' else None
            if cname == 'Pontuação':
                if f['Pneus Aplicados'] == 0:
                    exp = 0
                else:
                    mmv = f['Menor MM'] if M.isnum(f['Menor MM']) else 0
                    exp = f['Críticos'] * 1e6 + f['Em Demanda Ressolagem'] * 1e4 + f['Em Demanda Compra'] * 100 + max(0, 20 - mmv) + r / 1e7
                f['Pontuação'] = exp
            put(ws, f'{col(2+jj)}{r}', '=' + fml[cname], f=font(9), a=al('center'),
                nf='0.00' if cname == 'Menor MM' else ('0.0000000' if cname == 'Pontuação' else '0'), expected=exp)
    last = r
    add_table(ws, 'tbFrotas', f'B{hr}:{col(1+len(cols))}{last}', cols, calc=fml)
    table_header_style(ws, hr, 2, 1 + len(cols), h=34)
    for ci in range(2 + len(FROTA_IN), 2 + len(cols)):
        ws.cell(hr, ci).fill = fill(C['blue_deep'])
    # entradas em amarelo-claro
    for rr in range(hr + 1, last + 1):
        for ci in range(2, 2 + len(FROTA_IN)):
            ws.cell(rr, ci).fill = INPUT_FILL
    ws.column_dimensions['R'].hidden = True
    put(ws, 'B5', 'Colunas em amarelo: cadastro (editável). Colunas azul-escuro: indicadores calculados a partir da Base Tratada. '
        'Frota deve ser digitada como texto (ex.: VA113, 3897).', f=font(8, False, C['muted'], italic=True))
    ws.freeze_panes = f'C{hr+1}'
    # alerta visual: veículos com pneus em demanda
    red = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['danger_fg']), fill=fill(C['danger_soft']))
    amb = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['warning_fg']), fill=fill(C['warning_soft']))
    from openpyxl.formatting.rule import Rule
    rng = f'M{hr+1}:M{hr+400}'
    ws.conditional_formatting.add(rng, Rule(type='cellIs', operator='greaterThan', formula=['0'], dxf=red))
    ws.conditional_formatting.add(f'N{hr+1}:N{hr+400}', Rule(type='cellIs', operator='greaterThan', formula=['0'], dxf=red))
    ws.conditional_formatting.add(f'O{hr+1}:O{hr+400}', Rule(type='cellIs', operator='greaterThan', formula=['0'], dxf=amb))
    return last

# ------------------------------------------------------------------ FLUXO RESSOLAGEM
FLUXO_IN = ['Lote', 'Referência do Lote', 'N.Fogo', 'Status', 'Recapadora', 'Data do Status', 'Observação']
FLUXO_CALC = ['Medida', 'Vida Atual', 'Menor MM Atual', 'Situação Atual', 'Placa Atual', 'Etapa', 'Situação no Fluxo',
              'Último Registro', 'Chave']

def fluxo_formulas():
    F = lambda c: R(c, 'tbFluxo')
    m = f'MATCH({F("N.Fogo")},tbPneus[N.Fogo],0)'
    return {
        'Medida': f'IFERROR(INDEX(tbPneus[Dimensão],{m})&"","Fogo não localizado")',
        'Vida Atual': f'IFERROR(INDEX(tbPneus[N. Vida],{m}),"")',
        'Menor MM Atual': f'IFERROR(INDEX(tbPneus[Menor Milimetragem],{m}),"")',
        'Situação Atual': f'IFERROR(INDEX(tbPneus[Situação Pneu],{m})&"","")',
        'Placa Atual': f'IFERROR(INDEX(tbPneus[Placa],{m})&"","")',
        'Etapa': f'IFERROR(INDEX(tbFluxoStatus[Etapa],MATCH({F("Status")},tbFluxoStatus[Status],0)),"")',
        'Situação no Fluxo': f'IFERROR(INDEX(tbFluxoStatus[Situação no Fluxo],MATCH({F("Status")},tbFluxoStatus[Status],0)),"Status inválido")',
        'Último Registro': f'IF(_xlfn.MAXIFS(tbFluxo[Lote],tbFluxo[N.Fogo],{F("N.Fogo")})={F("Lote")},"SIM","NÃO")',
        'Chave': f'{F("N.Fogo")}&"|"&{F("Lote")}',
    }

def build_fluxo(wb, ws, fluxo_rows, base_rows):
    sheet_base(ws, C['gold'])
    widths(ws, {'B': 7, 'C': 15, 'D': 9, 'E': 22, 'F': 17, 'G': 12, 'H': 36, 'I': 12, 'J': 8, 'K': 9, 'L': 11,
                'M': 10, 'N': 7, 'O': 12, 'P': 10, 'Q': 12, 'R': 3, 'S': 22, 'T': 7, 'U': 12, 'V': 24, 'W': 84})
    header(ws, 'Fluxo de ressolagem — controle de lotes', 'Registro dos pneus enviados (ou decididos) em cada lote de ressolagem. '
           'Inclua uma linha por pneu e atualize o Status a cada etapa. Histórico de 569 registros migrado da planilha anterior.', 23)
    byfogo = {b['N.Fogo']: b for b in base_rows}
    st = {s[1]: s for s in FLUXO_STATUS}
    # KPIs
    hr = 11
    n = len(fluxo_rows)
    lote_max = max(f['Lote'] for f in fluxo_rows)
    em_and = sum(1 for f in fluxo_rows if st.get(f['Status'], (0, 0, 'Aberto'))[2] == 'Aberto')
    kpi(ws, 6, 2, 3, 'REGISTROS NO FLUXO', '=ROWS(tbFluxo[N.Fogo])', n)
    kpi(ws, 6, 5, 2, 'LOTE MAIS RECENTE', '=MAX(tbFluxo[Lote])', lote_max, tone='blue')
    kpi(ws, 6, 7, 2, 'EM ANDAMENTO', '=COUNTIF(tbFluxo[Situação no Fluxo],"Aberto")', em_and, tone='warning')
    ref_lote = next(f['Referência do Lote'] for f in fluxo_rows if f['Lote'] == lote_max)
    kpi(ws, 6, 9, 3, 'REFERÊNCIA DO ÚLTIMO LOTE', '=INDEX(tbFluxo[Referência do Lote],MATCH(MAX(tbFluxo[Lote]),tbFluxo[Lote],0))&""',
        ref_lote, nf='@', tone='default')
    ws['I7'].font = font(14, True, C['text'])
    cols = FLUXO_IN + FLUXO_CALC
    for i, h in enumerate(cols):
        put(ws, f'{col(2+i)}{hr}', h)
    fml = fluxo_formulas()
    r = hr
    lote_by_fogo = {}
    for f in fluxo_rows:
        lote_by_fogo[f['N.Fogo']] = max(lote_by_fogo.get(f['N.Fogo'], -1), f['Lote'])
    for f in fluxo_rows:
        r += 1
        for j, cname in enumerate(FLUXO_IN):
            v = f.get(cname)
            c = put(ws, f'{col(2+j)}{r}', v, f=font(9, False, C['blue_deep']),
                    a=al('left' if cname in ('Observação', 'Recapadora', 'Status', 'Referência do Lote') else 'center'))
            if cname == 'Data do Status': c.number_format = 'dd/mm/yyyy'
        b = byfogo.get(f['N.Fogo'])
        def idx(v):
            return 0 if v is None else v
        exp = {
            'Medida': (b['Dimensão'] or '') if b else 'Fogo não localizado',
            'Vida Atual': idx(b['N. Vida']) if b else '',
            'Menor MM Atual': idx(b['Menor Milimetragem']) if b else '',
            'Situação Atual': (b['Situação Pneu'] or '') if b else '',
            'Placa Atual': (b['Placa'] or '') if b else '',
            'Etapa': st[f['Status']][0] if f['Status'] in st else '',
            'Situação no Fluxo': st[f['Status']][2] if f['Status'] in st else 'Status inválido',
            'Último Registro': 'SIM' if lote_by_fogo[f['N.Fogo']] == f['Lote'] else 'NÃO',
            'Chave': f"{f['N.Fogo']}|{f['Lote']}",
        }
        for j, cname in enumerate(FLUXO_CALC):
            jj = len(FLUXO_IN) + j
            put(ws, f'{col(2+jj)}{r}', '=' + fml[cname], f=font(9), a=al('center'),
                nf='0.00' if cname == 'Menor MM Atual' else 'General', expected=exp[cname])
    last = r
    add_table(ws, 'tbFluxo', f'B{hr}:{col(1+len(cols))}{last}', cols, calc=fml)
    table_header_style(ws, hr, 2, 1 + len(cols), h=34)
    for ci in range(2 + len(FLUXO_IN), 2 + len(cols)):
        ws.cell(hr, ci).fill = fill(C['blue_deep'])
    for rr in range(hr + 1, last + 1):
        for ci in range(2, 2 + len(FLUXO_IN)):
            ws.cell(rr, ci).fill = INPUT_FILL
    ws.column_dimensions['Q'].hidden = True
    # validações
    dv1 = DataValidation(type='list', formula1='=lstStatusFluxo', allow_blank=False, showErrorMessage=True,
                         errorTitle='Status inválido', error='Escolha um status da lista (tabela de status ao lado).')
    dv2 = DataValidation(type='list', formula1='=lstRecapadoras', allow_blank=True, showErrorMessage=False)
    dv3 = DataValidation(type='whole', operator='greaterThan', formula1='0', allow_blank=False, showErrorMessage=True,
                         errorTitle='Lote inválido', error='Informe o número do lote (inteiro).')
    for dv in (dv1, dv2, dv3): ws.add_data_validation(dv)
    dv1.add(f'E{hr+1}:E{last+1000}'); dv2.add(f'F{hr+1}:F{last+1000}'); dv3.add(f'B{hr+1}:B{last+1000}')
    # destaque de status
    from openpyxl.formatting.rule import Rule
    for txt, tone in [('Aberto', 'warning')]:
        ds = DifferentialStyle(font=Font(name=FONT, bold=True, color=C['warning_fg']), fill=fill(C['warning_soft']))
        ws.conditional_formatting.add(f'P{hr+1}:P{last+1000}', Rule(type='cellIs', operator='equal', formula=['"Aberto"'], dxf=ds))
    # tabela de status (referência das listas)
    put(ws, 'S10', 'STATUS DO FLUXO (LISTA)', f=font(9, True, C['blue']))
    scol = ['Status', 'Etapa', 'Situação no Fluxo', 'Equivalência anterior', 'Descrição']
    for i, h in enumerate(scol):
        put(ws, f'{col(19+i)}{hr}', h)
    rr = hr
    for e, s, sit, desc, eq in FLUXO_STATUS:
        rr += 1
        put(ws, f'S{rr}', s, f=font(9, True)); put(ws, f'T{rr}', e, f=font(9), a=al('center'))
        put(ws, f'U{rr}', sit, f=font(9), a=al('center')); put(ws, f'V{rr}', eq, f=font(9, False, C['muted']))
        put(ws, f'W{rr}', desc, f=font(9), a=al('left', 'center'))
    add_table(ws, 'tbFluxoStatus', f'S{hr}:W{rr}', scol, style='TableStyleLight9')
    table_header_style(ws, hr, 19, 23, h=34, bg=C['neutral'])
    rr += 2
    put(ws, f'S{rr}', 'RECAPADORAS (LISTA)', f=font(9, True, C['blue']))
    rr += 1
    put(ws, f'S{rr}', 'Recapadora')
    rh = rr
    nomes_recap, variacoes = fontes.recapadoras(fluxo_rows)
    for rc in nomes_recap:
        rr += 1
        put(ws, f'S{rr}', rc, f=INPUT_FONT, fl=INPUT_FILL)
    add_table(ws, 'tbRecapadoras', f'S{rh}:S{rr}', ['Recapadora'], style='TableStyleLight9')
    table_header_style(ws, rh, 19, 19, h=22, bg=C['neutral'])
    nota = 'Histórico: "Realizado" → Ressolado; "Recusado Recapagem" → Reprovado; "Descartar" mantido.'
    if variacoes:
        nota += ' ' + '; '.join(f'"{v}" foi preservado como no original (provável grafia de {f})' for v, f in variacoes) + '.'
    put(ws, f'S{rr+2}', nota, f=font(8, False, C['muted'], italic=True), a=al('left', 'top', wrap=True))
    ws.merge_cells(f'S{rr+2}:W{rr+4}')
    defname(wb, 'lstStatusFluxo', 'tbFluxoStatus[Status]')
    defname(wb, 'lstRecapadoras', 'tbRecapadoras[Recapadora]')
    ws.freeze_panes = f'A{hr+1}'
    return last

# ------------------------------------------------------------------ BASE TRATADA (tbPneus)
def pneus_formulas():
    mm = R('Menor Milimetragem')
    S = 'tbPneus[[#This Row],[Sulco 1]:[Sulco 4]]'
    V = R('N. Vida')
    sit = R('Situação Pneu')
    medm = f'MATCH({R("Dimensão")},tbMedidas[Medida],0)'
    frm = f'MATCH({R("N. Frota")},tbFrotas[Frota],0)'
    lc, lr, lp, lf = 'pLimCritico', 'pLimRessolagem', 'pLimCompra', 'pLimFaixa'
    fx = lambda p: f'FIXED({p},2)'
    f = {}
    f['Ativo'] = f'IF(COUNTIFS(tbSituacoes[Situação Pneu],{sit},tbSituacoes[Inativo],"SIM")>0,"NÃO","SIM")'
    f['Utilização'] = f'IFERROR(INDEX(tbMedidas[Utilização],{medm}),"Não cadastrada")'
    f['Placa'] = f'IF({R("N. Frota")}="","",IFERROR(INDEX(tbFrotas[Placa],{frm}),"(sem cadastro)"))'
    f['Tipo Veículo'] = f'IF({R("N. Frota")}="","",IFERROR(INDEX(tbFrotas[Tipo],{frm})&"",""))'
    f['Local de Operação'] = (f'IF({sit}="USO",IFERROR(INDEX(tbFrotas[Local de Operação],{frm})&"","(sem cadastro)"),'
                              f'IF({sit}="ESTOQUE","Estoque FL "&{R("Filial Pneu")},PROPER({sit})))')
    f['Operação'] = (f'IFERROR(INDEX(tbCentrosCusto[Operação],MATCH({R("Filial Pneu")}&"-"&{R("Cód Unidade")},'
                     f'tbCentrosCusto[Chave],0)),"Não mapeada")')
    p = R('Posição')
    f['Eixo'] = (f'IF({p}="","",IF(LEFT({p},2)="ED","Dianteiro",IF(LEFT({p},3)="EST","Estepe",'
                 f'IF(LEFT({p},2)="ET","Traseiro","Outro"))))')
    f['Faixa MM'] = (f'IF(NOT(ISNUMBER({mm})),"Sem medição",IF({mm}<{lc},"< "&{fx(lc)}&" mm",'
                     f'IF({mm}<{lr},{fx(lc)}&" a < "&{fx(lr)}&" mm",IF({mm}<{lp},{fx(lr)}&" a < "&{fx(lp)}&" mm",'
                     f'IF({mm}<{lf},{fx(lp)}&" a < "&{fx(lf)}&" mm","≥ "&{fx(lf)}&" mm")))))')
    spec = f'INDEX(tbMedidas[Limite Ressolagem Específico (mm)],{medm})'
    f['Limite Ressolagem Aplicado'] = f'IFERROR(IF(ISNUMBER({spec}),{spec},{lr}),{lr})'
    f['Demanda Ressolagem'] = f'IF(AND({R("Ativo")}="SIM",ISNUMBER({mm}),{mm}<{R("Limite Ressolagem Aplicado")}),"SIM","NÃO")'
    f['Classificação Desgaste'] = (f'IF({R("Ativo")}="NÃO","INATIVO",IF(NOT(ISNUMBER({mm})),"SEM AFERIÇÃO",'
                                   f'IF({R("Demanda Ressolagem")}="SIM",IF(AND({sit}="USO",{mm}<{lc}),"CRÍTICO",'
                                   f'"ENVIAR PARA RESSOLAGEM"),IF({mm}<{lp},"ACOMPANHAR","NORMAL"))))')
    f['Destino Sugerido'] = (f'IF({R("Demanda Ressolagem")}<>"SIM","",IF(N({V})>=IFERROR(INDEX(tbMedidas[Vida Máxima],{medm}),999),'
                             f'"DESCARTAR (LIMITE DE VIDA)","RESSOLAR"))')
    f['Demanda Compra'] = (f'IF(AND({R("Ativo")}="SIM",ISNUMBER({mm}),{mm}<{lp},{V}=pVidaCompra,'
                           f'OR(pIncluirEstoque="SIM",{sit}<>"ESTOQUE")),"SIM","NÃO")')
    f['Estoque Utilizável'] = f'IF(AND({sit}="ESTOQUE",ISNUMBER({mm}),{mm}>={R("Limite Ressolagem Aplicado")}),"SIM","NÃO")'
    f['Cobertura Compra'] = f'IF(AND({sit}="ESTOQUE",{V}=pVidaCompra,ISNUMBER({mm}),{mm}>={lp}),"SIM","NÃO")'
    f['Dias desde Medição'] = f'IF(ISNUMBER({R("Dt. Medição")}),INT({R("Atualizado em")})-INT({R("Dt. Medição")}),"")'
    d = R('Dias desde Medição')
    f['Status Aferição'] = (f'IF({R("Ativo")}="NÃO","",IF({d}="","Sem medição",IF({d}>pAfPrazo,"Vencida",'
                            f'IF({d}>pAfAviso,"Próx. do vencimento","Em dia"))))')
    f['Lote Ressolagem'] = f'IF(COUNTIF(tbFluxo[N.Fogo],{R("N.Fogo")})=0,"",_xlfn.MAXIFS(tbFluxo[Lote],tbFluxo[N.Fogo],{R("N.Fogo")}))'
    f['Status Fluxo'] = (f'IF({R("Lote Ressolagem")}="","",IFERROR(INDEX(tbFluxo[Status],MATCH({R("N.Fogo")}&"|"&{R("Lote Ressolagem")},'
                         f'tbFluxo[Chave],0))&"",""))')
    stf = R('Status Fluxo')
    f['Situação Fluxo'] = (f'IF({stf}="",IF({R("Demanda Ressolagem")}="SIM","Não programado",""),'
                           f'IF(IFERROR(INDEX(tbFluxoStatus[Situação no Fluxo],MATCH({stf},tbFluxoStatus[Status],0)),"Aberto")="Aberto",'
                           f'"Em andamento",IF(OR({stf}="Reprovado",{stf}="Descartar"),IF({R("Ativo")}="SIM",'
                           f'"Encerrado – baixar no Rodopar","Encerrado"),IF({R("Demanda Ressolagem")}="SIM","Não programado","Concluído"))))')
    cls = '{' + ','.join(f'"{c}"' for c in M.CLASSES) + '}'
    f['Prioridade'] = f'IFERROR(MATCH({R("Classificação Desgaste")},{cls},0),9)'
    a = {}
    a['AUD Sem medida'] = f'IF({R("Dimensão")}="",1,0)'
    a['AUD Medida não cadastrada'] = f'IF(AND({R("Dimensão")}<>"",COUNTIF(tbMedidas[Medida],{R("Dimensão")})=0),1,0)'
    a['AUD Sem Menor Milimetragem'] = f'IF(ISNUMBER({mm}),0,1)'
    a['AUD Milimetragem inválida'] = f'IF(OR(AND(ISNUMBER({mm}),OR({mm}<=0,{mm}>pMMMax)),MAX({S})>pMMMax),1,0)'
    a['AUD Menor MM diferente do menor sulco'] = f'IF(AND(ISNUMBER({mm}),COUNT({S})>0),IF(ABS({mm}-MIN({S}))>0.005,1,0),0)'
    a['AUD Variação entre sulcos'] = f'IF(COUNT({S})>1,IF(AND(MAX({S})<=pMMMax,MAX({S})-MIN({S})>pVarSulcos),1,0),0)'
    a['AUD Situação não cadastrada'] = f'IF(OR({sit}="",COUNTIF(tbSituacoes[Situação Pneu],{sit})=0),1,0)'
    a['AUD N. Vida inválido'] = f'IF(ISNUMBER({V}),IF(OR({V}<1,{V}<>INT({V})),1,0),1)'
    a['AUD N.Fogo vazio ou duplicado'] = f'IF({R("N.Fogo")}="",1,IF(COUNTIF(tbPneus[N.Fogo],{R("N.Fogo")})>1,1,0))'
    a['AUD Em uso sem frota ou posição'] = f'IF(AND({sit}="USO",OR({R("N. Frota")}="",{R("Posição")}="")),1,0)'
    a['AUD Frota sem cadastro'] = f'IF({R("N. Frota")}="",0,IF(COUNTIF(tbFrotas[Frota],{R("N. Frota")})=0,1,0))'
    a['AUD Fora de uso com frota'] = f'IF(AND({sit}<>"USO",{R("N. Frota")}<>""),1,0)'
    a['AUD Km Real negativo'] = f'IF(AND(ISNUMBER({R("Km Real")}),{R("Km Real")}<0),1,0)'
    dot = R('Dot')
    a['AUD DOT inválido'] = (f'IF(ISNUMBER({dot}),IF(OR({dot}<101,{dot}>5399,INT({dot}/100)<1,INT({dot}/100)>53,'
                             f'MOD({dot},100)>MOD(YEAR({R("Atualizado em")}),100)),1,0),IF({dot}="",0,1))')
    a['AUD Pneu novo com MM baixa'] = f'IF(AND({R("Status")}="Novo",ISNUMBER({mm}),{mm}<{lp}),1,0)'
    a['AUD Vida 1 com dados de recapagem'] = f'IF(AND({V}=1,OR({R("Borracha")}<>"",{R("Desenho")}<>"")),1,0)'
    a['AUD Filial ou unidade sem De-Para'] = f'IF({R("Operação")}="Não mapeada",1,0)'
    f['Qtd Inconsistências'] = f'SUM(tbPneus[[#This Row],[{M.AUD_COLS[0]}]:[{M.AUD_COLS[-1]}]])'
    parts = '&'.join(f'IF({R(c)}=1,"; {lab}","")' for c, lab, *_ in M.AUD)
    f['Inconsistências'] = f'MID({parts},3,600)'
    f.update(a)
    return f

PQ_FMT = {'Data Compra': 'dd/mm/yyyy', 'Dt. Medição': 'dd/mm/yyyy hh:mm', 'Dt. Calibragem': 'dd/mm/yyyy hh:mm',
          'Data Cadastro': 'dd/mm/yyyy hh:mm', 'Data Última Alteração': 'dd/mm/yyyy hh:mm', 'Atualizado em': 'dd/mm/yyyy hh:mm',
          'Menor Milimetragem': '0.00', 'Sulco 1': '0.00', 'Sulco 2': '0.00', 'Sulco 3': '0.00', 'Sulco 4': '0.00',
          'Calibragem': '0.0', 'Km Rodado': '#,##0', 'Km Real': '#,##0'}
CALC_FMT = {'Limite Ressolagem Aplicado': '0.00', 'Dias desde Medição': '0', 'Prioridade': '0', 'Qtd Inconsistências': '0'}
PQ_W = {'N.Fogo': 9, 'Filial Pneu': 7, 'Cód Unidade': 7, 'Cód Custo': 7, 'Data Compra': 11, 'Situação Pneu': 10,
        'N. Frota': 8, 'Filial Frota': 7, 'Marca': 13, 'Modelo Pneu': 20, 'Dimensão': 12, 'Posição': 8,
        'Menor Milimetragem': 10, 'Sulco 1': 7, 'Sulco 2': 7, 'Sulco 3': 7, 'Sulco 4': 7, 'Dt. Medição': 15,
        'Calibragem': 8, 'Dt. Calibragem': 15, 'Km Rodado': 10, 'Km Real': 10, 'Dot': 7, 'N. Vida': 7, 'Condição': 12,
        'Classificação': 11, 'Status': 17, 'Data Cadastro': 15, 'Número De Série': 9, 'Usuário De Inclusão': 17,
        'Usuário Última Alteração': 17, 'Data Última Alteração': 15, 'Desenho': 9, 'Borracha': 10,
        'Tratamento Aplicado': 30, 'Atualizado em': 15, 'Arquivo Fonte': 40}

def build_base(wb, ws, rows):
    sheet_base(ws, C['neutral'], zoom=85)
    cols = M.PQ_COLS + M.CALC_COLS
    hr = 5
    put(ws, 'B2', 'Base tratada — Rodopar 10', f=font(15, True, C['blue']))
    put(ws, 'B3', 'Gerada pela consulta Power Query "Pneus_Rodopar10" (colunas cinza) + colunas de regra calculadas (azul) e '
        'de auditoria (AUD). Não digite nesta tabela: ela é recriada a cada "Atualizar Tudo".', f=font(9, False, C['muted']))
    for i, h in enumerate(cols):
        put(ws, f'{col(2+i)}{hr}', h)
    fml = pneus_formulas()
    nf_map = dict(PQ_FMT); nf_map.update(CALC_FMT)
    r = hr
    for row in rows:
        r += 1
        for j, cname in enumerate(cols):
            coord = f'{col(2+j)}{r}'
            if cname in M.PQ_COLS:
                c = ws[coord]
                c.value = row[cname]
            else:
                c = put(ws, coord, '=' + fml[cname], expected=row[cname])
            c.font = font(9)
            if cname in nf_map: c.number_format = nf_map[cname]
    last = r
    add_table(ws, 'tbPneus', f'B{hr}:{col(1+len(cols))}{last}', cols, calc=fml, style='TableStyleLight1')
    ws.row_dimensions[hr].height = 42
    for j, cname in enumerate(cols):
        c = ws.cell(hr, 2 + j)
        c.font = font(9, True, 'FFFFFF')
        c.alignment = al('center', 'center', wrap=True)
        if cname in M.PQ_COLS:
            c.fill = fill(C['neutral'])
        elif cname.startswith('AUD'):
            c.fill = fill('8F2323')
        else:
            c.fill = fill(C['blue'])
        w = PQ_W.get(cname, 12 if not cname.startswith('AUD') else 10)
        if cname in ('Classificação Desgaste', 'Destino Sugerido', 'Situação Fluxo'): w = 22
        if cname == 'Inconsistências': w = 40
        if cname in ('Local de Operação', 'Operação'): w = 18
        ws.column_dimensions[col(2 + j)].width = w
    ws.freeze_panes = f'C{hr+1}'
    # agrupar colunas de auditoria
    first_aud = 2 + cols.index(M.AUD_COLS[0]); last_aud = 2 + cols.index(M.AUD_COLS[-1])
    ws.column_dimensions.group(col(first_aud), col(last_aud), hidden=False, outline_level=1)
    for j, cname in enumerate(cols):
        if cname.startswith('AUD'):
            for rr in range(hr + 1, last + 1):
                ws.cell(rr, 2 + j).number_format = '"●";;""'
                ws.cell(rr, 2 + j).alignment = al('center')
                ws.cell(rr, 2 + j).font = font(9, True, C['danger'])
    return last
