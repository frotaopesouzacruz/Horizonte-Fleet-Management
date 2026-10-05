"""Modelo de dados: emulação fiel do Power Query e de todas as fórmulas da planilha.

Serve a dois propósitos:
  1. gerar a carga inicial (dados pré-carregados) e os valores em cache das fórmulas;
  2. validar, célula a célula, o resultado recalculado pelo LibreOffice.
"""
import datetime as dt, math, re, warnings
import collections
import openpyxl
from config import *
import fontes
warnings.filterwarnings('ignore')

RODOPAR_COLS = ["N.Fogo", "Filial Pneu", "Cód Unidade", "Cód Custo", "Data Compra", "Situação Pneu", "N. Frota",
                "Filial Frota", "Marca", "Modelo Pneu", "Dimensão", "Posição", "Menor Milimetragem", "Sulco 1",
                "Sulco 2", "Sulco 3", "Sulco 4", "Dt. Medição", "Calibragem", "Dt. Calibragem", "Km Rodado",
                "Km Real", "Dot", "N. Vida", "Condição", "Classificação", "Status", "Data Cadastro",
                "Número De Série", "Usuário De Inclusão", "Usuário Última Alteração", "Data Última Alteração",
                "Desenho", "Borracha"]
PQ_EXTRA = ["Tratamento Aplicado", "Atualizado em", "Arquivo Fonte"]
PQ_COLS = RODOPAR_COLS + PQ_EXTRA

INT_COLS = {"N.Fogo", "Filial Pneu", "Cód Unidade", "Cód Custo", "Filial Frota", "Dot", "N. Vida"}
NUM_COLS = {"Menor Milimetragem", "Sulco 1", "Sulco 2", "Sulco 3", "Sulco 4", "Calibragem", "Km Rodado", "Km Real",
            "Número De Série"}
DATE_COLS = {"Data Compra"}
DATETIME_COLS = {"Dt. Medição", "Dt. Calibragem", "Data Cadastro", "Data Última Alteração"}
TEXT_COLS = {"Marca", "Modelo Pneu", "Posição", "Condição", "Classificação", "Status", "Usuário De Inclusão",
             "Usuário Última Alteração", "Desenho", "Borracha"}

AUD = [  # coluna, rótulo curto, severidade, escopo, como corrigir
    ("AUD Sem medida", "Sem medida", "Alta", "Todos", "Preencher a Dimensão do pneu no Rodopar."),
    ("AUD Medida não cadastrada", "Medida não cadastrada", "Alta", "Todos",
     "Incluir a medida na tabela de Medidas (Parâmetros) ou corrigir a grafia no Rodopar."),
    ("AUD Sem Menor Milimetragem", "Sem Menor Milimetragem", "Alta", "Todos",
     "Registrar a aferição de sulcos no Rodopar."),
    ("AUD Milimetragem inválida", "Milimetragem inválida", "Alta", "Todos",
     "Corrigir sulco ou menor milimetragem ≤ 0 ou acima do máximo plausível."),
    ("AUD Menor MM diferente do menor sulco", "Menor MM diferente do menor sulco", "Média", "Todos",
     "Conferir a aferição: a Menor Milimetragem deve ser o menor dos 4 sulcos."),
    ("AUD Variação entre sulcos", "Variação entre sulcos acima do limite", "Média", "Todos",
     "Conferir a aferição (possível erro de digitação) ou avaliar desgaste irregular/alinhamento."),
    ("AUD Situação não cadastrada", "Situação vazia ou não cadastrada", "Alta", "Todos",
     "Corrigir a Situação no Rodopar ou cadastrar a nova situação em Parâmetros."),
    ("AUD N. Vida inválido", "N. Vida vazio ou inválido", "Alta", "Todos", "Corrigir o número de vidas no Rodopar."),
    ("AUD N.Fogo vazio ou duplicado", "N.Fogo vazio ou duplicado", "Alta", "Todos",
     "Remover a duplicidade ou preencher o número de fogo no Rodopar."),
    ("AUD Em uso sem frota ou posição", "Em uso sem frota ou posição", "Média", "USO",
     "Vincular o pneu à frota/posição no Rodopar."),
    ("AUD Frota sem cadastro", "Frota sem cadastro de veículo", "Média", "Todos",
     "Incluir a frota na aba Frotas (placa, tipo, base)."),
    ("AUD Fora de uso com frota", "Fora de uso com frota vinculada", "Baixa", "Não USO",
     "Remover o vínculo de frota do pneu no Rodopar."),
    ("AUD Km Real negativo", "Km Real negativo", "Baixa", "Todos",
     "Revisar hodômetros de montagem/desmontagem no Rodopar."),
    ("AUD DOT inválido", "DOT inválido", "Baixa", "Todos",
     "Corrigir o DOT (semana 01–53 + ano com 2 dígitos, não futuro)."),
    ("AUD Pneu novo com MM baixa", "Pneu 'Novo' com milimetragem baixa", "Média", "Todos",
     "Conferir a aferição: pneu com status Novo abaixo do limite de compra."),
    ("AUD Vida 1 com dados de recapagem", "Vida 1 com dados de recapagem", "Média", "Todos",
     "Conferir o N. Vida: pneu de 1ª vida com Borracha/Desenho de recapagem preenchidos."),
    ("AUD Filial ou unidade sem De-Para", "Filial/Unidade sem De-Para de operação", "Baixa", "Todos",
     "Incluir a combinação Filial/Unidade na tabela de Centros de Custo (Parâmetros)."),
]
AUD_COLS = [a[0] for a in AUD]

CALC_COLS = ["Ativo", "Utilização", "Placa", "Tipo Veículo", "Local de Operação", "Operação", "Eixo", "Faixa MM",
             "Limite Ressolagem Aplicado", "Demanda Ressolagem", "Classificação Desgaste", "Destino Sugerido",
             "Demanda Compra", "Estoque Utilizável", "Cobertura Compra", "Dias desde Medição", "Status Aferição",
             "Lote Ressolagem", "Status Fluxo", "Situação Fluxo", "Prioridade", "Qtd Inconsistências",
             "Inconsistências"] + AUD_COLS
CLASSES = ["CRÍTICO", "ENVIAR PARA RESSOLAGEM", "ACOMPANHAR", "SEM AFERIÇÃO", "NORMAL", "INATIVO"]


# ------------------------------------------------------------------ helpers (semântica do Excel)
def blank(v):
    return v is None or v == ''


def isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and not (isinstance(v, float) and math.isnan(v))


def oadate(d):
    base = dt.datetime(1899, 12, 30)
    if isinstance(d, dt.date) and not isinstance(d, dt.datetime):
        d = dt.datetime(d.year, d.month, d.day)
    delta = d - base
    return delta.days + delta.seconds / 86400 + delta.microseconds / 86400e6


def fixed_pt(x, n=2):
    s = f"{x:,.{n}f}"
    return s.replace(',', 'X').replace('.', ',').replace('X', '.')


# ------------------------------------------------------------------ Power Query (emulação)
def norm_header(h):
    if h is None:
        return None
    s = str(h).strip().replace('\r', ' ').replace('\n', ' ')
    return ' '.join([p for p in s.split(' ') if p != ''])


def fn_texto(v):
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        v = v.isoformat()
    t = ''.join(ch for ch in str(v) if ch >= ' ' or ch == '\t').strip()
    # Text.Clean remove caracteres de controle; Text.Trim remove espaços nas pontas
    t = ''.join(ch for ch in t if ord(ch) >= 32).strip()
    return None if t == '' else t


def fn_numero(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, (dt.datetime, dt.date)):
        return oadate(v)
    t = str(v).replace(',', '.').strip()
    if t == '':
        return None
    try:
        return float(t)
    except ValueError:
        return None


def fn_inteiro(v):
    n = fn_numero(v)
    if n is None:
        return None
    # Number.Round usa arredondamento bancário (ToEven) por padrão no M
    return int(round(n))


def fn_data(v):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)):
        return dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))
    t = str(v).strip()
    for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y'):
        try:
            return dt.datetime.strptime(t, fmt)
        except ValueError:
            pass
    return None


def fn_medida(v):
    t = fn_texto(v)
    if t is None:
        return None
    return t.upper().replace(' ', '').replace('R', ' R')


def fn_situacao(v):
    t = fn_texto(v)
    return None if t is None else t.upper()


def fn_frota(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return str(int(v))
    t = fn_texto(v)
    return None if t is None else t.upper()


def fn_log(r):
    itens = []
    cal = r.get('Calibragem')
    if isinstance(cal, (dt.datetime, dt.date)):
        itens.append('Calibragem convertida de data para número')
    dim = r.get('Dimensão')
    if dim is not None and fn_medida(dim) != str(dim):
        itens.append('Medida padronizada')
    sit = r.get('Situação Pneu')
    if sit is not None and fn_situacao(sit) != str(sit):
        itens.append('Situação padronizada')
    if any(isinstance(r.get(c), str) for c in ['Menor Milimetragem', 'Sulco 1', 'Sulco 2', 'Sulco 3', 'Sulco 4', 'N. Vida']):
        itens.append('Número armazenado como texto convertido')
    if any(isinstance(r.get(c), str) and r.get(c).strip() != r.get(c) for c in ['Marca', 'Modelo Pneu', 'Posição', 'Status', 'N. Frota']):
        itens.append('Espaços extras removidos')
    return '; '.join(itens) if itens else None


def read_rodopar_raw(path=None):
    wb = openpyxl.load_workbook(path or fontes.arquivo('rodopar'), data_only=True)
    ws = wb.worksheets[0]
    if 'Planilha1' in wb.sheetnames:
        ws = wb['Planilha1']
    rows = list(ws.iter_rows(values_only=True))
    hdr = [norm_header(h) for h in rows[0]]
    recs = []
    for row in rows[1:]:
        rec = {}
        for h, v in zip(hdr, row):
            if h is not None:
                rec[h] = v
        recs.append(rec)
    return recs


def emulate_pq(raw, atualizado=CARGA_INICIAL, fonte=None):
    fonte = fonte or fontes.arquivo_fonte_inicial()
    out = []
    for r in raw:
        sel = {c: r.get(c) for c in RODOPAR_COLS}
        vals = [None if (isinstance(v, str) and v.strip() == '') else v for v in sel.values()]
        if sum(v is not None for v in vals) == 0:
            continue
        log = fn_log(sel)
        n = {}
        for c in RODOPAR_COLS:
            v = sel[c]
            if c in INT_COLS:
                n[c] = fn_inteiro(v)
            elif c in NUM_COLS:
                n[c] = fn_numero(v)
            elif c in DATE_COLS:
                d = fn_data(v)
                n[c] = None if d is None else dt.datetime(d.year, d.month, d.day)
            elif c in DATETIME_COLS:
                n[c] = fn_data(v)
            elif c == 'Situação Pneu':
                n[c] = fn_situacao(v)
            elif c == 'N. Frota':
                n[c] = fn_frota(v)
            elif c == 'Dimensão':
                n[c] = fn_medida(v)
            elif c == 'Posição':
                t = fn_texto(v)
                n[c] = None if t is None else t.upper()
            else:
                n[c] = fn_texto(v)
        n['Tratamento Aplicado'] = log
        n['Atualizado em'] = atualizado
        n['Arquivo Fonte'] = fonte
        out.append(n)
    out.sort(key=lambda d: (d['N.Fogo'] is not None, d['N.Fogo'] if d['N.Fogo'] is not None else 0))
    return out


def data_base(base):
    """Data da base Rodopar 10: última alteração registrada (datas posteriores à carga são ignoradas)."""
    datas = [b['Data Última Alteração'] for b in base
             if isinstance(b['Data Última Alteração'], dt.datetime) and b['Data Última Alteração'] <= CARGA_INICIAL]
    return max(datas).date() if datas else CARGA_INICIAL.date()


# ------------------------------------------------------------------ tabelas de apoio
def read_frotas():
    """Cadastro de veículos: Base Cadastro Frotas (CPK) + alocação mais recente da Base de Fidelização."""
    cpk = fontes._wb('cpk')
    ws = cpk['Base Cadastro Frotas']
    cad = {}
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        fr, pl, tp, ma, mo, ca = row[1:7]
        if fr is None:
            continue
        key = str(int(fr)) if isinstance(fr, (int, float)) else str(fr).strip().upper()
        cad[key] = dict(Frota=key, Placa=str(pl).strip().upper(), Tipo=tp, Marca=ma, Modelo=str(mo).strip(), Carroceria=ca)
    ctrl = fontes._wb('controle', False)
    wsf = ctrl['Base_Dados_Fidelização']
    dia = fontes.data_alocacao()
    alloc = {}
    for r in range(10, wsf.max_row + 1):
        ref, fr, pl, top, tp, mo, data, st, loc, ges = [wsf.cell(r, c).value for c in range(2, 12)]
        if data is None or fr is None:
            continue
        if not isinstance(data, dt.datetime) or data.date() != dia:
            continue
        key = str(int(fr)) if isinstance(fr, (int, float)) else str(fr).strip().upper()
        alloc[key] = dict(Placa=pl, Tipo=tp, Carroceria=mo, TipoOp=top, Local=loc, Gestor=ges, Status=st)
    rows = []
    for key in sorted(set(cad) | set(alloc), key=lambda k: (not k.isdigit(), k)):
        c = cad.get(key, {})
        a = alloc.get(key, {})
        rows.append({
            'Frota': key,
            'Placa': c.get('Placa') or a.get('Placa'),
            'Tipo': c.get('Tipo') or a.get('Tipo'),
            'Marca': c.get('Marca'),
            'Modelo': c.get('Modelo'),
            'Carroceria': c.get('Carroceria') or a.get('Carroceria'),
            'Tipo Operação': a.get('TipoOp'),
            'Local de Operação': a.get('Local'),
            'Gestor': a.get('Gestor'),
            'Status Frota': a.get('Status') or 'Sem alocação',
        })
    return rows


def read_historico():
    ctrl = fontes._wb('controle', False)
    ws = ctrl['Ressolagem']
    obs = {}
    rec = fontes._wb('recapagem', False).active
    for r in range(16, rec.max_row + 1):
        f, w = rec.cell(r, 2).value, rec.cell(r, 23).value
        if f is not None and w:
            obs[(int(f), int(rec.cell(r, 3).value))] = str(w).strip()
    rows = []
    for r in range(14, ws.max_row + 1):
        fogo, dem, mes, recap, st = [ws.cell(r, c).value for c in (4, 5, 6, 7, 8)]
        if fogo is None:
            continue
        rows.append({
            'Lote': int(dem),
            'Referência do Lote': mes,
            'N.Fogo': int(fogo),
            'Status': STATUS_MAP_ANTIGO.get(str(st).strip(), str(st).strip()),
            'Recapadora': recap,
            'Data do Status': None,
            'Observação': obs.get((int(fogo), int(dem))),
        })
    return rows


# ------------------------------------------------------------------ regras (emulação das fórmulas)
def compute(base, frotas, fluxo, P=P):
    med = {m[0]: dict(util=m[1], vmax=m[2], lim=m[3], custo=m[4], custo_r=m[5]) for m in fontes.medidas()}
    sit_inat = {s[0] for s in SITUACOES if s[2] == 'SIM'}
    sit_all = {s[0] for s in SITUACOES}
    cc = {f"{a}-{b}": o for a, b, o in fontes.centros_custo()}
    fr = {f['Frota']: f for f in frotas}
    st_sit = {s[1]: s[2] for s in FLUXO_STATUS}
    fogos = [b['N.Fogo'] for b in base]
    # fluxo: chave e maxifs
    lote_max, chave_status = {}, {}
    for f in fluxo:
        k = f['N.Fogo']
        lote_max[k] = max(lote_max.get(k, -1e18), f['Lote'])
        ch = f"{f['N.Fogo']}|{f['Lote']}"
        chave_status.setdefault(ch, f['Status'])
    out = []
    for b in base:
        r = dict(b)
        sit = b['Situação Pneu'] or ''
        mm = b['Menor Milimetragem']
        vida = b['N. Vida']
        dim = b['Dimensão']
        frota = b['N. Frota']
        r['Ativo'] = 'NÃO' if sit in sit_inat else 'SIM'
        r['Utilização'] = med[dim]['util'] if dim in med else 'Não cadastrada'
        if blank(frota):
            r['Placa'] = ''
            r['Tipo Veículo'] = ''
        else:
            r['Placa'] = fr[frota]['Placa'] if frota in fr else '(sem cadastro)'
            r['Tipo Veículo'] = (fr[frota]['Tipo'] or '') if frota in fr else ''
        if sit == 'USO':
            r['Local de Operação'] = ((fr[frota]['Local de Operação'] or '') if frota in fr else '(sem cadastro)')
        elif sit == 'ESTOQUE':
            r['Local de Operação'] = 'Estoque FL ' + ('' if b['Filial Pneu'] is None else str(b['Filial Pneu']))
        else:
            r['Local de Operação'] = sit.title() if sit else ''
        key = f"{'' if b['Filial Pneu'] is None else b['Filial Pneu']}-{'' if b['Cód Unidade'] is None else b['Cód Unidade']}"
        r['Operação'] = cc.get(key, 'Não mapeada')
        pos = b['Posição'] or ''
        r['Eixo'] = '' if pos == '' else ('Dianteiro' if pos[:2] == 'ED' else ('Estepe' if pos[:3] == 'EST' else ('Traseiro' if pos[:2] == 'ET' else 'Outro')))
        lc, lr, lp, lf = P['pLimCritico'], P['pLimRessolagem'], P['pLimCompra'], P['pLimFaixa']
        if not isnum(mm):
            r['Faixa MM'] = 'Sem medição'
        elif mm < lc:
            r['Faixa MM'] = f"< {fixed_pt(lc)} mm"
        elif mm < lr:
            r['Faixa MM'] = f"{fixed_pt(lc)} a < {fixed_pt(lr)} mm"
        elif mm < lp:
            r['Faixa MM'] = f"{fixed_pt(lr)} a < {fixed_pt(lp)} mm"
        elif mm < lf:
            r['Faixa MM'] = f"{fixed_pt(lp)} a < {fixed_pt(lf)} mm"
        else:
            r['Faixa MM'] = f"≥ {fixed_pt(lf)} mm"
        spec = med[dim]['lim'] if dim in med else None
        lim_ap = spec if isnum(spec) else lr
        r['Limite Ressolagem Aplicado'] = lim_ap
        r['Demanda Ressolagem'] = 'SIM' if (r['Ativo'] == 'SIM' and isnum(mm) and mm < lim_ap) else 'NÃO'
        if r['Ativo'] == 'NÃO':
            cls = 'INATIVO'
        elif not isnum(mm):
            cls = 'SEM AFERIÇÃO'
        elif r['Demanda Ressolagem'] == 'SIM':
            cls = 'CRÍTICO' if (sit == 'USO' and mm < lc) else 'ENVIAR PARA RESSOLAGEM'
        else:
            cls = 'ACOMPANHAR' if mm < lp else 'NORMAL'
        r['Classificação Desgaste'] = cls
        if r['Demanda Ressolagem'] != 'SIM':
            r['Destino Sugerido'] = ''
        else:
            vmax = med[dim]['vmax'] if dim in med else 999
            r['Destino Sugerido'] = 'DESCARTAR (LIMITE DE VIDA)' if (vida if isnum(vida) else 0) >= vmax else 'RESSOLAR'
        inc_est = P['pIncluirEstoque'] == 'SIM'
        r['Demanda Compra'] = 'SIM' if (r['Ativo'] == 'SIM' and isnum(mm) and mm < lp and vida == P['pVidaCompra']
                                        and (inc_est or sit != 'ESTOQUE')) else 'NÃO'
        r['Estoque Utilizável'] = 'SIM' if (sit == 'ESTOQUE' and isnum(mm) and mm >= lim_ap) else 'NÃO'
        r['Cobertura Compra'] = 'SIM' if (sit == 'ESTOQUE' and vida == P['pVidaCompra'] and isnum(mm) and mm >= lp) else 'NÃO'
        med_dt = b['Dt. Medição']
        r['Dias desde Medição'] = ((b['Atualizado em'].date() - med_dt.date()).days if isinstance(med_dt, dt.datetime) else '')
        if r['Ativo'] == 'NÃO':
            r['Status Aferição'] = ''
        elif r['Dias desde Medição'] == '':
            r['Status Aferição'] = 'Sem medição'
        elif r['Dias desde Medição'] > P['pAfPrazo']:
            r['Status Aferição'] = 'Vencida'
        elif r['Dias desde Medição'] > P['pAfAviso']:
            r['Status Aferição'] = 'Próx. do vencimento'
        else:
            r['Status Aferição'] = 'Em dia'
        fogo = b['N.Fogo']
        if fogo in lote_max:
            r['Lote Ressolagem'] = lote_max[fogo]
            r['Status Fluxo'] = chave_status.get(f"{fogo}|{lote_max[fogo]}", '')
        else:
            r['Lote Ressolagem'] = ''
            r['Status Fluxo'] = ''
        stf = r['Status Fluxo']
        if stf == '':
            r['Situação Fluxo'] = 'Não programado' if r['Demanda Ressolagem'] == 'SIM' else ''
        elif st_sit.get(stf, 'Aberto') == 'Aberto':
            r['Situação Fluxo'] = 'Em andamento'
        elif stf in ('Reprovado', 'Descartar'):
            r['Situação Fluxo'] = 'Encerrado – baixar no Rodopar' if r['Ativo'] == 'SIM' else 'Encerrado'
        else:
            r['Situação Fluxo'] = 'Não programado' if r['Demanda Ressolagem'] == 'SIM' else 'Concluído'
        r['Prioridade'] = CLASSES.index(cls) + 1 if cls in CLASSES else 9
        # auditoria
        sul = [b[f'Sulco {i}'] for i in range(1, 5)]
        sul_n = [s for s in sul if isnum(s)]
        a = {}
        a['AUD Sem medida'] = blank(dim)
        a['AUD Medida não cadastrada'] = (not blank(dim)) and dim not in med
        a['AUD Sem Menor Milimetragem'] = not isnum(mm)
        a['AUD Milimetragem inválida'] = (isnum(mm) and (mm <= 0 or mm > P['pMMMax'])) or (max(sul_n) if sul_n else 0) > P['pMMMax']
        a['AUD Menor MM diferente do menor sulco'] = isnum(mm) and len(sul_n) > 0 and abs(mm - min(sul_n)) > 0.005
        a['AUD Variação entre sulcos'] = len(sul_n) > 1 and max(sul_n) <= P['pMMMax'] and (max(sul_n) - min(sul_n)) > P['pVarSulcos']
        a['AUD Situação não cadastrada'] = sit == '' or sit not in sit_all
        a['AUD N. Vida inválido'] = (not isnum(vida)) or vida < 1 or vida != int(vida)
        a['AUD N.Fogo vazio ou duplicado'] = blank(fogo) or fogos.count(fogo) > 1
        a['AUD Em uso sem frota ou posição'] = sit == 'USO' and (blank(frota) or blank(b['Posição']))
        a['AUD Frota sem cadastro'] = (not blank(frota)) and frota not in fr
        a['AUD Fora de uso com frota'] = sit != 'USO' and not blank(frota)
        a['AUD Km Real negativo'] = isnum(b['Km Real']) and b['Km Real'] < 0
        dot = b['Dot']
        yy = b['Atualizado em'].year % 100
        if isnum(dot):
            a['AUD DOT inválido'] = dot < 101 or dot > 5399 or int(dot / 100) < 1 or int(dot / 100) > 53 or (dot % 100) > yy
        else:
            a['AUD DOT inválido'] = not blank(dot)
        a['AUD Pneu novo com MM baixa'] = b['Status'] == 'Novo' and isnum(mm) and mm < lp
        a['AUD Vida 1 com dados de recapagem'] = vida == 1 and (not blank(b['Borracha']) or not blank(b['Desenho']))
        a['AUD Filial ou unidade sem De-Para'] = r['Operação'] == 'Não mapeada'
        for k, v in a.items():
            r[k] = 1 if v else 0
        r['Qtd Inconsistências'] = sum(r[k] for k in AUD_COLS)
        labels = [lab for col, lab, *_ in AUD if r[col] == 1]
        r['Inconsistências'] = '; '.join(labels)
        out.append(r)
    return out


def frotas_calc(frotas, rows):
    res = []
    for i, f in enumerate(frotas):
        k = f['Frota']
        uso = [r for r in rows if r['N. Frota'] == k and r['Situação Pneu'] == 'USO']
        d = dict(f)
        d['Pneus Aplicados'] = len(uso)
        d['Em Demanda Ressolagem'] = sum(r['Demanda Ressolagem'] == 'SIM' for r in uso)
        d['Críticos'] = sum(r['Classificação Desgaste'] == 'CRÍTICO' for r in uso)
        d['Em Demanda Compra'] = sum(r['Demanda Compra'] == 'SIM' for r in uso)
        mms = [r['Menor Milimetragem'] for r in uso if isnum(r['Menor Milimetragem'])]
        d['Menor MM'] = min(mms) if (len(uso) > 0 and mms) else ('' if len(uso) == 0 else 0)
        d['Aferições Vencidas'] = sum(r['Status Aferição'] == 'Vencida' for r in uso)
        res.append(d)
    return res


if __name__ == '__main__':  # diagnóstico rápido das regras sobre as fontes atuais
    raw = read_rodopar_raw()
    base = emulate_pq(raw)
    frotas = read_frotas()
    fluxo = read_historico()
    rows = compute(base, frotas, fluxo)
    conta = lambda col: dict(collections.Counter(r[col] for r in rows))
    print(len(rows), 'pneus;', len(frotas), 'frotas;', len(fluxo), 'registros de fluxo')
    print(conta('Classificação Desgaste'))
    print(conta('Destino Sugerido'))
    print(conta('Situação Fluxo'))
    print({c: sum(r[c] == 'SIM' for r in rows)
           for c in ('Demanda Ressolagem', 'Demanda Compra', 'Cobertura Compra', 'Estoque Utilizável')})
    print({c: sum(r[c] for r in rows) for c in AUD_COLS})
    print('qualidade', round(1 - sum(r['Qtd Inconsistências'] > 0 for r in rows) / len(rows), 4))
    print(conta('Tratamento Aplicado'))
    print(conta('Status Aferição'))
    print(conta('Faixa MM'))
    fc = frotas_calc(frotas, rows)
    print(sorted([(f['Frota'], f['Em Demanda Ressolagem'], f['Em Demanda Compra'], f['Menor MM']) for f in fc if f['Pneus Aplicados']], key=lambda x: (-x[1], -x[2]))[:12])
