"""Abas de planejamento e análise: Compra, Ressolagem, Estoque, Análises, Auditoria."""
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import DataBarRule, FormulaRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from views_common import *
import fontes

MED = [m[0] for m in MEDIDAS_REGRAS]
MEDINFO = {m[0]: m for m in MEDIDAS_REGRAS}  # Medida, Utilização, Vida Máxima, Limite Específico

def info_medida(m):
    """(Medida, Utilização, Vida Máxima, Limite Específico, Custo Novo, Custo Ressolagem, Referência)."""
    return {x[0]: x for x in fontes.medidas()}[m]

def med_row_formula(tbl, hdr_row):
    return f'=INDEX(tbMedidas[Medida],ROW()-ROW({tbl}[[#Headers],[Medida]]))'

def util_formula(tbl):
    return f'=INDEX(tbMedidas[Utilização],ROW()-ROW({tbl}[[#Headers],[Medida]]))'

def T(tbl, c):
    return f'{tbl}[[#This Row],[{c}]]'

# ======================================================================= COMPRA
def build_compra(wb, ws, rows):
    sheet_base(ws, C['blue'])
    widths(ws, {'B': 11, 'C': 13, 'D': 11, 'E': 12, 'F': 25, 'G': 11, 'H': 9, 'I': 10, 'J': 13, 'K': 9, 'L': 10,
                'M': 12, 'N': 10, 'O': 17, 'P': 12, 'Q': 7, 'R': 17, 'S': 17, 'T': 21})
    header(ws, 'Planejamento de compra de pneus', 'Demanda projetada de pneus de 1ª vida chegando ao limite, '
           'cobertura pelo estoque e necessidade líquida de compra por medida.', 20)
    lp, vc, inc = P['pLimCompra'], P['pVidaCompra'], P['pIncluirEstoque']
    put(ws, 'B5', '="Regra: Menor Milimetragem < "&FIXED(pLimCompra,2)&" mm   ·   N. Vida = "&pVidaCompra&"   ·   Situação diferente de Baixado e Descarte   ·   Pneus em estoque incluídos: "&pIncluirEstoque&"   ·   Cobertura: estoque com N. Vida = "&pVidaCompra&" e Menor Milimetragem ≥ "&FIXED(pLimCompra,2)&" mm"',
        f=font(9, True, C['blue_deep']), a=al('left', 'center'),
        expected=f"Regra: Menor Milimetragem < {M.fixed_pt(lp)} mm   ·   N. Vida = {vc}   ·   Situação diferente de Baixado e Descarte   ·   Pneus em estoque incluídos: {inc}   ·   Cobertura: estoque com N. Vida = {vc} e Menor Milimetragem ≥ {M.fixed_pt(lp)} mm")
    ws.row_dimensions[5].height = 22
    # ---------------- resumo por medida (tabela)
    hr = 11
    cols = ['Utilização', 'Medida', 'Pneus Ativos', 'Demanda Projetada', 'Em Uso', 'Em Estoque', 'Estoque Disponível',
            'Necessidade Líquida', 'Cobertura', 'Pressão de Reposição', 'Custo Unitário Ref.', 'Investimento Estimado',
            'Ressolagem em Andamento', 'Ordem']
    tb = 'tbResCompra'
    fml = {
        'Utilização': f'INDEX(tbMedidas[Utilização],ROW()-ROW({tb}[[#Headers],[Medida]]))',
        'Medida': f'INDEX(tbMedidas[Medida],ROW()-ROW({tb}[[#Headers],[Medida]]))',
        'Pneus Ativos': f'COUNTIFS(tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Ativo],"SIM")',
        'Demanda Projetada': f'COUNTIFS(tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Demanda Compra],"SIM")',
        'Em Uso': f'COUNTIFS(tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Demanda Compra],"SIM",tbPneus[Situação Pneu],"USO")',
        'Em Estoque': f'COUNTIFS(tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Demanda Compra],"SIM",tbPneus[Situação Pneu],"ESTOQUE")',
        'Estoque Disponível': f'COUNTIFS(tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Cobertura Compra],"SIM")',
        'Necessidade Líquida': f'MAX(0,{T(tb,"Demanda Projetada")}-{T(tb,"Estoque Disponível")})',
        'Cobertura': f'IF({T(tb,"Demanda Projetada")}=0,"",MIN(1,{T(tb,"Estoque Disponível")}/{T(tb,"Demanda Projetada")}))',
        'Pressão de Reposição': f'IF({T(tb,"Pneus Ativos")}=0,"",{T(tb,"Demanda Projetada")}/{T(tb,"Pneus Ativos")})',
        'Custo Unitário Ref.': f'IFERROR(INDEX(tbMedidas[Custo Pneu Novo (R$)],MATCH({T(tb,"Medida")},tbMedidas[Medida],0)),0)',
        'Investimento Estimado': f'{T(tb,"Necessidade Líquida")}*{T(tb,"Custo Unitário Ref.")}',
        'Ressolagem em Andamento': f'COUNTIFS(tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Situação Fluxo],"Em andamento")',
        'Ordem': f'{T(tb,"Necessidade Líquida")}*1000000+{T(tb,"Demanda Projetada")}*1000+(1000-ROW())/1000',
    }
    for i, h in enumerate(cols):
        put(ws, f'{col(2+i)}{hr}', h)
    summary = []
    for k, m in enumerate(MED):
        r = hr + 1 + k
        info = info_medida(m)
        dem = cnt(rows, Dimensão=m, **{'Demanda__Compra': 'SIM'})
        est = cnt(rows, Dimensão=m, **{'Cobertura__Compra': 'SIM'})
        atv = cnt(rows, Dimensão=m, Ativo='SIM')
        nec = max(0, dem - est)
        e = {'Utilização': info[1], 'Medida': m, 'Pneus Ativos': atv, 'Demanda Projetada': dem,
             'Em Uso': cnt(rows, Dimensão=m, **{'Demanda__Compra': 'SIM', 'Situação__Pneu': 'USO'}),
             'Em Estoque': cnt(rows, Dimensão=m, **{'Demanda__Compra': 'SIM', 'Situação__Pneu': 'ESTOQUE'}),
             'Estoque Disponível': est, 'Necessidade Líquida': nec,
             'Cobertura': '' if dem == 0 else min(1, est / dem),
             'Pressão de Reposição': '' if atv == 0 else dem / atv,
             'Custo Unitário Ref.': info[4] or 0, 'Investimento Estimado': nec * (info[4] or 0),
             'Ressolagem em Andamento': cnt(rows, Dimensão=m, **{'Situação__Fluxo': 'Em andamento'})}
        e['Ordem'] = nec * 1e6 + dem * 1e3 + (1000 - r) / 1000
        summary.append(e)
        for i, h in enumerate(cols):
            nf = {'Cobertura': '0%', 'Pressão de Reposição': '0%', 'Custo Unitário Ref.': '"R$" #,##0.00',
                  'Investimento Estimado': '"R$" #,##0', 'Ordem': '0.000'}.get(h, '#,##0')
            put(ws, f'{col(2+i)}{r}', '=' + fml[h], f=font(10, h in ('Medida', 'Necessidade Líquida')),
                a=al('left' if i < 2 else 'center'), nf=nf, expected=e[h])
    last = hr + len(MED)
    add_table(ws, tb, f'B{hr}:{col(1+len(cols))}{last}', cols, calc=fml)
    table_header_style(ws, hr, 2, 1 + len(cols), h=34)
    # total
    tr = last + 1
    put(ws, f'B{tr}', 'Total', f=font(10, True), a=al('left'))
    tot = {}
    for i, h in enumerate(cols[2:], start=2):
        cc = col(2 + i)
        if h in ('Cobertura',):
            v = (min(1, sum(s['Estoque Disponível'] for s in summary) / sum(s['Demanda Projetada'] for s in summary))
                 if sum(s['Demanda Projetada'] for s in summary) else '')
            put(ws, f'{cc}{tr}', f'=IF(SUM({tb}[Demanda Projetada])=0,"",MIN(1,SUM({tb}[Estoque Disponível])/SUM({tb}[Demanda Projetada])))',
                f=font(10, True), a=al('center'), nf='0%', expected=v)
        elif h == 'Pressão de Reposição':
            v = sum(s['Demanda Projetada'] for s in summary) / sum(s['Pneus Ativos'] for s in summary)
            put(ws, f'{cc}{tr}', f'=IF(SUM({tb}[Pneus Ativos])=0,"",SUM({tb}[Demanda Projetada])/SUM({tb}[Pneus Ativos]))',
                f=font(10, True), a=al('center'), nf='0%', expected=v)
        elif h in ('Custo Unitário Ref.', 'Ordem'):
            continue
        else:
            v = sum(s[h] for s in summary)
            tot[h] = v
            put(ws, f'{cc}{tr}', f'=SUM({tb}[{h}])', f=font(10, True), a=al('center'),
                nf='"R$" #,##0' if h == 'Investimento Estimado' else '#,##0', expected=v)
    ws.column_dimensions[col(1 + len(cols))].hidden = True
    for ci in range(2, 1 + len(cols)):
        c = ws.cell(tr, ci)
        c.fill = fill(C['primary_soft'])
        c.border = Border(top=Side(style='thin', color=C['blue']), bottom=Side(style='thin', color=C['blue']))
    medida_warning(ws, f'B{tr+1}', tb)
    ws.conditional_formatting.add(f'I{hr+1}:I{last}', DataBarRule(start_type='num', start_value=0, end_type='max',
                                                                    color='E9A3A3', showValue=True))
    ws.conditional_formatting.add(f'K{hr+1}:K{last}', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1,
                                                                    color='9ED3EC', showValue=True))
    # ---------------- KPIs (dependem do resumo)
    dem_t, est_t, nec_t, inv_t = tot['Demanda Projetada'], tot['Estoque Disponível'], tot['Necessidade Líquida'], tot['Investimento Estimado']
    kpi(ws, 7, 2, 2, 'DEMANDA PROJETADA', f'=SUM({tb}[Demanda Projetada])', dem_t, tone='blue', note='pneus de 1ª vida < limite')
    kpi(ws, 7, 4, 2, 'ESTOQUE DISPONÍVEL', f'=SUM({tb}[Estoque Disponível])', est_t, tone='success', note='cobertura (1ª vida ≥ limite)')
    kpi(ws, 7, 6, 2, 'NECESSIDADE LÍQUIDA', f'=SUM({tb}[Necessidade Líquida])', nec_t, tone='danger', note='demanda − estoque, por medida')
    kpi(ws, 7, 8, 3, 'INVESTIMENTO ESTIMADO', f'=SUM({tb}[Investimento Estimado])', inv_t, nf='"R$" #,##0', tone='gold',
        note='necessidade × custo de referência')
    nmed = sum(1 for s in summary if s['Necessidade Líquida'] > 0)
    kpi(ws, 7, 11, 2, 'MEDIDAS COM NECESSIDADE', f'=COUNTIF({tb}[Necessidade Líquida],">0")', nmed, tone='warning',
        note=f'de {len(MED)} medidas cadastradas')
    ws['K9'].value = f'="de "&ROWS({tb}[Medida])&" medidas cadastradas"'
    EXPECTED[(ws.title, 'K9')] = f'de {len(MED)} medidas cadastradas'
    cobg = (min(1, est_t / dem_t) if dem_t else '')
    kpi(ws, 7, 13, 2, 'COBERTURA GERAL', f'=IF(SUM({tb}[Demanda Projetada])=0,"",MIN(1,SUM({tb}[Estoque Disponível])/SUM({tb}[Demanda Projetada])))',
        cobg, nf='0%', tone='default', note='estoque ÷ demanda')
    section(ws, 'B10', 'Resumo por medida', 14)
    ws.cell(hr, 1 + len(cols)).fill = fill(C['blue'])
    # ---------------- por utilização
    ur = tr + 2
    section(ws, f'B{ur}', 'Por utilização', 8)
    ucols = ['Utilização', 'Demanda Projetada', 'Estoque Disponível', 'Necessidade Líquida', 'Investimento Estimado']
    put(ws, f'B{ur+1}', ucols[0]); put(ws, f'C{ur+1}', ucols[1]); put(ws, f'D{ur+1}', ucols[2]); put(ws, f'E{ur+1}', ucols[3])
    put(ws, f'F{ur+1}', ucols[4])
    table_header_style(ws, ur + 1, 2, 6, h=30)
    for k, u in enumerate(['Van', 'Fiorino', 'Caminhão']):
        r = ur + 2 + k
        put(ws, f'B{r}', u, f=font(10, True), a=al('left'))
        for j, h in enumerate(ucols[1:]):
            v = sum(s[h] for s in summary if s['Utilização'] == u)
            put(ws, f'{col(3+j)}{r}', f'=SUMIFS({tb}[{h}],{tb}[Utilização],$B{r})', f=font(10), a=al('center'),
                nf='"R$" #,##0' if h == 'Investimento Estimado' else '#,##0', expected=v)
        for ci in range(2, 7):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border']))
    put(ws, f'H{ur+1}', 'Por que o estoque ressolado não abate a compra?', f=font(9, True, C['blue_deep']))
    put(ws, f'H{ur+2}', 'Na base, 100% das posições dianteiras (EDE/EDD) usam pneus de 1ª vida. A demanda projetada repõe '
        'pneus de 1ª vida; por isso só o estoque de 1ª vida acima do limite cobre a compra (mesma lógica da lista manual '
        'anterior). Pneus ressolados abastecem posições traseiras e estepe e aparecem na aba Estoque.',
        f=font(8.5, False, C['muted']), a=al('left', 'top', wrap=True))
    ws.merge_cells(f'H{ur+2}:N{ur+4}')
    # ---------------- lista
    lr = ur + 7
    section(ws, f'B{lr}', 'Pneus em demanda de compra (rastreabilidade)', 20)
    filter_cell(ws, f'C{lr+1}', f'D{lr+1}', 'Medida:', '(Todas)', '=lstFiltroMedida', 'fCompraMedida', wb)
    filter_cell(ws, f'F{lr+1}', f'G{lr+1}', 'Situação:', '(Todas)', '"(Todas),USO,ESTOQUE"', 'fCompraSit', wb)
    put(ws, f'I{lr+1}', '=COUNTIFS(tbPneus[Demanda Compra],"SIM",tbPneus[Dimensão],IF(fCompraMedida="(Todas)","*",fCompraMedida),tbPneus[Situação Pneu],IF(fCompraSit="(Todas)","*",fCompraSit))&" pneus listados"',
        f=font(9, True, C['blue_deep']), a=al('left'), expected=f"{dem_t} pneus listados")
    link(ws, f'M{lr+1}', 'Ver estoque que cobre a demanda ▸', "'Estoque'!A1")
    hdrs = ['N.Fogo', 'Medida', 'Utilização', 'Marca', 'Modelo', 'Situação', 'Frota', 'Placa', 'Tipo Veículo', 'Posição',
            'Eixo', 'Menor MM', 'N. Vida', 'Status Rodopar', 'Dt. Medição', 'Dias', 'Aferição', 'Operação', 'Local de Operação']
    srcs = ['N.Fogo', 'Dimensão', 'Utilização', 'Marca', 'Modelo Pneu', 'Situação Pneu', 'N. Frota', 'Placa', 'Tipo Veículo',
            'Posição', 'Eixo', 'Menor Milimetragem', 'N. Vida', 'Status', 'Dt. Medição', 'Dias desde Medição', 'Status Aferição',
            'Operação', 'Local de Operação']
    fmts = [('0', 'center'), ('@', 'center'), ('@', 'center'), ('@', 'left'), ('@', 'left'), ('@', 'center'), ('0;-0;;@', 'center'),
            ('@', 'center'), ('@', 'left'), ('0;-0;;@', 'center'), ('@', 'center'), ('0.00', 'center'), ('0', 'center'),
            ('@', 'left'), ('dd/mm/yyyy', 'center'), ('0', 'center'), ('@', 'center'), ('@', 'left'), ('@', 'left')]
    cond = ('(tbPneus[Demanda Compra]="SIM")*(((fCompraMedida="(Todas)")+(tbPneus[Dimensão]=fCompraMedida))>0)'
            '*(((fCompraSit="(Todas)")+(tbPneus[Situação Pneu]=fCompraSit))>0)')
    prow = sel(rows, **{'Demanda__Compra': 'SIM'})
    vals, rng = list_block(ws, lr + 3, 2, hdrs, fmts, cond, srcs, [('Dimensão', 1), ('Menor Milimetragem', 1)], prow,
                           'Nenhum pneu atende aos filtros selecionados.')
    mm_cf(ws, f'M{lr+4}:M{lr+1500}', f'M{lr+4}')
    text_cf(ws, f'R{lr+4}:R{lr+1500}', [('Vencida', 'danger'), ('Próx. do vencimento', 'warning')])
    ws.freeze_panes = 'A5'
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return dict(summary=summary, tot=tot, list_row=lr + 3)

# ======================================================================= RESSOLAGEM
def build_ressolagem(wb, ws, rows):
    sheet_base(ws, C['blue'])
    widths(ws, {'B': 24, 'C': 26, 'D': 9, 'E': 12, 'F': 13, 'G': 24, 'H': 10, 'I': 9, 'J': 10, 'K': 11, 'L': 11,
                'M': 10, 'N': 7, 'O': 18, 'P': 12, 'Q': 7, 'R': 16, 'S': 20, 'T': 7, 'U': 19, 'V': 28})
    header(ws, 'Planejamento de ressolagem', 'Pneus abaixo do limite de ressolagem, prioridade, destino sugerido '
           '(ressolar ou descartar por limite de vida) e situação no fluxo de lotes.', 22)
    put(ws, 'B5', '="Regra: Menor Milimetragem < "&FIXED(pLimRessolagem,2)&" mm (ou limite específico da medida)   ·   Situação diferente de Baixado e Descarte   ·   CRÍTICO: em uso abaixo de "&FIXED(pLimCritico,2)&" mm   ·   ACOMPANHAR: entre "&FIXED(pLimRessolagem,2)&" e "&FIXED(pLimCompra,2)&" mm"',
        f=font(9, True, C['blue_deep']), a=al('left', 'center'),
        expected=f"Regra: Menor Milimetragem < {M.fixed_pt(P['pLimRessolagem'])} mm (ou limite específico da medida)   ·   Situação diferente de Baixado e Descarte   ·   CRÍTICO: em uso abaixo de {M.fixed_pt(P['pLimCritico'])} mm   ·   ACOMPANHAR: entre {M.fixed_pt(P['pLimRessolagem'])} e {M.fixed_pt(P['pLimCompra'])} mm")
    ws.row_dimensions[5].height = 22
    dr = sel(rows, **{'Demanda__Ressolagem': 'SIM'})
    n_dem = len(dr)
    n_res = sum(r['Destino Sugerido'] == 'RESSOLAR' for r in dr)
    n_des = sum(r['Destino Sugerido'] == 'DESCARTAR (LIMITE DE VIDA)' for r in dr)
    n_cri = sum(r['Classificação Desgaste'] == 'CRÍTICO' for r in dr)
    n_uso = sum(r['Situação Pneu'] == 'USO' for r in dr)
    n_est = sum(r['Situação Pneu'] == 'ESTOQUE' for r in dr)
    n_np = sum(r['Situação Fluxo'] == 'Não programado' for r in dr)
    custo = {m[0]: (m[5] or 0) for m in fontes.medidas()}
    c_tot = sum(custo.get(r['Dimensão'], 0) for r in dr if r['Destino Sugerido'] == 'RESSOLAR')
    kpi(ws, 7, 2, 1, 'DEMANDA DE RESSOLAGEM', '=COUNTIF(tbPneus[Demanda Ressolagem],"SIM")', n_dem, tone='danger', note='pneus abaixo do limite')
    kpi(ws, 7, 3, 1, 'A RESSOLAR', '=COUNTIF(tbPneus[Destino Sugerido],"RESSOLAR")', n_res, tone='blue', note='dentro do limite de vida')
    kpi(ws, 7, 4, 3, 'DESCARTAR (LIMITE DE VIDA)', '=COUNTIF(tbPneus[Destino Sugerido],"DESCARTAR (LIMITE DE VIDA)")', n_des, tone='neutral',
        note='vida máxima atingida')
    kpi(ws, 7, 7, 2, 'CRÍTICOS EM USO', '=COUNTIF(tbPneus[Classificação Desgaste],"CRÍTICO")', n_cri, tone='danger',
        note_formula='="abaixo de "&FIXED(pLimCritico,2)&" mm"', note_expected=f"abaixo de {M.fixed_pt(P['pLimCritico'])} mm")
    kpi(ws, 7, 9, 2, 'EM USO / EM ESTOQUE', '=COUNTIFS(tbPneus[Demanda Ressolagem],"SIM",tbPneus[Situação Pneu],"USO")&" / "&COUNTIFS(tbPneus[Demanda Ressolagem],"SIM",tbPneus[Situação Pneu],"ESTOQUE")',
        f'{n_uso} / {n_est}', nf='@', tone='default', note='retirar do veículo / enviar')
    kpi(ws, 7, 11, 3, 'NÃO PROGRAMADOS', '=COUNTIFS(tbPneus[Demanda Ressolagem],"SIM",tbPneus[Situação Fluxo],"Não programado")', n_np,
        tone='warning', note='sem lote aberto no fluxo')
    kpi(ws, 7, 14, 3, 'CUSTO ESTIMADO (RESSOLAR)', '=SUM(tbResRessolagem[Custo Estimado])', c_tot, nf='"R$" #,##0', tone='gold',
        note='a ressolar × custo de referência')
    # resumo por medida
    section(ws, 'B10', 'Resumo por medida', 13)
    hr = 11
    tb = 'tbResRessolagem'
    cols = ['Utilização', 'Medida', 'Limite Aplicado (mm)', 'Demanda', 'Ressolar', 'Descartar', 'Em Uso', 'Em Estoque',
            'Críticos', 'Não Programados', 'Em Andamento', 'Custo Estimado']
    base = f'tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Demanda Ressolagem],"SIM"'
    fml = {
        'Utilização': f'INDEX(tbMedidas[Utilização],ROW()-ROW({tb}[[#Headers],[Medida]]))',
        'Medida': f'INDEX(tbMedidas[Medida],ROW()-ROW({tb}[[#Headers],[Medida]]))',
        'Limite Aplicado (mm)': f'IFERROR(IF(ISNUMBER(INDEX(tbMedidas[Limite Ressolagem Específico (mm)],MATCH({T(tb,"Medida")},tbMedidas[Medida],0))),INDEX(tbMedidas[Limite Ressolagem Específico (mm)],MATCH({T(tb,"Medida")},tbMedidas[Medida],0)),pLimRessolagem),pLimRessolagem)',
        'Demanda': f'COUNTIFS({base})',
        'Ressolar': f'COUNTIFS({base},tbPneus[Destino Sugerido],"RESSOLAR")',
        'Descartar': f'COUNTIFS({base},tbPneus[Destino Sugerido],"DESCARTAR (LIMITE DE VIDA)")',
        'Em Uso': f'COUNTIFS({base},tbPneus[Situação Pneu],"USO")',
        'Em Estoque': f'COUNTIFS({base},tbPneus[Situação Pneu],"ESTOQUE")',
        'Críticos': f'COUNTIFS({base},tbPneus[Classificação Desgaste],"CRÍTICO")',
        'Não Programados': f'COUNTIFS({base},tbPneus[Situação Fluxo],"Não programado")',
        'Em Andamento': f'COUNTIFS(tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Situação Fluxo],"Em andamento")',
        'Custo Estimado': f'{T(tb,"Ressolar")}*IFERROR(INDEX(tbMedidas[Custo Ressolagem (R$)],MATCH({T(tb,"Medida")},tbMedidas[Medida],0)),0)',
    }
    for i, h in enumerate(cols):
        put(ws, f'{col(2+i)}{hr}', h)
    summ = []
    for k, m in enumerate(MED):
        r = hr + 1 + k
        info = info_medida(m)
        d = [x for x in dr if x['Dimensão'] == m]
        e = {'Utilização': info[1], 'Medida': m, 'Limite Aplicado (mm)': info[3] if M.isnum(info[3]) else P['pLimRessolagem'],
             'Demanda': len(d), 'Ressolar': sum(x['Destino Sugerido'] == 'RESSOLAR' for x in d),
             'Descartar': sum(x['Destino Sugerido'] == 'DESCARTAR (LIMITE DE VIDA)' for x in d),
             'Em Uso': sum(x['Situação Pneu'] == 'USO' for x in d), 'Em Estoque': sum(x['Situação Pneu'] == 'ESTOQUE' for x in d),
             'Críticos': sum(x['Classificação Desgaste'] == 'CRÍTICO' for x in d),
             'Não Programados': sum(x['Situação Fluxo'] == 'Não programado' for x in d),
             'Em Andamento': cnt(rows, Dimensão=m, **{'Situação__Fluxo': 'Em andamento'})}
        e['Custo Estimado'] = e['Ressolar'] * (info[5] or 0)
        summ.append(e)
        for i, h in enumerate(cols):
            nf = {'Limite Aplicado (mm)': '0.00', 'Custo Estimado': '"R$" #,##0'}.get(h, '#,##0')
            put(ws, f'{col(2+i)}{r}', '=' + fml[h], f=font(10, h in ('Medida', 'Demanda')), a=al('left' if i < 2 else 'center'),
                nf=nf, expected=e[h])
    last = hr + len(MED)
    add_table(ws, tb, f'B{hr}:{col(1+len(cols))}{last}', cols, calc=fml)
    table_header_style(ws, hr, 2, 1 + len(cols), h=34)
    tr = last + 1
    put(ws, f'B{tr}', 'Total', f=font(10, True))
    for i, h in enumerate(cols[3:], start=3):
        v = sum(s[h] for s in summ)
        put(ws, f'{col(2+i)}{tr}', f'=SUM({tb}[{h}])', f=font(10, True), a=al('center'),
            nf='"R$" #,##0' if h == 'Custo Estimado' else '#,##0', expected=v)
    for ci in range(2, 2 + len(cols)):
        c = ws.cell(tr, ci); c.fill = fill(C['primary_soft'])
        c.border = Border(top=Side(style='thin', color=C['blue']), bottom=Side(style='thin', color=C['blue']))
    ws.conditional_formatting.add(f'E{hr+1}:E{last}', DataBarRule(start_type='num', start_value=0, end_type='max', color='E9A3A3', showValue=True))
    medida_warning(ws, f'B{tr+1}', tb)
    # lista
    lr = tr + 2
    section(ws, f'B{lr}', 'Pneus em demanda de ressolagem (rastreabilidade)', 22)
    filter_cell(ws, f'B{lr+1}', f'C{lr+1}', 'Medida:', '(Todas)', '=lstFiltroMedida', 'fResMedida', wb)
    filter_cell(ws, f'E{lr+1}', f'F{lr+1}', 'Destino:', '(Todos)', '"(Todos),RESSOLAR,DESCARTAR (LIMITE DE VIDA)"', 'fResDestino', wb)
    ws[f'F{lr+1}'].alignment = al('center', 'center', wrap=True)
    filter_cell(ws, f'H{lr+1}', f'I{lr+1}', 'Situação:', '(Todas)', '"(Todas),USO,ESTOQUE"', 'fResSit', wb)
    ws.column_dimensions['I'].width = 10
    put(ws, f'K{lr+1}', '=COUNTIFS(tbPneus[Demanda Ressolagem],"SIM",tbPneus[Dimensão],IF(fResMedida="(Todas)","*",fResMedida),tbPneus[Destino Sugerido],IF(fResDestino="(Todos)","*",fResDestino),tbPneus[Situação Pneu],IF(fResSit="(Todas)","*",fResSit))&" pneus listados"',
        f=font(9, True, C['blue_deep']), expected=f'{n_dem} pneus listados')
    link(ws, f'O{lr+1}', 'Registrar lote no Fluxo de Ressolagem ▸', "'Fluxo Ressolagem'!A1")
    hdrs = ['Classificação', 'Destino Sugerido', 'N.Fogo', 'Medida', 'Marca', 'Modelo', 'Situação', 'Frota', 'Placa', 'Posição',
            'Eixo', 'Menor MM', 'N. Vida', 'Status Rodopar', 'Dt. Medição', 'Dias', 'Operação', 'Local de Operação', 'Lote',
            'Status no Fluxo', 'Situação no Fluxo']
    srcs = ['Classificação Desgaste', 'Destino Sugerido', 'N.Fogo', 'Dimensão', 'Marca', 'Modelo Pneu', 'Situação Pneu',
            'N. Frota', 'Placa', 'Posição', 'Eixo', 'Menor Milimetragem', 'N. Vida', 'Status', 'Dt. Medição',
            'Dias desde Medição', 'Operação', 'Local de Operação', 'Lote Ressolagem', 'Status Fluxo', 'Situação Fluxo']
    fmts = [('@', 'center'), ('@', 'center'), ('0', 'center'), ('@', 'center'), ('@', 'left'), ('@', 'left'), ('@', 'center'),
            ('0;-0;;@', 'center'), ('@', 'center'), ('0;-0;;@', 'center'), ('@', 'center'), ('0.00', 'center'), ('0', 'center'),
            ('@', 'left'), ('dd/mm/yyyy', 'center'), ('0', 'center'), ('@', 'left'), ('@', 'left'), ('0;-0;;@', 'center'),
            ('@', 'left'), ('@', 'left')]
    cond = ('(tbPneus[Demanda Ressolagem]="SIM")*(((fResMedida="(Todas)")+(tbPneus[Dimensão]=fResMedida))>0)'
            '*(((fResDestino="(Todos)")+(tbPneus[Destino Sugerido]=fResDestino))>0)'
            '*(((fResSit="(Todas)")+(tbPneus[Situação Pneu]=fResSit))>0)')
    vals, rng = list_block(ws, lr + 3, 2, hdrs, fmts, cond, srcs, [('Prioridade', 1), ('Menor Milimetragem', 1)], dr,
                           'Nenhum pneu atende aos filtros selecionados.')
    class_cf(ws, f'B{lr+4}:B{lr+1500}')
    text_cf(ws, f'C{lr+4}:C{lr+1500}', [('RESSOLAR', 'blue'), ('DESCARTAR (LIMITE DE VIDA)', 'neutral')])
    text_cf(ws, f'V{lr+4}:V{lr+1500}', [('Não programado', 'warning'), ('Em andamento', 'blue'), ('Encerrado – baixar no Rodopar', 'neutral')])
    mm_cf(ws, f'M{lr+4}:M{lr+1500}', f'M{lr+4}')
    ws.freeze_panes = 'A5'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return dict(summary=summ, n_dem=n_dem, n_res=n_res, n_des=n_des, n_cri=n_cri, n_np=n_np, custo=c_tot)

# ======================================================================= ESTOQUE
def build_estoque(wb, ws, rows):
    sheet_base(ws, C['blue'])
    widths(ws, {'B': 11, 'C': 13, 'D': 13, 'E': 24, 'F': 18, 'G': 9, 'H': 10, 'I': 16, 'J': 11, 'K': 11, 'L': 24, 'M': 27,
                'N': 15, 'O': 12, 'P': 7, 'Q': 7, 'R': 19})
    header(ws, 'Estoque de pneus', 'Composição do estoque por medida e condição: novos, usados de 1ª vida, ressolados, '
           'utilizáveis, comprometidos com ressolagem/descarte e cobertura da compra.', 18)
    est = sel(rows, **{'Situação__Pneu': 'ESTOQUE'})
    n_tot = len(est)
    n_novo = sum(r['Status'] == 'Novo' for r in est)
    n_us1 = sum((r['N. Vida'] == 1 and r['Status'] != 'Novo') for r in est)
    n_res = sum((M.isnum(r['N. Vida']) and r['N. Vida'] >= 2) for r in est)
    n_util = sum(r['Estoque Utilizável'] == 'SIM' for r in est)
    n_comp = sum(r['Demanda Ressolagem'] == 'SIM' for r in est)
    n_cob = sum(r['Cobertura Compra'] == 'SIM' for r in est)
    kpi(ws, 6, 2, 2, 'ESTOQUE TOTAL', '=COUNTIF(tbPneus[Situação Pneu],"ESTOQUE")', n_tot, tone='blue', note='situação ESTOQUE')
    kpi(ws, 6, 4, 2, 'NOVOS', '=COUNTIFS(tbPneus[Situação Pneu],"ESTOQUE",tbPneus[Status],"Novo")', n_novo, tone='default', note='status Novo')
    kpi(ws, 6, 6, 2, 'USADOS 1ª VIDA', '=COUNTIFS(tbPneus[Situação Pneu],"ESTOQUE",tbPneus[N. Vida],1,tbPneus[Status],"<>Novo")', n_us1,
        tone='default', note='vida 1, já rodados')
    kpi(ws, 6, 8, 2, 'RESSOLADOS', '=COUNTIFS(tbPneus[Situação Pneu],"ESTOQUE",tbPneus[N. Vida],">=2")', n_res, tone='default', note='N. Vida ≥ 2')
    kpi(ws, 6, 10, 2, 'UTILIZÁVEL', '=COUNTIF(tbPneus[Estoque Utilizável],"SIM")', n_util, tone='success',
        note_formula='="≥ "&FIXED(pLimRessolagem,2)&" mm"', note_expected=f"≥ {M.fixed_pt(P['pLimRessolagem'])} mm")
    kpi(ws, 6, 12, 1, 'COMPROMETIDO', '=COUNTIFS(tbPneus[Situação Pneu],"ESTOQUE",tbPneus[Demanda Ressolagem],"SIM")', n_comp,
        tone='danger', note='aguarda ressolagem/descarte')
    kpi(ws, 6, 13, 2, 'COBERTURA DE COMPRA', '=COUNTIF(tbPneus[Cobertura Compra],"SIM")', n_cob, tone='success',
        note_formula='="1ª vida ≥ "&FIXED(pLimCompra,2)&" mm"', note_expected=f"1ª vida ≥ {M.fixed_pt(P['pLimCompra'])} mm")
    section(ws, 'B10', 'Estoque por medida', 11)
    hr = 11
    tb = 'tbResEstoque'
    cols = ['Utilização', 'Medida', 'Total', 'Novos', 'Usados 1ª Vida', 'Ressolados', 'Utilizável', 'Comprometido',
            'Cobertura de Compra', 'Em Uso (referência)']
    b = f'tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Situação Pneu],"ESTOQUE"'
    fml = {
        'Utilização': f'INDEX(tbMedidas[Utilização],ROW()-ROW({tb}[[#Headers],[Medida]]))',
        'Medida': f'INDEX(tbMedidas[Medida],ROW()-ROW({tb}[[#Headers],[Medida]]))',
        'Total': f'COUNTIFS({b})',
        'Novos': f'COUNTIFS({b},tbPneus[Status],"Novo")',
        'Usados 1ª Vida': f'COUNTIFS({b},tbPneus[N. Vida],1,tbPneus[Status],"<>Novo")',
        'Ressolados': f'COUNTIFS({b},tbPneus[N. Vida],">=2")',
        'Utilizável': f'COUNTIFS(tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Estoque Utilizável],"SIM")',
        'Comprometido': f'COUNTIFS({b},tbPneus[Demanda Ressolagem],"SIM")',
        'Cobertura de Compra': f'COUNTIFS(tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Cobertura Compra],"SIM")',
        'Em Uso (referência)': f'COUNTIFS(tbPneus[Dimensão],{T(tb,"Medida")},tbPneus[Situação Pneu],"USO")',
    }
    for i, h in enumerate(cols):
        put(ws, f'{col(2+i)}{hr}', h)
    summ = []
    for k, m in enumerate(MED):
        r = hr + 1 + k
        e_ = [x for x in est if x['Dimensão'] == m]
        e = {'Utilização': MEDINFO[m][1], 'Medida': m, 'Total': len(e_), 'Novos': sum(x['Status'] == 'Novo' for x in e_),
             'Usados 1ª Vida': sum((x['N. Vida'] == 1 and x['Status'] != 'Novo') for x in e_),
             'Ressolados': sum((M.isnum(x['N. Vida']) and x['N. Vida'] >= 2) for x in e_),
             'Utilizável': sum(x['Estoque Utilizável'] == 'SIM' for x in e_),
             'Comprometido': sum(x['Demanda Ressolagem'] == 'SIM' for x in e_),
             'Cobertura de Compra': sum(x['Cobertura Compra'] == 'SIM' for x in e_),
             'Em Uso (referência)': cnt(rows, Dimensão=m, **{'Situação__Pneu': 'USO'})}
        summ.append(e)
        for i, h in enumerate(cols):
            put(ws, f'{col(2+i)}{r}', '=' + fml[h], f=font(10, h in ('Medida', 'Total')), a=al('left' if i < 2 else 'center'),
                nf='#,##0', expected=e[h])
    last = hr + len(MED)
    add_table(ws, tb, f'B{hr}:{col(1+len(cols))}{last}', cols, calc=fml)
    table_header_style(ws, hr, 2, 1 + len(cols), h=34)
    tr = last + 1
    put(ws, f'B{tr}', 'Total', f=font(10, True))
    for i, h in enumerate(cols[2:], start=2):
        put(ws, f'{col(2+i)}{tr}', f'=SUM({tb}[{h}])', f=font(10, True), a=al('center'), nf='#,##0', expected=sum(s[h] for s in summ))
    for ci in range(2, 2 + len(cols)):
        c = ws.cell(tr, ci); c.fill = fill(C['primary_soft'])
        c.border = Border(top=Side(style='thin', color=C['blue']), bottom=Side(style='thin', color=C['blue']))
    medida_warning(ws, f'M{tr}', tb)
    put(ws, f'B{tr+1}', 'Utilizável: em estoque com Menor Milimetragem ≥ limite de ressolagem.  Comprometido: em estoque abaixo do limite '
        '(parte da demanda de ressolagem).  Cobertura de compra: 1ª vida com Menor Milimetragem ≥ limite de compra (abate a necessidade de compra).',
        f=font(8, False, C['muted'], italic=True), a=al('left', 'top', wrap=True))
    ws.merge_cells(f'B{tr+1}:L{tr+2}')
    # lista
    lr = tr + 4
    section(ws, f'B{lr}', 'Pneus em estoque (rastreabilidade)', 18)
    filter_cell(ws, f'B{lr+1}', f'C{lr+1}', 'Medida:', '(Todas)', '=lstFiltroMedida', 'fEstMedida', wb)
    filter_cell(ws, f'E{lr+1}', f'F{lr+1}', 'Somente cobertura de compra:', 'NÃO', '"SIM,NÃO"', 'fEstCob', wb)
    ws[f'E{lr+1}'].alignment = al('right', 'center', wrap=True)
    put(ws, f'H{lr+1}', '=COUNTIFS(tbPneus[Situação Pneu],"ESTOQUE",tbPneus[Dimensão],IF(fEstMedida="(Todas)","*",fEstMedida),tbPneus[Cobertura Compra],IF(fEstCob="SIM","SIM","*"))&" pneus listados"',
        f=font(9, True, C['blue_deep']), expected=f'{n_tot} pneus listados')
    hdrs = ['N.Fogo', 'Medida', 'Marca', 'Modelo', 'Status Rodopar', 'N. Vida', 'Menor MM', 'Faixa', 'Utilizável',
            'Cobertura Compra', 'Classificação', 'Destino Sugerido', 'Local / Filial', 'Dt. Medição', 'Dias', 'Lote', 'Status no Fluxo']
    srcs = ['N.Fogo', 'Dimensão', 'Marca', 'Modelo Pneu', 'Status', 'N. Vida', 'Menor Milimetragem', 'Faixa MM',
            'Estoque Utilizável', 'Cobertura Compra', 'Classificação Desgaste', 'Destino Sugerido', 'Local de Operação',
            'Dt. Medição', 'Dias desde Medição', 'Lote Ressolagem', 'Status Fluxo']
    fmts = [('0', 'center'), ('@', 'center'), ('@', 'left'), ('@', 'left'), ('@', 'left'), ('0', 'center'), ('0.00', 'center'),
            ('@', 'center'), ('@', 'center'), ('@', 'center'), ('@', 'center'), ('@', 'center'), ('@', 'center'),
            ('dd/mm/yyyy', 'center'), ('0', 'center'), ('0;-0;;@', 'center'), ('@', 'left')]
    cond = ('(tbPneus[Situação Pneu]="ESTOQUE")*(((fEstMedida="(Todas)")+(tbPneus[Dimensão]=fEstMedida))>0)'
            '*(((fEstCob="NÃO")+(tbPneus[Cobertura Compra]="SIM"))>0)')
    vals, rng = list_block(ws, lr + 3, 2, hdrs, fmts, cond, srcs, [('Dimensão', 1), ('N. Vida', 1), ('Menor Milimetragem', -1)],
                           est, 'Nenhum pneu atende aos filtros selecionados.')
    class_cf(ws, f'L{lr+4}:L{lr+1500}')
    text_cf(ws, f'J{lr+4}:K{lr+1500}', [('SIM', 'success')])
    mm_cf(ws, f'H{lr+4}:H{lr+1500}', f'H{lr+4}')
    ws.freeze_panes = 'A5'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return dict(summary=summ, n_tot=n_tot, n_novo=n_novo, n_res=n_res, n_util=n_util, n_comp=n_comp, n_cob=n_cob, n_us1=n_us1)
