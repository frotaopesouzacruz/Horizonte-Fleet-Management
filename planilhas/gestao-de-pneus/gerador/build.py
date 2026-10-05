"""Orquestra a construção da planilha (etapa 1: openpyxl)."""
import sys, json, copy, datetime as dt
import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.writer.theme import theme_xml
import model as M
import fontes
from config import *
from xlh import *
import xlh
import build_ref as BR
import build_views as BV
import build_views2 as BV2
import build_front as BF

SHEETS = ['Início', 'Dashboard', 'Compra', 'Ressolagem', 'Fluxo Ressolagem', 'Estoque', 'Análises', 'Auditoria',
          'Parâmetros', 'Frotas', 'Base Tratada', 'Documentação']

def brand_theme():
    import re
    t = theme_xml
    cmap = {'dk2': C['navy'], 'lt2': C['bg'], 'accent1': C['blue'], 'accent2': C['cyan'], 'accent3': C['gold'],
            'accent4': C['success'], 'accent5': C['danger'], 'accent6': C['neutral'], 'hlink': C['blue'], 'folHlink': C['blue_deep']}
    for tag, hexc in cmap.items():
        t, n = re.subn(rf'(<a:{tag}>\s*<a:srgbClr val=")[0-9A-F]{{6}}("/>\s*</a:{tag}>)', rf'\g<1>{hexc}\g<2>', t)
        assert n == 1, tag
    t = t.replace('<a:clrScheme name="Office">', '<a:clrScheme name="Horizonte">')
    t = t.replace('<a:fontScheme name="Office">', '<a:fontScheme name="Horizonte">')
    t = t.replace('<a:latin typeface="Cambria"/>', '<a:latin typeface="Arial"/>').replace('<a:latin typeface="Calibri"/>', '<a:latin typeface="Arial"/>')
    return t.encode('utf-8')

def agosto_trace(rows, fluxo):
    import openpyxl
    R = {r['N.Fogo']: r for r in rows}
    rec = fontes._wb('recapagem', False).active
    h22 = {f['N.Fogo']: f['Status'] for f in fluxo if f['Lote'] == 22}
    t = dict(total=0, ressolado=0, baixado=0, ativo_desc=0, nao_reg=[], a_enviar=0, a_desc=0)
    for i in range(16, rec.max_row + 1):
        f = rec.cell(i, 2).value
        if f is None: continue
        f = int(f); st = rec.cell(i, 6).value; r = R[f]
        t['total'] += 1
        t['a_enviar'] += st == 'Há Enviar'; t['a_desc'] += st == 'Descartar'
        if f not in h22: t['nao_reg'].append(f)
        elif h22[f] == 'Ressolado': t['ressolado'] += 1
        elif r['Ativo'] == 'NÃO': t['baixado'] += 1
        else: t['ativo_desc'] += 1
    return t

def reconciliation(rows, fluxo):
    dem_c = sum(r['Demanda Compra'] == 'SIM' for r in rows)
    dem_c_uso = sum(r['Demanda Compra'] == 'SIM' and r['Situação Pneu'] == 'USO' for r in rows)
    cob = sum(r['Cobertura Compra'] == 'SIM' for r in rows)
    dem_r = sum(r['Demanda Ressolagem'] == 'SIM' for r in rows)
    l23 = [f for f in fluxo if f['Lote'] == 23]
    R = {r['N.Fogo']: r for r in rows}
    l23_est = sum(R[f['N.Fogo']]['Situação Pneu'] == 'ESTOQUE' for f in l23)
    l23_uso = sum(R[f['N.Fogo']]['Situação Pneu'] == 'USO' for f in l23)
    pend = sum(r['Situação Fluxo'] == 'Encerrado – baixar no Rodopar' for r in rows)
    ag = agosto_trace(rows, fluxo)
    db = M.data_base(rows)
    return [
        ('Demanda de compra (pneus listados)', '72', str(dem_c), f'+{dem_c-72}',
         '7 pneus em ESTOQUE com 1ª vida < 4 mm entram pela regra solicitada (a lista manual só tinha pneus em USO): 66895, 66970, '
         '67014, 75317, 76286, 76287 e 81717. 2 pneus com exatamente 4,00 mm (76274, 76280) estavam na lista manual e saem, pois a '
         'regra é estritamente < 4.',
         f'Regra aplicada como especificada. Com o parâmetro "Incluir pneus em ESTOQUE" = NÃO a demanda fica em {dem_c_uso} '
         '(os 72 pneus em uso da lista manual menos os 2 com 4,00 mm). 81717 ("Novo" com 1,00 mm) é provável erro de aferição — '
         'ver Auditoria.'),
        ('Demanda por medida (quadro "Previsionamento")', '225/75: 16 · 205/75: 23 · 225/65: 2 · 175/70: 16 · 275/80: 16',
         '225/75: 19 · 205/75: 24 · 225/65: 2 · 175/70: 18 · 275/80: 14', '—',
         'Erro de fórmula no arquivo anterior: as contagens de Fiorino (F16) e Caminhão (F19) apontavam para o cabeçalho F$12 '
         '(225/75 R16) e repetiam 16. Pela própria lista, eram 17 (Fiorino) e 14 (Caminhão).',
         'Corrigido: cada medida é contada pelo próprio rótulo (tabela de resumo).'),
        ('Necessidade de compra 205/75 R16', 'célula vazia (G25)', '14', '—',
         'A célula G25 do arquivo anterior não tinha fórmula; o total de Vans (5) omitia 13 pneus de 205/75 R16 (23 − 10).',
         'Corrigido: todas as medidas cadastradas têm necessidade calculada.'),
        ('Estoque de 175/70 R14 (cobertura)', '1', '10', '+9',
         'No arquivo anterior 9 pneus estavam como "175/70R14" (sem espaço) e não eram contados pelo CONT.SES de "175/70 R14".',
         'Corrigido: o Power Query padroniza a medida antes das regras (registrado em Tratamento Aplicado).'),
        ('Estoque disponível (cobertura) — total', '35', str(cob), str(cob - 35),
         'Mesmos 35 pneus: estoque, 1ª vida, Menor MM ≥ 4,00 mm (31 novos + 4 usados).', 'Regra de cobertura confirmada e mantida.'),
        ('Necessidade líquida — total', '34 (5 Vans + 15 Fiorino + 14 Caminhão)', '42', '+8',
         'Efeito combinado das correções acima e dos 7 pneus em estoque incluídos na demanda.',
         'Valor da nova planilha é o consistente com a regra; ver quadro por medida na aba Compra.'),
        ('Demanda de ressolagem', '37 (lista "Recapar" da aba Ressolagem)', str(dem_r), f'+{dem_r-37}',
         'Os 37 estão contidos nos 51. Em 10 pneus o menor sulco é o Sulco 4, que a fórmula anterior ignorava (usava o menor entre '
         'Sulco 1–3); 4 Fiorinos (67526, 67535, 75307, 77212) estão entre 2,10 e 2,75 mm — a fórmula anterior usava 2,10 mm para '
         'Fiorino. A regra anterior era "≤ 2,75"; o único pneu ativo com exatamente 2,75 mm (74635) não entra em nenhuma das duas.',
         'Regra aplicada como especificada (Menor Milimetragem < 2,75 para todas as medidas). Para manter 2,10 mm em Fiorino, informe '
         'o "Limite Ressolagem Específico" da medida 175/70 R14 em Parâmetros. Vários casos novos têm um sulco muito diferente dos '
         'demais (ex.: 67097: 9/9/9/2,44) — conferir em Auditoria › Variação entre sulcos.'),
        ('Lote 23 (Outubro_2026) — pneus "Descartar"', f'{len(l23)} decididos para descarte', f'{l23_est + l23_uso} ativos no Rodopar', '—',
         f'Os {len(l23)} pneus continuam em ESTOQUE ({l23_est}) ou USO ({l23_uso}) no Rodopar 10 de {db:%d/%m/%Y}.',
         f'Aparecem como "Encerrado – baixar no Rodopar" na Ressolagem ({pend} pneus no total com descarte/reprovação pendente de baixa).'),
        ('Demanda Recapagem Agosto/2026 — FL 87', f"{ag['total']} pneus ({ag['a_enviar']} a enviar, {ag['a_desc']} a descartar)", 'rastreados', '—',
         f"Hoje: {ag['ressolado']} ressolados (de volta em uso/estoque), {ag['baixado']} baixados (descarte/reprovação), "
         f"{ag['ativo_desc']} decididos para descarte ainda ativos no Rodopar e {len(ag['nao_reg'])} que não foram registrados no "
         f"controle de lotes ({', '.join(str(x) for x in ag['nao_reg'])}) e seguem abaixo do limite.",
         f"Os {len(ag['nao_reg'])} não registrados aparecem como \"Não programado\" na lista de ressolagem."),
    ]

def test_cases(rows):
    R = {r['N.Fogo']: r for r in rows}
    spec = [
        ('1. Menor MM 2,50 · Situação Aplicado (USO) → Ressolagem = SIM', 76219, 'Demanda Ressolagem', 'SIM'),
        ('2. Menor MM 2,48 · Situação BAIXADO → Ressolagem = NÃO', 66965, 'Demanda Ressolagem', 'NÃO'),
        ('3. Menor MM 3,54 · Vida 1 · USO → Compra = SIM', 78004, 'Demanda Compra', 'SIM'),
        ('4. Menor MM 3,39 · Vida 2 · USO → Compra = NÃO', 67088, 'Demanda Compra', 'NÃO'),
        ('5. Menor MM 3,62 · Vida 1 · DESCARTE → Compra = NÃO', 77205, 'Demanda Compra', 'NÃO'),
        ('Fronteira: Menor MM exatamente 2,75 · USO → Ressolagem = NÃO (estritamente menor)', 74635, 'Demanda Ressolagem', 'NÃO'),
        ('Fronteira: Menor MM exatamente 4,00 · Vida 1 · USO → Compra = NÃO (estritamente menor)', 76274, 'Demanda Compra', 'NÃO'),
        ('Estoque: Menor MM 3,59 · Vida 1 · ESTOQUE → Compra = SIM (estoque incluído)', 75317, 'Demanda Compra', 'SIM'),
        ('Limite de vida: Van vida 3 · 2,51 mm · USO → Destino = DESCARTAR', 66911, 'Destino Sugerido', 'DESCARTAR (LIMITE DE VIDA)'),
        ('Limite de vida: Fiorino vida 1 · 1,76 mm · USO → Destino = DESCARTAR', 75306, 'Destino Sugerido', 'DESCARTAR (LIMITE DE VIDA)'),
    ]
    return [(c, f, campo, esp, R[f]) for c, f, campo, esp in spec]

def main(out_path):
    raw = M.read_rodopar_raw()
    base = M.emulate_pq(raw)
    frotas = M.read_frotas()
    fluxo = M.read_historico()
    rows = M.compute(base, frotas, fluxo)
    frotas_c = M.frotas_calc(frotas, rows)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = {name: wb.create_sheet(name) for name in SHEETS}
    wb.loaded_theme = brand_theme()
    # referência e dados
    BR.build_parametros(wb, ws['Parâmetros'])
    # lista auxiliar de filtro de medidas (Parâmetros, coluna J oculta)
    wp = ws['Parâmetros']
    put(wp, 'J7', 'Filtro de medidas (auxiliar)', f=font(8, False, C['muted']))
    put(wp, 'J8', '(Todas)', f=font(8, False, C['muted']))
    for k in range(1, 21):
        exp = BV.MED[k - 1] if k <= len(BV.MED) else ''
        put(wp, f'J{8+k}', f'=IFERROR(INDEX(tbMedidas[Medida],{k})&"","")', f=font(8, False, C['muted']), expected=exp)
    wp.column_dimensions['J'].hidden = True
    wb.defined_names['lstFiltroMedida'] = DefinedName('lstFiltroMedida', attr_text="OFFSET('Parâmetros'!$J$8,0,0,1+ROWS(tbMedidas[Medida]),1)")
    BR.build_frotas(wb, ws['Frotas'], frotas_c)
    BR.build_fluxo(wb, ws['Fluxo Ressolagem'], fluxo, rows)
    BR.build_base(wb, ws['Base Tratada'], rows)
    ctx = {}
    ctx['compra'] = BV.build_compra(wb, ws['Compra'], rows)
    ctx['res'] = BV.build_ressolagem(wb, ws['Ressolagem'], rows)
    ctx['est'] = BV.build_estoque(wb, ws['Estoque'], rows)
    ctx['an'] = BV2.build_analises(wb, ws['Análises'], rows, frotas_c)
    ctx['aud'] = BV2.build_auditoria(wb, ws['Auditoria'], rows)
    ctx['frotas'] = frotas_c
    BF.build_dashboard(wb, ws['Dashboard'], rows, ctx)
    BF.build_inicio(wb, ws['Início'], rows, ctx)
    BF.build_doc(wb, ws['Documentação'], rows, ctx, reconciliation(rows, fluxo), test_cases(rows))
    for name in SHEETS:
        w = ws[name]
        w.page_setup.orientation = 'landscape'
        w.page_setup.paperSize = w.PAPERSIZE_A4
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = (name != 'Base Tratada')
        w.print_options.horizontalCentered = True
        w.page_margins.left = w.page_margins.right = 0.4
        w.page_margins.top = w.page_margins.bottom = 0.5
        w.oddFooter.center.text = '&8Gestão de Pneus — Horizonte · &A · página &P de &N'
    wb.active = 0
    for name in SHEETS:
        ws[name].sheet_view.tabSelected = (name == 'Início')
    wb.properties.title = 'Gestão de Pneus - Horizonte'
    wb.properties.subject = 'Planejamento de ressolagem, compra e estoque de pneus — Operações Souza Cruz'
    wb.properties.creator = 'Horizonte Logística — Gestão de Frota'
    wb.properties.keywords = 'pneus; ressolagem; compra; estoque; Rodopar 10'
    wb.save(out_path)
    meta = dict(n_formulas=len(xlh.EXPECTED), n_dynamic=len(xlh.DYNAMIC))
    return wb, rows, ctx, meta

if __name__ == '__main__':
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(OUT_DIR, 'stage1.xlsx')
    wb, rows, ctx, meta = main(out)
    import pickle
    with open(os.path.join(OUT_DIR, 'stage1_meta.pkl'), 'wb') as fh:
        pickle.dump(dict(EXPECTED=xlh.EXPECTED, DYNAMIC=xlh.DYNAMIC, TABLE_CALC=xlh.TABLE_CALC), fh)
    print('saved', out, meta)
