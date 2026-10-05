"""Utilitários de estilo e registro de fórmulas (fórmula + valor esperado)."""
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, NamedStyle
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn, TableFormula
from openpyxl.worksheet.hyperlink import Hyperlink
from config import C, FONT, CARGA_INICIAL

EXPECTED = {}      # (sheet, coord) -> valor esperado (fórmulas clássicas)
DYNAMIC = []       # fórmulas de matriz dinâmica a injetar no pós-processamento
TABLE_CALC = {}    # tabela -> {coluna: fórmula}

def font(size=10, bold=False, color=None, italic=False, underline=None):
    return Font(name=FONT, size=size, bold=bold, color=color or C['text'], italic=italic, underline=underline)

def fill(hexcolor):
    return PatternFill('solid', start_color=hexcolor, end_color=hexcolor)

def side(color=None, style='thin'):
    return Side(style=style, color=color or C['border'])

def border(l=None, r=None, t=None, b=None):
    return Border(left=l or Side(), right=r or Side(), top=t or Side(), bottom=b or Side())

BOX = Border(left=side(), right=side(), top=side(), bottom=side())
HAIR_B = Border(bottom=side())

def al(h='left', v='center', wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)

def put(ws, coord, value, f=None, fl=None, a=None, b=None, nf=None, expected=None):
    c = ws[coord]
    c.value = value
    if f is not None: c.font = f
    if fl is not None: c.fill = fl
    if a is not None: c.alignment = a
    if b is not None: c.border = b
    if nf is not None: c.number_format = nf
    if isinstance(value, str) and value.startswith('='):
        EXPECTED[(ws.title, coord)] = expected
    return c

def dyn(ws, coord, formula, values, nf=None, f=None, a=None):
    """Registra uma fórmula de matriz dinâmica (FILTRO/CLASSIFICAR/ÚNICO) e seu resultado esperado (lista de linhas)."""
    DYNAMIC.append(dict(sheet=ws.title, coord=coord, formula=formula, values=values))
    c = ws[coord]
    if f is not None: c.font = f
    if a is not None: c.alignment = a
    if nf is not None: c.number_format = nf
    return c

def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w

def col(n):
    return get_column_letter(n)

def cidx(letter):
    return column_index_from_string(letter)

def link(ws, coord, text, target, f=None, a=None):
    c = ws[coord]
    c.value = text
    c.hyperlink = Hyperlink(ref=coord, location=target, display=text)
    c.font = f or font(9, color=C['blue'], underline='single')
    if a is not None: c.alignment = a
    return c

def sheet_base(ws, tab_color, zoom=90, grid=False):
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = grid
    ws.sheet_view.zoomScale = zoom
    ws.column_dimensions['A'].width = 2.2

def header(ws, title, subtitle, last_col, data_ref=True, home=True):
    """Cabeçalho padrão: título, subtítulo, data dos dados e filete dourado."""
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 6
    put(ws, 'B2', title, f=font(17, True, C['blue']), a=al('left', 'center'))
    put(ws, 'B3', subtitle, f=font(9, False, C['muted']), a=al('left', 'center'))
    lc = col(last_col)
    if data_ref:
        fc = col(max(2, last_col - 4))
        put(ws, f'{fc}2', '=MAX(tbPneus[Atualizado em])', f=font(9, False, C['muted']), a=al('right', 'center'),
            nf='"Dados Rodopar 10 atualizados em "dd/mm/yyyy hh:mm', expected=CARGA_INICIAL)
        ws.merge_cells(f'{fc}2:{lc}2')
    if home:
        link(ws, f'{lc}3', '◂ Início', "'Início'!A1", f=font(9, color=C['blue'], underline='single'), a=al('right'))
    for ci in range(2, last_col + 1):
        ws.cell(4, ci).border = Border(bottom=Side(style='medium', color=C['gold']))

def section(ws, coord, text, last_col=None, row_h=20):
    r = ws[coord].row
    ws.row_dimensions[r].height = row_h
    put(ws, coord, text.upper(), f=font(9, True, C['blue']), a=al('left', 'bottom'))
    if last_col:
        c0 = ws[coord].column
        for ci in range(c0, last_col + 1):
            ws.cell(r, ci).border = Border(bottom=Side(style='thin', color=C['border_strong']))

def kpi(ws, row, col0, span, label, formula, expected, nf='#,##0', tone='default', note=None, note_formula=None, note_expected=None):
    """Cartão de indicador: rótulo / valor / nota, centralizado na seleção (sem mesclar)."""
    tones = {
        'default': (C['surface2'], C['text']),
        'blue': (C['primary_soft'], C['blue_deep']),
        'danger': (C['danger_soft'], C['danger_fg']),
        'warning': (C['warning_soft'], C['warning_fg']),
        'success': (C['success_soft'], C['success_fg']),
        'neutral': (C['neutral_soft'], C['neutral_fg']),
        'gold': (C['gold_soft'], C['gold_soft_fg']),
    }
    bg, fg = tones[tone]
    cols = list(range(col0, col0 + span))
    for r in (row, row + 1, row + 2):
        for ci in cols:
            c = ws.cell(r, ci)
            c.fill = fill(bg)
            left = Side(style='thin', color='FFFFFF') if ci == cols[0] else Side()
            right = Side(style='thin', color='FFFFFF') if ci == cols[-1] else Side()
            c.border = Border(left=left, right=right,
                              top=Side(style='medium', color=fg) if r == row else Side(),
                              bottom=Side(style='thin', color='FFFFFF') if r == row + 2 else Side())
    first = col(cols[0])
    put(ws, f'{first}{row}', label, f=font(8, True, C['muted']), a=Alignment(horizontal='centerContinuous', vertical='bottom', wrap_text=False))
    for ci in cols[1:]:
        ws.cell(row, ci).alignment = Alignment(horizontal='centerContinuous', vertical='bottom')
    put(ws, f'{first}{row+1}', formula, f=font(19, True, fg), a=Alignment(horizontal='centerContinuous', vertical='center'), nf=nf, expected=expected)
    for ci in cols[1:]:
        ws.cell(row + 1, ci).alignment = Alignment(horizontal='centerContinuous', vertical='center')
    if note_formula is not None:
        put(ws, f'{first}{row+2}', note_formula, f=font(8, False, C['muted']), a=Alignment(horizontal='centerContinuous', vertical='top'), expected=note_expected)
    elif note:
        put(ws, f'{first}{row+2}', note, f=font(8, False, C['muted']), a=Alignment(horizontal='centerContinuous', vertical='top'))
    for ci in cols[1:]:
        ws.cell(row + 2, ci).alignment = Alignment(horizontal='centerContinuous', vertical='top')
    ws.row_dimensions[row].height = 18
    ws.row_dimensions[row + 1].height = 30
    ws.row_dimensions[row + 2].height = 16

def table_header_style(ws, row, c1, c2, h=30, bg=None, fg='FFFFFF', dividers=True):
    ws.row_dimensions[row].height = h
    for ci in range(c1, c2 + 1):
        c = ws.cell(row, ci)
        c.font = font(9, True, fg)
        c.fill = fill(bg or C['blue'])
        c.alignment = al('center', 'center', wrap=True)
        if dividers:
            c.border = Border(left=Side(style='thin', color='FFFFFF'), right=Side(style='thin', color='FFFFFF'))

def body_style(ws, r1, r2, c1, c2, band=True, size=9):
    for r in range(r1, r2 + 1):
        for ci in range(c1, c2 + 1):
            c = ws.cell(r, ci)
            c.font = font(size)
            if band and (r - r1) % 2 == 1:
                c.fill = fill(C['surface2'])
            c.border = Border(bottom=Side(style='thin', color=C['border']))
            if c.alignment is None or c.alignment.horizontal is None:
                c.alignment = al('center', 'center')

def add_table(ws, name, ref, columns, style='TableStyleMedium2', calc=None, stripes=True):
    """Cria tabela estruturada com colunas calculadas (fórmula em nível de coluna)."""
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showFirstColumn=False, showLastColumn=False,
                                      showRowStripes=stripes, showColumnStripes=False)
    cols = []
    for i, cname in enumerate(columns, start=1):
        tc = TableColumn(id=i, name=cname)
        if calc and cname in calc:
            tc.calculatedColumnFormula = TableFormula(attr_text=calc[cname])
        cols.append(tc)
    t.tableColumns = cols
    t._initialise_columns = lambda: None
    ws.add_table(t)
    if calc:
        TABLE_CALC[name] = dict(calc)
    return t
