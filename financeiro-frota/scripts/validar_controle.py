#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
validar_controle.py - Conferência independente do Controle_Financeiro_Gerencial_Frota.xlsx

Recalcula, com pandas, os totais da Frota diretamente a partir das abas de origem
(Relação de NF's, Fornecedores, Rodopar 108 e 1015) e compara com os valores que a
planilha apresenta na aba "Conciliação e Qualidade" e em "Dados Gráficos".

Uso:
    python validar_controle.py "Controle_Financeiro_Gerencial_Frota.xlsx"

Pré-requisito: o arquivo deve ter sido aberto e salvo no Excel (para que os valores
calculados estejam gravados). Requer: pandas, openpyxl.
"""
import sys
import pandas as pd
import openpyxl

S_NF = "Relação de NF's"
S_FOR = "Fornecedores"
S_108 = "Doc's Pagos | Rodopar 108"
S_1015 = "Doc Há Pagar | Rodopar 1015"
S_CQ = "Conciliação e Qualidade"
S_DG = "Dados Gráficos"


def main(path):
    nf = pd.read_excel(path, sheet_name=S_NF, usecols="B,E,F,I,J")
    fo = pd.read_excel(path, sheet_name=S_FOR, usecols="A,C")
    pg = pd.read_excel(path, sheet_name=S_108, usecols="C:M")
    ap = pd.read_excel(path, sheet_name=S_1015, usecols="C:K")
    nf = nf.dropna(subset=["Documento"])
    nf["doc"] = nf["Documento"].astype(str).str.strip()
    nf["forn"] = nf["Razão Social"].astype(str).str.strip()
    nf["tipo"] = nf["Tipo"].astype(str).str.strip().str.upper()
    nf["em"] = pd.to_datetime(nf["Emissão"], errors="coerce").dt.normalize()
    nf["k1"] = nf["doc"] + "|" + nf["forn"] + "|" + nf["tipo"] + "|" + nf["em"].dt.strftime("%Y-%m-%d")
    nf["k2"] = nf["doc"] + "|" + nf["forn"] + "|" + nf["tipo"]
    frota = set(fo.loc[fo["Area"] == "Frota", "Fornecedor"].astype(str).str.strip())
    nf["frota"] = nf["forn"].isin(frota)
    nf["dup"] = nf.duplicated("k1")
    nf["ok"] = nf["frota"] & ~nf["dup"] & nf["Valor Doc"].notna() & nf["em"].notna()
    nf["comp"] = nf["em"].dt.strftime("%Y-%m")

    pg = pg[pg["Fornecedor"].notna()].copy()
    pg["J"] = pd.to_numeric(pg["Vlr. Parc."], errors="coerce")
    pg["H"] = pd.to_datetime(pg["Vencto"], errors="coerce").dt.normalize()
    pg = pg[pg["J"].notna() & pg["H"].notna()].copy()
    pg["doc"] = pg["Documento"].astype(str).str.strip()
    pg["forn"] = pg["Fornecedor"].astype(str).str.strip()
    pg["tipo"] = pg["Tipo"].astype(str).str.strip().str.upper()
    pg["em"] = pd.to_datetime(pg["Emissão"], errors="coerce").dt.normalize()
    pg["k1"] = pg["doc"] + "|" + pg["forn"] + "|" + pg["tipo"] + "|" + pg["em"].dt.strftime("%Y-%m-%d").fillna("s/data")
    pg["k2"] = pg["doc"] + "|" + pg["forn"] + "|" + pg["tipo"]
    pg["dup"] = pg.duplicated(["k2", "H", "J"])
    nfk1, nfk2 = set(nf["k1"]), set(nf["k2"])
    pg["linked"] = pg["k1"].isin(nfk1) | pg["k2"].isin(nfk2)
    pg["ok"] = pg["linked"] & ~pg["dup"]
    pg["tot"] = pd.to_numeric(pg["Vlr. Total"], errors="coerce").fillna(pg["J"])

    ap["J"] = pd.to_numeric(ap["Vlr. Parcela"], errors="coerce")
    ap["G"] = pd.to_datetime(ap["Vencimento"], errors="coerce").dt.normalize()
    ap = ap[ap["Doc."].notna() & ap["J"].notna() & ap["G"].notna()].copy()
    ap["k2"] = ap["Doc."].astype(str).str.strip() + "|" + ap["Fornecedor"].astype(str).str.strip() + "|" + ap["Tipo"].astype(str).str.strip().str.upper()
    ap["dup"] = ap.duplicated(["k2", "G", "J", "Parcela"])
    ap["ok"] = ap["k2"].isin(nfk2) & ~ap["dup"]

    wb = openpyxl.load_workbook(path, data_only=True)
    cq = wb[S_CQ]
    # localiza as linhas da conciliação de totais pelo rótulo
    labels = {cq.cell(r, 2).value: r for r in range(1, 120) if cq.cell(r, 2).value}
    def cell(label, col):
        r = labels.get(label)
        return cq.cell(r, col).value if r else None

    rows_cons = [r for r in range(1, 120) if cq.cell(r, 2).value == "   das quais: Frota com NF vinculada (consideradas)"]
    r108 = rows_cons[0] if rows_cons else None
    r1015 = rows_cons[1] if len(rows_cons) > 1 else None
    comps = [
        ("108 - parcelas consideradas (qtd)", int(pg["ok"].sum()), cq.cell(r108, 3).value if r108 else None),
        ("108 - parcelas consideradas (valor bruto)", round(pg.loc[pg["ok"], "J"].sum(), 2), cq.cell(r108, 4).value if r108 else None),
        ("1015 - parcelas consideradas (qtd)", int(ap["ok"].sum()), cq.cell(r1015, 3).value if r1015 else None),
        ("1015 - parcelas consideradas (valor)", round(ap.loc[ap["ok"], "J"].sum(), 2), cq.cell(r1015, 4).value if r1015 else None),
        ("NF - documentos considerados (valor)", round(nf.loc[nf["ok"], "Valor Doc"].sum(), 2), cell("   dos quais: considerados na DRE", 4)),
        ("NF - documentos considerados (qtd)", int(nf["ok"].sum()), cell("   dos quais: considerados na DRE", 3)),
    ]
    print(f"{'Verificação':55s} {'pandas':>18s} {'planilha':>18s}  status")
    ok_all = True
    for name, a, b in comps:
        st = "OK" if b is not None and abs(float(a) - float(b)) < 0.01 else "DIVERGENTE"
        ok_all &= st == "OK"
        print(f"{name:55s} {a:>18,.2f} {(b if b is not None else float('nan')):>18,.2f}  {st}")
    # custo por competência mensal x Dados Gráficos (ano do Painel)
    dg = wb[S_DG]
    print("\nCusto por competência (mês) - pandas x Dados Gráficos (ano do Painel, sem filtros):")
    for m in range(1, 13):
        comp = dg.cell(5 + m, 3).value
        exp = round(nf.loc[nf["ok"] & (nf["comp"] == comp), "Valor Doc"].sum(), 2)
        got = dg.cell(5 + m, 5).value or 0
        st = "OK" if abs(exp - float(got)) < 0.01 else "DIVERGENTE (verifique filtros/ajustes manuais de competência)"
        print(f"  {comp}: {exp:>14,.2f}  x  {float(got):>14,.2f}  {st}")
    print("\nResultado geral:", "OK" if ok_all else "HÁ DIVERGÊNCIAS - revise a aba Conciliação e Qualidade")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1])
