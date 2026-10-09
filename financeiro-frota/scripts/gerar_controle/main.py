# -*- coding: utf-8 -*-
import sys, time, openpyxl
from openpyxl.workbook.properties import CalcProperties
from common import *
import sources, params, reports, painel

SRC = sys.argv[1]
OUT = sys.argv[2]
t0 = time.time()
wb = openpyxl.load_workbook(SRC)

def last_data_row(ws, col):
    last = 1
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, col).value not in (None, ""):
            last = r
            break
    return last

n_nf = last_data_row(wb[S_NF], 2)       # Documento
n_for = last_data_row(wb[S_FOR], 1)
n_108 = last_data_row(wb[S_108], 3)
n_1015 = last_data_row(wb[S_1015], 4)
print("rows", n_nf, n_for, n_108, n_1015)

# limpa fórmulas antigas das colunas de tratamento / meses nas abas de origem (serão reescritas)
ws = wb[S_NF]
for r in range(1, ws.max_row + 1):
    for c in range(16, ws.max_column + 1):
        ws.cell(r, c).value = None
ws1015 = wb[S_1015]
ws1015["B1"].value = "Mês Venc."

# novas abas
for name in [S_PAINEL, S_REM, S_DRE, S_FC, S_CC, S_CP, S_PO, S_AF, S_CQ, S_LEIA, S_PAR, S_DG]:
    wb.create_sheet(name)

params.build_parametros(wb, wb[S_PAR])
params.build_remuneracao(wb, wb[S_REM])
dg = params.build_dados_graficos(wb, wb[S_DG])
sources.build_nf(wb[S_NF], n_nf)
sources.build_fornecedores(wb[S_FOR], n_for)
sources.build_108(wb[S_108], n_108)
sources.build_1015(wb[S_1015], n_1015)
reports.build_dre(wb, wb[S_DRE])
reports.build_fluxo(wb, wb[S_FC])
reports.build_comp_caixa(wb, wb[S_CC])
reports.build_custos_pacote(wb, wb[S_CP], dg)
reports.build_pagamentos(wb, wb[S_PO])
reports.build_fornecedores_analise(wb, wb[S_AF])
reports.build_conciliacao(wb, wb[S_CQ])
painel.build_painel(wb, wb[S_PAINEL])
painel.build_leiame(wb[S_LEIA])

order = [S_PAINEL, S_REM, S_DRE, S_FC, S_CC, S_CP, S_PO, S_AF, S_CQ, S_LEIA, S_PAR, S_DG, S_NF, S_FOR, S_108, S_1015]
wb._sheets = [wb[n] for n in order]
wb.active = 0
for ws in wb.worksheets:
    ws.sheet_view.tabSelected = (ws.title == S_PAINEL)
# cores das guias
for n in [S_PAINEL]:
    wb[n].sheet_properties.tabColor = NAVY
for n in [S_REM]:
    wb[n].sheet_properties.tabColor = GOLD
for n in [S_DRE, S_FC, S_CC, S_CP, S_PO, S_AF]:
    wb[n].sheet_properties.tabColor = BLUE
for n in [S_CQ, S_LEIA, S_PAR, S_DG]:
    wb[n].sheet_properties.tabColor = GRAY4
for n in [S_NF, S_FOR, S_108, S_1015]:
    wb[n].sheet_properties.tabColor = GRAPHITE
wb.calculation = CalcProperties(fullCalcOnLoad=True, calcMode="auto")
wb.properties.title = "Controle Financeiro Gerencial - Frota"
wb.properties.creator = "Controladoria Frota"
wb.save(OUT)
print("saved", OUT, round(time.time() - t0, 1), "s")
