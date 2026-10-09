# -*- coding: utf-8 -*-
"""Painel Executivo e Leia-me."""
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, DoughnutChart, Reference, Series
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.data_source import StrRef
from common import *
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.styles import Font
from sources import NF, P, A, F, MESES, MESES_ABREV, _whole, SCOPE_NF, SCOPE_SEMNF, SCOPE_OUTRA, SCOPE_NC, SEM_CLASS, NAO_CAD
from params import N_PAC, N_CAT, add_name
from reports import (style_chart, color_series, nfI, nfAQ, nfComp, nfPac, nfCat, nfFor, nfAno, nfMes, nfSt, nfBase, p_val, p_liq, p_cons, p_cx, p_cv, p_pac, p_cat, p_for,
                     p_jur, p_des, p_anocx, p_esc, a_val, a_cons, a_cv, a_st, a_pac, a_cat, a_for, a_esc, a_fx, FILT_NF, FILT_P, FILT_A, remC, remV, remA)


def build_painel(wb, ws):
    widths = {"A": 2}
    for c in range(2, 18):
        widths[get_column_letter(c)] = 13.5
    widths["R"] = 2
    set_widths(ws, widths)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    title_block(ws, "PAINEL EXECUTIVO FINANCEIRO | SETOR DE FROTA", "Horizonte Logística - Controle Financeiro Gerencial: quanto a Frota custa por competência, quanto paga por mês e quanto ainda compromete", "B", "Q")
    # ---- filtros
    labels = [("B", "Ano"), ("D", "Mês"), ("F", "Visão"), ("H", "Pacote"), ("K", "Categoria"), ("N", "Fornecedor")]
    for col, t in labels:
        write(ws, f"{col}4", t, font(True, NAVY, 9), align=ALIGN_L)
    input_cell(ws, "B5", 2026, "0", "Ano de referência do painel (competência, caixa e vencimentos).")
    ws.merge_cells("B5:C5")
    input_cell(ws, "D5", "Setembro", None, "Mês de referência dos indicadores mensais.")
    ws.merge_cells("D5:E5")
    input_cell(ws, "F5", "Mês", None, "Mês = indicadores do mês selecionado; Acumulado no ano = janeiro até o mês selecionado.")
    ws.merge_cells("F5:G5")
    input_cell(ws, "H5", "(Todos)", None, "Filtro de pacote: afeta custos, pagamentos e obrigações. Não afeta a remuneração.")
    ws.merge_cells("H5:J5")
    input_cell(ws, "K5", "(Todos)")
    ws.merge_cells("K5:M5")
    input_cell(ws, "N5", "(Todos)")
    ws.merge_cells("N5:Q5")
    add_list_validation(ws, "B5", "=ListaAnos")
    add_list_validation(ws, "D5", "=Meses")
    add_list_validation(ws, "F5", "=ListaVisao")
    add_list_validation(ws, "H5", "=ListaFiltroPacote")
    add_list_validation(ws, "K5", "=ListaFiltroCategoria")
    add_list_validation(ws, "N5", "=ListaFiltroFornecedor")
    for nm, ref in [("AnoSel", "$B$5"), ("MesNome", "$D$5"), ("Visao", "$F$5"), ("FiltroPacote", "$H$5"), ("FiltroCategoria", "$K$5"), ("FiltroFornecedor", "$N$5")]:
        add_name(wb, nm, f"{R_PAINEL}!{ref}")
    write(ws, "B6", '="Período: "&IF(VisaoAcum=1,"Janeiro a ","")&INDEX(Meses,MesSel)&"/"&AnoSel&"   |   Data de referência (vencido x a vencer): "&TEXT(DataRef,"dd/mm/yyyy")&"   |   Caixa realizado por: "&BaseCaixaTxt&IF(FiltroAtivo=1,"   |   FILTRO DE CUSTO ATIVO - remuneração e resultado não são filtrados","")',
          font(True, AMBER, 8), align=ALIGN_L)
    ws.merge_cells("B6:Q6")
    ws.row_dimensions[6].height = 16

    # helpers
    YTD = "VisaoAcum=1"
    def nf_sum(crit_month):
        return f"SUMIFS({nfI},{nfAQ},1,{crit_month},{FILT_NF})"
    cost_mes = nf_sum(f"{nfComp},CompSel")
    cost_ant = nf_sum(f"{nfComp},CompAnt")
    cost_ytd = nf_sum(f'{nfAno},AnoSel,{nfMes},"<="&MesSel')
    cost_sel = f"IF({YTD},{cost_ytd},{cost_mes})"
    rem_mes = f"(SUMIFS({remV},{remC},CompSel)+SUMIFS({remA},{remC},CompSel))"
    rem_ytd = f'(SUMIFS({remV},{R_REM}!$B:$B,AnoSel,{R_REM}!$C:$C,"<="&MesSel)+SUMIFS({remA},{R_REM}!$B:$B,AnoSel,{R_REM}!$C:$C,"<="&MesSel))'
    rem_sel = f"IF({YTD},{rem_ytd},{rem_mes})"
    jur_mes = f"SUMIFS({p_jur},{p_cons},1,{p_cx},CompSel,{FILT_P})"
    des_mes = f"SUMIFS({p_des},{p_cons},1,{p_cx},CompSel,{FILT_P})"
    jur_ytd = f'SUMIF({R_DG}!$N$6:$N$17,"<="&MesSel,{R_DG}!$J$6:$J$17)'
    des_ytd = f'SUMIF({R_DG}!$N$6:$N$17,"<="&MesSel,{R_DG}!$K$6:$K$17)'
    res_mes = f"({rem_mes}-{cost_mes}-{jur_mes}+{des_mes})"
    res_ytd = f"({rem_ytd}-{cost_ytd}-{jur_ytd}+{des_ytd})"
    res_sel = f"IF({YTD},{res_ytd},{res_mes})"
    pago_mes = f"SUMIFS({p_liq},{p_cons},1,{p_cx},CompSel,{FILT_P})"
    pago_ytd = f'SUMIF({R_DG}!$N$6:$N$17,"<="&MesSel,{R_DG}!$G$6:$G$17)'
    pago_sel = f"IF({YTD},{pago_ytd},{pago_mes})"
    prev_mes = f"SUMIFS({a_val},{a_cons},1,{a_cv},CompSel,{FILT_A})"
    prev_ant = f"SUMIFS({a_val},{a_cons},1,{a_cv},CompAnt,{FILT_A})"
    compr_mes = f"(SUMIFS({a_val},{a_cons},1,{a_cv},CompSel,{FILT_A})+SUMIFS({p_val},{p_cons},1,{p_cv},CompSel,{FILT_P}))"
    compr_ant = f"(SUMIFS({a_val},{a_cons},1,{a_cv},CompAnt,{FILT_A})+SUMIFS({p_val},{p_cons},1,{p_cv},CompAnt,{FILT_P}))"
    prev_ytd = f'SUMIF({R_DG}!$N$6:$N$17,"<="&MesSel,{R_DG}!$H$6:$H$17)'
    prev_sel = f"IF({YTD},{prev_ytd},{prev_mes})"
    aberto = f"SUMIFS({a_val},{a_cons},1,{FILT_A})"
    vencido = f'SUMIFS({a_val},{a_cons},1,{a_st},"Vencido",{FILT_A})'
    avencer = f'SUMIFS({a_val},{a_cons},1,{a_st},"A vencer",{FILT_A})'
    prox30 = f'SUMIFS({a_val},{a_cons},1,{a_st},"A vencer",{_whole(R_1015,A,"Dias para vencer")},"<=30",{FILT_A})'
    na_filtro = '"n/a (filtro ativo)"'

    # ---- bloco 1: perguntas-chave (adendo)
    section(ws, 8, "1. QUANTO CUSTA, QUANTO PAGA, QUANTO COMPROMETE - competência x caixa (período selecionado)", "B", "Q", NAVY)
    cards = [
        ("CUSTO TOTAL POR COMPETÊNCIA", f"={cost_sel}", f'="Acumulado no ano: R$ "&TEXT({cost_ytd},"#,##0")'),
        ("CUSTO EM FLUXO DE CAIXA - REALIZADO", f"={pago_sel}", f'=IF(FlagDataPgto=1,"Data efetiva de pagamento","Títulos pagos por mês de vencimento (data não validada)")'),
        ("CUSTO EM FLUXO DE CAIXA - PREVISTO", f"={prev_sel}", '="Obrigações em aberto com vencimento no período"'),
        ("TOTAL DE OBRIGAÇÕES EM ABERTO", f"={aberto}", f'="Vencido: R$ "&TEXT({vencido},"#,##0")&"  |  A vencer: R$ "&TEXT({avencer},"#,##0")'),
        ("RESULTADO GERENCIAL", f'=IF(FiltroAtivo=1,{na_filtro},{res_sel})', f'=IF(FiltroAtivo=1,"Remova os filtros de custo",IF({rem_sel}>0,"Margem: "&TEXT({res_sel}/{rem_sel},"0.0%"),"Remuneração não informada"))'),
    ]
    for i, (t, f, sub) in enumerate(cards):
        kpi_card(ws, 9, 2 + i * 3, t, f, FMT_BRL0, sub, None, width_cols=3)
    # ---- bloco 2: receitas e resultados
    section(ws, 13, "2. RECEITAS E RESULTADOS (regime de competência - sem filtros de custo)", "B", "Q", NAVY2)
    cards = [
        ("Remuneração do mês", f"={rem_mes}", f'=IF({rem_mes}=0,"Não informada em Remuneração","Competência "&CompSel)'),
        ("Remuneração acumulada no ano", f"={rem_ytd}", '="Janeiro a "&INDEX(Meses,MesSel)'),
        ("Resultado gerencial do mês", f'=IF(FiltroAtivo=1,{na_filtro},{res_mes})', f'=IF(FiltroAtivo=1,"",IF({rem_mes}>0,"Margem gerencial: "&TEXT({res_mes}/{rem_mes},"0.0%"),"Margem n/d (sem remuneração)"))'),
        ("Resultado acumulado no ano", f'=IF(FiltroAtivo=1,{na_filtro},{res_ytd})', f'=IF(FiltroAtivo=1,"",IF({rem_ytd}>0,"Margem acumulada: "&TEXT({res_ytd}/{rem_ytd},"0.0%"),"Margem n/d (sem remuneração)"))'),
        ("Custo sobre remuneração (período)", f'=IF(OR(FiltroAtivo=1,{rem_sel}<=0),"n/d",{cost_sel}/{rem_sel})', '="Custos operacionais / receita líquida"'),
    ]
    for i, (t, f, sub) in enumerate(cards):
        kpi_card(ws, 14, 2 + i * 3, t, f, FMT_PCT if i == 4 else FMT_BRL0, sub, None, width_cols=3, color=NAVY2)
    # ---- bloco 3: compromissos
    section(ws, 18, "3. COMPROMISSOS FINANCEIROS (parcelas de documentos da Frota - data de referência)", "B", "Q", NAVY2)
    cards = [
        ("Documentos registrados (ano)", f"=COUNTIFS({nfAQ},1,{nfAno},AnoSel,{FILT_NF})", f'="Valor: R$ "&TEXT(SUMIFS({nfI},{nfAQ},1,{nfAno},AnoSel,{FILT_NF}),"#,##0")'),
        ("Parcelas pagas identificadas (ano)", f"=COUNTIFS({p_cons},1,{p_anocx},AnoSel,{FILT_P})", f'="Líquido pago: R$ "&TEXT(SUMIFS({p_liq},{p_cons},1,{p_anocx},AnoSel,{FILT_P}),"#,##0")'),
        ("Valor vencido (em aberto)", f"={vencido}", f'=COUNTIFS({a_cons},1,{a_st},"Vencido",{FILT_A})&" parcelas vencidas"'),
        ("A vencer - próximos 30 dias", f"={prox30}", f'="Total a vencer: R$ "&TEXT({avencer},"#,##0")'),
        ("Variação dos compromissos (títulos por vencimento) vs mês anterior", f'=IF({compr_ant}>0,({compr_mes}-{compr_ant})/{compr_ant},"n/d")', f'="Mês: R$ "&TEXT({compr_mes},"#,##0")&" | Anterior: R$ "&TEXT({compr_ant},"#,##0")&" (pagos + em aberto)"'),
    ]
    for i, (t, f, sub) in enumerate(cards):
        kpi_card(ws, 19, 2 + i * 3, t, f, FMT_INT if i < 2 else (FMT_PCT if i == 4 else FMT_BRL0), sub, None, width_cols=3, color=NAVY2)
    # ---- bloco 4: eficiência
    section(ws, 23, "4. EFICIÊNCIA FINANCEIRA (custos por competência)", "B", "Q", NAVY2)
    top_pac = f'INDEX(DG_PacNomes,MATCH(MAX(DG_PacValores),DG_PacValores,0))'
    top_pac_pct = f'IF(SUM(DG_PacValores)>0,MAX(DG_PacValores)/SUM(DG_PacValores),0)'
    media_3 = f'IF(COUNTIFS({R_DG}!$N$6:$N$17,">"&(MesSel-3),{R_DG}!$N$6:$N$17,"<="&MesSel,{R_DG}!$E$6:$E$17,">0")>0,SUMIFS({R_DG}!$E$6:$E$17,{R_DG}!$N$6:$N$17,">"&(MesSel-3),{R_DG}!$N$6:$N$17,"<="&MesSel)/COUNTIFS({R_DG}!$N$6:$N$17,">"&(MesSel-3),{R_DG}!$N$6:$N$17,"<="&MesSel,{R_DG}!$E$6:$E$17,">0"),0)'
    media_6 = media_3.replace("MesSel-3", "MesSel-6")
    media_ano = f'IF(COUNTIFS({R_DG}!$N$6:$N$17,"<="&MesSel,{R_DG}!$E$6:$E$17,">0")>0,{cost_ytd}/COUNTIFS({R_DG}!$N$6:$N$17,"<="&MesSel,{R_DG}!$E$6:$E$17,">0"),0)'
    cards = [
        ("Custo médio mensal (ano, meses com custo)", f"={media_ano}", f'="Média 3M: R$ "&TEXT({media_3},"#,##0")&" | 6M: R$ "&TEXT({media_6},"#,##0")'),
        ("Variação do custo vs mês anterior", f'=IF({cost_ant}>0,({cost_mes}-{cost_ant})/{cost_ant},"n/d")', f'="Mês: R$ "&TEXT({cost_mes},"#,##0")&" | Anterior: R$ "&TEXT({cost_ant},"#,##0")'),
        ("Maior pacote (participação no ano)", f"={top_pac_pct}", f'=IF(SUM(DG_PacValores)>0,{top_pac},"-")'),
        ("Maior fornecedor (participação no ano)", f"={R_DG}!$D$86", f"={R_DG}!$B$86"),
        ("Custo do mês vs média 3 meses", f'=IF({media_3}>0,({cost_mes}-{media_3})/{media_3},"n/d")', '="Desvio do mês em relação à média dos 3 últimos meses"'),
    ]
    for i, (t, f, sub) in enumerate(cards):
        kpi_card(ws, 24, 2 + i * 3, t, f, FMT_BRL0 if i == 0 else FMT_PCT, sub, None, width_cols=3, color=NAVY2)
    for rng in ("K10", "H15", "K15", "N15", "N20", "E25", "N25"):
        pass
    # cores condicionais nos resultados
    for ref in ("N10:P10", "H15:J15", "K15:M15"):
        cf_negative_red_positive_green(ws, ref)
    ws.conditional_formatting.add("N20:P20", CellIsRule(operator="greaterThan", formula=["LimVar"], font=Font(name=FONT, color=RED, bold=True)))
    ws.conditional_formatting.add("E25:G25", CellIsRule(operator="greaterThan", formula=["LimVar"], font=Font(name=FONT, color=RED, bold=True)))
    ws.conditional_formatting.add("H25:J25", CellIsRule(operator="greaterThan", formula=["LimPacote"], font=Font(name=FONT, color=RED, bold=True)))
    ws.conditional_formatting.add("K25:M25", CellIsRule(operator="greaterThan", formula=["LimConc"], font=Font(name=FONT, color=RED, bold=True)))
    ws.conditional_formatting.add("H10:J10", CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, color=RED, bold=True)))

    # ---- destaques gerenciais automáticos
    section(ws, 28, "5. DESTAQUES GERENCIAIS AUTOMÁTICOS (fatos observáveis nos dados - sem justificativas inferidas)", "B", "Q", GRAPHITE)
    mesN = "INDEX(Meses,MesSel)"
    destaques = [
        f'=IF(FiltroAtivo=1,"Filtro de custo ativo: resultado e margem não são apresentados nesta seleção.",IF({rem_mes}=0,"Remuneração de "&{mesN}&"/"&AnoSel&" não informada: resultado e margem do mês indisponíveis.",IF({res_mes}<0,"Resultado gerencial NEGATIVO em "&{mesN}&": R$ "&TEXT({res_mes},"#,##0")&" (margem "&TEXT({res_mes}/{rem_mes},"0.0%")&").","Resultado gerencial positivo em "&{mesN}&": R$ "&TEXT({res_mes},"#,##0")&" (margem "&TEXT({res_mes}/{rem_mes},"0.0%")&").")))',
        f'=IF({cost_ant}>0,IF(ABS(({cost_mes}-{cost_ant})/{cost_ant})>LimVar,"Variação relevante de custo: "&TEXT(({cost_mes}-{cost_ant})/{cost_ant},"+0.0%;-0.0%")&" em "&{mesN}&" vs mês anterior (R$ "&TEXT({cost_mes},"#,##0")&" x R$ "&TEXT({cost_ant},"#,##0")&").","Custo de "&{mesN}&" variou "&TEXT(({cost_mes}-{cost_ant})/{cost_ant},"+0.0%;-0.0%")&" vs mês anterior (dentro do limite de "&TEXT(LimVar,"0%")&")."),"Sem custo no mês anterior para comparação.")',
        f'=IF(SUM(DG_PacValores)>0,IF({top_pac_pct}>LimPacote,"Concentração de custo: o pacote "&{top_pac}&" representa "&TEXT({top_pac_pct},"0.0%")&" do custo do ano.","Maior pacote do ano: "&{top_pac}&" ("&TEXT({top_pac_pct},"0.0%")&" do custo)."),"Sem custos no ano selecionado.")',
        f'=IF({R_DG}!$D$86>LimConc,"Concentração em fornecedor: "&{R_DG}!$B$86&" responde por "&TEXT({R_DG}!$D$86,"0.0%")&" do custo do ano (limite "&TEXT(LimConc,"0%")&").","Maior fornecedor: "&{R_DG}!$B$86&" ("&TEXT({R_DG}!$D$86,"0.0%")&" do custo do ano); dez maiores somam "&TEXT(SUM({R_DG}!$D$86:$D$95),"0.0%")&".")',
        f'=IF({vencido}>0,"Existem R$ "&TEXT({vencido},"#,##0")&" em "&COUNTIFS({a_cons},1,{a_st},"Vencido",{FILT_A})&" parcela(s) VENCIDA(S) e não liquidadas na data de referência.","Não há parcelas vencidas em aberto na data de referência.")',
        f'="Compromissos dos próximos 30 dias: R$ "&TEXT({prox30},"#,##0")&"; total a vencer: R$ "&TEXT({avencer},"#,##0")&" ("&COUNTIFS({a_cons},1,{a_st},"A vencer",{FILT_A})&" parcelas)."',
        f'=IF(FlagDataPgto=1,"Caixa realizado apurado pela data efetiva de pagamento do Rodopar 108.","ATENÇÃO: a extração 108 não preserva os grupos Doc. Pago em na ordem original; o caixa realizado é apresentado por mês de VENCIMENTO dos títulos pagos.")',
        f'="Mês de maior custo por competência no ano: "&INDEX(DG_Meses,MATCH(MAX({R_DG}!$E$6:$E$17),{R_DG}!$E$6:$E$17,0))&" (R$ "&TEXT(MAX({R_DG}!$E$6:$E$17),"#,##0")&"); meses com resultado negativo: "&IF(FiltroAtivo=1,"n/a",COUNTIF({R_DG}!$F$6:$F$17,"<0"))&"."',
        f'="Qualidade dos dados: "&COUNTIF({_whole(R_NF,NF,"Alertas")},"?*")&" documento(s) com alertas e "&(MAX({_whole(R_108,P,"Cum auditoria")})+MAX({_whole(R_1015,A,"Cum auditoria")}))&" parcela(s) com alertas - ver Conciliação e Qualidade."',
        f'=IF(COUNTIFS({nfSt},"Divergência*")>0,"Divergências de conciliação: "&COUNTIFS({nfSt},"Divergência*")&" documento(s), diferença de R$ "&TEXT(SUMIFS({_whole(R_NF,NF,"Diferença")},{nfSt},"Divergência*"),"#,##0.00")&".","Sem divergências de valor entre documentos e parcelas.")',
    ]
    for i, f in enumerate(destaques):
        r = 29 + i
        write(ws, f"B{r}", f"{i+1}.", font(True, NAVY, 9), None, ALIGN_C)
        write(ws, f"C{r}", f, font(False, GRAPHITE, 9), None, ALIGN_LW)
        ws.merge_cells(f"C{r}:Q{r}")
        ws.row_dimensions[r].height = 16
        ws.conditional_formatting.add(f"C{r}", FormulaRule(formula=[f'OR(ISNUMBER(SEARCH("NEGATIVO",C{r})),ISNUMBER(SEARCH("VENCIDA",C{r})),ISNUMBER(SEARCH("ATENÇÃO",C{r})),ISNUMBER(SEARCH("relevante",C{r})),ISNUMBER(SEARCH("Concentração",C{r})))'], font=Font(name=FONT, color=RED, bold=True)))

    # ---- gráficos
    section(ws, 40, "6. GRÁFICOS EXECUTIVOS (ano selecionado; filtros do Painel aplicados aos custos)", "B", "Q", NAVY)
    dg = wb[S_DG]
    cats = Reference(dg, min_col=2, min_row=6, max_row=17)
    # 1 remuneração x custo
    ch = LineChart()
    ch.add_data(Reference(dg, min_col=4, min_row=5, max_row=17), titles_from_data=True)
    ch.add_data(Reference(dg, min_col=5, min_row=5, max_row=17), titles_from_data=True)
    ch.set_categories(cats)
    style_chart(ch, "1. Remuneração x Despesas por competência", 17, 8)
    color_series(ch, [GREEN, NAVY], line=True)
    ws.add_chart(ch, "B41")
    # 2 resultado mensal
    ch = BarChart(); ch.type = "col"
    ch.add_data(Reference(dg, min_col=6, min_row=5, max_row=17), titles_from_data=True)
    ch.set_categories(cats)
    style_chart(ch, "2. Resultado gerencial mensal", 17, 8)
    color_series(ch, [NAVY])
    ch.series[0].invertIfNegative = False
    ws.add_chart(ch, "J41")
    # 3 pacotes
    pie = DoughnutChart()
    pie.add_data(Reference(dg, min_col=3, min_row=22, max_row=37), titles_from_data=False)
    pie.set_categories(Reference(dg, min_col=2, min_row=22, max_row=37))
    style_chart(pie, "3. Distribuição dos custos por pacote (ano)", 17, 8, "r")
    pie.holeSize = 50
    ws.add_chart(pie, "B58")
    # 4 categorias
    ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
    for k in range(1, 10):
        col = 2 + k
        s = Series(Reference(dg, min_col=col, min_row=70, max_row=81), title=None)
        s.tx = SeriesLabel(strRef=StrRef(f"{R_DG}!${get_column_letter(col)}$69"))
        ch.series.append(s)
    ch.set_categories(Reference(dg, min_col=2, min_row=70, max_row=81))
    style_chart(ch, "4. Evolução das despesas por categoria (8 maiores + outras)", 17, 8)
    color_series(ch)
    ws.add_chart(ch, "J58")
    # 5 top 10 fornecedores
    ch = BarChart(); ch.type = "bar"
    ch.add_data(Reference(dg, min_col=3, min_row=85, max_row=95), titles_from_data=True)
    ch.set_categories(Reference(dg, min_col=2, min_row=86, max_row=95))
    style_chart(ch, "5. Dez maiores fornecedores por custo (ano)", 17, 8)
    ch.x_axis.scaling.orientation = "maxMin"
    color_series(ch, [NAVY])
    ws.add_chart(ch, "B75")
    # 6 títulos pagos x a pagar (por vencimento)
    ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"
    ch.add_data(Reference(dg, min_col=9, min_row=5, max_row=17), titles_from_data=True)
    ch.add_data(Reference(dg, min_col=8, min_row=5, max_row=17), titles_from_data=True)
    ch.set_categories(cats)
    style_chart(ch, "6. Títulos pagos x a pagar, por mês de vencimento", 17, 8)
    color_series(ch, [GREEN, RED])
    ws.add_chart(ch, "J75")
    # 7 compromissos futuros
    ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
    ch.add_data(Reference(dg, min_col=4, max_col=5, min_row=99, max_row=112), titles_from_data=True)
    ch.set_categories(Reference(dg, min_col=2, min_row=100, max_row=112))
    style_chart(ch, "7. Composição dos compromissos futuros por mês (a partir da data de referência)", 17, 8)
    color_series(ch, [NAVY, RED])
    ws.add_chart(ch, "B92")
    # 8 competência x caixa
    ch = LineChart()
    ch.add_data(Reference(dg, min_col=5, min_row=5, max_row=17), titles_from_data=True)
    ch.add_data(Reference(dg, min_col=7, min_row=5, max_row=17), titles_from_data=True)
    ch.add_data(Reference(dg, min_col=8, min_row=5, max_row=17), titles_from_data=True)
    ch.set_categories(cats)
    style_chart(ch, "8. Competência x Caixa: custo reconhecido, pagos realizados e a pagar em aberto", 17, 8)
    color_series(ch, [NAVY, GREEN, RED], line=True)
    ws.add_chart(ch, "J92")
    ws.freeze_panes = "A7"
    setup_print(ws, area="B1:Q108", header_text="Painel Executivo Financeiro | Frota")
    ws.page_setup.fitToHeight = 0


LEIA_ME = [
    ("PROPÓSITO", [
        "Esta planilha transforma o antigo 'Fluxo de caixa.xlsx' em um controle gerencial do setor de Frota com quatro camadas integradas: "
        "(1) dados de origem - Relação de NF's, Fornecedores, Rodopar 108 e 1015; (2) tratamento e conciliação - colunas calculadas dentro das próprias tabelas de origem; "
        "(3) modelo financeiro - Remuneração, DRE, Fluxo de Caixa, Competência x Caixa; (4) apresentação - Painel Executivo e análises.",
        "Três perguntas respondidas no Painel: quanto a Frota custou por competência; quanto pagou efetivamente em cada mês; quanto ainda terá que pagar e quando.",
    ]),
    ("ROTINA MENSAL (5 PASSOS)", [
        "1) Rodopar 108 (documentos pagos): abra a aba 'Doc's Pagos | Rodopar 108', selecione o conteúdo antigo de C2 até a última linha da coluna M e exclua (Delete). Cole a nova extração completa a partir de C2, NA ORDEM ORIGINAL do relatório (não classifique!). As linhas 'Doc. Pago em : dd/mm/aaaa' devem permanecer: elas fornecem a data efetiva de pagamento.",
        "2) Rodopar 1015 (contas a pagar): mesma rotina na aba 'Doc Há Pagar | Rodopar 1015', colando a partir de C2 a extração completa com TODAS as parcelas em aberto (inclusive vencimentos de meses futuros e de anos seguintes).",
        "3) Relação de NF's: acrescente os novos documentos da Frota nas linhas seguintes (colunas B, E a O). As colunas A, C, D e todas as colunas cinza são calculadas. Use as colunas amarelas P, Q e R apenas para ajustes por documento (pacote, categoria, competência).",
        "4) Remuneração: informe o valor consolidado da competência (coluna F) na aba Remuneração. Ajustes (G) e observações são opcionais.",
        "5) Pressione F9 se o cálculo estiver manual (Fórmulas > Opções de Cálculo > Automático é o recomendado) e consulte Painel Executivo, DRE, Fluxo de Caixa e demais relatórios. Para apresentar, imprima/exporte em PDF o Painel e a DRE (áreas de impressão já configuradas).",
        "As tabelas (tblNF, tbl108, tbl1015, tblFornecedores) se expandem automaticamente ao colar/digitar abaixo da última linha e replicam as fórmulas das colunas calculadas. Se, ao colar, a tabela não expandir, use Design da Tabela > Redimensionar Tabela até a última linha colada.",
    ]),
    ("COMO A FROTA É IDENTIFICADA NAS BASES RODOPAR", [
        "As bases 108 e 1015 são da empresa inteira. Uma parcela só entra nos números da Frota quando é conciliada a um documento da Relação de NF's (escopo 'Frota - NF vinculada'). "
        "Parcelas de fornecedores cadastrados como Frota mas sem NF na relação ficam em fila de revisão ('Frota - sem NF') e NÃO entram em custos, pagamentos ou obrigações até serem tratadas. Fornecedores com Area '-' são de outras áreas.",
        "Chaves de conciliação: N1 = Documento + Fornecedor + Tipo + Data de emissão (Rodopar 108 x NF); N2 = Documento + Fornecedor + Tipo (Rodopar 1015, que não traz emissão). Correspondências ambíguas (mesma chave N2 em mais de uma NF) são sinalizadas.",
    ]),
    ("COMPETÊNCIA x CAIXA (REGRAS)", [
        "DRE (competência): cada documento entra uma única vez pelo Valor Doc, no mês de competência = 'Competência (ajuste manual)' ou, se vazio, mês de emissão (estimativa). Pagar parcelas nunca cria custo novo.",
        "Pagamentos realizados: parcelas do Rodopar 108 conciliadas; valor líquido pago = Vlr. Total. O mês de caixa é a data efetiva de pagamento (grupo 'Doc. Pago em') quando a extração está na ordem original; caso contrário, o mês de vencimento, sinalizado em todo o modelo.",
        "Compromissos futuros: parcelas do Rodopar 1015 conciliadas, por vencimento. Vencido = vencimento anterior à data de referência (Parâmetros) e ainda em aberto. A ausência de uma parcela no 1015 não é tratada como pagamento.",
        "Remuneração é receita gerencial por competência e NÃO é entrada de caixa. Entradas de caixa só existem se cadastradas em Remuneração > Recebimentos efetivos.",
    ]),
    ("INTERPRETAÇÃO DA DRE", [
        "Receita Gerencial Líquida = remuneração + ajustes. Custos operacionais por pacote (ordem definida em Parâmetros). Resultado operacional = receita - custos. Juros pagos e descontos obtidos (Rodopar 108) ajustam o resultado consolidado no mês de caixa.",
        "Margem gerencial = resultado / receita líquida; custo sobre remuneração = custos / receita líquida. Quando não há remuneração informada o indicador aparece como 'n/d'. Não há tributos, depreciação, provisões, metas ou orçamento na base - nada disso foi inventado.",
    ]),
    ("CONFERÊNCIA DE COMPROMISSOS E INCONSISTÊNCIAS", [
        "Pagamentos e Obrigações: cartões de aging (vencido, 7/15/30/60 dias), maiores credores e tabela analítica por parcela com filtro de status.",
        "Conciliação e Qualidade: verificações (fornecedor sem cadastro, sem pacote, duplicidades, divergências de valor, parcelas sem correspondência, registros de outras áreas, competência estimada) e três filas de revisão. "
        "Divergências de valor típicas: retenções na fonte (parcela líquida menor que o documento) e parcelas com vencimento fora do período extraído no 1015 - extraia sempre o 1015 completo.",
        "A Relação de NF's, o 108 e o 1015 trazem colunas 'Alertas' e 'Status' em cada linha, com filtros (AutoFiltro) para investigação direta.",
    ]),
    ("ANO SEGUINTE", [
        "Mantenha o histórico nas abas de origem (as tabelas podem conter vários exercícios). Para 2027: acrescente as linhas de 2027 na aba Remuneração (já pré-criadas), continue colando as extrações completas do Rodopar e mude o Ano no Painel Executivo. "
        "DRE, Fluxo de Caixa, Competência x Caixa, Custos por Pacote e Análise de Fornecedores têm célula de ano própria (padrão = ano do Painel).",
        "Novos pacotes ou categorias: inclua em Parâmetros (listas H e J) e no cadastro de Fornecedores. Novos fornecedores: inclua na aba Fornecedores com Area = Frota, Pacote e Categoria.",
    ]),
    ("LIMITAÇÕES CONHECIDAS", [
        "O Rodopar 108 do arquivo original foi classificado após a colagem; por isso a data efetiva de pagamento não pôde ser validada e o caixa realizado está por mês de vencimento. Colar a próxima extração na ordem original resolve automaticamente.",
        "O Rodopar 1015 não informa data de emissão nem valor original do documento; o valor da parcela pode vir líquido de retenções. O campo 'Acumulado' é apenas auxiliar.",
        "Capacidade: as filas de revisão exibem até 300 itens, a tabela analítica de parcelas até 1500 e o detalhamento de documentos por pacote até 300 (o total é sempre informado). As tabelas de origem não têm limite.",
        "Compatível com Excel Desktop (Windows). Sem macros, sem Power Query e sem vínculos externos: tudo é fórmula nativa, recalculada ao abrir o arquivo.",
    ]),
]


def build_leiame(ws):
    set_widths(ws, {"A": 2, "B": 150})
    title_block(ws, "LEIA-ME | INSTRUÇÕES DE UTILIZAÇÃO", "Controle Financeiro Gerencial - Setor de Frota (Horizonte Logística)", "B", "B")
    r = 5
    for titulo, paras in LEIA_ME:
        section(ws, r, titulo, "B", "B")
        r += 1
        for p in paras:
            write(ws, f"B{r}", p, font(False, GRAPHITE, 9), None, ALIGN_LW)
            ws.row_dimensions[r].height = max(18, 13 * (len(p) // 140 + 1))
            r += 1
        r += 1
    ws.sheet_view.showGridLines = False
    setup_print(ws, landscape=False, area=f"B1:B{r}", header_text="Leia-me")
