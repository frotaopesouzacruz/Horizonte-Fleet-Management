# -*- coding: utf-8 -*-
"""Estilos, constantes e utilitários compartilhados pelo gerador da planilha."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

# ---------------------------------------------------------------- nomes de abas
S_PAINEL = "Painel Executivo"
S_REM = "Remuneração"
S_DRE = "DRE Gerencial"
S_FC = "Fluxo de Caixa"
S_CC = "Custos | Competência x Caixa"
S_CP = "Custos por Pacote"
S_PO = "Pagamentos e Obrigações"
S_AF = "Análise de Fornecedores"
S_CQ = "Conciliação e Qualidade"
S_LEIA = "Leia-me"
S_PAR = "Parâmetros"
S_DG = "Dados Gráficos"
S_NF = "Relação de NF's"
S_FOR = "Fornecedores"
S_108 = "Doc's Pagos | Rodopar 108"
S_1015 = "Doc Há Pagar | Rodopar 1015"


def q(sheet):
    """Referência de aba com aspas (escapa apóstrofo)."""
    return "'" + sheet.replace("'", "''") + "'"


R_NF = q(S_NF)
R_108 = q(S_108)
R_1015 = q(S_1015)
R_FOR = q(S_FOR)
R_PAR = q(S_PAR)
R_REM = q(S_REM)
R_DG = q(S_DG)
R_CP = q(S_CP)
R_PO = q(S_PO)
R_AF = q(S_AF)
R_CQ = q(S_CQ)
R_PAINEL = q(S_PAINEL)

# ---------------------------------------------------------------- paleta
NAVY = "1F3864"
NAVY2 = "2F5597"
BLUE = "2E75B6"
LBLUE = "9DC3E6"
GRAPHITE = "3A3A3A"
GRAY1 = "F2F2F2"
GRAY2 = "D9D9D9"
GRAY3 = "BFBFBF"
GRAY4 = "7F7F7F"
WHITE = "FFFFFF"
GREEN = "2E7D32"
GREEN_L = "E2EFDA"
RED = "C00000"
RED_L = "FBE5E5"
AMBER = "BF8F00"
AMBER_L = "FFF2CC"
INPUT_FILL = "FFF9E6"
INPUT_FONT = "0000FF"
GOLD = "F2A516"

FONT = "Arial"

# ---------------------------------------------------------------- formatos
FMT_BRL = '"R$ "#,##0.00;[Red]"-R$ "#,##0.00;"R$ -"'
FMT_BRL0 = '"R$ "#,##0;[Red]"-R$ "#,##0;"R$ -"'
FMT_NUM = '#,##0.00;[Red]-#,##0.00;"-"'
FMT_INT = '#,##0;[Red]-#,##0;"-"'
FMT_PCT = '0.0%;[Red]-0.0%;"-"'
FMT_DATE = 'dd/mm/yyyy'
FMT_DATETIME = 'dd/mm/yyyy hh:mm'
FMT_TXT = '@'


def fill(hexcolor):
    return PatternFill("solid", start_color=hexcolor, end_color=hexcolor)


def font(bold=False, color=GRAPHITE, size=10, italic=False, name=FONT):
    return Font(name=name, bold=bold, color=color, size=size, italic=italic)


THIN = Side(style="thin", color=GRAY3)
HAIR = Side(style="hair", color=GRAY3)
MED = Side(style="medium", color=NAVY)
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_BOTTOM = Border(bottom=THIN)
BORDER_TOP_MED = Border(top=MED)

ALIGN_C = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_L = Alignment(horizontal="left", vertical="center", wrap_text=False)
ALIGN_LW = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_R = Alignment(horizontal="right", vertical="center")


def set_widths(ws, widths):
    """widths: dict letra->largura ou lista a partir de A."""
    if isinstance(widths, dict):
        for k, v in widths.items():
            ws.column_dimensions[k].width = v
    else:
        for i, v in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = v


def style_range(ws, rng, font_=None, fill_=None, align=None, border=None, fmt=None):
    for row in ws[rng]:
        for c in row:
            if font_ is not None:
                c.font = font_
            if fill_ is not None:
                c.fill = fill_
            if align is not None:
                c.alignment = align
            if border is not None:
                c.border = border
            if fmt is not None:
                c.number_format = fmt


def write(ws, ref, value, font_=None, fill_=None, align=None, fmt=None, border=None):
    c = ws[ref]
    c.value = value
    if font_ is not None:
        c.font = font_
    if fill_ is not None:
        c.fill = fill_
    if align is not None:
        c.alignment = align
    if fmt is not None:
        c.number_format = fmt
    if border is not None:
        c.border = border
    return c


def title_block(ws, title, subtitle, first_col="B", last_col="P", row=1):
    ws.row_dimensions[row].height = 30
    ws.row_dimensions[row + 1].height = 18
    ws.merge_cells(f"{first_col}{row}:{last_col}{row}")
    ws.merge_cells(f"{first_col}{row+1}:{last_col}{row+1}")
    write(ws, f"{first_col}{row}", title, font(True, WHITE, 16), fill(NAVY), Alignment(horizontal="left", vertical="center", indent=1))
    write(ws, f"{first_col}{row+1}", subtitle, font(False, WHITE, 10, italic=True), fill(NAVY2), Alignment(horizontal="left", vertical="center", indent=1))
    # faixa dourada fina
    ws.row_dimensions[row + 2].height = 4
    for col in range(column_index_from_string(first_col), column_index_from_string(last_col) + 1):
        ws.cell(row + 2, col).fill = fill(GOLD)


def section(ws, row, title, first_col="B", last_col="P", color=NAVY2):
    ws.merge_cells(f"{first_col}{row}:{last_col}{row}")
    ws.row_dimensions[row].height = 20
    write(ws, f"{first_col}{row}", title, font(True, WHITE, 11), fill(color), Alignment(horizontal="left", vertical="center", indent=1))


def header_row(ws, row, col_start, headers, fill_color=NAVY, font_color=WHITE, height=30, widths=None):
    for i, h in enumerate(headers):
        c = ws.cell(row, col_start + i, h)
        c.font = font(True, font_color, 9)
        c.fill = fill(fill_color)
        c.alignment = ALIGN_C
        c.border = BORDER_ALL
    ws.row_dimensions[row].height = height


def input_cell(ws, ref, value=None, fmt=None, comment=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.fill = fill(INPUT_FILL)
    c.font = font(True, INPUT_FONT, 10)
    c.border = BORDER_ALL
    c.alignment = ALIGN_C
    if fmt:
        c.number_format = fmt
    if comment:
        c.comment = Comment(comment, "Controle Financeiro")
    return c


def label(ws, ref, text, bold=False, color=GRAPHITE, size=9, align=ALIGN_L, italic=False):
    return write(ws, ref, text, font(bold, color, size, italic=italic), align=align)


def kpi_card(ws, row, col, title_, formula, fmt=FMT_BRL0, sub_formula=None, sub_fmt=None, width_cols=2, color=NAVY):
    """Cartão de indicador ocupando `width_cols` colunas e 3 linhas (título, valor, subinfo)."""
    c1 = get_column_letter(col)
    c2 = get_column_letter(col + width_cols - 1)
    ws.merge_cells(f"{c1}{row}:{c2}{row}")
    ws.merge_cells(f"{c1}{row+1}:{c2}{row+1}")
    ws.merge_cells(f"{c1}{row+2}:{c2}{row+2}")
    write(ws, f"{c1}{row}", title_, font(True, WHITE, 8), fill(color), ALIGN_C)
    write(ws, f"{c1}{row+1}", formula, font(True, NAVY, 14), fill(GRAY1), ALIGN_C, fmt)
    write(ws, f"{c1}{row+2}", sub_formula if sub_formula else "", font(False, GRAY4, 8, italic=True), fill(GRAY1), ALIGN_C, sub_fmt)
    for r in range(row, row + 3):
        for cc in range(col, col + width_cols):
            ws.cell(r, cc).border = BORDER_ALL
    ws.row_dimensions[row].height = 24
    ws.row_dimensions[row + 1].height = 26
    ws.row_dimensions[row + 2].height = 16


def add_list_validation(ws, sqref, formula1, allow_blank=True, error="Selecione um valor da lista."):
    dv = DataValidation(type="list", formula1=formula1, allow_blank=allow_blank, showErrorMessage=True,
                        errorTitle="Valor inválido", error=error)
    dv.add(sqref)
    ws.add_data_validation(dv)
    return dv


def cf_negative_red_positive_green(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color=RED, bold=True)))
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=Font(name=FONT, color=GREEN, bold=True)))


def cf_text(ws, rng, text, fill_color, font_color=None, contains=True):
    first = rng.split(":")[0]
    if contains:
        f = f'ISNUMBER(SEARCH("{text}",{first}))'
    else:
        f = f'{first}="{text}"'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f], fill=fill(fill_color), font=Font(name=FONT, color=font_color or GRAPHITE, bold=True)))


def comp_formula(date_ref):
    """AAAA-MM independente de idioma."""
    return f'YEAR({date_ref})&"-"&TEXT(MONTH({date_ref}),"00")'


def date_key_formula(date_ref):
    return f'YEAR({date_ref})&"-"&TEXT(MONTH({date_ref}),"00")&"-"&TEXT(DAY({date_ref}),"00")'


def col_map(start_letter, names):
    """Retorna dict nome->letra a partir de uma letra inicial."""
    idx = column_index_from_string(start_letter)
    return {n: get_column_letter(idx + i) for i, n in enumerate(names)}


def setup_print(ws, landscape=True, fit_width=True, area=None, title_rows=None, header_text=None, footer_text=None):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if area:
        ws.print_area = area
    if title_rows:
        ws.print_title_rows = title_rows
    if header_text:
        ws.oddHeader.center.text = header_text
        ws.oddHeader.center.size = 9
    ws.oddFooter.left.text = footer_text or "Horizonte Logística | Controle Financeiro Gerencial - Frota"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Página &P de &N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.center.text = "Impresso em &D"
    ws.oddFooter.center.size = 8
