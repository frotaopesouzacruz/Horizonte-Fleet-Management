# -*- coding: utf-8 -*-
"""Testes de validação (seção 12): injeta dados em cópias da planilha, recalcula com LibreOffice e verifica resultados."""
import sys, os, shutil, subprocess, json, datetime, copy
sys.path.insert(0, "build")
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
from sources import NF, P, A, F
from common import *

RECALC = [p for p in __import__("glob").glob("/root/.claude/skills/synced/*/xlsx/scripts/recalc.py")][0]
BASE = "out/Controle_Financeiro_Gerencial_Frota.xlsx"

def append_row(ws, table_name, raw_values):
    """Simula a expansão de tabela do Excel: nova linha com valores brutos + fórmulas copiadas da linha anterior."""
    t = ws.tables[table_name]
    first, last = t.ref.split(":")
    last_row = int(''.join(ch for ch in last if ch.isdigit()))
    last_col = ''.join(ch for ch in last if ch.isalpha())
    new = last_row + 1
    for c in range(1, ci(last_col) + 1):
        src = ws.cell(last_row, c)
        if isinstance(src.value, str) and src.value.startswith("="):
            ws.cell(new, c).value = Translator(src.value, origin=f"{gl(c)}{last_row}").translate_formula(f"{gl(c)}{new}")
        ws.cell(new, c).number_format = src.number_format
        ws.cell(new, c).font = copy.copy(src.font)
    for col, val in raw_values.items():
        ws[f"{col}{new}"] = val
    t.ref = f"{first}:{last_col}{new}"
    if t.autoFilter is not None:
        t.autoFilter.ref = t.ref
    return new

def recalc(path):
    out = subprocess.run([sys.executable, RECALC, path, "900"], capture_output=True, text=True)
    j = json.loads(out.stdout)
    assert j.get("status") == "success", j
    return j

def val(wb, sheet, ref):
    return wb[sheet][ref].value

def approx(a, b, tol=0.01):
    return abs((a or 0) - (b or 0)) <= tol

results = []
def check(name, cond, info=""):
    results.append((name, bool(cond), info))
    print(("OK   " if cond else "FALHA"), name, info)

# ------------------------------------------------------------------ T1: remuneração, novo doc, parcelamento, pagamento, duplicidade, ano seguinte, fornecedor sem cadastro
t1 = "test/t1.xlsx"
shutil.copy(BASE, t1)
wb = openpyxl.load_workbook(t1)
rem = wb[S_REM]
rem["F8"] = 300000; rem["F16"] = 350000; rem["G16"] = -10000; rem["J16"] = datetime.datetime(2026, 10, 9)
rem["F20"] = 10000  # 2027-01
nf = wb[S_NF]
r_new = append_row(nf, "tblNF", {"B": "TESTE-001", "E": "CASA DA SPRINTER LTDA", "F": "NSE", "G": "49.353.383/0001-70", "H": 87, "I": 12000,
                                 "J": datetime.datetime(2026, 9, 15), "K": datetime.datetime(2026, 9, 16), "L": "CONTAGEM", "M": "MG", "N": 1556, "O": "TESTE"})
r_2027 = append_row(nf, "tblNF", {"B": "TESTE-2027", "E": "LOCALIZA FLEET S/A", "F": "FT", "G": "02.286.479/0001-08", "H": 87, "I": 5000,
                                  "J": datetime.datetime(2027, 1, 10), "K": datetime.datetime(2027, 1, 12), "L": "BELO HORIZONTE", "M": "MG", "N": 1556, "O": "TESTE"})
r_nc = append_row(nf, "tblNF", {"B": "TESTE-NC", "E": "FORNECEDOR TESTE SEM CADASTRO LTDA", "F": "NFE", "G": "00.000.000/0001-00", "H": 87, "I": 777,
                                "J": datetime.datetime(2026, 9, 20), "K": datetime.datetime(2026, 9, 20), "L": "X", "M": "MG", "N": 1556, "O": "TESTE"})
ap = wb[S_1015]
for i, venc in enumerate([datetime.datetime(2026, 11, 15), datetime.datetime(2026, 12, 15)], start=2):
    append_row(ap, "tbl1015", {"C": 87, "D": "TESTE-001", "E": "1", "F": "CASA DA SPRINTER LTDA", "G": venc, "H": "NSE", "I": i, "J": 4000, "K": 4000})
append_row(ap, "tbl1015", {"C": 87, "D": "TESTE-2027", "E": "1", "F": "LOCALIZA FLEET S/A", "G": datetime.datetime(2027, 2, 10), "H": "FT", "I": 1, "J": 5000, "K": 5000})
pg = wb[S_108]
# parcela 1 do TESTE-001 paga (vencimento 15/10, com juros e desconto para testar)
append_row(pg, "tbl108", {"C": "TESTE-001", "D": "1", "E": "CASA DA SPRINTER LTDA", "F": datetime.datetime(2026, 9, 15), "G": "NSE", "H": datetime.datetime(2026, 10, 15),
                          "I": datetime.datetime(2026, 9, 15), "J": 4000, "K": 50, "L": 10, "M": 3960})
# reimportação em duplicidade: copia 5 parcelas Frota já existentes (fornecedor LOCALIZA FLEET S/A)
dups = []
for r in range(2, pg.max_row + 1):
    if pg.cell(r, 5).value == "LOCALIZA FLEET S/A" and isinstance(pg.cell(r, 10).value, (int, float)):
        dups.append({gl(c): pg.cell(r, c).value for c in range(3, 14)})
    if len(dups) == 5:
        break
dup_total = sum(d["J"] for d in dups)
for d in dups:
    append_row(pg, "tbl108", d)
wb[S_DRE]["C4"] = 2027  # DRE em 2027 enquanto o Painel permanece em 2026
wb.save(t1)
recalc(t1)
w = openpyxl.load_workbook(t1, data_only=True)
# Teste 02 - remuneração
check("T02 Remuneração Jan alimenta Painel/Dados (300.000)", approx(val(w, S_DG, "D6"), 300000), str(val(w, S_DG, "D6")))
check("T02 Receita líquida Set = 350.000 - 10.000 ajuste", approx(val(w, S_DG, "D14"), 340000), str(val(w, S_DG, "D14")))
custo_set = val(w, S_DG, "E14")
check("T03/T04 Custo Set inclui novo doc 12.000 uma única vez", approx(custo_set, 272144.62 + 12000), str(custo_set))
res_set = val(w, S_PAINEL, "N10")
check("T02 Resultado gerencial Set = receita - custo - juros + descontos (filtro inativo)", approx(res_set, 340000 - custo_set), f"{res_set}")
check("T02 Fluxo de caixa: entradas de Set = 0 (sem recebimento informado)", approx(val(w, S_FC, "K10"), 0), str(val(w, S_FC, "K10")))
check("T02 Painel margem exibida", "Margem" in str(val(w, S_PAINEL, "N11")), str(val(w, S_PAINEL, "N11")))
# Teste 05 - pagamento
nfst = w[S_NF]
st = nfst.cell(r_new, ci(NF["Status documento"])).value
check("T05 Documento TESTE-001 status 'Parcialmente pago'", st == "Parcialmente pago", st)
check("T05 Pago bruto 4.000 / em aberto 8.000", approx(nfst.cell(r_new, ci(NF["Valor pago bruto"])).value, 4000) and approx(nfst.cell(r_new, ci(NF["Valor em aberto"])).value, 8000),
      f"{nfst.cell(r_new, ci(NF['Valor pago bruto'])).value} / {nfst.cell(r_new, ci(NF['Valor em aberto'])).value}")
check("T05 Juros 10 e desconto 50 reconhecidos na DRE (mês de caixa Out)", approx(val(w, S_DG, "J15"), 10) and approx(val(w, S_DG, "K15"), 50), f"{val(w, S_DG, 'J15')} / {val(w, S_DG, 'K15')}")
check("T05 Pagos realizados Out (líquido) inclui 3.960", approx(val(w, S_DG, "G15"), 126361.60 + 3960), str(val(w, S_DG, "G15")))
check("T04 Obrigações em aberto total = 169.311,14 + 8.000 + 5.000", approx(val(w, S_PAINEL, "K10"), 169311.14 + 8000 + 5000), str(val(w, S_PAINEL, "K10")))
check("T04 Projeção Nov inclui 4.000", approx(val(w, S_DG, "D102"), 19948.36 + 4000), str(val(w, S_DG, "D102")))
# Teste 06 - duplicidade
check("T06 Parcelas 108 consideradas = 527 (526 + 1 nova; 5 duplicadas ignoradas)", val(w, S_PAINEL, "E20") == 527, str(val(w, S_PAINEL, "E20")))
check("T06 Auditoria aponta 5 parcelas duplicadas", val(w, S_CQ, "C17") == 5 and approx(val(w, S_CQ, "D17"), dup_total), f"{val(w, S_CQ, 'C17')} / {val(w, S_CQ, 'D17')} (esperado {dup_total})")
check("T06 Total pago bruto 2026 inalterado pelas duplicatas", approx(val(w, S_DG, "I18"), 2385181.48 + 4000), str(val(w, S_DG, "I18")))
# Teste 07 - ano seguinte
check("T07 DRE 2027 Jan: remuneração 10.000 e custo 5.000", approx(val(w, S_DRE, "C10"), 10000) and approx(val(w, S_DRE, "C30"), 5000), f"{val(w, S_DRE, 'C10')} / {val(w, S_DRE, 'C30')}")
check("T07 DRE 2027 resultado Jan = 5.000 e margem 50%", approx(val(w, S_DRE, "C39"), 5000) and approx(val(w, S_DRE, "C42"), 0.5), f"{val(w, S_DRE, 'C39')} / {val(w, S_DRE, 'C42')}")
check("T07 Painel 2026 não contaminado pelo doc de 2027 (acumulado)", approx(val(w, S_DG, "E18"), 2558436.66 + 12000), str(val(w, S_DG, "E18")))
check("T07 Projeção 12 meses inclui Fev/2027 = 5.000", approx(val(w, S_DG, "D105"), 5000), str(val(w, S_DG, "D105")))
# Teste 08 - conciliação
check("T08 Fornecedor sem cadastro aparece na auditoria (1 doc, R$ 777) e fora da DRE", val(w, S_CQ, "C11") == 1 and approx(val(w, S_CQ, "D11"), 777) and nfst.cell(r_nc, ci(NF["Considerar DRE"])).value == 0,
      f"{val(w, S_CQ, 'C11')} / {val(w, S_CQ, 'D11')} / considerar={nfst.cell(r_nc, ci(NF['Considerar DRE'])).value}")
check("T08 Divergências de valor listadas (13)", val(w, S_CQ, "C20") == 13, str(val(w, S_CQ, "C20")))
check("T08 Fila de revisão 1 lista o fornecedor sem cadastro", any("TESTE-NC" in str(w[S_CQ].cell(r, 2).value) for r in range(59, 120)))
json.dump(results, open("test/results_t1.json", "w"), ensure_ascii=False, indent=1)
