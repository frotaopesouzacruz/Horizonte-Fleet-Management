# -*- coding: utf-8 -*-
"""Camada 1 (dados de origem) e Camada 2 (tratamento) - colunas calculadas nas tabelas de origem."""
from openpyxl.worksheet.table import Table, TableStyleInfo, TableColumn, TableFormula
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.comments import Comment
from common import *

# ---------------------------------------------------------------- definição de colunas
NF_COLS = col_map("A", [
    "Area", "Documento", "Pacote", "Categoria", "Razão Social", "Tipo", "CNPJ/CPF", "Filial", "Valor Doc",
    "Emissão", "Referência", "Munícipio", "UF", "CFOP", "Atualizado Por",
    "Pacote (ajuste manual)", "Categoria (ajuste manual)", "Competência (ajuste manual)",
    "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
    "Total conciliado", "Diferença",
    "Chave N1", "Chave N2", "Pacote final", "Categoria final", "Origem classificação", "Competência", "Base competência",
    "Registro válido", "Ocorrência", "Escopo Frota", "Considerar DRE",
    "Qtd parc. pagas", "Valor pago bruto", "Valor pago líquido", "Qtd parc. abertas", "Valor em aberto",
    "Juros", "Descontos", "Retenções/dif. pagas", "Status documento", "Ambíguo N2", "Alertas",
    "Cum drill", "Cum auditoria", "Ano comp.", "Mês comp.", "Vencido (R$)",
])
C108_COLS = col_map("A", [
    "Ref Pesquisa", "Mês", "Documento", "Série", "Fornecedor", "Emissão", "Tipo", "Vencto", "Data Ref.",
    "Vlr. Parc.", "Vlr. Desco.", "Vlr. Juros", "Vlr. Total",
    "Registro válido", "Data Pgto (derivada)", "Documento (norm.)", "Fornecedor (norm.)", "Tipo (norm.)",
    "Emissão (data)", "Vencimento (data)", "Chave N1", "Chave N2", "Ocorrência", "Linha NF (N1)", "Linha NF (N2)",
    "Nível conciliação", "Linha NF", "Área fornecedor", "Escopo", "Considerar", "Pacote", "Categoria", "Competência",
    "Comp. Vencimento", "Data caixa validada", "Comp. Caixa", "Base data caixa", "Valor bruto", "Desconto", "Juros",
    "Valor líquido pago", "Retenção/diferença", "Status", "Alertas", "Cum vista", "Cum auditoria", "Ano comp.", "Ano venc.", "Ano caixa",
])
C1015_COLS = col_map("A", [
    "Ref Pesquisa", "Mês Venc.", "Filial", "Doc.", "Série", "Fornecedor", "Vencimento", "Tipo", "Parcela",
    "Vlr. Parcela", "Acumulado",
    "Registro válido", "Documento (norm.)", "Fornecedor (norm.)", "Tipo (norm.)", "Vencimento (data)", "Chave N2",
    "Ocorrência", "Linha NF", "Nível conciliação", "Área fornecedor", "Escopo", "Considerar", "Pacote", "Categoria",
    "Competência", "Comp. Vencimento", "Valor parcela", "Dias para vencer", "Status", "Dias em atraso",
    "Faixa de vencimento", "Alertas", "Cum vista", "Cum auditoria", "Ano comp.", "Ano venc.",
])
FOR_COLS = col_map("A", [
    "Fornecedor", "CNPJ", "Area", "Pacote", "Categoria",
    "Qtd docs (ano)", "Custo competência (ano)", "Participação %", "Pago líquido (ano, caixa)", "Em aberto (total)",
    "Vencido (total)", "Qtd parcelas abertas", "Movimento sem NF (revisar)", "Alertas cadastro", "Rank custo",
    "Rank em aberto", "Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez",
    "Meses com custo", "Cum alertas",
])

NF = NF_COLS
P = C108_COLS
A = C1015_COLS
F = FOR_COLS

MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
MESES_ABREV = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]

SCOPE_NF = "Frota - NF vinculada"
SCOPE_SEMNF = "Frota - sem NF (revisar)"
SCOPE_OUTRA = "Outras áreas"
SCOPE_NC = "Fornecedor não cadastrado"
SEM_CLASS = "SEM CLASSIFICAÇÃO"
NAO_CAD = "NÃO CADASTRADO"


def _ref(cols, name, row, absolute=False):
    L = cols[name]
    return f"${L}${row}" if absolute else f"{L}{row}"


def _whole(sheet_ref, cols, name):
    L = cols[name]
    return f"{sheet_ref}!${L}:${L}"


def make_table(ws, name, last_col_letter, last_row, style="TableStyleMedium2", first_row=1):
    ref = f"A{first_row}:{last_col_letter}{last_row}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    # colunas calculadas: usa a fórmula da primeira linha de dados
    cols = []
    ncols = column_index_from_string(last_col_letter)
    for i in range(1, ncols + 1):
        hdr = ws.cell(first_row, i).value
        tc = TableColumn(id=i, name=str(hdr))
        v = ws.cell(first_row + 1, i).value
        if isinstance(v, str) and v.startswith("="):
            tc.calculatedColumnFormula = TableFormula(attr_text=v[1:])
        cols.append(tc)
    t.tableColumns = cols
    from openpyxl.worksheet.filters import AutoFilter
    t.autoFilter = AutoFilter(ref=ref)
    ws.add_table(t)
    return t


def _hdr(ws, cols, names, fill_color, start_name=None):
    for n in names:
        c = ws[f"{cols[n]}1"]
        c.value = n
        c.font = font(True, WHITE, 9)
        c.fill = fill(fill_color)
        c.alignment = ALIGN_C
        c.border = BORDER_ALL


# ====================================================================== RELAÇÃO DE NF's
def build_nf(ws, n_rows):
    """n_rows = última linha com dados (inclusive)."""
    last = n_rows
    # cabeçalhos
    base_names = list(NF_COLS.keys())
    _hdr(ws, NF, base_names[:15], NAVY)
    _hdr(ws, NF, base_names[15:18], AMBER)          # colunas manuais
    _hdr(ws, NF, base_names[18:32], NAVY2)          # meses + conciliação
    _hdr(ws, NF, base_names[32:], GRAPHITE)         # tratamento
    ws.row_dimensions[1].height = 42
    ws["S1"].comment = Comment("Parcelas (108 pagas + 1015 em aberto) vinculadas ao documento, por mês de VENCIMENTO, "
                               "no ano selecionado no Painel Executivo. Não representa competência.", "Controle Financeiro")
    ws[f"{NF['Competência (ajuste manual)']}1"].comment = Comment(
        "Opcional. Informe AAAA-MM (ex.: 2026-09) ou uma data do mês de competência quando a competência "
        "do custo for diferente do mês de emissão. Em branco = competência estimada pelo mês de emissão.", "Controle Financeiro")
    ws[f"{NF['Pacote (ajuste manual)']}1"].comment = Comment(
        "Opcional. Preencha somente quando a classificação deste documento for diferente da classificação padrão do fornecedor.", "Controle Financeiro")

    for r in range(2, last + 1):
        R = r
        E = f"$E{R}"
        # Área: lookup com sinalização de não cadastrado
        ws[f"A{R}"] = f'=IF({E}="","",IFERROR(VLOOKUP({E},{R_FOR}!$A:$C,3,0),"{NAO_CAD}"))'
        ws[f"C{R}"] = f'=IF({E}="","",IFERROR(VLOOKUP({E},{R_FOR}!$A:$E,4,0),""))'
        ws[f"D{R}"] = f'=IF({E}="","",IFERROR(VLOOKUP({E},{R_FOR}!$A:$E,5,0),""))'
        # colunas manuais: estilo de entrada
        for n in ("Pacote (ajuste manual)", "Categoria (ajuste manual)", "Competência (ajuste manual)"):
            c = ws[f"{NF[n]}{R}"]
            c.fill = fill(INPUT_FILL); c.font = font(False, INPUT_FONT, 9)
        # meses (vencimento no ano selecionado) - parcelas conciliadas ao documento
        for i, m in enumerate(MESES, start=1):
            comp = f'AnoSel&"-"&TEXT({i},"00")'
            ws[f"{NF[m]}{R}"] = (f'=IF($B{R}="","",SUMIFS({_whole(R_108,P,"Valor bruto")},{_whole(R_108,P,"Linha NF")},ROW(),'
                                 f'{_whole(R_108,P,"Considerar")},1,{_whole(R_108,P,"Comp. Vencimento")},{comp})'
                                 f'+SUMIFS({_whole(R_1015,A,"Valor parcela")},{_whole(R_1015,A,"Linha NF")},ROW(),'
                                 f'{_whole(R_1015,A,"Considerar")},1,{_whole(R_1015,A,"Comp. Vencimento")},{comp}))')
            ws[f"{NF[m]}{R}"].number_format = FMT_NUM
        ws[f"{NF['Total conciliado']}{R}"] = f'=IF($B{R}="","",{NF["Valor pago bruto"]}{R}+{NF["Valor em aberto"]}{R})'
        ws[f"{NF['Diferença']}{R}"] = f'=IF($B{R}="","",ROUND($I{R}-{NF["Total conciliado"]}{R},2))'
        # chaves
        ws[f"{NF['Chave N1']}{R}"] = (f'=IF($B{R}="","",TRIM($B{R}&"")&"|"&TRIM({E})&"|"&UPPER(TRIM($F{R}))&"|"&'
                                      f'IF(ISNUMBER($J{R}),{date_key_formula(f"$J{R}")},"s/data"))')
        ws[f"{NF['Chave N2']}{R}"] = f'=IF($B{R}="","",TRIM($B{R}&"")&"|"&TRIM({E})&"|"&UPPER(TRIM($F{R})))'
        pa, ca, co = NF["Pacote (ajuste manual)"], NF["Categoria (ajuste manual)"], NF["Competência (ajuste manual)"]
        ws[f"{NF['Pacote final']}{R}"] = f'=IF($B{R}="","",IF({pa}{R}<>"",{pa}{R},IF($C{R}<>"",$C{R},"{SEM_CLASS}")))'
        ws[f"{NF['Categoria final']}{R}"] = f'=IF($B{R}="","",IF({ca}{R}<>"",{ca}{R},IF($D{R}<>"",$D{R},"{SEM_CLASS}")))'
        ws[f"{NF['Origem classificação']}{R}"] = (f'=IF($B{R}="","",IF(OR({pa}{R}<>"",{ca}{R}<>""),"Documento (ajuste manual)",'
                                                  f'IF($C{R}<>"","Cadastro de fornecedor","{SEM_CLASS}")))')
        ws[f"{NF['Competência']}{R}"] = (f'=IF($B{R}="","",IF({co}{R}<>"",IF(ISNUMBER({co}{R}),{comp_formula(f"{co}{R}")},TRIM({co}{R}&"")),'
                                         f'IF(ISNUMBER($J{R}),{comp_formula(f"$J{R}")},"s/data")))')
        ws[f"{NF['Base competência']}{R}"] = f'=IF($B{R}="","",IF({co}{R}<>"","Informada (ajuste manual)","Estimada (mês de emissão)"))'
        ws[f"{NF['Registro válido']}{R}"] = f'=IF(AND($B{R}<>"",ISNUMBER($I{R}),ISNUMBER($J{R}),{E}<>""),1,0)'
        k1 = NF["Chave N1"]
        ws[f"{NF['Ocorrência']}{R}"] = f'=IF($B{R}="","",COUNTIF(${k1}$2:${k1}{R},{k1}{R}))'
        ws[f"{NF['Escopo Frota']}{R}"] = f'=IF($A{R}="Frota",1,0)'
        ws[f"{NF['Considerar DRE']}{R}"] = f'=IF(AND({NF["Registro válido"]}{R}=1,{NF["Escopo Frota"]}{R}=1,{NF["Ocorrência"]}{R}=1),1,0)'
        ws[f"{NF['Qtd parc. pagas']}{R}"] = f'=IF($B{R}="","",COUNTIFS({_whole(R_108,P,"Linha NF")},ROW(),{_whole(R_108,P,"Considerar")},1))'
        ws[f"{NF['Valor pago bruto']}{R}"] = f'=IF($B{R}="","",SUMIFS({_whole(R_108,P,"Valor bruto")},{_whole(R_108,P,"Linha NF")},ROW(),{_whole(R_108,P,"Considerar")},1))'
        ws[f"{NF['Valor pago líquido']}{R}"] = f'=IF($B{R}="","",SUMIFS({_whole(R_108,P,"Valor líquido pago")},{_whole(R_108,P,"Linha NF")},ROW(),{_whole(R_108,P,"Considerar")},1))'
        ws[f"{NF['Qtd parc. abertas']}{R}"] = f'=IF($B{R}="","",COUNTIFS({_whole(R_1015,A,"Linha NF")},ROW(),{_whole(R_1015,A,"Considerar")},1))'
        ws[f"{NF['Valor em aberto']}{R}"] = f'=IF($B{R}="","",SUMIFS({_whole(R_1015,A,"Valor parcela")},{_whole(R_1015,A,"Linha NF")},ROW(),{_whole(R_1015,A,"Considerar")},1))'
        ws[f"{NF['Vencido (R$)']}{R}"] = f'=IF($B{R}="","",SUMIFS({_whole(R_1015,A,"Valor parcela")},{_whole(R_1015,A,"Linha NF")},ROW(),{_whole(R_1015,A,"Considerar")},1,{_whole(R_1015,A,"Status")},"Vencido"))'
        ws[f"{NF['Juros']}{R}"] = f'=IF($B{R}="","",SUMIFS({_whole(R_108,P,"Juros")},{_whole(R_108,P,"Linha NF")},ROW(),{_whole(R_108,P,"Considerar")},1))'
        ws[f"{NF['Descontos']}{R}"] = f'=IF($B{R}="","",SUMIFS({_whole(R_108,P,"Desconto")},{_whole(R_108,P,"Linha NF")},ROW(),{_whole(R_108,P,"Considerar")},1))'
        ws[f"{NF['Retenções/dif. pagas']}{R}"] = f'=IF($B{R}="","",SUMIFS({_whole(R_108,P,"Retenção/diferença")},{_whole(R_108,P,"Linha NF")},ROW(),{_whole(R_108,P,"Considerar")},1))'
        qp, qa, dif, va = NF["Qtd parc. pagas"], NF["Qtd parc. abertas"], NF["Diferença"], NF["Valor em aberto"]
        ws[f"{NF['Status documento']}{R}"] = (
            f'=IF($B{R}="","",IF({NF["Registro válido"]}{R}=0,"Registro inválido",IF({NF["Ocorrência"]}{R}>1,"Duplicado (desconsiderado)",'
            f'IF($A{R}<>"Frota","Fora do escopo Frota",IF(AND({qp}{R}=0,{qa}{R}=0),"Sem correspondência nas bases",'
            f'IF(ABS({dif}{R})>0.01,"Divergência de conciliação",IF({qa}{R}=0,"Pago",IF({qp}{R}>0,"Parcialmente pago","Em aberto"))))))))')
        k2 = NF["Chave N2"]
        ws[f"{NF['Ambíguo N2']}{R}"] = f'=IF($B{R}="","",IF(COUNTIF(${k2}:${k2},{k2}{R})>1,1,0))'
        st = NF["Status documento"]
        ws[f"{NF['Alertas']}{R}"] = (
            f'=IF($B{R}="","",IF({NF["Registro válido"]}{R}=0,"Registro inválido (valor/data/fornecedor); ","")'
            f'&IF({NF["Ocorrência"]}{R}>1,"Documento duplicado; ","")'
            f'&IF($A{R}="{NAO_CAD}","Fornecedor sem cadastro; ",IF($A{R}<>"Frota","Fornecedor de outra área; ",""))'
            f'&IF(AND($A{R}="Frota",$C{R}="",{pa}{R}=""),"Fornecedor sem pacote; ","")'
            f'&IF(AND($A{R}="Frota",$D{R}="",{ca}{R}=""),"Fornecedor sem categoria; ","")'
            f'&IF({st}{R}="Sem correspondência nas bases","Sem correspondência no Rodopar 108/1015; ","")'
            f'&IF({st}{R}="Divergência de conciliação","Diferença de R$ "&TEXT({dif}{R},"#,##0.00")&" entre documento e parcelas (verificar retenções/parcelas fora da extração); ","")'
            f'&IF({NF["Ambíguo N2"]}{R}=1,"Chave Doc+Forn+Tipo repetida (conciliação N2 ambígua); ","")'
            f'&IF({NF["Vencido (R$)"]}{R}>0,"Parcela(s) vencida(s); ",""))')
        cd, ca_ = NF["Cum drill"], NF["Cum auditoria"]
        ws[f"{cd}{R}"] = (f'=N({cd}{R-1})+IF(AND({NF["Considerar DRE"]}{R}=1,{NF["Pacote final"]}{R}={R_CP}!$C$5,'
                          f'{NF["Ano comp."]}{R}={R_CP}!$C$4),1,0)')
        ws[f"{ca_}{R}"] = f'=N({ca_}{R-1})+IF(AND($B{R}<>"",{NF["Alertas"]}{R}<>""),1,0)'
        ws[f"{NF['Ano comp.']}{R}"] = f'=IF($B{R}="","",IFERROR(VALUE(LEFT({NF["Competência"]}{R},4)),""))'
        ws[f"{NF['Mês comp.']}{R}"] = f'=IF($B{R}="","",IFERROR(VALUE(RIGHT({NF["Competência"]}{R},2)),""))'
        # formatos
        for n in ("Valor Doc", "Total conciliado", "Diferença", "Valor pago bruto", "Valor pago líquido", "Valor em aberto", "Juros", "Descontos", "Retenções/dif. pagas", "Vencido (R$)"):
            ws[f"{NF[n]}{R}"].number_format = FMT_NUM
        for n in ("Emissão", "Referência"):
            ws[f"{NF[n]}{R}"].number_format = FMT_DATE
        for n in base_names:
            c = ws[f"{NF[n]}{R}"]
            if c.font.name != FONT or c.font.color is None or (c.font.color.rgb or "") != f"FF{INPUT_FONT}":
                c.font = font(False, GRAPHITE, 9)
    # larguras
    widths = {"A": 9, "B": 16, "C": 16, "D": 15, "E": 34, "F": 6, "G": 18, "H": 6, "I": 14, "J": 11, "K": 11, "L": 16, "M": 5, "N": 7, "O": 16,
              "P": 16, "Q": 16, "R": 14}
    for m in MESES:
        widths[NF[m]] = 11
    for n in base_names[30:]:
        widths[NF[n]] = 14
    widths[NF["Alertas"]] = 50
    widths[NF["Chave N1"]] = 30
    widths[NF["Chave N2"]] = 26
    set_widths(ws, widths)
    ws.freeze_panes = "F2"
    make_table(ws, "tblNF", NF["Vencido (R$)"], last, "TableStyleLight1")
    cf_text(ws, f"{NF['Status documento']}2:{NF['Status documento']}{last}", "Divergência", RED_L, RED)
    cf_text(ws, f"{NF['Status documento']}2:{NF['Status documento']}{last}", "Pago", GREEN_L, GREEN, contains=False)
    cf_text(ws, f"{NF['Alertas']}2:{NF['Alertas']}{last}", ";", AMBER_L, AMBER)
    ws.sheet_view.zoomScale = 90


# ====================================================================== FORNECEDORES
def build_fornecedores(ws, n_rows):
    last = n_rows
    names = list(F.keys())
    _hdr(ws, F, names[:5], NAVY)
    _hdr(ws, F, names[5:], GRAPHITE)
    ws.row_dimensions[1].height = 42
    ws["C1"].comment = Comment("Area = 'Frota' para fornecedores do setor; '-' para fornecedores de outras áreas (ignorados no resultado da Frota).", "Controle Financeiro")
    ano = f"{R_AF}!$C$4"
    for r in range(2, last + 1):
        R = r
        ws[f"B{R}"] = f'=IF($A{R}="","",IFERROR(VLOOKUP($A{R},{R_NF}!$E:$G,3,0),""))'
        for n in ("Fornecedor", "CNPJ", "Area", "Pacote", "Categoria"):
            c = ws[f"{F[n]}{R}"]
            c.font = font(False, INPUT_FONT if n != "CNPJ" else GRAPHITE, 9)
            if n in ("Area", "Pacote", "Categoria"):
                c.fill = fill(INPUT_FILL)
        nfE, nfI, nfAQ, nfAno, nfComp = _whole(R_NF, NF, "Razão Social"), _whole(R_NF, NF, "Valor Doc"), _whole(R_NF, NF, "Considerar DRE"), _whole(R_NF, NF, "Ano comp."), _whole(R_NF, NF, "Competência")
        ws[f"{F['Qtd docs (ano)']}{R}"] = f'=IF($A{R}="","",COUNTIFS({nfE},$A{R},{nfAQ},1,{nfAno},{ano}))'
        ws[f"{F['Custo competência (ano)']}{R}"] = f'=IF($A{R}="","",SUMIFS({nfI},{nfE},$A{R},{nfAQ},1,{nfAno},{ano}))'
        ws[f"{F['Participação %']}{R}"] = f'=IF($A{R}="","",IF({R_AF}!$C$7>0,{F["Custo competência (ano)"]}{R}/{R_AF}!$C$7,""))'
        ws[f"{F['Pago líquido (ano, caixa)']}{R}"] = (f'=IF($A{R}="","",SUMIFS({_whole(R_108,P,"Valor líquido pago")},{_whole(R_108,P,"Fornecedor (norm.)")},$A{R},'
                                                      f'{_whole(R_108,P,"Considerar")},1,{_whole(R_108,P,"Comp. Caixa")},{ano}&"-*"))')
        ws[f"{F['Em aberto (total)']}{R}"] = f'=IF($A{R}="","",SUMIFS({_whole(R_1015,A,"Valor parcela")},{_whole(R_1015,A,"Fornecedor (norm.)")},$A{R},{_whole(R_1015,A,"Considerar")},1))'
        ws[f"{F['Vencido (total)']}{R}"] = f'=IF($A{R}="","",SUMIFS({_whole(R_1015,A,"Valor parcela")},{_whole(R_1015,A,"Fornecedor (norm.)")},$A{R},{_whole(R_1015,A,"Considerar")},1,{_whole(R_1015,A,"Status")},"Vencido"))'
        ws[f"{F['Qtd parcelas abertas']}{R}"] = f'=IF($A{R}="","",COUNTIFS({_whole(R_1015,A,"Fornecedor (norm.)")},$A{R},{_whole(R_1015,A,"Considerar")},1))'
        ws[f"{F['Movimento sem NF (revisar)']}{R}"] = (f'=IF($A{R}="","",SUMIFS({_whole(R_108,P,"Valor bruto")},{_whole(R_108,P,"Fornecedor (norm.)")},$A{R},{_whole(R_108,P,"Escopo")},"{SCOPE_SEMNF}")'
                                                       f'+SUMIFS({_whole(R_1015,A,"Valor parcela")},{_whole(R_1015,A,"Fornecedor (norm.)")},$A{R},{_whole(R_1015,A,"Escopo")},"{SCOPE_SEMNF}"))')
        ws[f"{F['Alertas cadastro']}{R}"] = (f'=IF($A{R}="","",IF(COUNTIF($A:$A,$A{R})>1,"Fornecedor duplicado no cadastro; ","")'
                                             f'&IF(AND($C{R}="Frota",$D{R}=""),"Frota sem pacote; ","")&IF(AND($C{R}="Frota",$E{R}=""),"Frota sem categoria; ","")'
                                             f'&IF(AND($C{R}<>"Frota",$C{R}<>"-"),"Área não reconhecida (use Frota ou -); ","")'
                                             f'&IF(AND($C{R}="Frota",$B{R}=""),"Sem CNPJ localizado na relação de NF; ","")'
                                             f'&IF({F["Movimento sem NF (revisar)"]}{R}>0,"Movimentos no Rodopar sem NF na relação: R$ "&TEXT({F["Movimento sem NF (revisar)"]}{R},"#,##0.00")&"; ",""))')
        g = F["Custo competência (ano)"]
        ws[f"{F['Rank custo']}{R}"] = f'=IF($A{R}="","",IF(N({g}{R})>0,RANK({g}{R},${g}:${g})+COUNTIF(${g}$2:{g}{R},{g}{R})-1,""))'
        j = F["Em aberto (total)"]
        ws[f"{F['Rank em aberto']}{R}"] = f'=IF($A{R}="","",IF(N({j}{R})>0,RANK({j}{R},${j}:${j})+COUNTIF(${j}$2:{j}{R},{j}{R})-1,""))'
        for i, m in enumerate(MESES_ABREV, start=1):
            ws[f"{F[m]}{R}"] = f'=IF($A{R}="","",SUMIFS({nfI},{nfE},$A{R},{nfAQ},1,{nfComp},{ano}&"-"&TEXT({i},"00")))'
            ws[f"{F[m]}{R}"].number_format = FMT_NUM
        ws[f"{F['Meses com custo']}{R}"] = f'=IF($A{R}="","",COUNTIF({F["Jan"]}{R}:{F["Dez"]}{R},">0"))'
        ca_ = F["Cum alertas"]
        ws[f"{ca_}{R}"] = f'=N({ca_}{R-1})+IF(AND($A{R}<>"",{F["Alertas cadastro"]}{R}<>""),1,0)'
        for n in ("Custo competência (ano)", "Pago líquido (ano, caixa)", "Em aberto (total)", "Vencido (total)", "Movimento sem NF (revisar)"):
            ws[f"{F[n]}{R}"].number_format = FMT_NUM
        ws[f"{F['Participação %']}{R}"].number_format = FMT_PCT
        for n in names[5:]:
            ws[f"{F[n]}{R}"].font = font(False, GRAPHITE, 9)
    widths = {"A": 44, "B": 20, "C": 8, "D": 18, "E": 18}
    for n in names[5:]:
        widths[F[n]] = 13
    widths[F["Alertas cadastro"]] = 48
    set_widths(ws, widths)
    ws.freeze_panes = "B2"
    make_table(ws, "tblFornecedores", F["Cum alertas"], last, "TableStyleLight1")
    cf_text(ws, f"{F['Alertas cadastro']}2:{F['Alertas cadastro']}{last}", ";", AMBER_L, AMBER)


# ====================================================================== RODOPAR 108
def build_108(ws, n_rows):
    last = n_rows
    names = list(P.keys())
    _hdr(ws, P, names[:2], NAVY2)
    _hdr(ws, P, names[2:13], NAVY)
    _hdr(ws, P, names[13:], GRAPHITE)
    ws.row_dimensions[1].height = 42
    ws["C1"].comment = Comment("ÁREA DE COLAGEM: cole aqui (C2) o relatório Rodopar 108 na ORDEM ORIGINAL (sem classificar), "
                               "incluindo as linhas 'Doc. Pago em : dd/mm/aaaa'. Elas fornecem a data efetiva de pagamento.", "Controle Financeiro")
    ws[f"{P['Data Pgto (derivada)']}1"].comment = Comment("Data do grupo 'Doc. Pago em :' imediatamente acima da linha. Só é utilizada quando a extração "
                                                            "está na ordem original (ver Parâmetros > Disponibilidade da data de pagamento).", "Controle Financeiro")
    nfK1 = _whole(R_NF, NF, "Chave N1")
    nfK2 = _whole(R_NF, NF, "Chave N2")
    for r in range(2, last + 1):
        R = r
        ws[f"A{R}"] = f'=IF($C{R}="","",$C{R}&$E{R}&$G{R}&$F{R})'
        ws[f"B{R}"] = f'=IF(ISNUMBER($H{R}),INDEX(Meses,MONTH($H{R})),"")'
        v = P["Registro válido"]
        ws[f"{v}{R}"] = f'=IF(AND($C{R}<>"",$C{R}<>"Documento",LEFT($C{R},12)<>"Doc. Pago em",ISNUMBER($J{R}),ISNUMBER($H{R})),1,0)'
        dp = P["Data Pgto (derivada)"]
        ws[f"{dp}{R}"] = (f'=IF(LEFT($C{R},12)="Doc. Pago em",IFERROR(IF(ISNUMBER($D{R}),INT($D{R}),'
                          f'DATE(VALUE(MID($D{R},7,4)),VALUE(MID($D{R},4,2)),VALUE(LEFT($D{R},2)))),""),IF(ISNUMBER({dp}{R-1}),{dp}{R-1},""))')
        ws[f"{P['Documento (norm.)']}{R}"] = f'=IF({v}{R}=1,TRIM($C{R}&""),"")'
        ws[f"{P['Fornecedor (norm.)']}{R}"] = f'=IF({v}{R}=1,TRIM($E{R}),"")'
        ws[f"{P['Tipo (norm.)']}{R}"] = f'=IF({v}{R}=1,UPPER(TRIM($G{R})),"")'
        ws[f"{P['Emissão (data)']}{R}"] = f'=IF(AND({v}{R}=1,ISNUMBER($F{R})),INT($F{R}),"")'
        ws[f"{P['Vencimento (data)']}{R}"] = f'=IF({v}{R}=1,INT($H{R}),"")'
        d, fo, t, em, ve = P["Documento (norm.)"], P["Fornecedor (norm.)"], P["Tipo (norm.)"], P["Emissão (data)"], P["Vencimento (data)"]
        ws[f"{P['Chave N1']}{R}"] = f'=IF({v}{R}=1,{d}{R}&"|"&{fo}{R}&"|"&{t}{R}&"|"&IF(ISNUMBER({em}{R}),{date_key_formula(f"{em}{R}")},"s/data"),"")'
        ws[f"{P['Chave N2']}{R}"] = f'=IF({v}{R}=1,{d}{R}&"|"&{fo}{R}&"|"&{t}{R},"")'
        k2 = P["Chave N2"]
        ws[f"{P['Ocorrência']}{R}"] = f'=IF({v}{R}=1,COUNTIFS(${k2}$2:${k2}{R},{k2}{R},${ve}$2:${ve}{R},{ve}{R},$J$2:$J{R},$J{R}),"")'
        k1 = P["Chave N1"]
        ws[f"{P['Linha NF (N1)']}{R}"] = f'=IF({v}{R}=1,IFERROR(MATCH({k1}{R},{nfK1},0),""),"")'
        ws[f"{P['Linha NF (N2)']}{R}"] = f'=IF(AND({v}{R}=1,{P["Linha NF (N1)"]}{R}=""),IFERROR(MATCH({k2}{R},{nfK2},0),""),"")'
        n1, n2 = P["Linha NF (N1)"], P["Linha NF (N2)"]
        ws[f"{P['Nível conciliação']}{R}"] = (f'=IF({v}{R}=0,"",IF({n1}{R}<>"","N1 - Doc+Forn+Tipo+Emissão",IF({n2}{R}<>"",'
                                              f'IF(COUNTIF({nfK2},{k2}{R})>1,"N2 ambíguo","N2 - Doc+Forn+Tipo"),"Sem NF")))')
        ws[f"{P['Linha NF']}{R}"] = f'=IF({n1}{R}<>"",{n1}{R},IF({n2}{R}<>"",{n2}{R},""))'
        ws[f"{P['Área fornecedor']}{R}"] = f'=IF({v}{R}=1,IFERROR(VLOOKUP({fo}{R},{R_FOR}!$A:$C,3,0),"{NAO_CAD}"),"")'
        ln, ar = P["Linha NF"], P["Área fornecedor"]
        ws[f"{P['Escopo']}{R}"] = (f'=IF({v}{R}=0,"",IF({ln}{R}<>"","{SCOPE_NF}",IF({ar}{R}="Frota","{SCOPE_SEMNF}",'
                                   f'IF({ar}{R}="{NAO_CAD}","{SCOPE_NC}","{SCOPE_OUTRA}"))))')
        es, oc = P["Escopo"], P["Ocorrência"]
        ws[f"{P['Considerar']}{R}"] = f'=IF(AND({v}{R}=1,{es}{R}="{SCOPE_NF}",{oc}{R}=1),1,0)'
        ws[f"{P['Pacote']}{R}"] = (f'=IF({v}{R}=0,"",IF({ln}{R}<>"",INDEX({_whole(R_NF,NF,"Pacote final")},{ln}{R}),'
                                   f'IF({ar}{R}="Frota",IF(IFERROR(VLOOKUP({fo}{R},{R_FOR}!$A:$E,4,0),"")="","{SEM_CLASS}",VLOOKUP({fo}{R},{R_FOR}!$A:$E,4,0)),"")))')
        ws[f"{P['Categoria']}{R}"] = (f'=IF({v}{R}=0,"",IF({ln}{R}<>"",INDEX({_whole(R_NF,NF,"Categoria final")},{ln}{R}),'
                                      f'IF({ar}{R}="Frota",IF(IFERROR(VLOOKUP({fo}{R},{R_FOR}!$A:$E,5,0),"")="","{SEM_CLASS}",VLOOKUP({fo}{R},{R_FOR}!$A:$E,5,0)),"")))')
        ws[f"{P['Competência']}{R}"] = (f'=IF({v}{R}=0,"",IF({ln}{R}<>"",INDEX({_whole(R_NF,NF,"Competência")},{ln}{R}),'
                                        f'IF(ISNUMBER({em}{R}),{comp_formula(f"{em}{R}")},"")))')
        ws[f"{P['Comp. Vencimento']}{R}"] = f'=IF(ISNUMBER({ve}{R}),{comp_formula(f"{ve}{R}")},"")'
        ws[f"{P['Data caixa validada']}{R}"] = f'=IF(AND({v}{R}=1,FlagDataPgto=1,ISNUMBER({dp}{R})),{dp}{R},"")'
        dc = P["Data caixa validada"]
        ws[f"{P['Comp. Caixa']}{R}"] = f'=IF({v}{R}=0,"",IF(ISNUMBER({dc}{R}),{comp_formula(f"{dc}{R}")},{P["Comp. Vencimento"]}{R}))'
        ws[f"{P['Base data caixa']}{R}"] = f'=IF({v}{R}=0,"",IF(ISNUMBER({dc}{R}),"Data de pagamento","Mês de vencimento (estimado)"))'
        ws[f"{P['Valor bruto']}{R}"] = f'=IF({v}{R}=1,$J{R},"")'
        ws[f"{P['Desconto']}{R}"] = f'=IF({v}{R}=1,N($K{R}),"")'
        ws[f"{P['Juros']}{R}"] = f'=IF({v}{R}=1,N($L{R}),"")'
        ws[f"{P['Valor líquido pago']}{R}"] = f'=IF({v}{R}=1,IF(ISNUMBER($M{R}),$M{R},$J{R}),"")'
        vb, de, ju, vl = P["Valor bruto"], P["Desconto"], P["Juros"], P["Valor líquido pago"]
        ws[f"{P['Retenção/diferença']}{R}"] = f'=IF({v}{R}=1,ROUND({vb}{R}-{de}{R}+{ju}{R}-{vl}{R},2),"")'
        ws[f"{P['Status']}{R}"] = (f'=IF({v}{R}=0,"",IF({oc}{R}>1,"Duplicado (desconsiderado)",IF({es}{R}="{SCOPE_NF}","Pago",'
                                   f'IF({es}{R}="{SCOPE_SEMNF}","Pago - pendente de classificação","Fora do escopo"))))')
        nv, rt, pc = P["Nível conciliação"], P["Retenção/diferença"], P["Pacote"]
        ws[f"{P['Alertas']}{R}"] = (
            f'=IF(OR({v}{R}=0,{es}{R}="{SCOPE_OUTRA}",{es}{R}="{SCOPE_NC}"),"",'
            f'IF({oc}{R}>1,"Parcela duplicada na extração; ","")'
            f'&IF({es}{R}="{SCOPE_SEMNF}","Fornecedor Frota sem NF na relação (revisar escopo/competência); ","")'
            f'&IF({nv}{R}="N2 ambíguo","Correspondência ambígua com a relação de NF; ","")'
            f'&IF({nv}{R}="N2 - Doc+Forn+Tipo","Emissão divergente da relação de NF; ","")'
            f'&IF(NOT(ISNUMBER({em}{R})),"Emissão inválida; ","")'
            f'&IF({rt}{R}<>0,"Retenção/diferença de R$ "&TEXT({rt}{R},"#,##0.00")&" entre parcela e valor pago; ","")'
            f'&IF({pc}{R}="{SEM_CLASS}","Sem pacote; ","")'
            f'&IF(COUNTIFS({_whole(R_1015,A,"Chave N2")},{k2}{R},{_whole(R_1015,A,"Vencimento (data)")},{ve}{R},{_whole(R_1015,A,"Valor parcela")},$J{R})>0,"Parcela consta também em aberto (1015); ",""))')
        cv, ca_ = P["Cum vista"], P["Cum auditoria"]
        st = P["Status"]
        flt = f"{R_PO}!$C$5"
        ws[f"{cv}{R}"] = (f'=N({cv}{R-1})+IF(AND({v}{R}=1,OR({es}{R}="{SCOPE_NF}",{es}{R}="{SCOPE_SEMNF}"),'
                          f'OR({flt}="(Todos)",{flt}={st}{R},AND({flt}="Pendente de classificação",{es}{R}="{SCOPE_SEMNF}"),'
                          f'AND({flt}="Duplicado",LEFT({st}{R},9)="Duplicado"))),1,0)')
        ws[f"{ca_}{R}"] = f'=N({ca_}{R-1})+IF(AND({v}{R}=1,{P["Alertas"]}{R}<>""),1,0)'
        ws[f"{P['Ano comp.']}{R}"] = f'=IF({P["Competência"]}{R}="","",IFERROR(VALUE(LEFT({P["Competência"]}{R},4)),""))'
        ws[f"{P['Ano venc.']}{R}"] = f'=IF(ISNUMBER({ve}{R}),YEAR({ve}{R}),"")'
        ws[f"{P['Ano caixa']}{R}"] = f'=IF({P["Comp. Caixa"]}{R}="","",VALUE(LEFT({P["Comp. Caixa"]}{R},4)))'
        for n in ("Vlr. Parc.", "Vlr. Desco.", "Vlr. Juros", "Vlr. Total", "Valor bruto", "Desconto", "Juros", "Valor líquido pago", "Retenção/diferença"):
            ws[f"{P[n]}{R}"].number_format = FMT_NUM
        for n in ("Emissão", "Vencto", "Data Ref.", "Data Pgto (derivada)", "Emissão (data)", "Vencimento (data)", "Data caixa validada"):
            ws[f"{P[n]}{R}"].number_format = FMT_DATE
        for n in names:
            ws[f"{P[n]}{R}"].font = font(False, GRAPHITE, 9)
    widths = {"A": 10, "B": 10, "C": 16, "D": 11, "E": 36, "F": 11, "G": 6, "H": 11, "I": 11, "J": 12, "K": 10, "L": 10, "M": 12}
    for n in names[13:]:
        widths[P[n]] = 13
    widths[P["Chave N1"]] = 34; widths[P["Chave N2"]] = 28; widths[P["Alertas"]] = 50; widths[P["Escopo"]] = 22; widths[P["Nível conciliação"]] = 22
    set_widths(ws, widths)
    ws.freeze_panes = "F2"
    make_table(ws, "tbl108", P["Ano caixa"], last, "TableStyleLight1")
    cf_text(ws, f"{P['Alertas']}2:{P['Alertas']}{last}", ";", AMBER_L, AMBER)
    cf_text(ws, f"{P['Escopo']}2:{P['Escopo']}{last}", SCOPE_NF, GREEN_L, GREEN, contains=False)
    cf_text(ws, f"{P['Escopo']}2:{P['Escopo']}{last}", SCOPE_SEMNF, AMBER_L, AMBER, contains=False)
    ws.sheet_view.zoomScale = 90


# ====================================================================== RODOPAR 1015
def build_1015(ws, n_rows):
    last = n_rows
    names = list(A.keys())
    _hdr(ws, A, names[:2], NAVY2)
    _hdr(ws, A, names[2:11], NAVY)
    _hdr(ws, A, names[11:], GRAPHITE)
    ws.row_dimensions[1].height = 42
    ws["C1"].comment = Comment("ÁREA DE COLAGEM: cole aqui (C2) o relatório Rodopar 1015 completo (todas as parcelas em aberto, incluindo vencimentos futuros de todos os meses).", "Controle Financeiro")
    ws["K1"].comment = Comment("Campo 'Acumulado' do Rodopar: controle auxiliar da extração. NÃO é somado como obrigação.", "Controle Financeiro")
    nfK2 = _whole(R_NF, NF, "Chave N2")
    for r in range(2, last + 1):
        R = r
        ws[f"A{R}"] = f'=IF($D{R}="","",$D{R}&$F{R}&$H{R})'
        ws[f"B{R}"] = f'=IF(ISNUMBER($G{R}),INDEX(Meses,MONTH($G{R})),"")'
        v = A["Registro válido"]
        ws[f"{v}{R}"] = f'=IF(AND($D{R}<>"",$D{R}<>"Doc.",ISNUMBER($J{R}),ISNUMBER($G{R})),1,0)'
        ws[f"{A['Documento (norm.)']}{R}"] = f'=IF({v}{R}=1,TRIM($D{R}&""),"")'
        ws[f"{A['Fornecedor (norm.)']}{R}"] = f'=IF({v}{R}=1,TRIM($F{R}),"")'
        ws[f"{A['Tipo (norm.)']}{R}"] = f'=IF({v}{R}=1,UPPER(TRIM($H{R})),"")'
        ws[f"{A['Vencimento (data)']}{R}"] = f'=IF({v}{R}=1,INT($G{R}),"")'
        d, fo, t, ve = A["Documento (norm.)"], A["Fornecedor (norm.)"], A["Tipo (norm.)"], A["Vencimento (data)"]
        ws[f"{A['Chave N2']}{R}"] = f'=IF({v}{R}=1,{d}{R}&"|"&{fo}{R}&"|"&{t}{R},"")'
        k2 = A["Chave N2"]
        ws[f"{A['Ocorrência']}{R}"] = f'=IF({v}{R}=1,COUNTIFS(${k2}$2:${k2}{R},{k2}{R},${ve}$2:${ve}{R},{ve}{R},$J$2:$J{R},$J{R},$I$2:$I{R},$I{R}),"")'
        ws[f"{A['Linha NF']}{R}"] = f'=IF({v}{R}=1,IFERROR(MATCH({k2}{R},{nfK2},0),""),"")'
        ln = A["Linha NF"]
        ws[f"{A['Nível conciliação']}{R}"] = f'=IF({v}{R}=0,"",IF({ln}{R}<>"",IF(COUNTIF({nfK2},{k2}{R})>1,"N2 ambíguo","N2 - Doc+Forn+Tipo"),"Sem NF"))'
        ws[f"{A['Área fornecedor']}{R}"] = f'=IF({v}{R}=1,IFERROR(VLOOKUP({fo}{R},{R_FOR}!$A:$C,3,0),"{NAO_CAD}"),"")'
        ar = A["Área fornecedor"]
        ws[f"{A['Escopo']}{R}"] = (f'=IF({v}{R}=0,"",IF({ln}{R}<>"","{SCOPE_NF}",IF({ar}{R}="Frota","{SCOPE_SEMNF}",'
                                   f'IF({ar}{R}="{NAO_CAD}","{SCOPE_NC}","{SCOPE_OUTRA}"))))')
        es, oc = A["Escopo"], A["Ocorrência"]
        ws[f"{A['Considerar']}{R}"] = f'=IF(AND({v}{R}=1,{es}{R}="{SCOPE_NF}",{oc}{R}=1),1,0)'
        ws[f"{A['Pacote']}{R}"] = (f'=IF({v}{R}=0,"",IF({ln}{R}<>"",INDEX({_whole(R_NF,NF,"Pacote final")},{ln}{R}),'
                                   f'IF({ar}{R}="Frota",IF(IFERROR(VLOOKUP({fo}{R},{R_FOR}!$A:$E,4,0),"")="","{SEM_CLASS}",VLOOKUP({fo}{R},{R_FOR}!$A:$E,4,0)),"")))')
        ws[f"{A['Categoria']}{R}"] = (f'=IF({v}{R}=0,"",IF({ln}{R}<>"",INDEX({_whole(R_NF,NF,"Categoria final")},{ln}{R}),'
                                      f'IF({ar}{R}="Frota",IF(IFERROR(VLOOKUP({fo}{R},{R_FOR}!$A:$E,5,0),"")="","{SEM_CLASS}",VLOOKUP({fo}{R},{R_FOR}!$A:$E,5,0)),"")))')
        ws[f"{A['Competência']}{R}"] = f'=IF({v}{R}=0,"",IF({ln}{R}<>"",INDEX({_whole(R_NF,NF,"Competência")},{ln}{R}),""))'
        ws[f"{A['Comp. Vencimento']}{R}"] = f'=IF(ISNUMBER({ve}{R}),{comp_formula(f"{ve}{R}")},"")'
        ws[f"{A['Valor parcela']}{R}"] = f'=IF({v}{R}=1,$J{R},"")'
        ws[f"{A['Dias para vencer']}{R}"] = f'=IF({v}{R}=1,{ve}{R}-DataRef,"")'
        dv = A["Dias para vencer"]
        ws[f"{A['Status']}{R}"] = (f'=IF({v}{R}=0,"",IF({oc}{R}>1,"Duplicado (desconsiderado)",'
                                   f'IF(COUNTIFS({_whole(R_108,P,"Chave N2")},{k2}{R},{_whole(R_108,P,"Vencimento (data)")},{ve}{R},{_whole(R_108,P,"Valor bruto")},$J{R})>0,'
                                   f'"Divergência (consta como pago no 108)",IF({dv}{R}<0,"Vencido","A vencer"))))')
        st = A["Status"]
        ws[f"{A['Dias em atraso']}{R}"] = f'=IF({st}{R}="Vencido",-{dv}{R},IF({v}{R}=1,0,""))'
        ws[f"{A['Faixa de vencimento']}{R}"] = (f'=IF({v}{R}=0,"",IF({dv}{R}<0,"Vencido",IF({dv}{R}<=7,"Até 7 dias",IF({dv}{R}<=15,"8 a 15 dias",'
                                                f'IF({dv}{R}<=30,"16 a 30 dias",IF({dv}{R}<=60,"31 a 60 dias","Acima de 60 dias"))))))')
        nv, pc = A["Nível conciliação"], A["Pacote"]
        ws[f"{A['Alertas']}{R}"] = (
            f'=IF(OR({v}{R}=0,{es}{R}="{SCOPE_OUTRA}",{es}{R}="{SCOPE_NC}"),"",'
            f'IF({oc}{R}>1,"Parcela duplicada na extração; ","")'
            f'&IF({es}{R}="{SCOPE_SEMNF}","Fornecedor Frota sem NF na relação (revisar escopo/competência); ","")'
            f'&IF({nv}{R}="N2 ambíguo","Correspondência ambígua com a relação de NF; ","")'
            f'&IF(LEFT({st}{R},11)="Divergência","Consta como pago no Rodopar 108 e em aberto no 1015; ","")'
            f'&IF({st}{R}="Vencido","Parcela vencida há "&{A["Dias em atraso"]}{R}&" dias; ","")'
            f'&IF({pc}{R}="{SEM_CLASS}","Sem pacote; ",""))')
        cv, ca_ = A["Cum vista"], A["Cum auditoria"]
        flt = f"{R_PO}!$C$5"
        ws[f"{cv}{R}"] = (f'=N({cv}{R-1})+IF(AND({v}{R}=1,OR({es}{R}="{SCOPE_NF}",{es}{R}="{SCOPE_SEMNF}"),'
                          f'OR({flt}="(Todos)",{flt}={st}{R},AND({flt}="Pendente de classificação",{es}{R}="{SCOPE_SEMNF}"),'
                          f'AND({flt}="Duplicado",LEFT({st}{R},9)="Duplicado"),AND({flt}="Divergência",LEFT({st}{R},11)="Divergência"))),1,0)')
        ws[f"{ca_}{R}"] = f'=N({ca_}{R-1})+IF(AND({v}{R}=1,{A["Alertas"]}{R}<>""),1,0)'
        ws[f"{A['Ano comp.']}{R}"] = f'=IF({A["Competência"]}{R}="","",IFERROR(VALUE(LEFT({A["Competência"]}{R},4)),""))'
        ws[f"{A['Ano venc.']}{R}"] = f'=IF(ISNUMBER({ve}{R}),YEAR({ve}{R}),"")'
        for n in ("Vlr. Parcela", "Acumulado", "Valor parcela"):
            ws[f"{A[n]}{R}"].number_format = FMT_NUM
        for n in ("Vencimento", "Vencimento (data)"):
            ws[f"{A[n]}{R}"].number_format = FMT_DATE
        for n in names:
            ws[f"{A[n]}{R}"].font = font(False, GRAPHITE, 9)
    widths = {"A": 10, "B": 10, "C": 7, "D": 14, "E": 7, "F": 36, "G": 11, "H": 6, "I": 8, "J": 12, "K": 12}
    for n in names[11:]:
        widths[A[n]] = 13
    widths[A["Chave N2"]] = 28; widths[A["Alertas"]] = 50; widths[A["Escopo"]] = 22; widths[A["Nível conciliação"]] = 22; widths[A["Status"]] = 22
    set_widths(ws, widths)
    ws.freeze_panes = "G2"
    make_table(ws, "tbl1015", A["Ano venc."], last, "TableStyleLight1")
    cf_text(ws, f"{A['Status']}2:{A['Status']}{last}", "Vencido", RED_L, RED, contains=False)
    cf_text(ws, f"{A['Status']}2:{A['Status']}{last}", "Divergência", AMBER_L, AMBER)
    cf_text(ws, f"{A['Alertas']}2:{A['Alertas']}{last}", ";", AMBER_L, AMBER)
    ws.sheet_view.zoomScale = 90
