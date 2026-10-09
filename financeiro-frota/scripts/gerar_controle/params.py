# -*- coding: utf-8 -*-
"""Parâmetros, Remuneração (entrada manual) e Dados Gráficos (base dos gráficos do Painel)."""
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from common import *
from sources import NF, P, A, F, MESES, MESES_ABREV, _whole, SCOPE_NF, SCOPE_SEMNF, SCOPE_OUTRA, SCOPE_NC, SEM_CLASS, NAO_CAD

PACOTES_INICIAIS = ["Manutenção", "Pneus", "OPE Merchandising", "Sinistro", "Lavagem", "Pedagios", "MTSR", "Telemetria", "Documentação", "Segurança"]
CATEGORIAS_INICIAIS = ["Vans", "Peças E Serviços", "Tecnologia", "Locação", "Frota Leve OPE", "Lavagem", "Caminhão", "Serviços Gerais",
                       "Lanternagem", "Ressolagem", "Pedagios", "Guincho", "Telemetria", "Parabrisa", "Compra", "Contrato", "Tacografos",
                       "Vistoria Inmetro", "Extintores"]
N_PAC = 15
N_CAT = 25
N_FOR_LIST = 400
STATUS_FILTRO = ["(Todos)", "Pago", "A vencer", "Vencido", "Divergência", "Duplicado", "Pendente de classificação"]


def add_name(wb, name, ref):
    dn = DefinedName(name, attr_text=ref)
    wb.defined_names[name] = dn


# ====================================================================== PARÂMETROS
def build_parametros(wb, ws):
    set_widths(ws, {"A": 2, "B": 46, "C": 34, "D": 3, "E": 14, "F": 8, "G": 3, "H": 24, "I": 3, "J": 22, "K": 3, "L": 8, "M": 3,
                    "N": 20, "O": 3, "P": 26, "Q": 3, "R": 24, "S": 3, "T": 24, "U": 3, "V": 44, "W": 3, "X": 14, "Y": 3, "Z": 8})
    title_block(ws, "PARÂMETROS E LISTAS DO MODELO", "Seleções ativas, limites de destaque e listas utilizadas nas validações (edite apenas as células amarelas)", "B", "V")
    section(ws, 4, "Seleção ativa (espelho dos filtros do Painel Executivo)", "B", "C")
    rows = [
        (5, "Ano selecionado (Painel)", "=AnoSel", "0", "AnoSelPar"),
        (6, "Mês selecionado (nº)", "=IFERROR(MATCH(MesNome,Meses,0),MONTH(TODAY()))", "0", "MesSel"),
        (7, "Competência selecionada (AAAA-MM)", '=AnoSel&"-"&TEXT(MesSel,"00")', "@", "CompSel"),
        (8, "Competência anterior (AAAA-MM)", '=IF(MesSel=1,(AnoSel-1)&"-12",AnoSel&"-"&TEXT(MesSel-1,"00"))', "@", "CompAnt"),
        (9, "Data de referência (sobrescrever, opcional)", None, FMT_DATE, "DataRefManual"),
        (10, "Data de referência efetiva (vencido x a vencer)", "=IF(ISNUMBER(DataRefManual),DataRefManual,TODAY())", FMT_DATE, "DataRef"),
        (11, "Critério de pacote aplicado", '=IF(OR(FiltroPacote="(Todos)",FiltroPacote=""),"*",FiltroPacote)', "@", "CritPacote"),
        (12, "Critério de categoria aplicado", '=IF(OR(FiltroCategoria="(Todos)",FiltroCategoria=""),"*",FiltroCategoria)', "@", "CritCategoria"),
        (13, "Critério de fornecedor aplicado", '=IF(OR(FiltroFornecedor="(Todos)",FiltroFornecedor=""),"*",FiltroFornecedor)', "@", "CritFornecedor"),
        (14, "Filtro de custo ativo (1 = sim)", '=IF(OR(CritPacote<>"*",CritCategoria<>"*",CritFornecedor<>"*"),1,0)', "0", "FiltroAtivo"),
        (15, "Visão acumulada no ano (1 = sim)", '=IF(Visao="Acumulado no ano",1,0)', "0", "VisaoAcum"),
    ]
    for r, lab, f, fmt, nm in rows:
        label(ws, f"B{r}", lab)
        if f is None:
            input_cell(ws, f"C{r}", None, fmt, "Deixe em branco para usar a data de hoje (TODAY).")
        else:
            write(ws, f"C{r}", f, font(True, NAVY, 10), fill(GRAY1), ALIGN_C, fmt, BORDER_ALL)
        add_name(wb, nm, f"{R_PAR}!$C${r}")

    section(ws, 17, "Disponibilidade da data efetiva de pagamento (Rodopar 108)", "B", "C")
    c108 = f"{R_108}!$C:$C"
    rows = [
        (18, "Grupos 'Doc. Pago em :' encontrados na extração", f'=COUNTIF({c108},"Doc. Pago em*")', "0", "QtdGruposPgto"),
        (19, "Grupos seguidos do cabeçalho (ordem original preservada)", f'=SUMPRODUCT((LEFT({R_108}!$C$2:$C$40000,12)="Doc. Pago em")*({R_108}!$C$3:$C$40001="Documento"))', "0", "QtdGruposOk"),
        (20, "Data de pagamento validada (1 = sim)", "=IF(AND(QtdGruposPgto>0,QtdGruposPgto=QtdGruposOk),1,0)", "0", "FlagDataPgto"),
        (21, "Base utilizada na visão de caixa realizado", '=IF(FlagDataPgto=1,"Data efetiva de pagamento (grupo Doc. Pago em)","Mês de VENCIMENTO dos títulos pagos (data de pagamento não disponível na extração)")', "@", "BaseCaixaTxt"),
        (22, "Parcelas pagas com data efetiva atribuída", f'=COUNTIFS({_whole(R_108,P,"Considerar")},1,{_whole(R_108,P,"Base data caixa")},"Data de pagamento")', "0", "QtdPagasComData"),
    ]
    for r, lab, f, fmt, nm in rows:
        label(ws, f"B{r}", lab)
        write(ws, f"C{r}", f, font(True, NAVY, 10), fill(GRAY1), Alignment(horizontal="center", vertical="center", wrap_text=True), fmt, BORDER_ALL)
        add_name(wb, nm, f"{R_PAR}!$C${r}")
    ws.row_dimensions[21].height = 40

    section(ws, 24, "Limites para destaques automáticos", "B", "C")
    rows = [(25, "Concentração por fornecedor (alerta acima de)", 0.25, FMT_PCT, "LimConc"),
            (26, "Variação mensal de custo relevante (acima de)", 0.20, FMT_PCT, "LimVar"),
            (27, "Pacote dominante (participação acima de)", 0.50, FMT_PCT, "LimPacote")]
    for r, lab, v, fmt, nm in rows:
        label(ws, f"B{r}", lab)
        input_cell(ws, f"C{r}", v, fmt)
        add_name(wb, nm, f"{R_PAR}!$C${r}")

    section(ws, 29, "Legenda", "B", "C")
    input_cell(ws, "B30", "Célula amarela / fonte azul = entrada manual")
    write(ws, "B31", "Célula cinza / fonte azul-marinho = calculada (não editar)", font(True, NAVY, 10), fill(GRAY1), ALIGN_C, None, BORDER_ALL)

    # ---- listas
    def list_col(col, hdr, values, n, name, first_row=4, editable=True):
        write(ws, f"{col}3", hdr, font(True, WHITE, 9), fill(NAVY2), ALIGN_C, None, BORDER_ALL)
        for i in range(n):
            r = first_row + i
            c = ws[f"{col}{r}"]
            c.value = values[i] if i < len(values) else None
            c.border = BORDER_ALL
            if editable:
                c.fill = fill(INPUT_FILL); c.font = font(False, INPUT_FONT, 9)
            else:
                c.font = font(False, GRAPHITE, 9)
        add_name(wb, name, f"{R_PAR}!${col}${first_row}:${col}${first_row + n - 1}")

    list_col("E", "Meses", MESES, 12, "Meses", editable=False)
    list_col("F", "Abrev.", MESES_ABREV, 12, "MesesAbrev", editable=False)
    list_col("H", "Pacotes (ordem da DRE)", PACOTES_INICIAIS, N_PAC, "ListaPacotes")
    list_col("J", "Categorias", CATEGORIAS_INICIAIS, N_CAT, "ListaCategorias")
    list_col("L", "Anos", list(range(2025, 2036)), 11, "ListaAnos", editable=False)
    list_col("N", "Visão do Painel", ["Mês", "Acumulado no ano"], 2, "ListaVisao", editable=False)
    list_col("P", "Status (filtro pagamentos)", STATUS_FILTRO, len(STATUS_FILTRO), "ListaStatus", editable=False)
    list_col("R", "Filtro de pacote", ["(Todos)"] + [f'=IF($H{4+i}="","",$H{4+i})' for i in range(N_PAC)], N_PAC + 1, "ListaFiltroPacote", editable=False)
    list_col("T", "Filtro de categoria", ["(Todos)"] + [f'=IF($J{4+i}="","",$J{4+i})' for i in range(N_CAT)], N_CAT + 1, "ListaFiltroCategoria", editable=False)
    list_col("V", "Filtro de fornecedor", ["(Todos)"] + [f'=IF({R_FOR}!$A{2+i}="","",{R_FOR}!$A{2+i})' for i in range(N_FOR_LIST)], N_FOR_LIST + 1, "ListaFiltroFornecedor", editable=False)
    list_col("X", "Matriz: mostrar", ["Total", "Pago", "Em aberto"], 3, "ListaMostrar", editable=False)
    list_col("Z", "Mês nº", list(range(1, 13)), 12, "ListaMesesNum", editable=False)
    ws["H3"].comment = Comment("Inclua aqui novos pacotes (até 15). A ordem define as linhas da DRE e das análises.", "Controle Financeiro")
    ws.freeze_panes = "B4"
    ws.sheet_view.zoomScale = 90


# ====================================================================== REMUNERAÇÃO
def build_remuneracao(wb, ws):
    set_widths(ws, {"A": 2, "B": 10, "C": 9, "D": 13, "E": 14, "F": 18, "G": 14, "H": 20, "I": 44, "J": 16, "K": 22})
    title_block(ws, "REMUNERAÇÃO MENSAL | COMPETÊNCIA", "Receita gerencial consolidada do setor de Frota por competência mensal (sem segregação por pacote, categoria ou fornecedor)", "B", "K")
    label(ws, "B4", "Como preencher: informe o Ano e o Mês de competência e o valor total da Remuneração (R$). Ajustes são opcionais. "
                    "A competência AAAA-MM é gerada automaticamente. Ao alterar um valor, a DRE, o Painel e os indicadores são recalculados.", size=9, align=ALIGN_LW)
    ws.merge_cells("B4:K4"); ws.row_dimensions[4].height = 30
    input_cell(ws, "B5", "Amarelo = entrada manual"); ws.merge_cells("B5:D5")
    write(ws, "E5", "Cinza = calculado automaticamente", font(True, NAVY, 9), fill(GRAY1), ALIGN_C, None, BORDER_ALL); ws.merge_cells("E5:G5")
    write(ws, "H5", "Importante: remuneração por competência NÃO é entrada de caixa. Recebimentos efetivos (se houver) são informados na tabela abaixo.",
          font(True, RED, 8), align=ALIGN_LW); ws.merge_cells("H5:K5")
    hdrs = ["Ano", "Mês", "Mês (nome)", "Competência", "Remuneração (R$)", "Ajustes (R$)", "Receita gerencial ajustada (R$)", "Observações", "Data de atualização", "Verificação"]
    header_row(ws, 7, 2, hdrs)
    first, nrows = 8, 24  # 2026 e 2027
    for i in range(nrows):
        r = first + i
        ano = 2026 + i // 12
        mes = i % 12 + 1
        input_cell(ws, f"B{r}", ano, "0")
        input_cell(ws, f"C{r}", mes, "0")
        write(ws, f"D{r}", f'=IF(AND(ISNUMBER($C{r}),$C{r}>=1,$C{r}<=12),INDEX(Meses,$C{r}),"")', font(False, GRAPHITE, 9), fill(GRAY1), ALIGN_C, None, BORDER_ALL)
        write(ws, f"E{r}", f'=IF(AND(ISNUMBER($B{r}),ISNUMBER($C{r})),$B{r}&"-"&TEXT($C{r},"00"),"")', font(True, NAVY, 9), fill(GRAY1), ALIGN_C, "@", BORDER_ALL)
        input_cell(ws, f"F{r}", None, FMT_BRL)
        input_cell(ws, f"G{r}", None, FMT_BRL)
        write(ws, f"H{r}", f'=IF(AND($F{r}="",$G{r}=""),"",N($F{r})+N($G{r}))', font(True, NAVY, 9), fill(GRAY1), ALIGN_R, FMT_BRL, BORDER_ALL)
        input_cell(ws, f"I{r}", None); ws[f"I{r}"].alignment = ALIGN_LW; ws[f"I{r}"].font = font(False, INPUT_FONT, 9)
        input_cell(ws, f"J{r}", None, FMT_DATE)
        write(ws, f"K{r}", f'=IF($E{r}="","",IF(COUNTIF($E$8:$E$2000,$E{r})>1,"COMPETÊNCIA DUPLICADA",IF($F{r}="","Remuneração não informada","OK")))',
              font(False, GRAPHITE, 8), fill(GRAY1), ALIGN_C, None, BORDER_ALL)
    last = first + nrows - 1
    cf_text(ws, f"K{first}:K{last}", "DUPLICADA", RED_L, RED)
    cf_text(ws, f"K{first}:K{last}", "OK", GREEN_L, GREEN, contains=False)
    t = Table(displayName="tblRemuneracao", ref=f"B7:K{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    add_list_validation(ws, f"C{first}:C{last + 200}", "=ListaMesesNum")
    add_list_validation(ws, f"B{first}:B{last + 200}", "=ListaAnos")
    label(ws, f"B{last + 1}", "Para novos exercícios, basta continuar digitando Ano/Mês nas linhas seguintes (a tabela se expande automaticamente).", size=8, italic=True)
    ws.merge_cells(f"B{last + 1}:K{last + 1}")

    r0 = last + 4
    section(ws, r0, "Saldo inicial de caixa (opcional) - utilizado somente no Fluxo de Caixa projetado", "B", "K")
    label(ws, f"B{r0+1}", "Data do saldo inicial"); ws.merge_cells(f"B{r0+1}:D{r0+1}")
    input_cell(ws, f"E{r0+1}", None, FMT_DATE)
    label(ws, f"B{r0+2}", "Valor do saldo inicial (R$)"); ws.merge_cells(f"B{r0+2}:D{r0+2}")
    input_cell(ws, f"E{r0+2}", None, FMT_BRL)
    add_name(wb, "SaldoInicialData", f"{R_REM}!$E${r0+1}")
    add_name(wb, "SaldoInicialValor", f"{R_REM}!$E${r0+2}")
    label(ws, f"F{r0+1}", "Sem saldo inicial informado, o Fluxo de Caixa não calcula saldo projetado (evita saldo fictício).", size=8, italic=True)
    ws.merge_cells(f"F{r0+1}:K{r0+2}")

    r1 = r0 + 5
    section(ws, r1, "Recebimentos efetivos de caixa (opcional) - somente entradas confirmadas, com data de recebimento", "B", "K")
    header_row(ws, r1 + 1, 2, ["Data do recebimento", "Valor recebido (R$)", "Competência de origem", "Descrição / referência", "Data de atualização"])
    ws.merge_cells(f"D{r1+1}:D{r1+1}")
    rf, nrec = r1 + 2, 60
    for i in range(nrec):
        r = rf + i
        input_cell(ws, f"B{r}", None, FMT_DATE)
        input_cell(ws, f"C{r}", None, FMT_BRL)
        input_cell(ws, f"D{r}", None, "@")
        input_cell(ws, f"E{r}", None); ws[f"E{r}"].alignment = ALIGN_LW
        input_cell(ws, f"F{r}", None, FMT_DATE)
    add_name(wb, "RecebData", f"{R_REM}!$B${rf}:$B${rf + nrec - 1 + 500}")
    add_name(wb, "RecebValor", f"{R_REM}!$C${rf}:$C${rf + nrec - 1 + 500}")
    t2 = Table(displayName="tblRecebimentos", ref=f"B{r1+1}:F{rf + nrec - 1}")
    t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t2)
    ws.freeze_panes = "B8"
    ws.sheet_view.showGridLines = False
    setup_print(ws, area=f"B1:K{rf + nrec - 1}", header_text="Remuneração mensal | Frota")
    return first, last


# ====================================================================== DADOS GRÁFICOS
def build_dados_graficos(wb, ws):
    """Blocos de dados (ano do Painel = AnoSel, filtros do Painel aplicados aos custos)."""
    set_widths(ws, {"A": 2, "B": 30, "C": 14, "D": 16, "E": 16, "F": 16, "G": 16, "H": 16, "I": 16, "J": 16, "K": 16, "L": 16, "M": 16, "N": 10, "O": 16, "P": 16})
    title_block(ws, "DADOS DOS GRÁFICOS DO PAINEL EXECUTIVO", "Camada auxiliar calculada automaticamente (não editar). Ano e filtros conforme Painel Executivo.", "B", "P")
    nfI, nfAQ, nfComp = _whole(R_NF, NF, "Valor Doc"), _whole(R_NF, NF, "Considerar DRE"), _whole(R_NF, NF, "Competência")
    nfPac, nfCat, nfFor = _whole(R_NF, NF, "Pacote final"), _whole(R_NF, NF, "Categoria final"), _whole(R_NF, NF, "Razão Social")
    nfAno = _whole(R_NF, NF, "Ano comp.")
    filt_nf = f"{nfPac},CritPacote,{nfCat},CritCategoria,{nfFor},CritFornecedor"
    p_val, p_liq, p_cons, p_cx, p_cv = _whole(R_108, P, "Valor bruto"), _whole(R_108, P, "Valor líquido pago"), _whole(R_108, P, "Considerar"), _whole(R_108, P, "Comp. Caixa"), _whole(R_108, P, "Comp. Vencimento")
    p_pac, p_cat, p_for, p_jur, p_des = _whole(R_108, P, "Pacote"), _whole(R_108, P, "Categoria"), _whole(R_108, P, "Fornecedor (norm.)"), _whole(R_108, P, "Juros"), _whole(R_108, P, "Desconto")
    filt_p = f"{p_pac},CritPacote,{p_cat},CritCategoria,{p_for},CritFornecedor"
    a_val, a_cons, a_cv, a_st = _whole(R_1015, A, "Valor parcela"), _whole(R_1015, A, "Considerar"), _whole(R_1015, A, "Comp. Vencimento"), _whole(R_1015, A, "Status")
    a_pac, a_cat, a_for = _whole(R_1015, A, "Pacote"), _whole(R_1015, A, "Categoria"), _whole(R_1015, A, "Fornecedor (norm.)")
    filt_a = f"{a_pac},CritPacote,{a_cat},CritCategoria,{a_for},CritFornecedor"
    remC, remV, remA = f"{R_REM}!$E:$E", f"{R_REM}!$F:$F", f"{R_REM}!$G:$G"

    # ---- Bloco 1: evolução mensal
    section(ws, 4, "Bloco 1 - Evolução mensal (ano selecionado)", "B", "P")
    hdr = ["Mês", "Competência", "Remuneração líquida", "Custo por competência", "Resultado gerencial", "Pagos realizados (líquido)",
           "A pagar em aberto (por venc.)", "Pagos (bruto, por venc.)", "Juros pagos", "Descontos obtidos", "Custo acumulado", "Pagos acumulados", "Mês nº", "Margem %", "A vencer (por venc.)"]
    header_row(ws, 5, 2, hdr)
    for m in range(1, 13):
        r = 5 + m
        comp = f'AnoSel&"-"&TEXT({m},"00")'
        ws[f"B{r}"] = f"=INDEX(Meses,{m})"
        ws[f"C{r}"] = f"={comp}"
        ws[f"D{r}"] = f"=SUMIFS({remV},{remC},$C{r})+SUMIFS({remA},{remC},$C{r})"
        ws[f"E{r}"] = f"=SUMIFS({nfI},{nfAQ},1,{nfComp},$C{r},{filt_nf})"
        ws[f"J{r}"] = f"=SUMIFS({p_jur},{p_cons},1,{p_cx},$C{r},{filt_p})"
        ws[f"K{r}"] = f"=SUMIFS({p_des},{p_cons},1,{p_cx},$C{r},{filt_p})"
        ws[f"F{r}"] = f"=IF(FiltroAtivo=1,0,$D{r}-$E{r}-$J{r}+$K{r})"
        ws[f"G{r}"] = f"=SUMIFS({p_liq},{p_cons},1,{p_cx},$C{r},{filt_p})"
        ws[f"H{r}"] = f"=SUMIFS({a_val},{a_cons},1,{a_cv},$C{r},{filt_a})"
        ws[f"I{r}"] = f"=SUMIFS({p_val},{p_cons},1,{p_cv},$C{r},{filt_p})"
        ws[f"L{r}"] = f"=SUM($E$6:$E{r})"
        ws[f"M{r}"] = f"=SUM($G$6:$G{r})"
        ws[f"N{r}"] = m
        ws[f"O{r}"] = f'=IF(OR(FiltroAtivo=1,$D{r}<=0),"",$F{r}/$D{r})'
        ws[f"P{r}"] = f'=SUMIFS({a_val},{a_cons},1,{a_cv},$C{r},{a_st},"A vencer",{filt_a})'
        for col in "DEFGHIJKLMP":
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"O{r}"].number_format = FMT_PCT
        style_range(ws, f"B{r}:P{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    write(ws, "B18", "Total", font(True, NAVY, 9), fill(GRAY2), ALIGN_L, None, BORDER_ALL)
    for col in "DEFGHIJKP":
        write(ws, f"{col}18", f"=SUM({col}6:{col}17)", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    add_name(wb, "DG_Meses", f"{R_DG}!$B$6:$B$17")

    # ---- Bloco 2: pacotes
    section(ws, 20, "Bloco 2 - Custo por pacote (ano selecionado, por competência)", "B", "P")
    header_row(ws, 21, 2, ["Pacote", "Custo (ano)", "Participação %", "Custo (mês sel.)", "Custo (mês ant.)", "Variação M-1 %", "Custo acumulado até mês sel."])
    for i in range(N_PAC):
        r = 22 + i
        ws[f"B{r}"] = f'=IF(INDEX(ListaPacotes,{i+1})="","",INDEX(ListaPacotes,{i+1}))'
        ws[f"C{r}"] = f'=IF($B{r}="",0,SUMIFS({nfI},{nfAQ},1,{nfAno},AnoSel,{nfPac},$B{r},{nfCat},CritCategoria,{nfFor},CritFornecedor))'
        ws[f"D{r}"] = f'=IF($C$38>0,$C{r}/$C$38,0)'
        ws[f"E{r}"] = f'=IF($B{r}="",0,SUMIFS({nfI},{nfAQ},1,{nfComp},CompSel,{nfPac},$B{r},{nfCat},CritCategoria,{nfFor},CritFornecedor))'
        ws[f"F{r}"] = f'=IF($B{r}="",0,SUMIFS({nfI},{nfAQ},1,{nfComp},CompAnt,{nfPac},$B{r},{nfCat},CritCategoria,{nfFor},CritFornecedor))'
        ws[f"G{r}"] = f'=IF($F{r}>0,($E{r}-$F{r})/$F{r},"")'
        ws[f"H{r}"] = f'=IF($B{r}="",0,SUMIFS({nfI},{nfAQ},1,{nfAno},AnoSel,{_whole(R_NF,NF,"Mês comp.")},"<="&MesSel,{nfPac},$B{r},{nfCat},CritCategoria,{nfFor},CritFornecedor))'
    r = 22 + N_PAC  # 37 - outros
    ws[f"B{r}"] = "Outros / sem classificação"
    ws[f"C{r}"] = f'=SUMIFS({nfI},{nfAQ},1,{nfAno},AnoSel,{nfCat},CritCategoria,{nfFor},CritFornecedor)-SUM($C$22:$C${r-1})'
    ws[f"D{r}"] = f'=IF($C$38>0,$C{r}/$C$38,0)'
    ws[f"E{r}"] = f'=SUMIFS({nfI},{nfAQ},1,{nfComp},CompSel,{nfCat},CritCategoria,{nfFor},CritFornecedor)-SUM($E$22:$E${r-1})'
    ws[f"F{r}"] = f'=SUMIFS({nfI},{nfAQ},1,{nfComp},CompAnt,{nfCat},CritCategoria,{nfFor},CritFornecedor)-SUM($F$22:$F${r-1})'
    ws[f"G{r}"] = f'=IF($F{r}>0,($E{r}-$F{r})/$F{r},"")'
    ws[f"H{r}"] = f'=SUMIFS({nfI},{nfAQ},1,{nfAno},AnoSel,{_whole(R_NF,NF,"Mês comp.")},"<="&MesSel,{nfCat},CritCategoria,{nfFor},CritFornecedor)-SUM($H$22:$H${r-1})'
    ws["B38"] = "Total"
    for col in "CEFH":
        ws[f"{col}38"] = f"=SUM({col}22:{col}37)"
    ws["D38"] = "=SUM(D22:D37)"
    ws["G38"] = '=IF($F$38>0,($E$38-$F$38)/$F$38,"")'
    for r in range(22, 39):
        for col in "CEFH":
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"D{r}"].number_format = FMT_PCT; ws[f"G{r}"].number_format = FMT_PCT
        style_range(ws, f"B{r}:H{r}", font(r == 38, NAVY if r == 38 else GRAPHITE, 9), fill(GRAY2) if r == 38 else None, None, BORDER_ALL)
    add_name(wb, "DG_PacNomes", f"{R_DG}!$B$22:$B$37")
    add_name(wb, "DG_PacValores", f"{R_DG}!$C$22:$C$37")

    # ---- Bloco 3: categorias (ano) + top 8 x meses
    section(ws, 40, "Bloco 3 - Custo por categoria (ano selecionado) e evolução mensal das 8 maiores", "B", "P")
    header_row(ws, 41, 2, ["Categoria", "Custo (ano)", "Rank"])
    for i in range(N_CAT):
        r = 42 + i
        ws[f"B{r}"] = f'=IF(INDEX(ListaCategorias,{i+1})="","",INDEX(ListaCategorias,{i+1}))'
        ws[f"C{r}"] = f'=IF($B{r}="",0,SUMIFS({nfI},{nfAQ},1,{nfAno},AnoSel,{nfCat},$B{r},{nfPac},CritPacote,{nfFor},CritFornecedor))'
        ws[f"D{r}"] = f'=IF($C{r}>0,RANK($C{r},$C$42:$C$66)+COUNTIF($C$42:$C{r},$C{r})-1,"")'
        ws[f"C{r}"].number_format = FMT_NUM
        style_range(ws, f"B{r}:D{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    header_row(ws, 68, 2, ["Mês"] + [f"Top {k}" for k in range(1, 9)] + ["Outras categorias", "Total"])
    for k in range(1, 9):
        col = get_column_letter(2 + k)
        ws[f"{col}69"] = f'=IFERROR(INDEX($B$42:$B$66,MATCH({k},$D$42:$D$66,0)),"-")'
        ws[f"{col}69"].font = font(True, NAVY, 9); ws[f"{col}69"].border = BORDER_ALL; ws[f"{col}69"].fill = fill(GRAY2)
    write(ws, "B69", "Categoria", font(True, NAVY, 9), fill(GRAY2), ALIGN_L, None, BORDER_ALL)
    write(ws, "K69", "Outras", font(True, NAVY, 9), fill(GRAY2), ALIGN_C, None, BORDER_ALL)
    write(ws, "L69", "Total", font(True, NAVY, 9), fill(GRAY2), ALIGN_C, None, BORDER_ALL)
    for m in range(1, 13):
        r = 69 + m
        ws[f"B{r}"] = f"=INDEX(Meses,{m})"
        for k in range(1, 9):
            col = get_column_letter(2 + k)
            ws[f"{col}{r}"] = f'=IF({col}$69="-",0,SUMIFS({nfI},{nfAQ},1,{nfComp},$C{5+m},{nfCat},{col}$69,{nfPac},CritPacote,{nfFor},CritFornecedor))'
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"L{r}"] = f"=$E{5+m}"
        ws[f"K{r}"] = f"=$L{r}-SUM($C{r}:$J{r})"
        ws[f"K{r}"].number_format = FMT_NUM; ws[f"L{r}"].number_format = FMT_NUM
        style_range(ws, f"B{r}:L{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)

    # ---- Bloco 4: top 10 fornecedores (ano selecionado - usa ranking da aba Fornecedores com AnoAF)
    section(ws, 84, "Bloco 4 - Dez maiores fornecedores por custo de competência (ano da Análise de Fornecedores)", "B", "P")
    header_row(ws, 85, 2, ["Fornecedor", "Custo (ano)", "Participação %"])
    fRank, fNome, fCusto, fPart = _whole(R_FOR, F, "Rank custo"), _whole(R_FOR, F, "Fornecedor"), _whole(R_FOR, F, "Custo competência (ano)"), _whole(R_FOR, F, "Participação %")
    for k in range(1, 11):
        r = 85 + k
        ws[f"B{r}"] = f'=IFERROR(INDEX({fNome},MATCH({k},{fRank},0)),"-")'
        ws[f"C{r}"] = f'=IFERROR(INDEX({fCusto},MATCH({k},{fRank},0)),0)'
        ws[f"D{r}"] = f'=IFERROR(INDEX({fPart},MATCH({k},{fRank},0)),0)'
        ws[f"C{r}"].number_format = FMT_NUM; ws[f"D{r}"].number_format = FMT_PCT
        style_range(ws, f"B{r}:D{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)

    # ---- Bloco 5: próximos 12 meses a partir da data de referência
    section(ws, 98, "Bloco 5 - Compromissos futuros: obrigações em aberto por mês de vencimento a partir da data de referência", "B", "P")
    header_row(ws, 99, 2, ["Período", "Competência venc.", "A vencer (R$)", "Vencido em aberto (R$)", "Total (R$)", "Acumulado (R$)", "Qtd parcelas"])
    a_ve = _whole(R_1015, A, "Vencimento (data)")
    ws["B100"] = "Vencidos (em aberto)"
    ws["C100"] = "-"
    ws["D100"] = 0
    ws["E100"] = f'=SUMIFS({a_val},{a_cons},1,{a_st},"Vencido",{filt_a})'
    ws["F100"] = "=D100+E100"
    ws["G100"] = "=F100"
    ws["H100"] = f'=COUNTIFS({a_cons},1,{a_st},"Vencido",{filt_a})'
    for k in range(12):
        r = 101 + k
        d0 = f"DATE(YEAR(DataRef),MONTH(DataRef)+{k},1)"
        ws[f"B{r}"] = f'=INDEX(MesesAbrev,MONTH({d0}))&"/"&YEAR({d0})'
        ws[f"C{r}"] = f'=YEAR({d0})&"-"&TEXT(MONTH({d0}),"00")'
        ws[f"D{r}"] = f'=SUMIFS({a_val},{a_cons},1,{a_cv},$C{r},{a_st},"A vencer",{filt_a})'
        ws[f"E{r}"] = 0
        ws[f"F{r}"] = f"=D{r}+E{r}"
        ws[f"G{r}"] = f"=G{r-1}+F{r}"
        ws[f"H{r}"] = f'=COUNTIFS({a_cons},1,{a_cv},$C{r},{a_st},"A vencer",{filt_a})'
    for r in range(100, 113):
        for col in "DEFG":
            ws[f"{col}{r}"].number_format = FMT_NUM
        style_range(ws, f"B{r}:H{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    write(ws, "B113", "Total em aberto", font(True, NAVY, 9), fill(GRAY2), ALIGN_L, None, BORDER_ALL)
    for col in "DEF":
        write(ws, f"{col}113", f"=SUM({col}100:{col}112)", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    write(ws, "H113", "=SUM(H100:H112)", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_INT, BORDER_ALL)
    ws["B114"] = f'=IF(F113<>SUMIFS({a_val},{a_cons},1,{filt_a}),"ATENÇÃO: existem vencimentos além de 12 meses (R$ "&TEXT(SUMIFS({a_val},{a_cons},1,{filt_a})-F113,"#,##0.00")&") não exibidos na projeção.","Projeção cobre 100% das obrigações em aberto.")'
    ws["B114"].font = font(False, AMBER, 8, italic=True)
    ws.merge_cells("B114:H114")

    # ---- Bloco 6: matriz fornecedor x pacote (ano Custos por Pacote) e categoria x pacote
    section(ws, 116, "Bloco 6 - Matriz Fornecedor x Pacote e Categoria x Pacote (ano da aba Custos por Pacote)", "B", "P")
    header_row(ws, 117, 2, ["Fornecedor"] + [f"=IF(INDEX(ListaPacotes,{i+1})=\"\",\"-\",INDEX(ListaPacotes,{i+1}))" for i in range(N_PAC)] + ["Pacote selecionado", "Rank (pacote sel.)"])
    anoCP = f"{R_CP}!$C$4"
    pacSel = f"{R_CP}!$C$5"
    for i in range(N_FOR_LIST):
        r = 118 + i
        ws[f"B{r}"] = f'=IF({R_FOR}!$A{2+i}="","",{R_FOR}!$A{2+i})'
        for j in range(N_PAC):
            col = get_column_letter(3 + j)
            ws[f"{col}{r}"] = f'=IF(OR($B{r}="",{col}$117="-"),0,SUMIFS({nfI},{nfAQ},1,{nfAno},{anoCP},{nfFor},$B{r},{nfPac},{col}$117))'
            ws[f"{col}{r}"].number_format = FMT_NUM
        selcol = get_column_letter(3 + N_PAC)
        rkcol = get_column_letter(4 + N_PAC)
        ws[f"{selcol}{r}"] = f'=IF($B{r}="",0,SUMIFS({nfI},{nfAQ},1,{nfAno},{anoCP},{nfFor},$B{r},{nfPac},{pacSel}))'
        ws[f"{selcol}{r}"].number_format = FMT_NUM
        ws[f"{rkcol}{r}"] = f'=IF({selcol}{r}>0,RANK({selcol}{r},${selcol}$118:${selcol}${118+N_FOR_LIST-1})+COUNTIF(${selcol}$118:{selcol}{r},{selcol}{r})-1,"")'
        style_range(ws, f"B{r}:{rkcol}{r}", font(False, GRAPHITE, 9))
    rr = 118 + N_FOR_LIST + 2  # 520
    header_row(ws, rr, 2, ["Categoria"] + [f"=IF(INDEX(ListaPacotes,{i+1})=\"\",\"-\",INDEX(ListaPacotes,{i+1}))" for i in range(N_PAC)] + ["Pacote selecionado"])
    for i in range(N_CAT):
        r = rr + 1 + i
        ws[f"B{r}"] = f'=IF(INDEX(ListaCategorias,{i+1})="","",INDEX(ListaCategorias,{i+1}))'
        for j in range(N_PAC):
            col = get_column_letter(3 + j)
            ws[f"{col}{r}"] = f'=IF(OR($B{r}="",{col}${rr}="-"),0,SUMIFS({nfI},{nfAQ},1,{nfAno},{anoCP},{nfCat},$B{r},{nfPac},{col}${rr}))'
            ws[f"{col}{r}"].number_format = FMT_NUM
        selcol = get_column_letter(3 + N_PAC)
        ws[f"{selcol}{r}"] = f'=IF($B{r}="",0,SUMIFS({nfI},{nfAQ},1,{nfAno},{anoCP},{nfCat},$B{r},{nfPac},{pacSel}))'
        ws[f"{selcol}{r}"].number_format = FMT_NUM
        style_range(ws, f"B{r}:{selcol}{r}", font(False, GRAPHITE, 9))
    ws.freeze_panes = "B4"
    ws.sheet_view.zoomScale = 85
    return {"forn_matrix_first": 118, "forn_matrix_last": 118 + N_FOR_LIST - 1, "cat_matrix_first": rr + 1, "cat_matrix_last": rr + N_CAT,
            "selcol": get_column_letter(3 + N_PAC), "rkcol": get_column_letter(4 + N_PAC), "pac_hdr_row": 117, "cat_hdr_row": rr}
