# -*- coding: utf-8 -*-
"""Testes de validação (seção 12): injeta dados em cópias da planilha, recalcula com LibreOffice e verifica resultados."""
import sys, os, shutil, subprocess, json, datetime, copy
sys.path.insert(0, "build")
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
from sources import NF, P, A, F
from common import *

RECALC = [p for p in __import__("glob").glob("/root/.claude/skills/synced/*/xlsx/scripts/recalc.py")][0]
BASE = "out/Controle_Financeiro_Gerencial_Frota.xlsx"

def append_row(ws, table_name, raw_values):
    """Simula a expansão de tabela do Excel: nova linha com valores brutos + fórmulas copiadas da linha anterior."""
    t = ws.tables[table_name]
    first, last = t.ref.split(":")
    last_row = int(''.join(ch for ch in last if ch.isdigit()))
    last_col = ''.join(ch for ch in last if ch.isalpha())
    new = last_row + 1
    for c in range(1, ci(last_col) + 1):
        src = ws.cell(last_row, c)
        if isinstance(src.value, str) and src.value.startswith("="):
            ws.cell(new, c).value = Translator(src.value, origin=f"{gl(c)}{last_row}").translate_formula(f"{gl(c)}{new}")
        ws.cell(new, c).number_format = src.number_format
        ws.cell(new, c).font = copy.copy(src.font)
    for col, val in raw_values.items():
        ws[f"{col}{new}"] = val
    t.ref = f"{first}:{last_col}{new}"
    if t.autoFilter is not None:
        t.autoFilter.ref = t.ref
    return new

def recalc(path):
    out = subprocess.run([sys.executable, RECALC, path, "900"], capture_output=True, text=True)
    j = json.loads(out.stdout)
    assert j.get("status") == "success", j
    return j

def val(wb, sheet, ref):
    return wb[sheet][ref].value

def approx(a, b, tol=0.01):
    return abs((a or 0) - (b or 0)) <= tol

results = []
def check(name, cond, info=""):
    results.append((name, bool(cond), info))
    print(("OK   " if cond else "FALHA"), name, info)

