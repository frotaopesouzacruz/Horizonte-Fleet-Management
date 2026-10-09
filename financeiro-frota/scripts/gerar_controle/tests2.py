# -*- coding: utf-8 -*-
"""T2: extração 108 na ordem original (grupos 'Doc. Pago em') -> data efetiva de pagamento; T1 integridade dos dados brutos."""
import sys, shutil, subprocess, json, datetime, copy
sys.path.insert(0, "build")
import openpyxl
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
from sources import NF, P, A, F
from common import *
from tests_lib import append_row, recalc, val, approx, check, results

BASE = "out/Controle_Financeiro_Gerencial_Frota.xlsx"
# ---------------------------------------------------------------- T01 integridade: dados brutos preservados
orig = openpyxl.load_workbook("src/Fluxo_de_caixa.xlsx", data_only=True)
new = openpyxl.load_workbook(BASE)
diffs = 0
for sheet, cols in [(S_NF, "BEFGHIJKLMNO"), (S_FOR, "ACDE"), (S_108, "CDEFGHIJKLM"), (S_1015, "CDEFGHIJK")]:
    o, n = orig[sheet], new[sheet]
    for r in range(2, o.max_row + 1):
        for c in cols:
            if (o[f"{c}{r}"].value or None) != (n[f"{c}{r}"].value or None):
                diffs += 1
                if diffs < 5: print("diff", sheet, c, r, o[f"{c}{r}"].value, n[f"{c}{r}"].value)
check("T01 Dados brutos das 4 abas de origem preservados célula a célula", diffs == 0, f"diferenças={diffs}")
tot_o = sum(v for v in (orig[S_NF][f"I{r}"].value for r in range(2, 495)) if isinstance(v, (int, float)))
check("T01 Total Valor Doc original = 2.562.627,16", approx(tot_o, 2562627.16), str(tot_o))

# ---------------------------------------------------------------- T2: 108 em ordem original
t2 = "test/t2.xlsx"
shutil.copy(BASE, t2)
wb = openpyxl.load_workbook(t2)
pg = wb[S_108]
# coleta linhas de dados e datas dos grupos
rows, dates = [], []
for r in range(2, pg.max_row + 1):
    c = pg.cell(r, 3).value
    if c is None: continue
    if str(c).startswith("Doc. Pago em"):
        dates.append(datetime.datetime.strptime(pg.cell(r, 4).value, "%d/%m/%Y"))
    elif c not in ("Documento",) and not str(c).startswith(("Filial", "Total do Dia")) and isinstance(pg.cell(r, 10).value, (int, float)):
        rows.append([pg.cell(r, k).value for k in range(3, 14)])
dates.sort()
# atribui cada parcela ao primeiro grupo cuja data >= vencimento (sintético, apenas para validar o mecanismo)
groups = {d: [] for d in dates}
for row in rows:
    venc = row[5]
    d = next((d for d in dates if d >= venc), dates[-1])
    groups[d].append(row)
# limpa C:M e reescreve na estrutura do relatório
for r in range(2, pg.max_row + 1):
    for k in range(3, 14):
        pg.cell(r, k).value = None
r = 2
pg.cell(r, 3).value = "Filial "; pg.cell(r, 4).value = "HORIZONTE LOGÍSTICA LTDA"; r += 1
expected_by_month = {}
for d in dates:
    pg.cell(r, 3).value = "Doc. Pago em :"; pg.cell(r, 4).value = d.strftime("%d/%m/%Y"); r += 1
    for k, h in enumerate(["Documento", "Série", "Fornecedor", "Emissão", "Tipo", "Vencto", "Data Ref.", "Vlr. Parc.", "Vlr. Desco.", "Vlr. Juros", "Vlr. Total"]):
        pg.cell(r, 3 + k).value = h
    r += 1
    for row in groups[d]:
        for k, v in enumerate(row):
            pg.cell(r, 3 + k).value = v
        r += 1
    pg.cell(r, 3).value = "Total do Dia :"; pg.cell(r, 4).value = sum(x[10] for x in groups[d]); r += 1
wb.save(t2)
recalc(t2)
w = openpyxl.load_workbook(t2, data_only=True)
check("T2 Flag de data de pagamento validada = 1 (ordem original)", val(w, S_PAR, "C20") == 1, str(val(w, S_PAR, "C20")))
check("T2 Parcelas Frota consideradas continuam 526", val(w, S_PAINEL, "E20") == 526, str(val(w, S_PAINEL, "E20")))
check("T2 Todas as 526 parcelas com data efetiva atribuída", val(w, S_PAR, "C22") == 526, str(val(w, S_PAR, "C22")))
check("T2 Fluxo de caixa: saídas com data validada = saídas totais (todos os meses)", all(approx(val(w, S_FC, f"{c}11"), val(w, S_FC, f"{c}12")) for c in "CDEFGHIJKLMN"))
check("T2 Painel sinaliza base = data efetiva", "Data efetiva" in str(val(w, S_PAR, "C21")), str(val(w, S_PAR, "C21")))
# conferência independente em pandas do pago por mês de pagamento
import pandas as pd
recs = []
for d in dates:
    for row in groups[d]:
        recs.append((d, row[2], row[0], row[4], row[10]))
df = pd.DataFrame(recs, columns=["pg", "forn", "doc", "tipo", "tot"])
ws = w[S_108]
liq = {}
for rr in range(2, ws.max_row + 1):
    if ws.cell(rr, ci(P["Considerar"])).value == 1:
        liq[ws.cell(rr, ci(P["Comp. Caixa"])).value] = liq.get(ws.cell(rr, ci(P["Comp. Caixa"])).value, 0) + ws.cell(rr, ci(P["Valor líquido pago"])).value
fc = {f"2026-{m:02d}": val(w, S_FC, f"{gl(2+m)}11") for m in range(1, 13)}
check("T2 Pagos por mês (Fluxo de Caixa) = soma das parcelas consideradas por Comp. Caixa", all(approx(fc[k], liq.get(k, 0)) for k in fc), str({k: round(v) for k, v in fc.items()}))
check("T2 Total pago líquido inalterado (2.362.314,66)", approx(sum(fc.values()), 2362314.66), str(sum(fc.values())))
json.dump(results, open("test/results_t2.json", "w"), ensure_ascii=False, indent=1)
