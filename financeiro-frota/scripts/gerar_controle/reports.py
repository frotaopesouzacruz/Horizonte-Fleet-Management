# -*- coding: utf-8 -*-
"""Camada 3/4 - DRE, Fluxo de Caixa, Competência x Caixa, Custos por Pacote, Pagamentos, Fornecedores, Conciliação."""
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, PieChart, DoughnutChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from common import *
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from sources import NF, P, A, F, MESES, MESES_ABREV, _whole, SCOPE_NF, SCOPE_SEMNF, SCOPE_OUTRA, SCOPE_NC, SEM_CLASS, NAO_CAD
from params import N_PAC, N_CAT, N_FOR_LIST, add_name

MCOLS = [get_column_letter(3 + i) for i in range(12)]  # C..N
nfI, nfAQ, nfComp = _whole(R_NF, NF, "Valor Doc"), _whole(R_NF, NF, "Considerar DRE"), _whole(R_NF, NF, "Competência")
nfPac, nfCat, nfFor, nfAno, nfMes = _whole(R_NF, NF, "Pacote final"), _whole(R_NF, NF, "Categoria final"), _whole(R_NF, NF, "Razão Social"), _whole(R_NF, NF, "Ano comp."), _whole(R_NF, NF, "Mês comp.")
nfBase, nfSt = _whole(R_NF, NF, "Base competência"), _whole(R_NF, NF, "Status documento")
p_val, p_liq, p_cons, p_cx, p_cv = _whole(R_108, P, "Valor bruto"), _whole(R_108, P, "Valor líquido pago"), _whole(R_108, P, "Considerar"), _whole(R_108, P, "Comp. Caixa"), _whole(R_108, P, "Comp. Vencimento")
p_pac, p_cat, p_for, p_jur, p_des = _whole(R_108, P, "Pacote"), _whole(R_108, P, "Categoria"), _whole(R_108, P, "Fornecedor (norm.)"), _whole(R_108, P, "Juros"), _whole(R_108, P, "Desconto")
p_comp, p_anoc, p_anov, p_anocx, p_esc, p_ret = _whole(R_108, P, "Competência"), _whole(R_108, P, "Ano comp."), _whole(R_108, P, "Ano venc."), _whole(R_108, P, "Ano caixa"), _whole(R_108, P, "Escopo"), _whole(R_108, P, "Retenção/diferença")
a_val, a_cons, a_cv, a_st = _whole(R_1015, A, "Valor parcela"), _whole(R_1015, A, "Considerar"), _whole(R_1015, A, "Comp. Vencimento"), _whole(R_1015, A, "Status")
a_pac, a_cat, a_for, a_comp, a_anoc, a_anov, a_esc, a_fx = _whole(R_1015, A, "Pacote"), _whole(R_1015, A, "Categoria"), _whole(R_1015, A, "Fornecedor (norm.)"), _whole(R_1015, A, "Competência"), _whole(R_1015, A, "Ano comp."), _whole(R_1015, A, "Ano venc."), _whole(R_1015, A, "Escopo"), _whole(R_1015, A, "Faixa de vencimento")
FILT_NF = f"{nfPac},CritPacote,{nfCat},CritCategoria,{nfFor},CritFornecedor"
FILT_P = f"{p_pac},CritPacote,{p_cat},CritCategoria,{p_for},CritFornecedor"
FILT_A = f"{a_pac},CritPacote,{a_cat},CritCategoria,{a_for},CritFornecedor"
remC, remV, remA = f"{R_REM}!$E:$E", f"{R_REM}!$F:$F", f"{R_REM}!$G:$G"

SERIES_COLORS = [NAVY, BLUE, GOLD, GREEN, RED, GRAY4, LBLUE, "8E44AD", "C9A227", "1ABC9C"]


def style_chart(ch, title, w=16, h=7.5, legend_pos="b", ymin=None):
    ch.title = title
    ch.style = 2
    ch.width = w
    ch.height = h
    if ch.legend is not None:
        ch.legend.position = legend_pos
    try:
        ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRAY2))
        ch.y_axis.number_format = '#,##0'
        ch.y_axis.delete = False
        ch.x_axis.delete = False
        if ymin is not None:
            ch.y_axis.scaling.min = ymin
    except Exception:
        pass
    return ch


def color_series(ch, colors=None, line=False, smooth=False, width=22000):
    colors = colors or SERIES_COLORS
    for i, s in enumerate(ch.series):
        col = colors[i % len(colors)]
        if line:
            s.graphicalProperties.line.solidFill = col
            s.graphicalProperties.line.width = width
            s.smooth = smooth
            s.marker.symbol = "circle"
            s.marker.size = 5
            s.marker.graphicalProperties = GraphicalProperties(solidFill=col)
            s.marker.graphicalProperties.line.solidFill = col
        else:
            s.graphicalProperties.solidFill = col
            s.graphicalProperties.line.solidFill = col


def year_cell(ws, wb, name, ref="C4", note="Padrão: ano selecionado no Painel Executivo. Pode ser sobrescrito aqui."):
    label(ws, "B4", "Ano de referência", True)
    input_cell(ws, ref, "=AnoSel", "0", note)
    add_list_validation(ws, ref, "=ListaAnos")
    add_name(wb, name, f"{q(ws.title)}!${ref[0]}${ref[1:]}")


# ====================================================================== DRE
def build_dre(wb, ws):
    set_widths(ws, {"A": 2, "B": 50, **{c: 14 for c in MCOLS}, "O": 16, "P": 11, "Q": 11, "R": 3})
    title_block(ws, "DRE GERENCIAL SIMPLIFICADA | SETOR DE FROTA", "Regime de competência - receita gerencial (remuneração) x custos operacionais reconhecidos no mês de competência do documento", "B", "Q")
    year_cell(ws, wb, "AnoDRE")
    write(ws, "D4", '="Documentos com competência estimada pelo mês de emissão: "&COUNTIFS(' + nfAQ + ',1,' + nfAno + ',AnoDRE,' + nfBase + ',"Estimada*")&" de "&COUNTIFS(' + nfAQ + ',1,' + nfAno + ',AnoDRE)&"  |  Competências com remuneração duplicada: "&SUMPRODUCT((' + f"{R_REM}!$K$8:$K$2000" + '="COMPETÊNCIA DUPLICADA")*1)',
          font(False, GRAY4, 8, italic=True), align=ALIGN_L)
    ws.merge_cells("D4:Q4")
    write(ws, "B5", '="Demonstrativo gerencial - não substitui a DRE contábil oficial. Sem tributos, depreciação ou provisões (ausentes da base). Ano "&AnoDRE', font(False, GRAY4, 8, italic=True))
    ws.merge_cells("B5:Q5")
    hdr = ["DRE Gerencial Simplificada (R$)"] + [f"=INDEX(Meses,{m})" for m in range(1, 13)] + ["Total ano", "% Receita", "% Custos"]
    header_row(ws, 7, 2, hdr, height=24)

    def comp(m):
        return f'AnoDRE&"-"&TEXT({m},"00")'

    def line(r, text, formulas, bold=False, fill_color=None, fmt=FMT_NUM, pct_rec=True, pct_cust=False, total="sum", color=GRAPHITE):
        write(ws, f"B{r}", text, font(bold, color, 9), fill(fill_color) if fill_color else None, ALIGN_L, None, BORDER_ALL)
        for i, col in enumerate(MCOLS):
            write(ws, f"{col}{r}", formulas(i + 1, col), font(bold, color, 9), fill(fill_color) if fill_color else None, ALIGN_R, fmt, BORDER_ALL)
        if total == "sum":
            write(ws, f"O{r}", f"=SUM(C{r}:N{r})", font(True, color, 9), fill(fill_color or GRAY1), ALIGN_R, fmt, BORDER_ALL)
        elif total is not None:
            write(ws, f"O{r}", total, font(True, color, 9), fill(fill_color or GRAY1), ALIGN_R, fmt, BORDER_ALL)
        else:
            write(ws, f"O{r}", "", font(True, color, 9), fill(fill_color or GRAY1), ALIGN_R, fmt, BORDER_ALL)
        write(ws, f"P{r}", f'=IF($O$12<>0,O{r}/$O$12,"")' if pct_rec else "", font(False, GRAY4, 8), None, ALIGN_R, FMT_PCT, BORDER_ALL)
        write(ws, f"Q{r}", f'=IF($O$30<>0,O{r}/$O$30,"")' if pct_cust else "", font(False, GRAY4, 8), None, ALIGN_R, FMT_PCT, BORDER_ALL)

    section(ws, 9, "1. RECEITA GERENCIAL (por competência)", "B", "Q", NAVY)
    line(10, "(+) Remuneração mensal do setor", lambda m, c: f"=SUMIFS({remV},{remC},{comp(m)})")
    line(11, "(+/-) Ajustes informados na aba Remuneração", lambda m, c: f"=SUMIFS({remA},{remC},{comp(m)})")
    line(12, "(=) Receita Gerencial Líquida", lambda m, c: f"={c}10+{c}11", True, GRAY2, color=NAVY)
    section(ws, 14, "2. CUSTOS OPERACIONAIS DA FROTA (por competência do documento)", "B", "Q", NAVY)
    pac_rows = []
    for i in range(N_PAC - 2):  # 13 slots
        r = 15 + i
        pac_rows.append(r)
        write(ws, f"B{r}", f'=IF(INDEX(ListaPacotes,{i+1})="","(-) Pacote não utilizado","(-) "&INDEX(ListaPacotes,{i+1}))', font(False, GRAPHITE, 9), None, ALIGN_L, None, BORDER_ALL)
        for j, col in enumerate(MCOLS):
            m = j + 1
            write(ws, f"{col}{r}", f'=IF(INDEX(ListaPacotes,{i+1})="",0,SUMIFS({nfI},{nfAQ},1,{nfComp},{comp(m)},{nfPac},INDEX(ListaPacotes,{i+1})))',
                  font(False, GRAPHITE, 9), None, ALIGN_R, FMT_NUM, BORDER_ALL)
        write(ws, f"O{r}", f"=SUM(C{r}:N{r})", font(True, GRAPHITE, 9), fill(GRAY1), ALIGN_R, FMT_NUM, BORDER_ALL)
        write(ws, f"P{r}", f'=IF($O$12<>0,O{r}/$O$12,"")', font(False, GRAY4, 8), None, ALIGN_R, FMT_PCT, BORDER_ALL)
        write(ws, f"Q{r}", f'=IF($O$30<>0,O{r}/$O$30,"")', font(False, GRAY4, 8), None, ALIGN_R, FMT_PCT, BORDER_ALL)
    r = 28
    line(r, "(-) Demais pacotes / sem classificação", lambda m, c: f"=SUMIFS({nfI},{nfAQ},1,{nfComp},{comp(m)})-SUM({c}15:{c}27)", pct_cust=True)
    line(29, "", lambda m, c: "", fmt=None, pct_rec=False, total=None)
    line(30, "(=) Total de Custos Operacionais", lambda m, c: f"=SUM({c}15:{c}28)", True, GRAY2, pct_cust=True, color=NAVY)
    section(ws, 32, "3. RESULTADO OPERACIONAL GERENCIAL", "B", "Q", NAVY)
    line(33, "(=) Receita Gerencial Líquida - Custos Operacionais", lambda m, c: f"={c}12-{c}30", True, GRAY2, color=NAVY)
    section(ws, 35, "4. AJUSTES E DESPESAS ADICIONAIS (reconhecidos no mês de caixa)", "B", "Q", NAVY)
    line(36, "(-) Despesas administrativas (não identificadas na base - linha reservada)", lambda m, c: "=0")
    line(37, "(-) Juros e encargos financeiros pagos (Rodopar 108)", lambda m, c: f"=SUMIFS({p_jur},{p_cons},1,{p_cx},{comp(m)})")
    line(38, "(+) Descontos financeiros obtidos (Rodopar 108)", lambda m, c: f"=SUMIFS({p_des},{p_cons},1,{p_cx},{comp(m)})")
    line(39, "(=) Resultado Gerencial Consolidado", lambda m, c: f"={c}33-{c}36-{c}37+{c}38", True, GRAY2, color=NAVY)
    ws.row_dimensions[39].height = 20
    section(ws, 41, "5. INDICADORES", "B", "Q", NAVY)
    line(42, "Margem gerencial (%) = Resultado / Receita líquida", lambda m, c: f'=IF({c}12>0,{c}39/{c}12,"n/d")', fmt=FMT_PCT, pct_rec=False, total='=IF(O12>0,O39/O12,"n/d")')
    line(43, "Custos operacionais / Receita líquida (%)", lambda m, c: f'=IF({c}12>0,{c}30/{c}12,"n/d")', fmt=FMT_PCT, pct_rec=False, total='=IF(O12>0,O30/O12,"n/d")')
    line(44, "Resultado acumulado no ano (YTD)", lambda m, c: f"=SUM($C$39:{c}39)", pct_rec=False, total="=O39")
    line(45, "Variação do resultado vs mês anterior (R$)", lambda m, c: (f"={c}39-{MCOLS[m-2]}39" if m > 1 else f'={c}39-(SUMIFS({remV},{remC},(AnoDRE-1)&"-12")+SUMIFS({remA},{remC},(AnoDRE-1)&"-12")-SUMIFS({nfI},{nfAQ},1,{nfComp},(AnoDRE-1)&"-12")-SUMIFS({p_jur},{p_cons},1,{p_cx},(AnoDRE-1)&"-12")+SUMIFS({p_des},{p_cons},1,{p_cx},(AnoDRE-1)&"-12"))'), pct_rec=False, total=None)
    line(46, "Custo mensal médio (meses com custo até o mês)", lambda m, c: f'=IF(COUNTIF($C$30:{c}30,">0")>0,SUM($C$30:{c}30)/COUNTIF($C$30:{c}30,">0"),"n/d")', pct_rec=False, total='=IF(COUNTIF(C30:N30,">0")>0,O30/COUNTIF(C30:N30,">0"),"n/d")')
    line(47, "Variação do custo vs mês anterior (%)", lambda m, c: (f'=IF({MCOLS[m-2]}30>0,({c}30-{MCOLS[m-2]}30)/{MCOLS[m-2]}30,"n/d")' if m > 1 else f'=IF(SUMIFS({nfI},{nfAQ},1,{nfComp},(AnoDRE-1)&"-12")>0,({c}30-SUMIFS({nfI},{nfAQ},1,{nfComp},(AnoDRE-1)&"-12"))/SUMIFS({nfI},{nfAQ},1,{nfComp},(AnoDRE-1)&"-12"),"n/d")'), fmt=FMT_PCT, pct_rec=False, total=None)
    line(48, "Receita acumulada no ano (YTD)", lambda m, c: f"=SUM($C$12:{c}12)", pct_rec=False, total="=O12")
    line(49, "Custo acumulado no ano (YTD)", lambda m, c: f"=SUM($C$30:{c}30)", pct_rec=False, total="=O30")
    section(ws, 51, "Participação de cada pacote nos custos operacionais (%)", "B", "Q", GRAPHITE)
    for i, pr in enumerate(pac_rows):
        r = 52 + i
        write(ws, f"B{r}", f"=B{pr}", font(False, GRAPHITE, 9), None, ALIGN_L, None, BORDER_ALL)
        for col in MCOLS:
            write(ws, f"{col}{r}", f'=IF({col}$30>0,{col}{pr}/{col}$30,"")', font(False, GRAPHITE, 9), None, ALIGN_R, FMT_PCT, BORDER_ALL)
        write(ws, f"O{r}", f'=IF($O$30>0,O{pr}/$O$30,"")', font(True, GRAPHITE, 9), fill(GRAY1), ALIGN_R, FMT_PCT, BORDER_ALL)
    r = 52 + len(pac_rows)
    write(ws, f"B{r}", "=B28", font(False, GRAPHITE, 9), None, ALIGN_L, None, BORDER_ALL)
    for col in MCOLS:
        write(ws, f"{col}{r}", f'=IF({col}$30>0,{col}28/{col}$30,"")', font(False, GRAPHITE, 9), None, ALIGN_R, FMT_PCT, BORDER_ALL)
    write(ws, f"O{r}", f'=IF($O$30>0,O28/$O$30,"")', font(True, GRAPHITE, 9), fill(GRAY1), ALIGN_R, FMT_PCT, BORDER_ALL)
    r += 2
    section(ws, r, "Critérios e limitações", "B", "Q", GRAPHITE)
    notas = [
        "Competência do custo = mês informado em 'Competência (ajuste manual)' na Relação de NF's; se em branco, mês de EMISSÃO do documento (estimativa documentada). O campo 'Referência' da NF não foi confirmado como competência (é a data de registro, poucos dias após a emissão) e não é utilizado.",
        "Cada documento é reconhecido uma única vez pelo valor total do documento (Valor Doc), independentemente do número de parcelas ou do status de pagamento. Pagamentos de parcelas nunca geram novo custo.",
        "Somente documentos da Relação de NF's cujo fornecedor está cadastrado como Área = Frota entram na DRE. Documentos duplicados ou inválidos são desconsiderados e listados em Conciliação e Qualidade.",
        "Juros e descontos vêm do Rodopar 108 (parcelas conciliadas a documentos da Frota) e são reconhecidos no mês de caixa da parcela (data de pagamento validada ou, na ausência, mês de vencimento).",
        "Não há tributos, depreciação, provisões, orçamento ou metas na base - nenhum desses itens foi inventado. Linhas reservadas permitem inclusão futura.",
        "Remuneração por competência não é entrada de caixa. Margens e percentuais dependentes de receita aparecem como 'n/d' quando não há remuneração informada.",
    ]
    for i, n in enumerate(notas):
        rr = r + 1 + i
        write(ws, f"B{rr}", f"{i+1}. {n}", font(False, GRAPHITE, 8), align=ALIGN_LW)
        ws.merge_cells(f"B{rr}:Q{rr}")
        ws.row_dimensions[rr].height = 26
    cf_negative_red_positive_green(ws, "C33:O33")
    cf_negative_red_positive_green(ws, "C39:O39")
    cf_negative_red_positive_green(ws, "C42:O42")
    cf_negative_red_positive_green(ws, "C44:O44")
    ws.freeze_panes = "C8"
    setup_print(ws, area=f"B1:Q{r + len(notas) + 1}", title_rows="7:7", header_text="DRE Gerencial Simplificada | Frota")
    ws.sheet_view.zoomScale = 90
    ws.sheet_view.showGridLines = False


# ====================================================================== FLUXO DE CAIXA
def build_fluxo(wb, ws):
    set_widths(ws, {"A": 2, "B": 52, **{c: 14 for c in MCOLS}, "O": 16, "P": 3})
    title_block(ws, "FLUXO DE CAIXA | REALIZADO E PROJETADO", "Entradas confirmadas, saídas pagas (Rodopar 108) e compromissos em aberto por vencimento (Rodopar 1015)", "B", "O")
    year_cell(ws, wb, "AnoFC")
    write(ws, "D4", '="Base das saídas realizadas: "&BaseCaixaTxt', font(True, AMBER, 9), align=ALIGN_L)
    ws.merge_cells("D4:O4")
    write(ws, "B5", "Remuneração por competência NÃO é entrada de caixa. Entradas somente quando cadastradas em Remuneração > Recebimentos efetivos. Sem saldo inicial informado, o saldo projetado não é calculado.", font(False, GRAY4, 8, italic=True), align=ALIGN_LW)
    ws.merge_cells("B5:O5"); ws.row_dimensions[5].height = 24
    hdr = ["Fluxo de caixa mensal (R$)"] + [f"=INDEX(Meses,{m})" for m in range(1, 13)] + ["Total ano"]
    header_row(ws, 7, 2, hdr, height=24)

    def comp(m):
        return f'AnoFC&"-"&TEXT({m},"00")'

    def line(r, text, formulas, bold=False, fill_color=None, fmt=FMT_NUM, total="sum", color=GRAPHITE):
        write(ws, f"B{r}", text, font(bold, color, 9), fill(fill_color) if fill_color else None, ALIGN_L, None, BORDER_ALL)
        for i, col in enumerate(MCOLS):
            write(ws, f"{col}{r}", formulas(i + 1, col), font(bold, color, 9), fill(fill_color) if fill_color else None, ALIGN_R, fmt, BORDER_ALL)
        if total == "sum":
            write(ws, f"O{r}", f"=SUM(C{r}:N{r})", font(True, color, 9), fill(fill_color or GRAY1), ALIGN_R, fmt, BORDER_ALL)
        else:
            write(ws, f"O{r}", total or "", font(True, color, 9), fill(fill_color or GRAY1), ALIGN_R, fmt, BORDER_ALL)

    section(ws, 9, "A. REALIZADO (somente movimentações confirmadas)", "B", "O", NAVY)
    line(10, "(+) Entradas efetivamente recebidas (cadastro manual)", lambda m, c: f'=SUMIFS(RecebValor,RecebData,">="&DATE(AnoFC,{m},1),RecebData,"<"&DATE(AnoFC,{m}+1,1))')
    line(11, "(-) Saídas pagas - parcelas de documentos da Frota (valor líquido pago)", lambda m, c: f"=SUMIFS({p_liq},{p_cons},1,{p_cx},{comp(m)})")
    line(12, "    das quais com data efetiva de pagamento validada", lambda m, c: f'=SUMIFS({p_liq},{p_cons},1,{p_cx},{comp(m)},{_whole(R_108,P,"Base data caixa")},"Data de pagamento")', color=GRAY4)
    line(13, "(=) Fluxo líquido realizado", lambda m, c: f"={c}10-{c}11", True, GRAY2, color=NAVY)
    section(ws, 15, "B. PROJETADO (obrigações em aberto no Rodopar 1015, por mês de vencimento)", "B", "O", NAVY)
    line(16, "(-) Obrigações a vencer", lambda m, c: f'=SUMIFS({a_val},{a_cons},1,{a_cv},{comp(m)},{a_st},"A vencer")')
    line(17, "(-) Obrigações vencidas e não pagas (no mês de vencimento original)", lambda m, c: f'=SUMIFS({a_val},{a_cons},1,{a_cv},{comp(m)},{a_st},"Vencido")')
    line(18, "(=) Total de contas a pagar previstas", lambda m, c: f"={c}16+{c}17", True, GRAY2, color=NAVY)
    section(ws, 20, "C. CONSOLIDADO", "B", "O", NAVY)
    line(21, "Fluxo projetado do mês = Entradas - Saídas pagas - Contas a pagar previstas", lambda m, c: f"={c}10-{c}11-{c}18", True)
    line(22, "Saldo inicial do mês (somente se saldo inicial informado)", lambda m, c: (
        f'=IF(AND(ISNUMBER(SaldoInicialValor),ISNUMBER(SaldoInicialData),YEAR(SaldoInicialData)=AnoFC,MONTH(SaldoInicialData)={m}),SaldoInicialValor,'
        + (f'IF(ISNUMBER({MCOLS[m-2]}23),{MCOLS[m-2]}23,"n/d"))' if m > 1 else '"n/d")')), total=None)
    line(23, "Saldo final projetado", lambda m, c: f'=IF(ISNUMBER({c}22),{c}22+{c}21,"n/d")', True, GRAY2, total='=IF(ISNUMBER(N23),N23,"n/d")', color=NAVY)
    section(ws, 25, "D. ACUMULADO NO ANO", "B", "O", NAVY)
    line(26, "Entradas recebidas acumuladas", lambda m, c: f"=SUM($C$10:{c}10)", total="=O10")
    line(27, "Saídas pagas acumuladas", lambda m, c: f"=SUM($C$11:{c}11)", total="=O11")
    line(28, "Contas a pagar previstas acumuladas", lambda m, c: f"=SUM($C$18:{c}18)", total="=O18")
    line(29, "Memo: Remuneração por competência (não é caixa)", lambda m, c: f"=SUMIFS({remV},{remC},{comp(m)})+SUMIFS({remA},{remC},{comp(m)})", color=GRAY4)
    cf_negative_red_positive_green(ws, "C13:O13"); cf_negative_red_positive_green(ws, "C21:O21"); cf_negative_red_positive_green(ws, "C23:O23")

    section(ws, 31, "E. PRÓXIMOS 12 MESES A PARTIR DA DATA DE REFERÊNCIA (obrigações em aberto, filtros do Painel aplicados)", "B", "O", NAVY2)
    header_row(ws, 32, 2, ["Período", "A vencer (R$)", "Vencido em aberto (R$)", "Total (R$)", "Acumulado (R$)", "Qtd parcelas"])
    for i in range(13):
        r = 33 + i
        src = 100 + i
        ws[f"B{r}"] = f"={R_DG}!B{src}"
        ws[f"C{r}"] = f"={R_DG}!D{src}"
        ws[f"D{r}"] = f"={R_DG}!E{src}"
        ws[f"E{r}"] = f"={R_DG}!F{src}"
        ws[f"F{r}"] = f"={R_DG}!G{src}"
        ws[f"G{r}"] = f"={R_DG}!H{src}"
        for col in "CDEF":
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"G{r}"].number_format = FMT_INT
        style_range(ws, f"B{r}:G{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    write(ws, "B46", "Total em aberto", font(True, NAVY, 9), fill(GRAY2), ALIGN_L, None, BORDER_ALL)
    for col in "CDE":
        write(ws, f"{col}46", f"={R_DG}!{chr(ord(col)+1)}113", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    write(ws, "G46", f"={R_DG}!H113", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_INT, BORDER_ALL)
    write(ws, "B47", f"={R_DG}!B114", font(False, AMBER, 8, italic=True)); ws.merge_cells("B47:O47")
    ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
    ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=32, max_row=45), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=33, max_row=45))
    style_chart(ch, "Compromissos em aberto por mês de vencimento (próximos 12 meses)", 24, 8)
    color_series(ch, [NAVY, RED])
    ws.add_chart(ch, "I32")
    ws.freeze_panes = "C8"
    setup_print(ws, area="B1:O47", header_text="Fluxo de Caixa | Frota")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90


# ====================================================================== COMPETÊNCIA x CAIXA
def build_comp_caixa(wb, ws):
    set_widths(ws, {"A": 2, "B": 30, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18, "H": 18, "I": 18, "J": 18, "K": 16, "L": 16, "M": 16, "N": 16, "O": 16, "P": 16, "Q": 16, "R": 16, "S": 16, "T": 16, "U": 14})
    title_block(ws, "CUSTOS | COMPETÊNCIA x FLUXO DE CAIXA", "Visão 1: custo reconhecido na competência  |  Visão 2: desembolsos realizados e compromissos por vencimento (filtros do Painel aplicados)", "B", "T")
    year_cell(ws, wb, "AnoCC")
    write(ws, "D4", '="Mês de referência dos indicadores mensais: "&INDEX(Meses,MesSel)&"/"&AnoCC&"   |   Base caixa realizado: "&BaseCaixaTxt', font(True, AMBER, 9))
    ws.merge_cells("D4:T4")

    def comp(m):
        return f'AnoCC&"-"&TEXT({m},"00")'

    cs = f'AnoCC&"-"&TEXT(MesSel,"00")'
    kpis = [
        ("Custo total por competência (mês)", f"=SUMIFS({nfI},{nfAQ},1,{nfComp},{cs},{FILT_NF})"),
        ("Custo total pago no mês (líquido)", f"=SUMIFS({p_liq},{p_cons},1,{p_cx},{cs},{FILT_P})"),
        ("Custo previsto para pagamento no mês (a vencer)", f'=SUMIFS({a_val},{a_cons},1,{a_cv},{cs},{a_st},"A vencer",{FILT_A})'),
        ("Obrigações vencidas e não pagas (total)", f'=SUMIFS({a_val},{a_cons},1,{a_st},"Vencido",{FILT_A})'),
        ("Obrigações a vencer (total)", f'=SUMIFS({a_val},{a_cons},1,{a_st},"A vencer",{FILT_A})'),
        ("Total pago acumulado no ano", f'=SUMIFS({p_liq},{p_cons},1,{p_anocx},AnoCC,{FILT_P})'),
        ("Custo por competência acumulado no ano", f"=SUMIFS({nfI},{nfAQ},1,{nfAno},AnoCC,{FILT_NF})"),
        ("Total de obrigações em aberto", f"=SUMIFS({a_val},{a_cons},1,{FILT_A})"),
        ("Comprometimento financeiro futuro (a vencer após a data de ref.)", f'=SUMIFS({a_val},{a_cons},1,{a_st},"A vencer",{FILT_A})'),
    ]
    for i, (t, f) in enumerate(kpis):
        row = 6 + (i // 5) * 3
        col = 2 + (i % 5) * 2
        kpi_card(ws, row, col, t, f, FMT_BRL0, width_cols=2)
    r0 = 13
    section(ws, r0, "Tabela de acompanhamento mensal - perspectivas temporais distintas: NÃO somar entre colunas", "B", "T", NAVY)
    hdr = ["Mês", "Custo por Competência", "Pagamentos Realizados (líquido)", "Pagamentos Previstos (a vencer, por venc.)", "Saldo em Aberto (obrigações pendentes, por competência do custo)",
           "Vencido em aberto (por venc.)", "Pagos (bruto, por venc.)", "Custo acumulado (competência)", "Pagos acumulados", "Previstos acumulados", "Em aberto por vencimento"]
    header_row(ws, r0 + 1, 2, hdr, height=44)
    first = r0 + 2
    for m in range(1, 13):
        r = first + m - 1
        ws[f"B{r}"] = f"=INDEX(Meses,{m})"
        ws[f"C{r}"] = f"=SUMIFS({nfI},{nfAQ},1,{nfComp},{comp(m)},{FILT_NF})"
        ws[f"D{r}"] = f"=SUMIFS({p_liq},{p_cons},1,{p_cx},{comp(m)},{FILT_P})"
        ws[f"E{r}"] = f'=SUMIFS({a_val},{a_cons},1,{a_cv},{comp(m)},{a_st},"A vencer",{FILT_A})'
        ws[f"F{r}"] = f"=SUMIFS({a_val},{a_cons},1,{a_comp},{comp(m)},{FILT_A})"
        ws[f"G{r}"] = f'=SUMIFS({a_val},{a_cons},1,{a_cv},{comp(m)},{a_st},"Vencido",{FILT_A})'
        ws[f"H{r}"] = f"=SUMIFS({p_val},{p_cons},1,{p_cv},{comp(m)},{FILT_P})"
        ws[f"I{r}"] = f"=SUM($C${first}:C{r})"
        ws[f"J{r}"] = f"=SUM($D${first}:D{r})"
        ws[f"K{r}"] = f"=SUM($E${first}:E{r})"
        ws[f"L{r}"] = f"=E{r}+G{r}"
        for col in "CDEFGHIJKL":
            ws[f"{col}{r}"].number_format = FMT_NUM
        style_range(ws, f"B{r}:L{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    tot = first + 12
    write(ws, f"B{tot}", "TOTAL", font(True, NAVY, 9), fill(GRAY2), ALIGN_L, None, BORDER_ALL)
    for col in "CDEFGHL":
        write(ws, f"{col}{tot}", f"=SUM({col}{first}:{col}{tot-1})", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    for col in "IJK":
        write(ws, f"{col}{tot}", f"={col}{tot-1}", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    write(ws, f"B{tot+1}", "Saldo em aberto total (todas as competências, inclusive anteriores ao ano):", font(False, GRAY4, 8, italic=True)); ws.merge_cells(f"B{tot+1}:E{tot+1}")
    write(ws, f"F{tot+1}", f"=SUMIFS({a_val},{a_cons},1,{FILT_A})", font(True, NAVY, 9), fill(GRAY1), ALIGN_R, FMT_NUM, BORDER_ALL)

    # ---- matriz competência x vencimento
    rm = tot + 4
    section(ws, rm, "Matriz: competência do custo (linhas) x mês de vencimento/pagamento das parcelas (colunas) - valores de parcelas conciliadas", "B", "T", NAVY)
    label(ws, f"B{rm+1}", "Mostrar:", True)
    input_cell(ws, f"C{rm+1}", "Total")
    add_list_validation(ws, f"C{rm+1}", "=ListaMostrar")
    write(ws, f"D{rm+1}", "Total = pago (108) + em aberto (1015). Pago = parcelas quitadas (bruto). Em aberto = parcelas pendentes (1015). Filtros de pacote/categoria/fornecedor do Painel aplicados.", font(False, GRAY4, 8, italic=True))
    ws.merge_cells(f"D{rm+1}:T{rm+1}")
    hdrm = ["Competência \\ Vencimento", "Venc. antes do ano"] + [f"=INDEX(MesesAbrev,{m})" for m in range(1, 13)] + ["Venc. após o ano", "Total", "Pago", "Em aberto", "% pago"]
    header_row(ws, rm + 2, 2, hdrm, height=30)
    show = f"$C${rm+1}"
    rows_def = [("Comp. anterior ao ano", "ant")] + [(f"=INDEX(Meses,{m})", m) for m in range(1, 13)] + [("Comp. posterior ao ano", "pos")]
    mfirst = rm + 3
    for i, (lab, key) in enumerate(rows_def):
        r = mfirst + i
        ws[f"B{r}"] = lab
        if key == "ant":
            crit_p = f'{p_anoc},"<"&AnoCC'; crit_a = f'{a_anoc},"<"&AnoCC'
        elif key == "pos":
            crit_p = f'{p_anoc},">"&AnoCC'; crit_a = f'{a_anoc},">"&AnoCC'
        else:
            crit_p = f"{p_comp},{comp(key)}"; crit_a = f"{a_comp},{comp(key)}"
        cols_def = [("ant", "C")] + [(m, get_column_letter(3 + m)) for m in range(1, 13)] + [("pos", "P")]
        for ckey, col in cols_def:
            if ckey == "ant":
                vp = f'{p_anov},"<"&AnoCC'; va = f'{a_anov},"<"&AnoCC'
            elif ckey == "pos":
                vp = f'{p_anov},">"&AnoCC'; va = f'{a_anov},">"&AnoCC'
            else:
                vp = f"{p_cv},{comp(ckey)}"; va = f"{a_cv},{comp(ckey)}"
            pago = f"SUMIFS({p_val},{p_cons},1,{crit_p},{vp},{FILT_P})"
            aberto = f"SUMIFS({a_val},{a_cons},1,{crit_a},{va},{FILT_A})"
            ws[f"{col}{r}"] = f'=IF({show}="Pago",{pago},IF({show}="Em aberto",{aberto},{pago}+{aberto}))'
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"Q{r}"] = f"=SUM(C{r}:P{r})"
        ws[f"R{r}"] = f"=SUMIFS({p_val},{p_cons},1,{crit_p},{FILT_P})"
        ws[f"S{r}"] = f"=SUMIFS({a_val},{a_cons},1,{crit_a},{FILT_A})"
        ws[f"T{r}"] = f'=IF(R{r}+S{r}>0,R{r}/(R{r}+S{r}),"")'
        for col in "QRS":
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"T{r}"].number_format = FMT_PCT
        style_range(ws, f"B{r}:T{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    mt = mfirst + len(rows_def)
    write(ws, f"B{mt}", "TOTAL", font(True, NAVY, 9), fill(GRAY2), ALIGN_L, None, BORDER_ALL)
    for col in [get_column_letter(c) for c in range(3, 20)]:
        write(ws, f"{col}{mt}", f"=SUM({col}{mfirst}:{col}{mt-1})", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    write(ws, f"T{mt}", f'=IF(R{mt}+S{mt}>0,R{mt}/(R{mt}+S{mt}),"")', font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_PCT, BORDER_ALL)
    write(ws, f"B{mt+1}", "Leitura: cada linha mostra em quais meses as parcelas dos custos daquela competência vencem/foram pagas; 'Pago' e 'Em aberto' mostram quanto da competência já foi liquidado e quanto está pendente.", font(False, GRAY4, 8, italic=True))
    ws.merge_cells(f"B{mt+1}:T{mt+1}")
    # conditional shading of heatmap-ish
    from openpyxl.formatting.rule import ColorScaleRule
    ws.conditional_formatting.add(f"C{mfirst}:P{mt-1}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", end_type="max", end_color=LBLUE))

    # ---- gráficos (5)
    rg = mt + 4
    section(ws, rg, "Gráficos - competência x caixa", "B", "T", NAVY)
    last = first + 11
    cats = Reference(ws, min_col=2, min_row=first, max_row=last)

    def line_chart(cols, title, anchor, colors):
        ch = LineChart()
        for c in cols:
            ch.add_data(Reference(ws, min_col=c, min_row=r0 + 1, max_row=last), titles_from_data=True)
        ch.set_categories(cats)
        style_chart(ch, title, 18, 8)
        color_series(ch, colors, line=True)
        ws.add_chart(ch, anchor)

    line_chart([3, 4], "1. Custo por competência x Pagamentos realizados", f"B{rg+1}", [NAVY, GREEN])
    line_chart([3, 5], "2. Custo por competência x Pagamentos previstos (a vencer)", f"H{rg+1}", [NAVY, GOLD])
    ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"
    for c in (3, 4, 12):
        ch.add_data(Reference(ws, min_col=c, min_row=r0 + 1, max_row=last), titles_from_data=True)
    ch.set_categories(cats)
    style_chart(ch, "3. Despesas reconhecidas x Pagamentos realizados x Obrigações em aberto (por venc.)", 18, 8)
    color_series(ch, [NAVY, GREEN, RED])
    ws.add_chart(ch, f"N{rg+1}")
    line_chart([9, 10, 11], "4. Acumulados: custo por competência x pagos x previstos", f"B{rg+18}", [NAVY, GREEN, GOLD])
    ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
    ch.add_data(Reference(wb[S_DG], min_col=4, max_col=5, min_row=99, max_row=112), titles_from_data=True)
    ch.set_categories(Reference(wb[S_DG], min_col=2, min_row=100, max_row=112))
    style_chart(ch, "5. Projeção de desembolsos - próximos 12 meses (a partir da data de referência)", 18, 8)
    color_series(ch, [NAVY, RED])
    ws.add_chart(ch, f"H{rg+18}")
    ws.freeze_panes = "C5"
    setup_print(ws, area=f"B1:T{rg+35}", header_text="Custos | Competência x Caixa")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85


# ====================================================================== CUSTOS POR PACOTE
def build_custos_pacote(wb, ws, dg):
    set_widths(ws, {"A": 2, "B": 30, **{get_column_letter(c): 12 for c in range(3, 15)}, "O": 14, "P": 13, "Q": 10, "R": 13, "S": 13, "T": 13, "U": 10, "V": 30, "W": 18, "X": 9})
    title_block(ws, "ANÁLISE DE CUSTOS | PACOTES", "Custos por competência (Valor Doc) por pacote - sem rateio de remuneração. Detalhamento até categoria, fornecedor e documento.", "B", "X")
    year_cell(ws, wb, "AnoCP")
    label(ws, "B5", "Pacote para detalhamento", True)
    input_cell(ws, "C5", "Manutenção")
    add_list_validation(ws, "C5", "=ListaPacotes")
    write(ws, "D4", '="Mês selecionado: "&INDEX(Meses,MesSel)&". Tabela abaixo sem filtros de pacote/categoria/fornecedor (visão completa do ano)."', font(False, GRAY4, 8, italic=True))
    ws.merge_cells("D4:X4")
    hdr = ["Pacote"] + [f"=INDEX(MesesAbrev,{m})" for m in range(1, 13)] + ["Total ano", "Média mensal", "Part. %", "Mês sel.", "Mês anterior", "Var. M-1 (R$)", "Var. M-1 %", "Principal fornecedor", "Principal categoria", "Meses c/ custo"]
    header_row(ws, 7, 2, hdr, height=30)
    first = 8
    fm_first, fm_last, cm_first, cm_last = dg["forn_matrix_first"], dg["forn_matrix_last"], dg["cat_matrix_first"], dg["cat_matrix_last"]

    def comp(m, ano="AnoCP"):
        return f'{ano}&"-"&TEXT({m},"00")'

    for i in range(N_PAC + 1):
        r = first + i
        is_other = i == N_PAC
        if is_other:
            ws[f"B{r}"] = "Outros / sem classificação"
        else:
            ws[f"B{r}"] = f'=IF(INDEX(ListaPacotes,{i+1})="","",INDEX(ListaPacotes,{i+1}))'
        for m in range(1, 13):
            col = get_column_letter(2 + m)
            if is_other:
                ws[f"{col}{r}"] = f"=SUMIFS({nfI},{nfAQ},1,{nfComp},{comp(m)})-SUM({col}{first}:{col}{r-1})"
            else:
                ws[f"{col}{r}"] = f'=IF($B{r}="",0,SUMIFS({nfI},{nfAQ},1,{nfComp},{comp(m)},{nfPac},$B{r}))'
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"O{r}"] = f"=SUM(C{r}:N{r})"
        ws[f"P{r}"] = f'=IF(X{r}>0,O{r}/X{r},0)'
        ws[f"Q{r}"] = f'=IF($O${first+N_PAC+1}>0,O{r}/$O${first+N_PAC+1},0)'
        ws[f"R{r}"] = f"=INDEX(C{r}:N{r},MesSel)"
        ws[f"S{r}"] = (f'=IF(MesSel>1,INDEX(C{r}:N{r},MesSel-1),' + (f'SUMIFS({nfI},{nfAQ},1,{nfComp},(AnoCP-1)&"-12")-SUM(S{first}:S{r-1})' if is_other else f'IF($B{r}="",0,SUMIFS({nfI},{nfAQ},1,{nfComp},(AnoCP-1)&"-12",{nfPac},$B{r}))') + ')')
        ws[f"T{r}"] = f"=R{r}-S{r}"
        ws[f"U{r}"] = f'=IF(S{r}>0,T{r}/S{r},"")'
        if is_other:
            ws[f"V{r}"] = "-"; ws[f"W{r}"] = "-"
        else:
            mcol = get_column_letter(3 + i)
            ws[f"V{r}"] = f'=IF(OR($B{r}="",MAX({R_DG}!{mcol}{fm_first}:{mcol}{fm_last})=0),"-",INDEX({R_DG}!$B${fm_first}:$B${fm_last},MATCH(MAX({R_DG}!{mcol}{fm_first}:{mcol}{fm_last}),{R_DG}!{mcol}{fm_first}:{mcol}{fm_last},0)))'
            ws[f"W{r}"] = f'=IF(OR($B{r}="",MAX({R_DG}!{mcol}{cm_first}:{mcol}{cm_last})=0),"-",INDEX({R_DG}!$B${cm_first}:$B${cm_last},MATCH(MAX({R_DG}!{mcol}{cm_first}:{mcol}{cm_last}),{R_DG}!{mcol}{cm_first}:{mcol}{cm_last},0)))'
        ws[f"X{r}"] = f'=COUNTIF(C{r}:N{r},">0")'
        for col in "OPRST":
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"Q{r}"].number_format = FMT_PCT; ws[f"U{r}"].number_format = FMT_PCT
        style_range(ws, f"B{r}:X{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    tot = first + N_PAC + 1
    write(ws, f"B{tot}", "TOTAL FROTA", font(True, NAVY, 9), fill(GRAY2), ALIGN_L, None, BORDER_ALL)
    for m in range(1, 13):
        col = get_column_letter(2 + m)
        write(ws, f"{col}{tot}", f"=SUM({col}{first}:{col}{tot-1})", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    write(ws, f"O{tot}", f"=SUM(O{first}:O{tot-1})", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    write(ws, f"P{tot}", f'=IF(X{tot}>0,O{tot}/X{tot},0)', font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    write(ws, f"Q{tot}", f"=SUM(Q{first}:Q{tot-1})", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_PCT, BORDER_ALL)
    write(ws, f"R{tot}", f"=SUM(R{first}:R{tot-1})", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    write(ws, f"S{tot}", f"=SUM(S{first}:S{tot-1})", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    write(ws, f"T{tot}", f"=R{tot}-S{tot}", font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_NUM, BORDER_ALL)
    write(ws, f"U{tot}", f'=IF(S{tot}>0,T{tot}/S{tot},"")', font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_PCT, BORDER_ALL)
    write(ws, f"X{tot}", f'=COUNTIF(C{tot}:N{tot},">0")', font(True, NAVY, 9), fill(GRAY2), ALIGN_R, FMT_INT, BORDER_ALL)
    for col in "VW":
        write(ws, f"{col}{tot}", "", None, fill(GRAY2), None, None, BORDER_ALL)
    cf_negative_red_positive_green(ws, f"U{first}:U{tot}")
    # gráfico evolução pacotes (stacked)
    ch = BarChart(); ch.type = "col"; ch.grouping = "stacked"; ch.overlap = 100
    for i in range(N_PAC + 1):
        r = first + i
        s = Series(Reference(ws, min_col=3, max_col=14, min_row=r), title_from_data=False, title=None)
        from openpyxl.chart.series import SeriesLabel
        from openpyxl.chart.data_source import StrRef
        s.tx = SeriesLabel(strRef=StrRef(f"{R_CP}!$B${r}"))
        ch.series.append(s)
    ch.set_categories(Reference(ws, min_col=3, max_col=14, min_row=7))
    style_chart(ch, "Evolução mensal dos custos por pacote (competência)", 30, 9)
    color_series(ch)
    ws.add_chart(ch, f"B{tot+2}")

    # ---- detalhamento
    rd = tot + 21
    section(ws, rd, "Detalhamento do pacote selecionado (célula C5)", "B", "X", NAVY)
    write(ws, f"B{rd+1}", '="Pacote: "&$C$5&"  |  Ano: "&AnoCP&"  |  Custo do pacote no ano: "&TEXT(SUMIFS(' + nfI + ',' + nfAQ + ',1,' + nfAno + ',AnoCP,' + nfPac + ',$C$5),"#,##0.00")&"  |  Documentos: "&COUNTIFS(' + nfAQ + ',1,' + nfAno + ',AnoCP,' + nfPac + ',$C$5)',
          font(True, NAVY, 10))
    ws.merge_cells(f"B{rd+1}:X{rd+1}")
    # categorias
    header_row(ws, rd + 3, 2, ["Categoria", "Custo (ano)", "Part. % no pacote", "Mês sel."])
    selcol = dg["selcol"]
    for i in range(N_CAT):
        r = rd + 4 + i
        src = cm_first + i
        ws[f"B{r}"] = f"={R_DG}!B{src}"
        ws[f"C{r}"] = f"={R_DG}!{selcol}{src}"
        ws[f"D{r}"] = f'=IF(SUM($C${rd+4}:$C${rd+3+N_CAT})>0,C{r}/SUM($C${rd+4}:$C${rd+3+N_CAT}),0)'
        ws[f"E{r}"] = f'=IF($B{r}="",0,SUMIFS({nfI},{nfAQ},1,{nfComp},{comp("MesSel")},{nfPac},$C$5,{nfCat},$B{r}))'
        ws[f"C{r}"].number_format = FMT_NUM; ws[f"E{r}"].number_format = FMT_NUM; ws[f"D{r}"].number_format = FMT_PCT
        style_range(ws, f"B{r}:E{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    # fornecedores top 10 do pacote
    header_row(ws, rd + 3, 7, ["Rank", "Fornecedor", "Custo (ano)", "Part. % no pacote", "Qtd docs"])
    rkcol = dg["rkcol"]
    for k in range(1, 11):
        r = rd + 3 + k
        ws[f"G{r}"] = k
        ws[f"H{r}"] = f'=IFERROR(INDEX({R_DG}!$B${fm_first}:$B${fm_last},MATCH({k},{R_DG}!${rkcol}${fm_first}:${rkcol}${fm_last},0)),"-")'
        ws[f"I{r}"] = f'=IFERROR(INDEX({R_DG}!${selcol}${fm_first}:${selcol}${fm_last},MATCH({k},{R_DG}!${rkcol}${fm_first}:${rkcol}${fm_last},0)),0)'
        ws[f"J{r}"] = f'=IF(SUM({R_DG}!${selcol}${fm_first}:${selcol}${fm_last})>0,I{r}/SUM({R_DG}!${selcol}${fm_first}:${selcol}${fm_last}),0)'
        ws[f"K{r}"] = f'=IF(H{r}="-",0,COUNTIFS({nfAQ},1,{nfAno},AnoCP,{nfPac},$C$5,{nfFor},H{r}))'
        ws[f"I{r}"].number_format = FMT_NUM; ws[f"J{r}"].number_format = FMT_PCT
        style_range(ws, f"G{r}:K{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    ws.merge_cells(f"H{rd+3}:H{rd+3}")
    # gráfico pizza categorias do pacote
    pie = DoughnutChart()
    pie.add_data(Reference(ws, min_col=3, min_row=rd + 4, max_row=rd + 3 + N_CAT), titles_from_data=False)
    pie.set_categories(Reference(ws, min_col=2, min_row=rd + 4, max_row=rd + 3 + N_CAT))
    style_chart(pie, "Composição do pacote por categoria (ano)", 14, 8, "r")
    pie.holeSize = 55
    ws.add_chart(pie, f"M{rd+3}")
    # documentos
    rdoc = rd + 4 + N_CAT + 2
    header_row(ws, rdoc, 2, ["Documento", "Fornecedor", "Tipo", "Emissão", "Competência", "Categoria", "Valor Doc", "Status documento", "Pago (bruto)", "Em aberto", "Base competência", "Alertas"])
    write(ws, f"B{rdoc-1}", '="Documentos do pacote no ano (até 300 exibidos) - total de documentos: "&MAX(' + _whole(R_NF, NF, "Cum drill") + ')', font(True, NAVY, 9))
    ws.merge_cells(f"B{rdoc-1}:M{rdoc-1}")
    cum = _whole(R_NF, NF, "Cum drill")
    fields = [("Documento", "Documento"), ("Razão Social", "Fornecedor"), ("Tipo", "Tipo"), ("Emissão", "Emissão"), ("Competência", "Competência"), ("Categoria final", "Categoria"),
              ("Valor Doc", "Valor"), ("Status documento", "Status"), ("Valor pago bruto", "Pago"), ("Valor em aberto", "Aberto"), ("Base competência", "Base"), ("Alertas", "Alertas")]
    for k in range(1, 301):
        r = rdoc + k
        for j, (src, _) in enumerate(fields):
            col = get_column_letter(2 + j)
            ws[f"{col}{r}"] = f'=IF({k}>MAX({cum}),"",INDEX({_whole(R_NF,NF,src)},MATCH({k},{cum},0)))'
            c = ws[f"{col}{r}"]
            c.font = font(False, GRAPHITE, 8); c.border = BORDER_ALL
            if src == "Emissão":
                c.number_format = FMT_DATE
            elif src in ("Valor Doc", "Valor pago bruto", "Valor em aberto"):
                c.number_format = FMT_NUM
    ws.freeze_panes = "C8"
    setup_print(ws, area=f"B1:X{rdoc+40}", header_text="Análise de Custos | Pacotes")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85


# ====================================================================== PAGAMENTOS E OBRIGAÇÕES
def build_pagamentos(wb, ws):
    widths = {"A": 2, "B": 34, "C": 16, "D": 9, "E": 7, "F": 7, "G": 8, "H": 16, "I": 16, "J": 13, "K": 13, "L": 11, "M": 11, "N": 13, "O": 11, "P": 11, "Q": 12, "R": 22, "S": 9, "T": 24, "U": 20, "V": 40}
    set_widths(ws, widths)
    title_block(ws, "GESTÃO DE PAGAMENTOS E OBRIGAÇÕES", "Contas a pagar (Rodopar 1015) e parcelas pagas (Rodopar 108) de documentos da Frota - cálculo no nível de parcela", "B", "V")
    label(ws, "B4", "Data de referência", True)
    write(ws, "C4", "=DataRef", font(True, NAVY, 10), fill(GRAY1), ALIGN_C, FMT_DATE, BORDER_ALL)
    label(ws, "B5", "Filtro de status da tabela analítica", True)
    input_cell(ws, "C5", "(Todos)")
    add_list_validation(ws, "C5", "=ListaStatus")
    write(ws, "D4", "Indicadores sem filtros do Painel (visão completa da Frota). Status 'Vencido' = vencimento anterior à data de referência e parcela ainda em aberto no Rodopar 1015.", font(False, GRAY4, 8, italic=True))
    ws.merge_cells("D4:V4")
    kp = [
        ("Total em aberto", f"=SUMIFS({a_val},{a_cons},1)", FMT_BRL0, f'=COUNTIFS({a_cons},1)&" parcelas"'),
        ("Total vencido", f'=SUMIFS({a_val},{a_cons},1,{a_st},"Vencido")', FMT_BRL0, f'=COUNTIFS({a_cons},1,{a_st},"Vencido")&" parcelas"'),
        ("A vencer até 7 dias", f'=SUMIFS({a_val},{a_cons},1,{a_fx},"Até 7 dias")', FMT_BRL0, f'=COUNTIFS({a_cons},1,{a_fx},"Até 7 dias")&" parcelas"'),
        ("A vencer de 8 a 15 dias", f'=SUMIFS({a_val},{a_cons},1,{a_fx},"8 a 15 dias")', FMT_BRL0, f'=COUNTIFS({a_cons},1,{a_fx},"8 a 15 dias")&" parcelas"'),
        ("A vencer de 16 a 30 dias", f'=SUMIFS({a_val},{a_cons},1,{a_fx},"16 a 30 dias")', FMT_BRL0, f'=COUNTIFS({a_cons},1,{a_fx},"16 a 30 dias")&" parcelas"'),
        ("A vencer de 31 a 60 dias", f'=SUMIFS({a_val},{a_cons},1,{a_fx},"31 a 60 dias")', FMT_BRL0, f'=COUNTIFS({a_cons},1,{a_fx},"31 a 60 dias")&" parcelas"'),
        ("A vencer acima de 60 dias", f'=SUMIFS({a_val},{a_cons},1,{a_fx},"Acima de 60 dias")', FMT_BRL0, f'=COUNTIFS({a_cons},1,{a_fx},"Acima de 60 dias")&" parcelas"'),
        ("Parcelas pagas identificadas (ano Painel)", f'=SUMIFS({p_liq},{p_cons},1,{p_anocx},AnoSel)', FMT_BRL0, f'=COUNTIFS({p_cons},1,{p_anocx},AnoSel)&" parcelas (líquido pago)"'),
        ("Documentos com divergência de conciliação", f'=COUNTIFS({nfSt},"Divergência*")', FMT_INT, f'="Diferença total R$ "&TEXT(SUMIFS({_whole(R_NF,NF,"Diferença")},{nfSt},"Divergência*"),"#,##0.00")'),
        ("Parcelas pendentes de classificação (sem NF)", f'=SUMIFS({p_val},{p_esc},"{SCOPE_SEMNF}")+SUMIFS({a_val},{a_esc},"{SCOPE_SEMNF}")', FMT_BRL0, f'=COUNTIFS({p_esc},"{SCOPE_SEMNF}")+COUNTIFS({a_esc},"{SCOPE_SEMNF}")&" parcelas de fornecedores Frota sem NF"'),
    ]
    for i, (t, f, fmt, sub) in enumerate(kp):
        row = 7 + (i // 5) * 3
        positions = [(2, 4), (6, 4), (10, 4), (14, 4), (18, 4)]
        c0, w = positions[i % 5]
        kpi_card(ws, row, c0, t, f, fmt, sub, None, width_cols=w)
    section(ws, 13, "Fornecedores com maiores obrigações em aberto", "B", "V", NAVY)
    header_row(ws, 14, 2, ["Fornecedor", "Em aberto (R$)", "Rank", "", "", "", "Vencido (R$)", "Qtd parcelas abertas"])
    ws.merge_cells("D14:G14")
    fRankA, fNome = _whole(R_FOR, F, "Rank em aberto"), _whole(R_FOR, F, "Fornecedor")
    for k in range(1, 11):
        r = 14 + k
        ws[f"D{r}"] = k
        ws.merge_cells(f"D{r}:G{r}")
        ws[f"B{r}"] = f'=IFERROR(INDEX({fNome},MATCH({k},{fRankA},0)),"-")'
        ws[f"C{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Em aberto (total)")},MATCH({k},{fRankA},0)),0)'
        ws[f"H{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Vencido (total)")},MATCH({k},{fRankA},0)),0)'
        ws[f"I{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Qtd parcelas abertas")},MATCH({k},{fRankA},0)),0)'
        ws[f"C{r}"].number_format = FMT_NUM; ws[f"H{r}"].number_format = FMT_NUM
        style_range(ws, f"B{r}:I{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    ch = BarChart(); ch.type = "bar"
    ch.add_data(Reference(ws, min_col=3, min_row=14, max_row=24), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=15, max_row=24))
    style_chart(ch, "Dez maiores obrigações em aberto por fornecedor", 16, 8)
    ch.x_axis.scaling.orientation = "maxMin"
    color_series(ch, [NAVY])
    ws.add_chart(ch, "K13")

    # ---- tabela analítica
    rt = 28
    section(ws, rt, "Tabela analítica de parcelas - documentos da Frota (em aberto no Rodopar 1015 e pagas no Rodopar 108) - até 1500 linhas", "B", "V", NAVY)
    write(ws, f"B{rt+1}", '="Parcelas exibidas: "&(MAX(' + _whole(R_1015, A, "Cum vista") + ')+MAX(' + _whole(R_108, P, "Cum vista") + '))&"  |  Filtro: "&$C$5&"  |  Data de pagamento preenchida somente quando validada pela extração em ordem original."', font(True, NAVY, 9))
    ws.merge_cells(f"B{rt+1}:V{rt+1}")
    hdr = ["Fornecedor", "Documento", "Série", "Tipo", "Filial", "Parcela", "Pacote", "Categoria", "Valor original doc.", "Valor da parcela", "Desconto", "Juros", "Valor líquido",
           "Emissão", "Vencimento", "Data de pagamento (validada)", "Status financeiro", "Dias em atraso", "Escopo / conciliação", "Competência", "Alertas"]
    header_row(ws, rt + 2, 2, hdr, height=36)
    cumA, cumP = _whole(R_1015, A, "Cum vista"), _whole(R_108, P, "Cum vista")
    nA = f"MAX({cumA})"
    # (campo 1015, campo 108) - None = vazio
    fields = [("Fornecedor (norm.)", "Fornecedor (norm.)", None), ("Documento (norm.)", "Documento (norm.)", None), ("Série", "Série", None), ("Tipo (norm.)", "Tipo (norm.)", None),
              ("Filial", None, None), ("Parcela", None, None), ("Pacote", "Pacote", None), ("Categoria", "Categoria", None),
              ("__docval", "__docval", FMT_NUM), ("Valor parcela", "Valor bruto", FMT_NUM), (None, "Desconto", FMT_NUM), (None, "Juros", FMT_NUM), ("Valor parcela", "Valor líquido pago", FMT_NUM),
              ("__emis", "Emissão (data)", FMT_DATE), ("Vencimento (data)", "Vencimento (data)", FMT_DATE), (None, "Data caixa validada", FMT_DATE), ("Status", "Status", None), ("Dias em atraso", None, FMT_INT),
              ("Escopo", "Escopo", None), ("Competência", "Competência", None), ("Alertas", "Alertas", None)]
    for k in range(1, 1501):
        r = rt + 2 + k
        ia = f"MATCH({k},{cumA},0)"
        ip = f"MATCH({k}-{nA},{cumP},0)"
        for j, (fa, fp, fmt) in enumerate(fields):
            col = get_column_letter(2 + j)
            if fa == "__docval":
                ea = f'IFERROR(INDEX({nfI},INDEX({_whole(R_1015,A,"Linha NF")},{ia})),"")'
                ep = f'IFERROR(INDEX({nfI},INDEX({_whole(R_108,P,"Linha NF")},{ip})),"")'
            elif fa == "__emis":
                ea = f'IFERROR(INDEX({_whole(R_NF,NF,"Emissão")},INDEX({_whole(R_1015,A,"Linha NF")},{ia})),"")'
                ep = f'INDEX({_whole(R_108,P,fp)},{ip})'
            else:
                ea = f'INDEX({_whole(R_1015,A,fa)},{ia})' if fa else '""'
                ep = f'INDEX({_whole(R_108,P,fp)},{ip})' if fp else '""'
            ws[f"{col}{r}"] = f'=IF({k}<={nA},{ea},IF({k}-{nA}<=MAX({cumP}),{ep},""))'
            c = ws[f"{col}{r}"]
            c.font = font(False, GRAPHITE, 8); c.border = BORDER_ALL
            if fmt:
                c.number_format = fmt
    rng = f"R{rt+3}:R{rt+1502}"
    cf_text(ws, rng, "Vencido", RED_L, RED, contains=False)
    cf_text(ws, rng, "Pago", GREEN_L, GREEN, contains=False)
    cf_text(ws, rng, "Divergência", AMBER_L, AMBER)
    cf_text(ws, rng, "pendente", AMBER_L, AMBER)
    ws.freeze_panes = f"C{rt+3}"
    ws.auto_filter.ref = f"B{rt+2}:V{rt+1502}"
    setup_print(ws, area=f"B1:V{rt+60}", header_text="Gestão de Pagamentos e Obrigações")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85


# ====================================================================== ANÁLISE DE FORNECEDORES
def build_fornecedores_analise(wb, ws):
    set_widths(ws, {"A": 2, "B": 7, "C": 42, "D": 16, "E": 16, "F": 15, "G": 9, "H": 10, "I": 8, "J": 14, "K": 14, "L": 13, **{get_column_letter(c): 11 for c in range(13, 25)}, "Y": 8, "Z": 26})
    title_block(ws, "ANÁLISE DE FORNECEDORES", "Concentração financeira por fornecedor - custos por competência (ano), pagamentos identificados e obrigações em aberto", "B", "Z")
    year_cell(ws, wb, "AnoAF")
    label(ws, "B6", "Custo total da Frota no ano (base de participação)", True)
    ws.merge_cells("B6:B6")
    write(ws, "C7", f"=SUMIFS({nfI},{nfAQ},1,{nfAno},AnoAF)", font(True, NAVY, 10), fill(GRAY1), ALIGN_R, FMT_BRL, BORDER_ALL)
    label(ws, "B7", "Custo total (ano)", True)
    fRank, fPart, fCusto = _whole(R_FOR, F, "Rank custo"), _whole(R_FOR, F, "Participação %"), _whole(R_FOR, F, "Custo competência (ano)")
    kp = [
        ("Fornecedores com custo no ano", f'=COUNTIF({fRank},">0")', FMT_INT, None),
        ("Participação do maior fornecedor", f'=IFERROR(INDEX({fPart},MATCH(1,{fRank},0)),0)', FMT_PCT, f'=IFERROR(INDEX({_whole(R_FOR,F,"Fornecedor")},MATCH(1,{fRank},0)),"-")'),
        ("Participação dos 3 maiores", f'=SUMPRODUCT(({R_FOR}!$O$2:$O$2000<=3)*({R_FOR}!$O$2:$O$2000>0)*N(+{R_FOR}!$H$2:$H$2000))', FMT_PCT, None),
        ("Participação dos 5 maiores", f'=SUMPRODUCT(({R_FOR}!$O$2:$O$2000<=5)*({R_FOR}!$O$2:$O$2000>0)*N(+{R_FOR}!$H$2:$H$2000))', FMT_PCT, None),
        ("Fornecedores acima do limite de concentração", f'=COUNTIF({fPart},">"&LimConc)', FMT_INT, '="Limite: "&TEXT(LimConc,"0%")&" (Parâmetros)"'),
    ]
    for i, (t, f, fmt, sub) in enumerate(kp):
        kpi_card(ws, 4 + (0 if i < 5 else 3), 4 + i * 3 if i < 5 else 4 + (i - 5) * 3, t, f, fmt, sub, None, width_cols=3)
    section(ws, 9, "Ranking por custo de competência no ano (50 maiores) - sinalizações: concentração acima do limite e custo recorrente", "B", "Z", NAVY)
    hdr = ["Rank", "Fornecedor", "Pacote", "Categoria", "Custo competência (ano)", "Part. %", "Part. acum.", "Qtd docs", "Pago líquido (ano)", "Em aberto", "Vencido"] + MESES_ABREV + ["Meses c/ custo", "Sinalização"]
    header_row(ws, 10, 2, hdr, height=30)
    for k in range(1, 51):
        r = 10 + k
        m = f"MATCH({k},{fRank},0)"
        ws[f"B{r}"] = k
        ws[f"C{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Fornecedor")},{m}),"")'
        ws[f"D{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Pacote")},{m}),"")'
        ws[f"E{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Categoria")},{m}),"")'
        ws[f"F{r}"] = f'=IFERROR(INDEX({fCusto},{m}),"")'
        ws[f"G{r}"] = f'=IFERROR(INDEX({fPart},{m}),"")'
        ws[f"H{r}"] = f'=IF(C{r}="","",SUM($G$11:G{r}))'
        ws[f"I{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Qtd docs (ano)")},{m}),"")'
        ws[f"J{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Pago líquido (ano, caixa)")},{m}),"")'
        ws[f"K{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Em aberto (total)")},{m}),"")'
        ws[f"L{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Vencido (total)")},{m}),"")'
        for j, ma in enumerate(MESES_ABREV):
            col = get_column_letter(13 + j)
            ws[f"{col}{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,ma)},{m}),"")'
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"Y{r}"] = f'=IFERROR(INDEX({_whole(R_FOR,F,"Meses com custo")},{m}),"")'
        ws[f"Z{r}"] = (f'=IF(C{r}="","",IF(N(G{r})>LimConc,"CONCENTRAÇÃO ELEVADA; ","")&IF(N(Y{r})>=6,"Custo recorrente; ","")'
                       f'&IF(N(L{r})>0,"Vencido em aberto; ",""))')
        for col in "FJKL":
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"G{r}"].number_format = FMT_PCT; ws[f"H{r}"].number_format = FMT_PCT
        style_range(ws, f"B{r}:Z{r}", font(False, GRAPHITE, 9), None, None, BORDER_ALL)
    cf_text(ws, "Z11:Z60", "CONCENTRAÇÃO", RED_L, RED)
    ch = BarChart(); ch.type = "bar"
    ch.add_data(Reference(ws, min_col=6, min_row=10, max_row=20), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=3, min_row=11, max_row=20))
    style_chart(ch, "Dez maiores fornecedores por custo de competência (ano)", 18, 9)
    ch.x_axis.scaling.orientation = "maxMin"
    color_series(ch, [NAVY])
    ws.add_chart(ch, "B63")
    pie = DoughnutChart()
    pie.add_data(Reference(ws, min_col=6, min_row=11, max_row=20), titles_from_data=False)
    pie.set_categories(Reference(ws, min_col=3, min_row=11, max_row=20))
    style_chart(pie, "Participação dos dez maiores no custo (ano)", 14, 9, "r")
    pie.holeSize = 50
    ws.add_chart(pie, "L63")
    ws.freeze_panes = "D11"
    setup_print(ws, area="B1:Z82", header_text="Análise de Fornecedores")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85


# ====================================================================== CONCILIAÇÃO E QUALIDADE
def build_conciliacao(wb, ws):
    set_widths(ws, {"A": 2, "B": 62, "C": 12, "D": 18, "E": 14, "F": 44, "G": 14, "H": 14, "I": 16, "J": 20, "K": 60})
    title_block(ws, "CONCILIAÇÃO E QUALIDADE DOS DADOS", "Auditoria das bases (Relação de NF's, Fornecedores, Rodopar 108 e 1015) - nenhuma divergência é mascarada", "B", "K")
    nfB = f"{R_NF}!$B:$B"
    nfA = f"{R_NF}!$A:$A"
    nfAl = _whole(R_NF, NF, "Alertas")
    nfVal = _whole(R_NF, NF, "Registro válido")
    nfOc = _whole(R_NF, NF, "Ocorrência")
    nfAmb = _whole(R_NF, NF, "Ambíguo N2")
    fA, fC, fD, fE = f"{R_FOR}!$A:$A", f"{R_FOR}!$C:$C", f"{R_FOR}!$D:$D", f"{R_FOR}!$E:$E"
    pC, pVal, pOc, pNv, pSt = f"{R_108}!$C:$C", _whole(R_108, P, "Registro válido"), _whole(R_108, P, "Ocorrência"), _whole(R_108, P, "Nível conciliação"), _whole(R_108, P, "Status")
    aD, aVal, aOc, aNv = f"{R_1015}!$D:$D", _whole(R_1015, A, "Registro válido"), _whole(R_1015, A, "Ocorrência"), _whole(R_1015, A, "Nível conciliação")
    section(ws, 4, "Resumo da qualidade dos dados", "B", "K", NAVY)
    kp = [
        ("Documentos na Relação de NF's", f"=COUNTA({nfB})-1", FMT_INT, f'="Valor: R$ "&TEXT(SUM({nfI}),"#,##0")'),
        ("Documentos íntegros (considerados, sem alertas)", f'=COUNTIFS({nfAQ},1,{nfAl},"")', FMT_INT, f'="R$ "&TEXT(SUMIFS({nfI},{nfAQ},1,{nfAl},""),"#,##0")'),
        ("Documentos com pendências/alertas", f'=COUNTIF({nfAl},"?*")', FMT_INT, f'="R$ "&TEXT(SUMIF({nfAl},"?*",{nfI}),"#,##0")'),
        ("Parcelas Frota com alertas (108 + 1015)", f'=MAX({_whole(R_108,P,"Cum auditoria")})+MAX({_whole(R_1015,A,"Cum auditoria")})', FMT_INT, None),
        ("Fornecedores com alertas de cadastro", f'=MAX({_whole(R_FOR,F,"Cum alertas")})', FMT_INT, None),
    ]
    for i, (t, f, fmt, sub) in enumerate(kp):
        kpi_card(ws, 5, 2 + i * 2 if i == 0 else 2 + i * 2, t, f, fmt, sub, None, width_cols=2)
    section(ws, 9, "Verificações de integridade e conciliação", "B", "K", NAVY)
    header_row(ws, 10, 2, ["Verificação", "Quantidade", "Valor (R$)", "Severidade", "Onde verificar / tratamento"])
    checks = [
        ("Documentos com fornecedor não cadastrado (Relação de NF's)", f'=COUNTIF({nfA},"{NAO_CAD}")', f'=SUMIFS({nfI},{nfA},"{NAO_CAD}")', "Alta", "Cadastrar o fornecedor na aba Fornecedores (Area, Pacote, Categoria)."),
        ("Documentos de fornecedores de outras áreas (excluídos do resultado)", f'=COUNTIFS({nfA},"<>Frota",{nfA},"<>{NAO_CAD}",{nfVal},1)', f'=SUMIFS({nfI},{nfA},"<>Frota",{nfA},"<>{NAO_CAD}",{nfVal},1)', "Média", "Confirmar área no cadastro de fornecedores."),
        ("Fornecedores Frota sem pacote", f'=COUNTIFS({fC},"Frota",{fD},"")', "", "Alta", "Aba Fornecedores - preencher Pacote."),
        ("Fornecedores Frota sem categoria", f'=COUNTIFS({fC},"Frota",{fE},"")', "", "Média", "Aba Fornecedores - preencher Categoria."),
        ("Documentos sem pacote (fornecedor e documento sem classificação)", f'=COUNTIFS({nfPac},"{SEM_CLASS}",{nfAQ},1)', f'=SUMIFS({nfI},{nfPac},"{SEM_CLASS}",{nfAQ},1)', "Alta", "Preencher 'Pacote (ajuste manual)' na NF ou o cadastro do fornecedor."),
        ("Documentos duplicados na Relação de NF's (mesma chave N1)", f'=COUNTIF({nfOc},">1")', f'=SUMIFS({nfI},{nfOc},">1")', "Alta", "Segunda ocorrência é desconsiderada; remover a linha duplicada."),
        ("Parcelas pagas duplicadas na extração 108 (Frota)", f'=COUNTIFS({pOc},">1",{p_esc},"Frota*")', f'=SUMIFS({p_val},{pOc},">1",{p_esc},"Frota*")', "Alta", "Reimportação em duplicidade - desconsideradas automaticamente."),
        ("Parcelas em aberto duplicadas na extração 1015 (Frota)", f'=COUNTIFS({aOc},">1",{a_esc},"Frota*")', f'=SUMIFS({a_val},{aOc},">1",{a_esc},"Frota*")', "Alta", "Reimportação em duplicidade - desconsideradas automaticamente."),
        ("Documentos da Frota sem correspondência nas bases 108/1015", f'=COUNTIF({nfSt},"Sem correspondência*")', f'=SUMIFS({nfI},{nfSt},"Sem correspondência*")', "Média", "Documento ainda não lançado no financeiro ou extração 1015 incompleta."),
        ("Documentos com divergência de valor (documento x soma das parcelas)", f'=COUNTIF({nfSt},"Divergência*")', f'=SUMIFS({_whole(R_NF,NF,"Diferença")},{nfSt},"Divergência*")', "Alta", "Verificar retenções na fonte e parcelas com vencimento fora do período extraído no 1015."),
        ("Parcelas pagas com retenção/diferença entre parcela e valor pago", f'=COUNTIFS({p_cons},1,{p_ret},"<>0")', f'=SUMIFS({p_ret},{p_cons},1)', "Info", "Diferença = retenções na fonte/ajustes; custo permanece pelo valor do documento."),
        ("Parcelas que constam como pagas (108) e em aberto (1015)", f'=COUNTIF({a_st},"Divergência*")', f'=SUMIFS({a_val},{a_st},"Divergência*")', "Alta", "Status não conclusivo - confirmar no Rodopar."),
        ("Correspondências ambíguas (mesma chave Doc+Forn+Tipo em mais de uma NF)", f'=COUNTIF({nfAmb},1)', f'=SUMIFS({nfI},{nfAmb},1)', "Média", "Conciliação por N2 pode associar parcela ao documento errado - revisar."),
        ("Registros inválidos na Relação de NF's (sem valor/data/fornecedor)", f'=COUNTIFS({nfB},"<>",{nfVal},0)', "", "Alta", "Corrigir o registro na Relação de NF's."),
        ("Linhas inválidas na extração 108 (exceto cabeçalhos e grupos)", f'=COUNTIFS({pC},"<>",{pC},"<>Documento",{pC},"<>Doc. Pago em*",{pC},"<>Filial*",{pC},"<>Total do Dia*",{pVal},0)', "", "Média", "Datas/valores não numéricos - verificar colagem (linhas 'Filial', 'Doc. Pago em', 'Documento' e 'Total do Dia' são estruturais e ignoradas)."),
        ("Linhas inválidas na extração 1015 (exceto cabeçalhos)", f'=COUNTIFS({aD},"<>",{aD},"<>Doc.",{aVal},0)', "", "Média", "Datas/valores não numéricos - verificar colagem."),
        ("Movimentos de fornecedores Frota SEM NF na relação (108)", f'=COUNTIF({p_esc},"{SCOPE_SEMNF}")', f'=SUMIFS({p_val},{p_esc},"{SCOPE_SEMNF}")', "Média", "Fila de revisão: incluir NF na relação (se da Frota) ou ajustar área do fornecedor. Não entram no resultado."),
        ("Movimentos de fornecedores Frota SEM NF na relação (1015)", f'=COUNTIF({a_esc},"{SCOPE_SEMNF}")', f'=SUMIFS({a_val},{a_esc},"{SCOPE_SEMNF}")', "Média", "Idem - obrigações não consideradas até classificação."),
        ("Registros de outras áreas nas bases Rodopar (108 + 1015)", f'=COUNTIF({p_esc},"{SCOPE_OUTRA}")+COUNTIF({a_esc},"{SCOPE_OUTRA}")', f'=SUMIFS({p_val},{p_esc},"{SCOPE_OUTRA}")+SUMIFS({a_val},{a_esc},"{SCOPE_OUTRA}")', "Info", "Excluídos do resultado da Frota (fornecedor com Area = '-')."),
        ("Registros de fornecedores não cadastrados nas bases Rodopar", f'=COUNTIF({p_esc},"{SCOPE_NC}")+COUNTIF({a_esc},"{SCOPE_NC}")', f'=SUMIFS({p_val},{p_esc},"{SCOPE_NC}")+SUMIFS({a_val},{a_esc},"{SCOPE_NC}")', "Info", "Excluídos; presumidos de outras áreas (folha, tributos etc.)."),
        ("Documentos com competência estimada pelo mês de emissão", f'=COUNTIFS({nfAQ},1,{nfBase},"Estimada*")', f'=SUMIFS({nfI},{nfAQ},1,{nfBase},"Estimada*")', "Info", "Informe 'Competência (ajuste manual)' quando a competência real for diferente."),
        ("Parcelas pagas sem data efetiva de pagamento (mês de vencimento usado)", f'=COUNTIFS({p_cons},1,{_whole(R_108,P,"Base data caixa")},"Mês de vencimento (estimado)")', f'=SUMIFS({p_liq},{p_cons},1,{_whole(R_108,P,"Base data caixa")},"Mês de vencimento (estimado)")', "Média", "Cole o Rodopar 108 na ordem original (com grupos 'Doc. Pago em') para habilitar a data efetiva."),
        ("Competências com remuneração duplicada", f'=SUMPRODUCT(({R_REM}!$K$8:$K$2000="COMPETÊNCIA DUPLICADA")*1)', "", "Alta", "Aba Remuneração - manter uma linha por competência."),
        ("Competências do ano com remuneração não informada (até o mês selecionado)", f'=MesSel-COUNTIFS({R_REM}!$B:$B,AnoSel,{R_REM}!$C:$C,"<="&MesSel,{R_REM}!$F:$F,"<>")', "", "Média", "Aba Remuneração - informar o valor do mês."),
    ]
    for i, (t, qf, vf, sev, onde) in enumerate(checks):
        r = 11 + i
        write(ws, f"B{r}", t, font(False, GRAPHITE, 9), None, ALIGN_LW, None, BORDER_ALL)
        write(ws, f"C{r}", qf, font(True, GRAPHITE, 9), None, ALIGN_R, FMT_INT, BORDER_ALL)
        write(ws, f"D{r}", vf if vf else "", font(False, GRAPHITE, 9), None, ALIGN_R, FMT_NUM, BORDER_ALL)
        write(ws, f"E{r}", sev, font(True, RED if sev == "Alta" else (AMBER if sev == "Média" else GRAY4), 9), None, ALIGN_C, None, BORDER_ALL)
        write(ws, f"F{r}", onde, font(False, GRAY4, 8), None, ALIGN_LW, None, BORDER_ALL)
        ws.merge_cells(f"F{r}:K{r}")
        ws.row_dimensions[r].height = 24
    rc = 11 + len(checks)
    ws.conditional_formatting.add(f"C11:C{rc-1}", FormulaRule(formula=[f'AND(C11>0,$E11="Alta")'], fill=fill(RED_L), font=Font(name=FONT, color=RED, bold=True)))
    ws.conditional_formatting.add(f"C11:C{rc-1}", FormulaRule(formula=[f'AND(C11>0,$E11="Média")'], fill=fill(AMBER_L), font=Font(name=FONT, color=AMBER, bold=True)))

    # ---- reconciliação de totais (Teste 10)
    rt = rc + 1
    section(ws, rt, "Conciliação dos totais de origem x totais considerados no modelo", "B", "K", NAVY)
    header_row(ws, rt + 1, 2, ["Base / recorte", "Linhas", "Valor (R$)", "", "Observação"])
    ws.merge_cells(f"F{rt+1}:K{rt+1}")
    recon = [
        ("Rodopar 108 - linhas totais coladas (inclui cabeçalhos e grupos)", f'=COUNTA({pC})-1', "", "Extração bruta."),
        ("Rodopar 108 - parcelas válidas", f'=COUNTIF({pVal},1)', f'=SUMIFS({p_val},{pVal},1)', "Linhas com documento, vencimento e valor numéricos."),
        ("   das quais: Frota com NF vinculada (consideradas)", f'=COUNTIF({p_cons},1)', f'=SUMIFS({p_val},{p_cons},1)', "Base dos pagamentos realizados da Frota."),
        ("   das quais: Frota sem NF (fila de revisão)", f'=COUNTIF({p_esc},"{SCOPE_SEMNF}")', f'=SUMIFS({p_val},{p_esc},"{SCOPE_SEMNF}")', "Não consideradas até classificação."),
        ("   das quais: outras áreas", f'=COUNTIF({p_esc},"{SCOPE_OUTRA}")', f'=SUMIFS({p_val},{p_esc},"{SCOPE_OUTRA}")', "Fora do escopo."),
        ("   das quais: fornecedor não cadastrado", f'=COUNTIF({p_esc},"{SCOPE_NC}")', f'=SUMIFS({p_val},{p_esc},"{SCOPE_NC}")', "Fora do escopo."),
        ("   das quais: duplicadas (Frota)", f'=COUNTIFS({pOc},">1",{p_esc},"Frota*")', f'=SUMIFS({p_val},{pOc},">1",{p_esc},"Frota*")', "Desconsideradas."),
        ("Rodopar 1015 - linhas totais coladas", f'=COUNTA({aD})-1', "", "Extração bruta."),
        ("Rodopar 1015 - parcelas válidas", f'=COUNTIF({aVal},1)', f'=SUMIFS({a_val},{aVal},1)', "Campo 'Acumulado' não é somado."),
        ("   das quais: Frota com NF vinculada (consideradas)", f'=COUNTIF({a_cons},1)', f'=SUMIFS({a_val},{a_cons},1)', "Base das obrigações em aberto da Frota."),
        ("   das quais: Frota sem NF (fila de revisão)", f'=COUNTIF({a_esc},"{SCOPE_SEMNF}")', f'=SUMIFS({a_val},{a_esc},"{SCOPE_SEMNF}")', "Não consideradas até classificação."),
        ("   das quais: outras áreas / não cadastrados", f'=COUNTIF({a_esc},"{SCOPE_OUTRA}")+COUNTIF({a_esc},"{SCOPE_NC}")', f'=SUMIFS({a_val},{a_esc},"{SCOPE_OUTRA}")+SUMIFS({a_val},{a_esc},"{SCOPE_NC}")', "Fora do escopo."),
        ("Relação de NF's - documentos", f'=COUNTA({nfB})-1', f'=SUM({nfI})', "Base dos custos por competência."),
        ("   dos quais: considerados na DRE", f'=COUNTIF({nfAQ},1)', f'=SUMIFS({nfI},{nfAQ},1)', "Válidos, Frota, sem duplicidade."),
        ("   dos quais: excluídos (inválidos, outras áreas, duplicados)", f'=COUNTA({nfB})-1-COUNTIF({nfAQ},1)', f'=SUM({nfI})-SUMIFS({nfI},{nfAQ},1)', "Listados acima por motivo."),
        ("Soma das parcelas conciliadas (108 pago + 1015 aberto) dos documentos considerados", "", f'=SUMIFS({_whole(R_NF,NF,"Total conciliado")},{nfAQ},1)', "Deve aproximar-se do valor dos documentos; diferença abaixo = retenções e parcelas fora da extração."),
        ("Diferença total (documentos considerados - parcelas conciliadas)", "", f'=SUMIFS({_whole(R_NF,NF,"Diferença")},{nfAQ},1)', "Detalhada por documento na lista de alertas."),
    ]
    for i, (t, qf, vf, obs) in enumerate(recon):
        r = rt + 2 + i
        write(ws, f"B{r}", t, font(t.startswith("   ") is False, GRAPHITE, 9), None, ALIGN_LW, None, BORDER_ALL)
        write(ws, f"C{r}", qf, font(False, GRAPHITE, 9), None, ALIGN_R, FMT_INT, BORDER_ALL)
        write(ws, f"D{r}", vf, font(False, GRAPHITE, 9), None, ALIGN_R, FMT_NUM, BORDER_ALL)
        write(ws, f"E{r}", "", None, None, None, None, BORDER_ALL)
        write(ws, f"F{r}", obs, font(False, GRAY4, 8), None, ALIGN_LW, None, BORDER_ALL)
        ws.merge_cells(f"F{r}:K{r}")

    # ---- listas de revisão
    rl = rt + 2 + len(recon) + 2
    section(ws, rl, "Fila de revisão 1 - Documentos com alertas (até 300)", "B", "K", NAVY2)
    header_row(ws, rl + 1, 2, ["Documento | Fornecedor", "Tipo", "Emissão", "Valor Doc", "Status documento", "Pago (bruto)", "Em aberto", "Diferença", "Pacote", "Alertas"])
    cum = _whole(R_NF, NF, "Cum auditoria")
    for k in range(1, 301):
        r = rl + 1 + k
        m = f"MATCH({k},{cum},0)"
        g = lambda src: f'INDEX({_whole(R_NF,NF,src)},{m})'
        ws[f"B{r}"] = f'=IF({k}>MAX({cum}),"",{g("Documento")}&" | "&{g("Razão Social")})'
        ws[f"C{r}"] = f'=IF({k}>MAX({cum}),"",{g("Tipo")})'
        ws[f"D{r}"] = f'=IF({k}>MAX({cum}),"",{g("Emissão")})'
        ws[f"E{r}"] = f'=IF({k}>MAX({cum}),"",{g("Valor Doc")})'
        ws[f"F{r}"] = f'=IF({k}>MAX({cum}),"",{g("Status documento")})'
        ws[f"G{r}"] = f'=IF({k}>MAX({cum}),"",{g("Valor pago bruto")})'
        ws[f"H{r}"] = f'=IF({k}>MAX({cum}),"",{g("Valor em aberto")})'
        ws[f"I{r}"] = f'=IF({k}>MAX({cum}),"",{g("Diferença")})'
        ws[f"J{r}"] = f'=IF({k}>MAX({cum}),"",{g("Pacote final")})'
        ws[f"K{r}"] = f'=IF({k}>MAX({cum}),"",{g("Alertas")})'
        for col in "BCDEFGHIJK":
            c = ws[f"{col}{r}"]; c.font = font(False, GRAPHITE, 8); c.border = BORDER_ALL
        ws[f"D{r}"].number_format = FMT_DATE
        for col in "EGHI":
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"K{r}"].alignment = ALIGN_LW
    rl2 = rl + 304
    section(ws, rl2, "Fila de revisão 2 - Parcelas com alertas: em aberto (1015) e pagas (108) de fornecedores Frota (até 300)", "B", "K", NAVY2)
    header_row(ws, rl2 + 1, 2, ["Documento | Fornecedor", "Origem", "Vencimento", "Valor parcela", "Status", "Escopo", "Nível conciliação", "Competência", "Pacote", "Alertas"])
    cumA, cumP = _whole(R_1015, A, "Cum auditoria"), _whole(R_108, P, "Cum auditoria")
    nA = f"MAX({cumA})"
    for k in range(1, 301):
        r = rl2 + 1 + k
        ia = f"MATCH({k},{cumA},0)"
        ip = f"MATCH({k}-{nA},{cumP},0)"
        ga = lambda src: f'INDEX({_whole(R_1015,A,src)},{ia})'
        gp = lambda src: f'INDEX({_whole(R_108,P,src)},{ip})'
        def both(ea, ep):
            return f'=IF({k}<={nA},{ea},IF({k}-{nA}<=MAX({cumP}),{ep},""))'
        ws[f"B{r}"] = both(f'{ga("Documento (norm.)")}&" | "&{ga("Fornecedor (norm.)")}', f'{gp("Documento (norm.)")}&" | "&{gp("Fornecedor (norm.)")}')
        ws[f"C{r}"] = both('"1015 - em aberto"', '"108 - pago"')
        ws[f"D{r}"] = both(ga("Vencimento (data)"), gp("Vencimento (data)"))
        ws[f"E{r}"] = both(ga("Valor parcela"), gp("Valor bruto"))
        ws[f"F{r}"] = both(ga("Status"), gp("Status"))
        ws[f"G{r}"] = both(ga("Escopo"), gp("Escopo"))
        ws[f"H{r}"] = both(ga("Nível conciliação"), gp("Nível conciliação"))
        ws[f"I{r}"] = both(ga("Competência"), gp("Competência"))
        ws[f"J{r}"] = both(ga("Pacote"), gp("Pacote"))
        ws[f"K{r}"] = both(ga("Alertas"), gp("Alertas"))
        for col in "BCDEFGHIJK":
            c = ws[f"{col}{r}"]; c.font = font(False, GRAPHITE, 8); c.border = BORDER_ALL
        ws[f"D{r}"].number_format = FMT_DATE; ws[f"E{r}"].number_format = FMT_NUM
        ws[f"K{r}"].alignment = ALIGN_LW
    rl3 = rl2 + 304
    section(ws, rl3, "Fila de revisão 3 - Fornecedores com alertas de cadastro (até 60)", "B", "K", NAVY2)
    header_row(ws, rl3 + 1, 2, ["Fornecedor", "Área", "Pacote", "Categoria", "Custo (ano)", "Em aberto", "Mov. sem NF", "", "", "Alertas"])
    cumF = _whole(R_FOR, F, "Cum alertas")
    for k in range(1, 61):
        r = rl3 + 1 + k
        m = f"MATCH({k},{cumF},0)"
        gf = lambda src: f'IF({k}>MAX({cumF}),"",INDEX({_whole(R_FOR,F,src)},{m}))'
        ws[f"B{r}"] = f"={gf('Fornecedor')}"
        ws[f"C{r}"] = f"={gf('Area')}"
        ws[f"D{r}"] = f"={gf('Pacote')}"
        ws[f"E{r}"] = f"={gf('Categoria')}"
        ws[f"F{r}"] = f"={gf('Custo competência (ano)')}"
        ws[f"G{r}"] = f"={gf('Em aberto (total)')}"
        ws[f"H{r}"] = f"={gf('Movimento sem NF (revisar)')}"
        ws[f"K{r}"] = f"={gf('Alertas cadastro')}"
        for col in "BCDEFGHIJK":
            c = ws[f"{col}{r}"]; c.font = font(False, GRAPHITE, 8); c.border = BORDER_ALL
        for col in "FGH":
            ws[f"{col}{r}"].number_format = FMT_NUM
        ws[f"K{r}"].alignment = ALIGN_LW
    ws.freeze_panes = "B4"
    setup_print(ws, area=f"B1:K{rt + 2 + len(recon)}", header_text="Conciliação e Qualidade dos Dados")
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
